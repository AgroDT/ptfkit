"""Manifest-aware audit of workbook predictors for retention candidates."""

from __future__ import annotations

import hashlib
import json
import math
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import asdict, dataclass
from numbers import Real
from pathlib import Path
from typing import TYPE_CHECKING, TypeGuard, cast
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from ptfkit import get_manifest

from .retention import LabelCount


if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook

    from .datasets import DatasetMapping, JsonScalar


class PredictorAuditError(ValueError):
    """Raised when predictor availability cannot be audited safely."""


@dataclass(frozen=True)
class CanonicalPredictorValue:
    """One manifest input candidate with raw-cell lineage and gating state."""

    sample_key: str
    sample_id: JsonScalar
    source: JsonScalar
    sample_number: JsonScalar
    function: str
    input: str
    raw_value: JsonScalar
    value: JsonScalar
    status: str
    blockers: tuple[str, ...]
    source_cells: tuple[str, ...]


@dataclass(frozen=True)
class FunctionSampleEligibility:
    """Predictor-only eligibility for one function and one master-sheet row."""

    sample_key: str
    sample_id: JsonScalar
    source: JsonScalar
    sample_number: JsonScalar
    function: str
    eligible: bool
    reasons: tuple[str, ...]


@dataclass(frozen=True)
class PredictorBindingAudit:
    """Count-only result for one declared function-input binding."""

    function: str
    input: str
    kind: str
    worksheet: str | None
    column: str | None
    binding_status: str
    candidate_sample_count: int
    populated_count: int
    compatible_count: int
    measurement_compatible_count: int
    particle_size_compatible_count: int | None
    missing_count: int
    excluded_count: int
    blockers: tuple[str, ...]


@dataclass(frozen=True)
class FunctionPredictorAudit:
    """Count-only predictor eligibility for one public ptfkit function."""

    function: str
    input_count: int
    candidate_sample_count: int
    complete_input_sample_count: int
    eligible_sample_count: int
    excluded_sample_count: int


@dataclass(frozen=True)
class PredictorAuditSummary:
    """Non-sensitive evidence for a manifest-aware predictor audit."""

    schema_version: int
    source_filename: str
    sha256: str
    manifest_schema_version: int
    public_function_count: int
    candidate_sample_count: int
    binding_audits: tuple[PredictorBindingAudit, ...]
    function_audits: tuple[FunctionPredictorAudit, ...]
    exclusion_reasons: tuple[LabelCount, ...]


@dataclass(frozen=True)
class PredictorAuditResult:
    """In-memory predictor records, eligibility, and count-only summary."""

    summary: PredictorAuditSummary
    values: tuple[CanonicalPredictorValue, ...]
    eligibility: tuple[FunctionSampleEligibility, ...]


@dataclass(frozen=True)
class _Sample:
    worksheet: str
    row: int
    sample_id: JsonScalar
    source: JsonScalar
    sample_number: JsonScalar

    @property
    def key(self) -> str:
        return f'{self.worksheet}!{self.row}'


@dataclass(frozen=True)
class _Binding:
    function: str
    input: str
    worksheet: str
    column: str
    status: str
    blockers: tuple[str, ...]
    category_map: Mapping[str, str]
    measurement_overrides: tuple[_MeasurementOverride, ...]


@dataclass(frozen=True)
class _MeasurementOverride:
    rows: frozenset[int]
    measurement: Mapping[str, object]
    source_headers: tuple[_SourceHeader, ...]


@dataclass(frozen=True)
class _SourceHeader:
    worksheet: str
    column: str
    row: int
    value: str

    @property
    def cell(self) -> str:
        return f'{self.worksheet}!{self.column}{self.row}'


def audit_predictor_workbook(
    source: str | Path,
    mapping: DatasetMapping,
    *,
    manifest: Mapping[str, object] | None = None,
) -> PredictorAuditResult:
    """Audit predictors only for rows declared as retention sample blocks."""
    path = Path(source)
    digest = _sha256(path)
    _validate_source(path, digest, mapping)
    manifest_document = get_manifest() if manifest is None else manifest
    schema_version, functions = _manifest_functions(manifest_document)
    bindings = _bindings(mapping.document)
    _validate_bindings_against_manifest(bindings, functions)
    binding_index = {(item.function, item.input): item for item in bindings}
    sample_rows = _candidate_rows(mapping.document)
    column_contracts = _worksheet_columns(mapping.document)

    values: list[CanonicalPredictorValue] = []
    eligibility: list[FunctionSampleEligibility] = []
    binding_audits: list[PredictorBindingAudit] = []
    function_audits: list[FunctionPredictorAudit] = []
    reason_counts: Counter[str] = Counter()

    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError) as error:
        message = f'Cannot audit workbook {path.name!r}: {error}'
        raise PredictorAuditError(message) from error

    try:
        _validate_binding_headers(workbook, bindings, column_contracts)
        samples = _samples(workbook, sample_rows, column_contracts)
        for function in functions:
            function_id = _string(function['id'])
            inputs = tuple(_mapping(item) for item in _sequence(function['inputs']))
            function_samples = _function_samples(samples, inputs, function_id, binding_index)
            sample_reasons: dict[str, set[str]] = {sample.key: set() for sample in function_samples}
            populated_inputs = Counter[str]()

            for input_contract in inputs:
                input_name = _string(input_contract['name'])
                kind = _string(input_contract['kind'])
                binding = binding_index.get((function_id, input_name))
                input_values, audit = _audit_binding(
                    workbook,
                    function_samples,
                    function_id,
                    input_name,
                    kind,
                    input_contract,
                    binding,
                    column_contracts,
                )
                values.extend(input_values)
                binding_audits.append(audit)
                for item in input_values:
                    if item.status not in ('missing_input', 'missing_input_binding'):
                        populated_inputs[item.sample_key] += 1
                    if item.status != 'compatible':
                        reasons = item.blockers or (item.status,)
                        sample_reasons[item.sample_key].update(reasons)

            for sample in function_samples:
                reasons = tuple(sorted(sample_reasons[sample.key]))
                reason_counts.update(reasons)
                eligibility.append(
                    FunctionSampleEligibility(
                        sample_key=sample.key,
                        sample_id=sample.sample_id,
                        source=sample.source,
                        sample_number=sample.sample_number,
                        function=function_id,
                        eligible=not reasons,
                        reasons=reasons,
                    )
                )
            eligible_count = sum(not sample_reasons[sample.key] for sample in function_samples)
            function_audits.append(
                FunctionPredictorAudit(
                    function=function_id,
                    input_count=len(inputs),
                    candidate_sample_count=len(function_samples),
                    complete_input_sample_count=sum(
                        populated_inputs[sample.key] == len(inputs) for sample in function_samples
                    ),
                    eligible_sample_count=eligible_count,
                    excluded_sample_count=len(function_samples) - eligible_count,
                )
            )
    finally:
        workbook.close()

    distinct_samples = {sample.key for sample in samples}
    summary = PredictorAuditSummary(
        schema_version=2,
        source_filename=path.name,
        sha256=digest,
        manifest_schema_version=schema_version,
        public_function_count=len(functions),
        candidate_sample_count=len(distinct_samples),
        binding_audits=tuple(binding_audits),
        function_audits=tuple(function_audits),
        exclusion_reasons=tuple(
            LabelCount(label=label, count=count) for label, count in sorted(reason_counts.items())
        ),
    )
    return PredictorAuditResult(summary, tuple(values), tuple(eligibility))


def write_predictor_audit_summary(result: PredictorAuditResult, destination: str | Path) -> None:
    """Write deterministic counts without redistributing workbook values."""
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = json.dumps(asdict(result.summary), ensure_ascii=False, indent=2, sort_keys=True)
    path.write_text(f'{document}\n', encoding='utf-8', newline='\n')


def _audit_binding(
    workbook: Workbook,
    samples: Sequence[_Sample],
    function: str,
    input_name: str,
    kind: str,
    input_contract: Mapping[str, object],
    binding: _Binding | None,
    column_contracts: Mapping[str, Mapping[str, Mapping[str, object]]],
) -> tuple[list[CanonicalPredictorValue], PredictorBindingAudit]:
    if binding is None:
        records = [
            _value(sample, function, input_name, None, None, 'missing_input_binding', (), ())
            for sample in samples
        ]
        return records, _binding_audit(
            function,
            input_name,
            kind,
            None,
            records,
            len(samples),
            (),
            0,
            None,
        )

    sheet = workbook[binding.worksheet]
    column_contract = column_contracts[binding.worksheet][binding.column]
    records: list[CanonicalPredictorValue] = []
    measurement_compatible_count = 0
    expects_particle_size = _expects_particle_size(kind, input_contract)
    particle_size_compatible_count = 0 if expects_particle_size else None
    for sample in samples:
        raw = sheet[f'{binding.column}{sample.row}'].value
        actual_contract, evidence_cells = _contract_for_sample(binding, column_contract, sample.row)
        contract_compatible = _measurement_compatible(kind, input_contract, actual_contract)
        particle_size_compatible = _particle_measurement_compatible(
            kind, input_contract, actual_contract
        )
        source_cells = (f'{binding.worksheet}!{binding.column}{sample.row}', *evidence_cells)
        if _missing(raw):
            records.append(
                _value(
                    sample,
                    function,
                    input_name,
                    raw,
                    None,
                    'missing_input',
                    (),
                    source_cells,
                )
            )
            continue
        measurement_compatible_count += contract_compatible
        if particle_size_compatible_count is not None:
            particle_size_compatible_count += particle_size_compatible is True
        if binding.status != 'confirmed':
            records.append(
                _value(
                    sample,
                    function,
                    input_name,
                    raw,
                    None,
                    'binding_blocked',
                    binding.blockers,
                    source_cells,
                )
            )
            continue
        if not contract_compatible:
            records.append(
                _value(
                    sample,
                    function,
                    input_name,
                    raw,
                    None,
                    'measurement_contract_mismatch',
                    (),
                    source_cells,
                )
            )
            continue
        normalized, status = _normalize_value(raw, kind, input_contract, binding)
        records.append(
            _value(
                sample,
                function,
                input_name,
                raw,
                normalized,
                status,
                (),
                source_cells,
            )
        )
    return records, _binding_audit(
        function,
        input_name,
        kind,
        binding,
        records,
        len(samples),
        binding.blockers,
        measurement_compatible_count,
        particle_size_compatible_count,
    )


def _binding_audit(
    function: str,
    input_name: str,
    kind: str,
    binding: _Binding | None,
    records: Sequence[CanonicalPredictorValue],
    sample_count: int,
    blockers: tuple[str, ...],
    measurement_compatible_count: int,
    particle_size_compatible_count: int | None,
) -> PredictorBindingAudit:
    compatible = sum(item.status == 'compatible' for item in records)
    missing = sum(item.status == 'missing_input' for item in records)
    populated = sample_count - missing if binding is not None else 0
    return PredictorBindingAudit(
        function=function,
        input=input_name,
        kind=kind,
        worksheet=None if binding is None else binding.worksheet,
        column=None if binding is None else binding.column,
        binding_status='missing' if binding is None else binding.status,
        candidate_sample_count=sample_count,
        populated_count=populated,
        compatible_count=compatible,
        measurement_compatible_count=measurement_compatible_count,
        particle_size_compatible_count=particle_size_compatible_count,
        missing_count=missing,
        excluded_count=sample_count - compatible,
        blockers=blockers,
    )


def _value(
    sample: _Sample,
    function: str,
    input_name: str,
    raw: object,
    normalized: object,
    status: str,
    blockers: tuple[str, ...],
    source_cells: tuple[str, ...],
) -> CanonicalPredictorValue:
    return CanonicalPredictorValue(
        sample_key=sample.key,
        sample_id=sample.sample_id,
        source=sample.source,
        sample_number=sample.sample_number,
        function=function,
        input=input_name,
        raw_value=_json_scalar(raw),
        value=_json_scalar(normalized),
        status=status,
        blockers=blockers,
        source_cells=source_cells,
    )


def _normalize_value(
    raw: object,
    kind: str,
    input_contract: Mapping[str, object],
    binding: _Binding,
) -> tuple[object, str]:
    if kind == 'categorical':
        text = str(raw).strip()
        mapped = binding.category_map.get(text, text)
        enum = _mapping(input_contract['enum'])
        allowed = {_string(_mapping(item)['value']) for item in _sequence(enum['values'])}
        if mapped not in allowed:
            return None, 'unknown_category'
        return mapped, 'compatible'
    if kind == 'numeric':
        if isinstance(raw, bool) or not isinstance(raw, Real):
            return None, 'invalid_numeric_input'
        value = float(raw)
        if not math.isfinite(value):
            return None, 'invalid_numeric_input'
        return value, 'compatible'
    return None, 'unsupported_input_kind'


def _measurement_compatible(
    kind: str,
    input_contract: Mapping[str, object],
    column_contract: Mapping[str, object],
) -> bool:
    if kind == 'categorical':
        return 'enum' in input_contract
    if kind != 'numeric':
        return False
    expected_value = input_contract.get('measurement')
    actual_value = column_contract.get('measurement')
    if not isinstance(expected_value, Mapping) or not isinstance(actual_value, Mapping):
        return False
    expected = cast('Mapping[str, object]', expected_value)
    actual = cast('Mapping[str, object]', actual_value)
    basis_value = expected.get('basis')
    basis = cast('Mapping[str, object]', basis_value) if isinstance(basis_value, Mapping) else {}
    checks = {
        'quantity': expected.get('quantity'),
        'unit': expected.get('unit'),
        'material_basis': basis.get('material'),
        'denominator_basis': basis.get('denominator'),
    }
    basic_compatible = all(
        actual.get(key) == value for key, value in checks.items() if value is not None
    )
    required_method = expected.get('method')
    method_compatible = required_method is None or actual.get('method') == required_method
    conditions_compatible = tuple(_sequence(expected.get('conditions', []))) == tuple(
        _sequence(actual.get('conditions', []))
    )
    return (
        basic_compatible
        and method_compatible
        and conditions_compatible
        and _particle_size_compatible(expected.get('particle_size'), actual.get('particle_size'))
    )


def _expects_particle_size(kind: str, input_contract: Mapping[str, object]) -> bool:
    if kind != 'numeric':
        return False
    measurement = input_contract.get('measurement')
    return isinstance(measurement, Mapping) and measurement.get('particle_size') is not None


def _particle_measurement_compatible(
    kind: str,
    input_contract: Mapping[str, object],
    column_contract: Mapping[str, object],
) -> bool | None:
    if not _expects_particle_size(kind, input_contract):
        return None
    expected = _mapping(input_contract['measurement'])
    actual_value = column_contract.get('measurement')
    if not isinstance(actual_value, Mapping):
        return False
    actual = cast('Mapping[str, object]', actual_value)
    return _particle_size_compatible(expected.get('particle_size'), actual.get('particle_size'))


def _contract_for_sample(
    binding: _Binding,
    column_contract: Mapping[str, object],
    row: int,
) -> tuple[Mapping[str, object], tuple[str, ...]]:
    for override in binding.measurement_overrides:
        if row in override.rows:
            contract = dict(column_contract)
            contract['measurement'] = override.measurement
            return contract, tuple(header.cell for header in override.source_headers)
    return column_contract, ()


def _particle_size_compatible(expected_value: object, actual_value: object) -> bool:
    if expected_value is None or actual_value is None:
        return expected_value is actual_value
    expected = _mapping(expected_value)
    actual = _mapping(actual_value)
    return (
        expected.get('unit') == actual.get('unit')
        and _bound_compatible(expected.get('lower'), actual, 'lower')
        and _bound_compatible(expected.get('upper'), actual, 'upper')
    )


def _bound_compatible(expected_value: object, actual: Mapping[str, object], prefix: str) -> bool:
    actual_value = actual.get(prefix)
    if expected_value is None:
        return actual_value is None
    expected = _mapping(expected_value)
    return expected.get('value') == actual_value and expected.get('inclusive') == actual.get(
        f'{prefix}_inclusive'
    )


def _function_samples(
    samples: Sequence[_Sample],
    inputs: Sequence[Mapping[str, object]],
    function: str,
    bindings: Mapping[tuple[str, str], _Binding],
) -> tuple[_Sample, ...]:
    worksheets = {
        binding.worksheet
        for input_contract in inputs
        if (binding := bindings.get((function, _string(input_contract['name'])))) is not None
    }
    if not worksheets:
        return tuple(samples)
    return tuple(sample for sample in samples if sample.worksheet in worksheets)


def _samples(
    workbook: Workbook,
    rows: Mapping[str, tuple[int, ...]],
    columns: Mapping[str, Mapping[str, Mapping[str, object]]],
) -> tuple[_Sample, ...]:
    result: list[_Sample] = []
    for worksheet, row_numbers in sorted(rows.items()):
        sheet = workbook[worksheet]
        role_columns = {
            _string(contract['role']): column for column, contract in columns[worksheet].items()
        }
        required = ('sample_id', 'source', 'sample_number')
        missing_roles = [role for role in required if role not in role_columns]
        if missing_roles:
            message = f'Worksheet {worksheet!r} lacks identity columns: {", ".join(missing_roles)}.'
            raise PredictorAuditError(message)
        result.extend(
            _Sample(
                worksheet=worksheet,
                row=row,
                sample_id=_json_scalar(sheet[f'{role_columns["sample_id"]}{row}'].value),
                source=_json_scalar(sheet[f'{role_columns["source"]}{row}'].value),
                sample_number=_json_scalar(sheet[f'{role_columns["sample_number"]}{row}'].value),
            )
            for row in row_numbers
        )
    return tuple(result)


def _candidate_rows(document: Mapping[str, object]) -> Mapping[str, tuple[int, ...]]:
    result: dict[str, set[int]] = {}
    for rule_value in _sequence(document['retention_reconstruction_rules']):
        rule = _mapping(rule_value)
        worksheet = _string(rule['master_worksheet'])
        rows = result.setdefault(worksheet, set())
        for block_value in _sequence(rule['blocks']):
            block = _mapping(block_value)
            rows.update(
                range(
                    _integer(block['first_sample_row']),
                    _integer(block['last_sample_row']) + 1,
                )
            )
        parameter_value = rule['parameter_block']
        if isinstance(parameter_value, Mapping):
            parameter = cast('Mapping[str, object]', parameter_value)
            for range_value in _sequence(parameter['sample_rows']):
                row_range = _mapping(range_value)
                rows.update(range(_integer(row_range['first']), _integer(row_range['last']) + 1))
    return {worksheet: tuple(sorted(rows)) for worksheet, rows in result.items()}


def _bindings(document: Mapping[str, object]) -> tuple[_Binding, ...]:
    result: list[_Binding] = []
    for value in _sequence(document.get('predictor_bindings', [])):
        binding = _mapping(value)
        categories = {
            _string(item['raw']): _string(item['value'])
            for category_value in _sequence(binding.get('category_map', []))
            for item in [_mapping(category_value)]
            if item['status'] == 'confirmed'
        }
        overrides = tuple(
            _measurement_override(_mapping(item))
            for item in _sequence(binding.get('measurement_overrides', []))
        )
        result.append(
            _Binding(
                function=_string(binding['function']),
                input=_string(binding['input']),
                worksheet=_string(binding['worksheet']),
                column=_string(binding['column']),
                status=_string(binding['status']),
                blockers=tuple(_string(item) for item in _sequence(binding['blockers'])),
                category_map=categories,
                measurement_overrides=overrides,
            )
        )
    return tuple(result)


def _measurement_override(value: Mapping[str, object]) -> _MeasurementOverride:
    rows = frozenset(
        row
        for range_value in _sequence(value['sample_rows'])
        for row_range in [_mapping(range_value)]
        for row in range(_integer(row_range['first']), _integer(row_range['last']) + 1)
    )
    return _MeasurementOverride(
        rows=rows,
        measurement=_mapping(value['measurement']),
        source_headers=tuple(
            _SourceHeader(
                worksheet=_string(header['worksheet']),
                column=_string(header['column']),
                row=_integer(header['row']),
                value=_string(header['value']),
            )
            for item in _sequence(value['source_headers'])
            for header in [_mapping(item)]
        ),
    )


def _worksheet_columns(
    document: Mapping[str, object],
) -> Mapping[str, Mapping[str, Mapping[str, object]]]:
    return {
        _string(worksheet['name']): {
            _string(column['column']): column
            for column_value in _sequence(worksheet['columns'])
            for column in [_mapping(column_value)]
        }
        for worksheet_value in _sequence(document['worksheets'])
        for worksheet in [_mapping(worksheet_value)]
    }


def _validate_binding_headers(
    workbook: Workbook,
    bindings: Sequence[_Binding],
    columns: Mapping[str, Mapping[str, Mapping[str, object]]],
) -> None:
    sheet_names = workbook.sheetnames
    for binding in bindings:
        if binding.worksheet not in sheet_names:
            message = f'Workbook has no predictor worksheet {binding.worksheet!r}.'
            raise PredictorAuditError(message)
        contract = columns[binding.worksheet][binding.column]
        expected = contract.get('source_header')
        if expected is not None:
            sheet = workbook[binding.worksheet]
            actual = sheet[f'{binding.column}1'].value
            if actual != expected:
                message = (
                    f'Predictor header {binding.worksheet}!{binding.column}1 does not match '
                    'the reviewed mapping.'
                )
                raise PredictorAuditError(message)
        for override in binding.measurement_overrides:
            for header in override.source_headers:
                actual = workbook[header.worksheet][f'{header.column}{header.row}'].value
                if actual != header.value:
                    message = (
                        f'Source-local predictor header {header.cell} does not match '
                        'the reviewed mapping.'
                    )
                    raise PredictorAuditError(message)


def _manifest_functions(
    manifest: Mapping[str, object],
) -> tuple[int, tuple[Mapping[str, object], ...]]:
    schema_version = manifest.get('schema_version')
    functions_value = manifest.get('functions')
    if not isinstance(schema_version, int) or isinstance(schema_version, bool):
        message = 'Manifest has no integer schema version.'
        raise PredictorAuditError(message)
    if not _is_sequence(functions_value):
        message = 'Manifest has no function list.'
        raise PredictorAuditError(message)
    functions = tuple(_mapping(item) for item in functions_value)
    return schema_version, functions


def _validate_bindings_against_manifest(
    bindings: Sequence[_Binding], functions: Sequence[Mapping[str, object]]
) -> None:
    manifest_inputs = {
        (_string(function['id']), _string(input_contract['name']))
        for function in functions
        for input_value in _sequence(function['inputs'])
        for input_contract in [_mapping(input_value)]
    }
    unknown = sorted(
        (binding.function, binding.input)
        for binding in bindings
        if (binding.function, binding.input) not in manifest_inputs
    )
    if unknown:
        labels = ', '.join(f'{function}.{input_name}' for function, input_name in unknown)
        message = f'Predictor bindings are absent from the public manifest: {labels}.'
        raise PredictorAuditError(message)


def _validate_source(path: Path, digest: str, mapping: DatasetMapping) -> None:
    if path.name != mapping.source_filename:
        message = (
            f'Workbook filename {path.name!r} does not match mapping source '
            f'{mapping.source_filename!r}.'
        )
        raise PredictorAuditError(message)
    if digest != mapping.source_sha256:
        message = 'Workbook checksum does not match the reviewed mapping.'
        raise PredictorAuditError(message)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open('rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                digest.update(chunk)
    except OSError as error:
        message = f'Cannot read workbook {path}: {error}'
        raise PredictorAuditError(message) from error
    return digest.hexdigest()


def _missing(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _json_scalar(value: object) -> JsonScalar:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def _mapping(value: object) -> Mapping[str, object]:
    if not isinstance(value, Mapping):
        message = f'Expected an object, got {type(value).__name__}.'
        raise PredictorAuditError(message)
    return cast('Mapping[str, object]', value)


def _sequence(value: object) -> Sequence[object]:
    if not _is_sequence(value):
        message = f'Expected an array, got {type(value).__name__}.'
        raise PredictorAuditError(message)
    return value


def _is_sequence(value: object) -> TypeGuard[Sequence[object]]:
    return isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray))


def _string(value: object) -> str:
    if not isinstance(value, str):
        message = f'Expected a string, got {type(value).__name__}.'
        raise PredictorAuditError(message)
    return value


def _integer(value: object) -> int:
    if not isinstance(value, int) or isinstance(value, bool):
        message = f'Expected an integer, got {type(value).__name__}.'
        raise PredictorAuditError(message)
    return value
