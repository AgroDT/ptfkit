"""Evidence-only profiling of long-form water-retention worksheets."""

from __future__ import annotations

import hashlib
import json
import math
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from numbers import Real
from pathlib import Path
from typing import TYPE_CHECKING
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


if TYPE_CHECKING:
    from collections.abc import Hashable

    from openpyxl.worksheet.worksheet import Worksheet


_OBSERVATION_ENCODINGS = frozenset({'W', 'Q', '?'})
_PARAMETER_ENCODING = 'VGN'
_EXPECTED_HEADERS = (
    'source',
    'sample_number',
    'Формат данных влажности почвы',
    'Формат данных давления влаги',
    'Переменная/Давление влаги',
    'Значение/Содержание влаги',
)
_MISSING_MARKERS = frozenset({'', '-', '\u2013', '\u2014', 'na', 'n/a', 'nan', 'none'})


class RetentionProfileError(ValueError):
    """Raised when a retention worksheet cannot be profiled without interpretation."""


@dataclass(frozen=True)
class NumericProfile:
    """Observed numeric-shape evidence without an inferred unit or scale."""

    numeric_count: int
    missing_count: int
    nonnumeric_count: int
    nonfinite_count: int
    minimum: float | None
    maximum: float | None
    negative_count: int
    zero_count: int
    positive_count: int
    at_most_one_count: int
    above_one_count: int


@dataclass(frozen=True)
class LabelCount:
    """Count of one exact workbook label."""

    label: str
    count: int


@dataclass(frozen=True)
class RetentionEncodingProfile:
    """Evidence for one exact water-content and pressure-label combination."""

    water_encoding: str
    pressure_encoding: str | None
    row_count: int
    sample_count: int
    first_row: int
    last_row: int
    pressure_values: NumericProfile
    water_values: NumericProfile


@dataclass(frozen=True)
class RetentionSourceProfile:
    """Retention evidence grouped by the source label exactly as stored."""

    source: str
    row_count: int
    sample_count: int
    observation_count: int
    fitted_parameter_count: int
    unclassified_count: int
    missing_sample_number_count: int
    encodings: tuple[RetentionEncodingProfile, ...]
    fitted_parameters: tuple[LabelCount, ...]
    unclassified_encodings: tuple[LabelCount, ...]


@dataclass(frozen=True)
class RetentionWorkbookProfile:
    """Checksum-bound source-level evidence for one long retention worksheet."""

    schema_version: int
    source_filename: str
    sha256: str
    sheet_name: str
    total_data_rows: int
    source_count: int
    observation_count: int
    fitted_parameter_count: int
    unclassified_count: int
    missing_source_count: int
    sources: tuple[RetentionSourceProfile, ...]


@dataclass
class _NumericAccumulator:
    numeric: list[float]
    missing_count: int = 0
    nonnumeric_count: int = 0
    nonfinite_count: int = 0

    @classmethod
    def empty(cls) -> _NumericAccumulator:
        return cls(numeric=[])

    def add(self, value: object) -> None:
        if value is None or (
            isinstance(value, str) and value.strip().casefold() in _MISSING_MARKERS
        ):
            self.missing_count += 1
            return
        if isinstance(value, Real) and not isinstance(value, bool):
            numeric = float(value)
            if not math.isfinite(numeric):
                self.nonfinite_count += 1
                return
            self.numeric.append(numeric)
            return
        self.nonnumeric_count += 1

    def finish(self) -> NumericProfile:
        return NumericProfile(
            numeric_count=len(self.numeric),
            missing_count=self.missing_count,
            nonnumeric_count=self.nonnumeric_count,
            nonfinite_count=self.nonfinite_count,
            minimum=min(self.numeric, default=None),
            maximum=max(self.numeric, default=None),
            negative_count=sum(value < 0 for value in self.numeric),
            zero_count=sum(value == 0 for value in self.numeric),
            positive_count=sum(value > 0 for value in self.numeric),
            at_most_one_count=sum(value <= 1 for value in self.numeric),
            above_one_count=sum(value > 1 for value in self.numeric),
        )


@dataclass
class _EncodingAccumulator:
    rows: list[int]
    samples: set[Hashable]
    pressure_values: _NumericAccumulator
    water_values: _NumericAccumulator

    @classmethod
    def empty(cls) -> _EncodingAccumulator:
        return cls(
            rows=[],
            samples=set(),
            pressure_values=_NumericAccumulator.empty(),
            water_values=_NumericAccumulator.empty(),
        )


@dataclass
class _SourceAccumulator:
    row_count: int
    samples: set[Hashable]
    observation_count: int
    fitted_parameter_count: int
    unclassified_count: int
    missing_sample_number_count: int
    encodings: dict[tuple[str, str | None], _EncodingAccumulator]
    fitted_parameters: Counter[str]
    unclassified_encodings: Counter[str]

    @classmethod
    def empty(cls) -> _SourceAccumulator:
        return cls(
            row_count=0,
            samples=set(),
            observation_count=0,
            fitted_parameter_count=0,
            unclassified_count=0,
            missing_sample_number_count=0,
            encodings=defaultdict(_EncodingAccumulator.empty),
            fitted_parameters=Counter(),
            unclassified_encodings=Counter(),
        )


@dataclass
class _ProfileAccumulator:
    sources: dict[str, _SourceAccumulator]
    total_data_rows: int = 0
    missing_source_count: int = 0

    @classmethod
    def empty(cls) -> _ProfileAccumulator:
        return cls(sources=defaultdict(_SourceAccumulator.empty))

    def add(self, row_number: int, row: tuple[object, ...]) -> None:
        if not any(value is not None and value != '' for value in row):
            return
        self.total_data_rows += 1
        source_value, sample_number, water_encoding, pressure_encoding, axis, water = row
        if source_value is None or str(source_value).strip() == '':
            self.missing_source_count += 1
            return

        accumulator = self.sources[str(source_value)]
        accumulator.row_count += 1
        sample_key = _sample_key(sample_number)
        if sample_key is None:
            accumulator.missing_sample_number_count += 1
        else:
            accumulator.samples.add(sample_key)

        water_label = _label(water_encoding)
        if water_label in _OBSERVATION_ENCODINGS:
            _add_observation(
                accumulator,
                row_number=row_number,
                sample_key=sample_key,
                water_encoding=water_label,
                pressure_encoding=_optional_label(pressure_encoding),
                pressure_value=axis,
                water_value=water,
            )
        elif water_label == _PARAMETER_ENCODING:
            accumulator.fitted_parameter_count += 1
            accumulator.fitted_parameters[_label(axis)] += 1
        else:
            accumulator.unclassified_count += 1
            accumulator.unclassified_encodings[water_label] += 1


def profile_retention_workbook(
    source: str | Path, *, sheet_name: str = 'ОГХ'
) -> RetentionWorkbookProfile:
    """Profile exact source labels and numeric shapes without inferring semantics."""
    path = Path(source)
    if path.suffix.casefold() != '.xlsx':
        message = f'Workbook must use the .xlsx extension: {path.name!r}.'
        raise RetentionProfileError(message)
    try:
        digest = _sha256(path)
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError) as error:
        message = f'Cannot profile workbook {path.name!r}: {error}'
        raise RetentionProfileError(message) from error

    try:
        if sheet_name not in workbook.sheetnames:
            message = f'Workbook has no retention worksheet {sheet_name!r}.'
            raise RetentionProfileError(message)
        sheet = workbook[sheet_name]
        _validate_headers(sheet_name, tuple(sheet.cell(1, column).value for column in range(2, 8)))

        profile = _profile_sheet(sheet)
    finally:
        workbook.close()

    source_profiles = tuple(
        _finish_source(label, accumulator)
        for label, accumulator in sorted(
            profile.sources.items(), key=lambda item: (item[0].casefold(), item[0])
        )
    )
    return RetentionWorkbookProfile(
        schema_version=1,
        source_filename=path.name,
        sha256=digest,
        sheet_name=sheet_name,
        total_data_rows=profile.total_data_rows,
        source_count=len(source_profiles),
        observation_count=sum(item.observation_count for item in source_profiles),
        fitted_parameter_count=sum(item.fitted_parameter_count for item in source_profiles),
        unclassified_count=sum(item.unclassified_count for item in source_profiles),
        missing_source_count=profile.missing_source_count,
        sources=source_profiles,
    )


def write_retention_profile(profile: RetentionWorkbookProfile, destination: str | Path) -> None:
    """Write a stable UTF-8 JSON profile without a machine-specific workbook path."""
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = json.dumps(asdict(profile), ensure_ascii=False, indent=2, sort_keys=True)
    path.write_text(f'{document}\n', encoding='utf-8')


def _profile_sheet(sheet: Worksheet) -> _ProfileAccumulator:
    profile = _ProfileAccumulator.empty()
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=2, min_col=2, max_col=7, values_only=True), start=2
    ):
        profile.add(row_number, row)
    return profile


def _add_observation(
    source: _SourceAccumulator,
    *,
    row_number: int,
    sample_key: Hashable | None,
    water_encoding: str,
    pressure_encoding: str | None,
    pressure_value: object,
    water_value: object,
) -> None:
    source.observation_count += 1
    encoding = source.encodings[(water_encoding, pressure_encoding)]
    encoding.rows.append(row_number)
    if sample_key is not None:
        encoding.samples.add(sample_key)
    encoding.pressure_values.add(pressure_value)
    encoding.water_values.add(water_value)


def _finish_source(label: str, source: _SourceAccumulator) -> RetentionSourceProfile:
    encodings = tuple(
        RetentionEncodingProfile(
            water_encoding=water_encoding,
            pressure_encoding=pressure_encoding,
            row_count=len(values.rows),
            sample_count=len(values.samples),
            first_row=min(values.rows),
            last_row=max(values.rows),
            pressure_values=values.pressure_values.finish(),
            water_values=values.water_values.finish(),
        )
        for (water_encoding, pressure_encoding), values in sorted(
            source.encodings.items(), key=lambda item: (item[0][0], item[0][1] or '')
        )
    )
    return RetentionSourceProfile(
        source=label,
        row_count=source.row_count,
        sample_count=len(source.samples),
        observation_count=source.observation_count,
        fitted_parameter_count=source.fitted_parameter_count,
        unclassified_count=source.unclassified_count,
        missing_sample_number_count=source.missing_sample_number_count,
        encodings=encodings,
        fitted_parameters=_finish_counts(source.fitted_parameters),
        unclassified_encodings=_finish_counts(source.unclassified_encodings),
    )


def _finish_counts(values: Counter[str]) -> tuple[LabelCount, ...]:
    return tuple(LabelCount(label, values[label]) for label in sorted(values))


def _sample_key(value: object) -> Hashable | None:
    if value is None or value == '':
        return None
    if isinstance(value, (bool, int, float, str)):
        return (type(value).__name__, value)
    return (type(value).__name__, str(value))


def _label(value: object) -> str:
    if value is None:
        return '<missing>'
    return str(value)


def _optional_label(value: object) -> str | None:
    if value is None or value == '':
        return None
    return str(value)


def _validate_headers(sheet_name: str, actual: tuple[object, ...]) -> None:
    if actual != _EXPECTED_HEADERS:
        message = (
            f'Retention worksheet {sheet_name!r} columns B:G do not match the '
            'reviewed meta_v3.3 layout.'
        )
        raise RetentionProfileError(message)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()
