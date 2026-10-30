"""Canonical, rule-gated normalization of long-form retention data."""

from __future__ import annotations

import hashlib
import json
import math
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import asdict, dataclass
from numbers import Real
from pathlib import Path
from typing import TYPE_CHECKING, cast
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException

from .retention import LabelCount


if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook

    from .datasets import DatasetMapping, JsonScalar


_OBSERVATION_ENCODINGS = frozenset({'W', 'Q', '?'})
_EXPECTED_HEADERS = (
    'source',
    'sample_number',
    'Формат данных влажности почвы',
    'Формат данных давления влаги',
    'Переменная/Давление влаги',
    'Значение/Содержание влаги',
)
_CENTIMETER_WATER_HEAD_TO_KILOPASCAL = 0.0980665
_FIRST_OBSERVATION_ROW = 2


class RetentionNormalizationError(ValueError):
    """Raised when a workbook cannot be normalized under its reviewed contract."""


@dataclass(frozen=True)
class NormalizedRetentionObservation:
    """One observed retention point in canonical units with cell-level lineage."""

    source_filename: str
    sheet_name: str
    row: int
    sample_id: JsonScalar
    source: str
    sample_number: JsonScalar
    water_encoding: str
    pressure_encoding: str | None
    raw_pressure: float
    raw_water_content: float
    pressure_quantity: str
    pressure_kilopascal: float
    pressure_unit: str
    pressure_convention: str
    water_quantity: str
    water_content: float
    water_unit: str
    water_basis: str
    source_cells: tuple[str, ...]


@dataclass(frozen=True)
class FittedRetentionParameter:
    """One workbook-supplied fitted parameter kept outside observations."""

    source_filename: str
    sheet_name: str
    row: int
    sample_id: JsonScalar
    source: str
    sample_number: JsonScalar
    model_family: str
    parameter_name: str
    value: JsonScalar
    source_cells: tuple[str, ...]


@dataclass(frozen=True)
class RetentionExclusion:
    """One source row excluded with an explicit reason and preserved raw values."""

    source_filename: str
    sheet_name: str
    row: int
    row_role: str
    reason: str
    blockers: tuple[str, ...]
    sample_id: JsonScalar
    source: str | None
    sample_number: JsonScalar
    water_encoding: str | None
    pressure_encoding: str | None
    raw_pressure: JsonScalar
    raw_water_content: JsonScalar
    source_cells: tuple[str, ...]


@dataclass(frozen=True)
class RetentionNormalizationSummary:
    """Non-sensitive row-accounting evidence for one normalization run."""

    schema_version: int
    source_filename: str
    sha256: str
    sheet_name: str
    total_data_rows: int
    input_observation_count: int
    normalized_observation_count: int
    excluded_observation_count: int
    fitted_parameter_count: int
    unclassified_exclusion_count: int
    missing_source_exclusion_count: int
    exclusion_reasons: tuple[LabelCount, ...]
    blocker_counts: tuple[LabelCount, ...]


@dataclass(frozen=True)
class RetentionNormalizationResult:
    """Complete in-memory normalization result and its auditable row balance."""

    summary: RetentionNormalizationSummary
    observations: tuple[NormalizedRetentionObservation, ...]
    fitted_parameters: tuple[FittedRetentionParameter, ...]
    exclusions: tuple[RetentionExclusion, ...]


@dataclass(frozen=True)
class _ObservationRule:
    status: str
    water_quantity: str | None
    water_unit: str | None
    canonical_water_unit: str | None
    pressure_representation: str | None
    pressure_convention: str | None
    canonical_pressure_unit: str | None
    blockers: tuple[str, ...]
    requires_row_review: bool


@dataclass(frozen=True)
class _SampleIdentity:
    sample_id: JsonScalar
    source: str
    sample_number: JsonScalar
    evidence_cells: tuple[str, ...]


@dataclass(frozen=True)
class _IdentityRules:
    source_aliases: Mapping[str, str]
    sample_crosswalks: Mapping[str, _SampleIdentity]


@dataclass(frozen=True)
class _RowReview:
    identity: _SampleIdentity | None
    reason: str | None
    evidence_cells: tuple[str, ...]


def normalize_retention_workbook(  # noqa: C901, PLR0912, PLR0915
    source: str | Path, mapping: DatasetMapping
) -> RetentionNormalizationResult:
    """Normalize only observations covered by confirmed source-specific rules."""
    path = Path(source)
    digest = _sha256(path)
    _validate_source_identity(path, digest, mapping)
    sheet_name = _profile_string(mapping, 'sheet_name')
    rules = _observation_rules(mapping.document)
    identity_rules = _identity_rules(mapping.document, sheet_name)
    row_reviews = _row_reviews(mapping.document, sheet_name)
    consumed_reviews: set[tuple[str, int]] = set()
    point_indexes: dict[str, tuple[int, dict[tuple[float, float], list[_SampleIdentity]]]] = {}
    released_points: set[tuple[str, float, float]] = set()

    observations: list[NormalizedRetentionObservation] = []
    parameters: list[FittedRetentionParameter] = []
    exclusions: list[RetentionExclusion] = []
    input_observation_count = 0
    total_data_rows = 0
    missing_source_count = 0
    unclassified_count = 0

    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError) as error:
        message = f'Cannot normalize workbook {path.name!r}: {error}'
        raise RetentionNormalizationError(message) from error

    try:
        if sheet_name not in workbook.sheetnames:
            message = f'Workbook has no retention worksheet {sheet_name!r}.'
            raise RetentionNormalizationError(message)
        sheet = workbook[sheet_name]
        actual_headers = tuple(sheet.cell(1, column).value for column in range(2, 8))
        if actual_headers != _EXPECTED_HEADERS:
            message = (
                f'Retention worksheet {sheet_name!r} columns B:G do not match the reviewed layout.'
            )
            raise RetentionNormalizationError(message)

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, min_col=1, max_col=7, values_only=True), start=2
        ):
            if not _has_value(row):
                continue
            total_data_rows += 1
            sample_id, source_value, sample_number, water_value, pressure_value, axis, water = row
            source_label = _optional_label(source_value)
            water_encoding = _optional_label(water_value)
            pressure_encoding = _optional_label(pressure_value)
            lineage = _lineage(sheet_name, row_number)

            if source_label is None:
                missing_source_count += 1
                exclusions.append(
                    _exclusion(
                        path.name,
                        sheet_name,
                        row_number,
                        'missing_source',
                        'missing_source',
                        (),
                        sample_id,
                        None,
                        sample_number,
                        water_encoding,
                        pressure_encoding,
                        axis,
                        water,
                        lineage,
                    )
                )
                continue

            if water_encoding == 'VGN':
                parameters.append(
                    FittedRetentionParameter(
                        source_filename=path.name,
                        sheet_name=sheet_name,
                        row=row_number,
                        sample_id=_json_scalar(sample_id),
                        source=source_label,
                        sample_number=_json_scalar(sample_number),
                        model_family='van_genuchten',
                        parameter_name=_required_label(axis, row_number, 'parameter name'),
                        value=_json_scalar(water),
                        source_cells=(
                            f'{sheet_name}!A{row_number}',
                            f'{sheet_name}!B{row_number}',
                            f'{sheet_name}!C{row_number}',
                            f'{sheet_name}!F{row_number}',
                            f'{sheet_name}!G{row_number}',
                        ),
                    )
                )
                continue

            if water_encoding not in _OBSERVATION_ENCODINGS:
                unclassified_count += 1
                exclusions.append(
                    _exclusion(
                        path.name,
                        sheet_name,
                        row_number,
                        'unclassified',
                        'unclassified_water_encoding',
                        (),
                        sample_id,
                        source_label,
                        sample_number,
                        water_encoding,
                        pressure_encoding,
                        axis,
                        water,
                        lineage,
                    )
                )
                continue

            input_observation_count += 1
            rule = rules.get((source_label, water_encoding, pressure_encoding))
            if rule is None:
                exclusions.append(
                    _exclusion(
                        path.name,
                        sheet_name,
                        row_number,
                        'observation',
                        'missing_source_rule',
                        (),
                        sample_id,
                        source_label,
                        sample_number,
                        water_encoding,
                        pressure_encoding,
                        axis,
                        water,
                        lineage,
                    )
                )
                continue
            if rule.requires_row_review and source_label not in row_reviews:
                message = 'This source requires exhaustive reviewed row decisions.'
                raise RetentionNormalizationError(message)
            review = None
            if source_label in row_reviews:
                review = row_reviews[source_label].get(row_number)
                if review is None:
                    message = f'Missing reviewed observation decision at {sheet_name}!{row_number}.'
                    raise RetentionNormalizationError(message)
                consumed_reviews.add((source_label, row_number))
                lineage = (*lineage, *review.evidence_cells)
                if review.reason is not None:
                    exclusions.append(
                        _exclusion(
                            path.name,
                            sheet_name,
                            row_number,
                            'observation',
                            review.reason,
                            ('retention_reconstruction',),
                            sample_id,
                            source_label,
                            sample_number,
                            water_encoding,
                            pressure_encoding,
                            axis,
                            water,
                            lineage,
                        )
                    )
                    continue
            if review is None and not _sample_identity_matches(
                sample_id, source_label, sample_number
            ):
                exclusions.append(
                    _exclusion(
                        path.name,
                        sheet_name,
                        row_number,
                        'observation',
                        'sample_identity_mismatch',
                        ('retention_reconstruction', *rule.blockers),
                        sample_id,
                        source_label,
                        sample_number,
                        water_encoding,
                        pressure_encoding,
                        axis,
                        water,
                        lineage,
                    )
                )
                continue
            if rule.status != 'confirmed':
                exclusions.append(
                    _exclusion(
                        path.name,
                        sheet_name,
                        row_number,
                        'observation',
                        'source_rule_blocked',
                        rule.blockers,
                        sample_id,
                        source_label,
                        sample_number,
                        water_encoding,
                        pressure_encoding,
                        axis,
                        water,
                        lineage,
                    )
                )
                continue

            identity = (
                review.identity
                if review is not None
                else _resolve_identity(sample_id, source_label, sample_number, identity_rules)
            )
            if identity is None:
                message = 'A released row requires a reviewed canonical identity.'
                raise RetentionNormalizationError(message)
            if review is not None:
                if source_label not in point_indexes:
                    point_indexes[source_label] = _review_point_index(
                        workbook, mapping.document, source_label
                    )
                places, point_index = point_indexes[source_label]
                point_key = _review_point_key(axis, water, places)
                match = point_index.get(point_key, [])
                release_key = (source_label, *point_key)
                if len(match) != 1:
                    message = (
                        f'Released row {row_number} requires exactly one master point; '
                        f'found {len(match)}.'
                    )
                    raise RetentionNormalizationError(message)
                master = match[0]
                if (master.sample_id, master.source, master.sample_number) != (
                    identity.sample_id,
                    identity.source,
                    identity.sample_number,
                ) or release_key in released_points:
                    message = f'Released row {row_number} requires an unused exact master identity.'
                    raise RetentionNormalizationError(message)
                if not set(master.evidence_cells).issubset(review.evidence_cells):
                    message = 'Reviewed row provenance must include its matched master cells.'
                    raise RetentionNormalizationError(message)
                released_points.add(release_key)
            normalized = _normalize_observation(
                path.name,
                sheet_name,
                row_number,
                identity.sample_id,
                identity.source,
                identity.sample_number,
                water_encoding,
                pressure_encoding,
                axis,
                water,
                (*lineage, *identity.evidence_cells),
                rule,
            )
            if isinstance(normalized, RetentionExclusion):
                exclusions.append(normalized)
            else:
                observations.append(normalized)
    finally:
        workbook.close()

    expected_reviews = {(label, row) for label, rows in row_reviews.items() for row in rows}
    if consumed_reviews != expected_reviews:
        message = 'Reviewed observation decisions must exactly cover the declared source rows.'
        raise RetentionNormalizationError(message)
    _validate_profile_counts(
        mapping,
        total_data_rows=total_data_rows,
        observation_count=input_observation_count,
        fitted_parameter_count=len(parameters),
        unclassified_count=unclassified_count,
        missing_source_count=missing_source_count,
    )
    observation_exclusions = tuple(item for item in exclusions if item.row_role == 'observation')
    if len(observations) + len(observation_exclusions) != input_observation_count:
        message = 'Observation row accounting is not balanced.'
        raise RetentionNormalizationError(message)

    reason_counts = Counter(item.reason for item in exclusions)
    blocker_counts = Counter(blocker for item in exclusions for blocker in item.blockers)
    summary = RetentionNormalizationSummary(
        schema_version=1,
        source_filename=path.name,
        sha256=digest,
        sheet_name=sheet_name,
        total_data_rows=total_data_rows,
        input_observation_count=input_observation_count,
        normalized_observation_count=len(observations),
        excluded_observation_count=len(observation_exclusions),
        fitted_parameter_count=len(parameters),
        unclassified_exclusion_count=unclassified_count,
        missing_source_exclusion_count=missing_source_count,
        exclusion_reasons=_counts(reason_counts),
        blocker_counts=_counts(blocker_counts),
    )
    return RetentionNormalizationResult(
        summary=summary,
        observations=tuple(observations),
        fitted_parameters=tuple(parameters),
        exclusions=tuple(exclusions),
    )


def write_retention_normalization_summary(
    result: RetentionNormalizationResult, destination: str | Path
) -> None:
    """Write stable row-accounting evidence without redistributing measurements."""
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = json.dumps(asdict(result.summary), ensure_ascii=False, indent=2, sort_keys=True)
    path.write_text(f'{document}\n', encoding='utf-8')


def _normalize_observation(  # noqa: PLR0917
    source_filename: str,
    sheet_name: str,
    row: int,
    sample_id: object,
    source: str,
    sample_number: object,
    water_encoding: str,
    pressure_encoding: str | None,
    raw_pressure: object,
    raw_water: object,
    lineage: tuple[str, ...],
    rule: _ObservationRule,
) -> NormalizedRetentionObservation | RetentionExclusion:
    semantics = (
        rule.water_quantity,
        rule.water_unit,
        rule.canonical_water_unit,
        rule.pressure_representation,
        rule.pressure_convention,
        rule.canonical_pressure_unit,
    )
    if any(value is None for value in semantics) or rule.blockers:
        message = f'Confirmed retention rule for {source!r} is incomplete.'
        raise RetentionNormalizationError(message)

    pressure = _finite_number(raw_pressure)
    water = _finite_number(raw_water)
    if pressure is None:
        return _exclusion(
            source_filename,
            sheet_name,
            row,
            'observation',
            'invalid_pressure_value',
            (),
            sample_id,
            source,
            sample_number,
            water_encoding,
            pressure_encoding,
            raw_pressure,
            raw_water,
            lineage,
        )
    if water is None:
        return _exclusion(
            source_filename,
            sheet_name,
            row,
            'observation',
            'invalid_water_content_value',
            (),
            sample_id,
            source,
            sample_number,
            water_encoding,
            pressure_encoding,
            raw_pressure,
            raw_water,
            lineage,
        )

    normalized_water = _normalize_water(water, cast('str', rule.water_unit))
    if normalized_water < 0 or (
        rule.water_quantity == 'volumetric_water_content' and normalized_water > 1
    ):
        return _exclusion(
            source_filename,
            sheet_name,
            row,
            'observation',
            'water_content_out_of_range',
            (),
            sample_id,
            source,
            sample_number,
            water_encoding,
            pressure_encoding,
            raw_pressure,
            raw_water,
            lineage,
        )
    normalized_pressure = _normalize_pressure(
        pressure,
        cast('str', rule.pressure_representation),
        cast('str', rule.pressure_convention),
    )
    if normalized_pressure is None:
        return _exclusion(
            source_filename,
            sheet_name,
            row,
            'observation',
            'pressure_convention_mismatch',
            (),
            sample_id,
            source,
            sample_number,
            water_encoding,
            pressure_encoding,
            raw_pressure,
            raw_water,
            lineage,
        )
    return NormalizedRetentionObservation(
        source_filename=source_filename,
        sheet_name=sheet_name,
        row=row,
        sample_id=_json_scalar(sample_id),
        source=source,
        sample_number=_json_scalar(sample_number),
        water_encoding=water_encoding,
        pressure_encoding=pressure_encoding,
        raw_pressure=pressure,
        raw_water_content=water,
        pressure_quantity='matric_potential',
        pressure_kilopascal=normalized_pressure,
        pressure_unit=cast('str', rule.canonical_pressure_unit),
        pressure_convention='magnitude',
        water_quantity=cast('str', rule.water_quantity),
        water_content=normalized_water,
        water_unit=cast('str', rule.canonical_water_unit),
        water_basis=_water_basis(cast('str', rule.water_quantity)),
        source_cells=lineage,
    )


def _normalize_water(value: float, unit: str) -> float:
    if unit in ('mass_fraction', 'volume_fraction'):
        return value
    if unit in ('mass_percent', 'volume_percent'):
        return value / 100.0
    message = f'Unsupported retention water unit {unit!r}.'
    raise RetentionNormalizationError(message)


def _water_basis(quantity: str) -> str:
    if quantity == 'gravimetric_water_content':
        return 'dry_soil_mass'
    if quantity == 'volumetric_water_content':
        return 'bulk_volume'
    message = f'Unsupported retention water quantity {quantity!r}.'
    raise RetentionNormalizationError(message)


def _normalize_pressure(value: float, representation: str, convention: str) -> float | None:
    if representation in ('pf_log10_centimeter_water_head', 'log10_kilopascal'):
        if convention != 'magnitude':
            return None
        try:
            scale = (
                _CENTIMETER_WATER_HEAD_TO_KILOPASCAL
                if representation == 'pf_log10_centimeter_water_head'
                else 1.0
            )
            normalized = 10.0**value * scale
        except OverflowError:
            return None
        return normalized if math.isfinite(normalized) else None
    if representation in ('kilopascal', 'megapascal'):
        scale = 1000.0 if representation == 'megapascal' else 1.0
        if convention == 'magnitude' and value >= 0:
            normalized = value * scale
        elif convention == 'negative' and value <= 0:
            normalized = abs(value) * scale
        else:
            return None
        return normalized if math.isfinite(normalized) else None
    message = f'Unsupported retention pressure representation {representation!r}.'
    raise RetentionNormalizationError(message)


def _observation_rules(
    document: Mapping[str, object],
) -> dict[tuple[str, str, str | None], _ObservationRule]:
    result: dict[tuple[str, str, str | None], _ObservationRule] = {}
    for source_value in _sequence(document['retention_source_rules']):
        source_rule = _mapping(source_value)
        source = _string(source_rule['source'])
        for observation_value in _sequence(source_rule['observations']):
            observation = _mapping(observation_value)
            key = (
                source,
                _string(observation['water_encoding']),
                cast('str | None', observation['pressure_encoding']),
            )
            if key in result:
                message = f'Duplicate retention source rule for {key!r}.'
                raise RetentionNormalizationError(message)
            result[key] = _ObservationRule(
                status=_string(observation['status']),
                requires_row_review=observation.get('requires_row_review') is True,
                water_quantity=cast('str | None', observation['water_quantity']),
                water_unit=cast('str | None', observation['water_unit']),
                canonical_water_unit=cast('str | None', observation['canonical_water_unit']),
                pressure_representation=cast('str | None', observation['pressure_representation']),
                pressure_convention=cast('str | None', observation['pressure_convention']),
                canonical_pressure_unit=cast('str | None', observation['canonical_pressure_unit']),
                blockers=tuple(_string(value) for value in _sequence(observation['blockers'])),
            )
    return result


def _identity_rules(document: Mapping[str, object], worksheet: str) -> _IdentityRules:
    identity_value = document.get('identity')
    if not isinstance(identity_value, Mapping):
        return _IdentityRules(source_aliases={}, sample_crosswalks={})

    aliases: dict[str, str] = {}
    for value in _sequence(identity_value.get('source_aliases', [])):
        alias = _mapping(value)
        if alias.get('worksheet') == worksheet and alias.get('status') == 'confirmed':
            aliases[_string(alias['raw'])] = _string(alias['canonical'])

    crosswalks: dict[str, _SampleIdentity] = {}
    for value in _sequence(identity_value.get('sample_crosswalks', [])):
        crosswalk = _mapping(value)
        if crosswalk.get('worksheet') != worksheet or crosswalk.get('status') != 'confirmed':
            continue
        raw_sample_id = _string(crosswalk['raw_sample_id'])
        crosswalks[raw_sample_id] = _SampleIdentity(
            sample_id=_json_scalar(crosswalk['canonical_sample_id']),
            source=_string(crosswalk['canonical_source']),
            sample_number=_json_scalar(crosswalk['canonical_sample_number']),
            evidence_cells=tuple(_string(cell) for cell in _sequence(crosswalk['evidence_cells'])),
        )
    return _IdentityRules(source_aliases=aliases, sample_crosswalks=crosswalks)


def _resolve_identity(
    sample_id: object,
    source: str,
    sample_number: object,
    rules: _IdentityRules,
) -> _SampleIdentity:
    if isinstance(sample_id, str) and sample_id in rules.sample_crosswalks:
        return rules.sample_crosswalks[sample_id]
    canonical_source = rules.source_aliases.get(source, source)
    canonical_number = _json_scalar(sample_number)
    return _SampleIdentity(
        sample_id=_json_scalar(f'{canonical_source}_{canonical_number}'),
        source=canonical_source,
        sample_number=canonical_number,
        evidence_cells=(),
    )


def _validate_source_identity(path: Path, digest: str, mapping: DatasetMapping) -> None:
    if path.suffix.casefold() != '.xlsx':
        message = f'Workbook must use the .xlsx extension: {path.name!r}.'
        raise RetentionNormalizationError(message)
    if path.name != mapping.source_filename:
        message = f'Workbook filename {path.name!r} does not match {mapping.source_filename!r}.'
        raise RetentionNormalizationError(message)
    if digest != mapping.source_sha256:
        message = 'Workbook checksum does not match the dataset mapping.'
        raise RetentionNormalizationError(message)


def _validate_profile_counts(
    mapping: DatasetMapping,
    *,
    total_data_rows: int,
    observation_count: int,
    fitted_parameter_count: int,
    unclassified_count: int,
    missing_source_count: int,
) -> None:
    actual = {
        'total_data_rows': total_data_rows,
        'observation_count': observation_count,
        'fitted_parameter_count': fitted_parameter_count,
        'unclassified_count': unclassified_count,
        'missing_source_count': missing_source_count,
    }
    mismatches = [
        field for field, value in actual.items() if mapping.retention_profile.get(field) != value
    ]
    if mismatches:
        message = (
            'Workbook row counts do not match the bound retention profile: '
            f'{", ".join(mismatches)}.'
        )
        raise RetentionNormalizationError(message)


def _exclusion(  # noqa: PLR0917
    source_filename: str,
    sheet_name: str,
    row: int,
    row_role: str,
    reason: str,
    blockers: tuple[str, ...],
    sample_id: object,
    source: str | None,
    sample_number: object,
    water_encoding: str | None,
    pressure_encoding: str | None,
    raw_pressure: object,
    raw_water_content: object,
    source_cells: tuple[str, ...],
) -> RetentionExclusion:
    return RetentionExclusion(
        source_filename=source_filename,
        sheet_name=sheet_name,
        row=row,
        row_role=row_role,
        reason=reason,
        blockers=blockers,
        sample_id=_json_scalar(sample_id),
        source=source,
        sample_number=_json_scalar(sample_number),
        water_encoding=water_encoding,
        pressure_encoding=pressure_encoding,
        raw_pressure=_json_scalar(raw_pressure),
        raw_water_content=_json_scalar(raw_water_content),
        source_cells=source_cells,
    )


def _lineage(sheet_name: str, row: int) -> tuple[str, ...]:
    return tuple(f'{sheet_name}!{column}{row}' for column in 'ABCDEFG')


def _counts(values: Counter[str]) -> tuple[LabelCount, ...]:
    return tuple(LabelCount(label, values[label]) for label in sorted(values))


def _profile_string(mapping: DatasetMapping, field: str) -> str:
    return _string(mapping.retention_profile[field])


def _finite_number(value: object) -> float | None:
    if not isinstance(value, Real) or isinstance(value, bool):
        return None
    number = float(value)
    return number if math.isfinite(number) else None


def _sample_identity_matches(sample_id: object, source: str, sample_number: object) -> bool:
    if sample_number is None or sample_number == '':
        return False
    expected = f'{source}_{sample_number}'
    return sample_id == expected


def _has_value(row: Sequence[object]) -> bool:
    return any(value is not None and value != '' for value in row)


def _optional_label(value: object) -> str | None:
    if value is None or value == '':
        return None
    return str(value)


def _required_label(value: object, row: int, label: str) -> str:
    result = _optional_label(value)
    if result is None:
        message = f'Row {row} has no fitted {label}.'
        raise RetentionNormalizationError(message)
    return result


def _json_scalar(value: object) -> JsonScalar:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open('rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                digest.update(chunk)
    except OSError as error:
        message = f'Cannot read workbook {path.name!r}: {error}'
        raise RetentionNormalizationError(message) from error
    return digest.hexdigest()


def _mapping(value: object) -> Mapping[str, object]:
    if not isinstance(value, Mapping):
        message = 'Dataset mapping contains a non-object rule.'
        raise RetentionNormalizationError(message)
    return cast('Mapping[str, object]', value)


def _sequence(value: object) -> Sequence[object]:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes, bytearray)):
        message = 'Dataset mapping contains a non-array rule list.'
        raise RetentionNormalizationError(message)
    return value


def _string(value: object) -> str:
    if not isinstance(value, str):
        message = 'Dataset mapping contains a non-string rule value.'
        raise RetentionNormalizationError(message)
    return value


def _row_reviews(
    document: Mapping[str, object], worksheet: str
) -> dict[str, dict[int, _RowReview]]:
    result: dict[str, dict[int, _RowReview]] = {}
    occupied: set[int] = set()
    for value in _sequence(document.get('retention_row_reviews', [])):
        policy = _mapping(value)
        source = _string(policy['source'])
        if policy['worksheet'] != worksheet or source in result:
            message = 'Row reviews require unique sources on the bound retention worksheet.'
            raise RetentionNormalizationError(message)
        rows: dict[int, _RowReview] = {}
        for item in _sequence(policy['decisions']):
            decision = _mapping(item)
            row = decision['row']
            if (
                isinstance(row, bool)
                or not isinstance(row, int)
                or row < _FIRST_OBSERVATION_ROW
                or row in occupied
            ):
                message = 'Row-review coordinates must be unique observation rows.'
                raise RetentionNormalizationError(message)
            occupied.add(row)
            cells = tuple(_string(cell) for cell in _sequence(decision['evidence_cells']))
            reason = decision.get('exclusion_reason')
            identity = (
                None
                if reason is not None
                else _SampleIdentity(
                    sample_id=_json_scalar(decision['canonical_sample_id']),
                    source=_string(decision['canonical_source']),
                    sample_number=_json_scalar(decision['canonical_sample_number']),
                    evidence_cells=(),
                )
            )
            if identity is not None and not _sample_identity_matches(
                identity.sample_id, identity.source, identity.sample_number
            ):
                message = 'Reviewed row identities must match their source and sample number.'
                raise RetentionNormalizationError(message)
            rows[row] = _RowReview(identity, None if reason is None else _string(reason), cells)
        if not rows:
            message = 'Row-review policies cannot be empty.'
            raise RetentionNormalizationError(message)
        result[source] = rows
    return result


def _review_point_key(pressure: object, water: object, places: int) -> tuple[float, float]:
    numeric_pressure, numeric_water = _finite_number(pressure), _finite_number(water)
    if numeric_pressure is None or numeric_water is None:
        message = 'Reviewed point matching requires finite pressure and water values.'
        raise RetentionNormalizationError(message)
    return round(numeric_pressure, places), round(numeric_water, places)


def _review_point_index(
    workbook: Workbook,
    document: Mapping[str, object],
    source: str,
) -> tuple[int, dict[tuple[float, float], list[_SampleIdentity]]]:
    rules = [
        _mapping(item)
        for item in _sequence(document.get('retention_reconstruction_rules', []))
        if _mapping(item)['long_source'] == source
    ]
    if len(rules) != 1 or rules[0]['water_scale_status'] != 'confirmed':
        message = 'Row releases require one reconstruction rule with confirmed water scaling.'
        raise RetentionNormalizationError(message)
    rule = rules[0]
    scale = _finite_number(rule['water_scale'])
    if scale is None or scale <= 0:
        message = 'Reviewed water scaling must be finite and positive.'
        raise RetentionNormalizationError(message)
    places = cast('int', rule['comparison_decimal_places'])
    sheet = workbook[_string(rule['master_worksheet'])]
    result: dict[tuple[float, float], list[_SampleIdentity]] = {}
    for value in _sequence(rule['blocks']):
        block = _mapping(value)
        for row in range(
            cast('int', block['first_sample_row']), cast('int', block['last_sample_row']) + 1
        ):
            identity = _SampleIdentity(
                _json_scalar(sheet.cell(row, 1).value),
                _string(rule['master_source']),
                _json_scalar(sheet.cell(row, 2).value),
                (),
            )
            if sheet.cell(row, 3).value != identity.source or not _sample_identity_matches(
                identity.sample_id, identity.source, identity.sample_number
            ):
                message = 'Reviewed master samples require consistent source/id/number cells.'
                raise RetentionNormalizationError(message)
            for column in range(
                column_index_from_string(_string(block['first_column'])),
                column_index_from_string(_string(block['last_column'])) + 1,
            ):
                water = _finite_number(sheet.cell(row, column).value)
                if water is None:
                    continue
                axis_cell = sheet.cell(cast('int', block['axis_row']), column)
                axis = _finite_number(axis_cell.value)
                if axis is None:
                    continue
                key = _review_point_key(axis, water * scale, places)
                evidence = (
                    f'{sheet.title}!{axis_cell.coordinate}',
                    f'{sheet.title}!{sheet.cell(row, column).coordinate}',
                    f'{sheet.title}!A{row}:C{row}',
                )
                result.setdefault(key, []).append(
                    _SampleIdentity(
                        identity.sample_id,
                        identity.source,
                        identity.sample_number,
                        evidence,
                    )
                )
    return places, result
