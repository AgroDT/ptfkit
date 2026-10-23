"""Read-only structural inventory of source XLSX workbooks."""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Literal, TypeAlias, cast
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException


if TYPE_CHECKING:
    from collections.abc import Iterable

    from openpyxl.worksheet.worksheet import Worksheet


JsonScalar: TypeAlias = bool | int | float | str | None
IdentifierReason: TypeAlias = Literal['header_match', 'formula_composite']

_FORMULA_EXAMPLE_LIMIT = 20
_MIN_HEADER_TEXT_CELLS = 2
_EXPLICIT_MISSING_MARKERS = frozenset(
    {
        '-',
        '\u2013',
        '\u2014',
        'na',
        'n/a',
        'n.a.',
        'nan',
        'none',
        'null',
        '\u043d\u0434',
        '\u043d/\u0434',
        '\u043d\u0435\u0442 \u0434\u0430\u043d\u043d\u044b\u0445',
    }
)
_IDENTIFIER_HEADERS = frozenset(
    {
        'id',
        'sample',
        'sample id',
        'sample number',
        'sample no',
        '\u043a\u043e\u0434 \u043e\u0431\u0440\u0430\u0437\u0446\u0430',
        '\u043d\u043e\u043c\u0435\u0440 \u043e\u0431\u0440\u0430\u0437\u0446\u0430',
        '\u043e\u0431\u0440\u0430\u0437\u0435\u0446',
    }
)


class WorkbookInspectionError(ValueError):
    """Raised when a workbook cannot be inventoried without modifying it."""


@dataclass(frozen=True)
class HeaderBlock:
    """One contiguous block of rows inferred to label following data."""

    rows: tuple[int, ...]
    nonempty_cells: int
    text_cells: int


@dataclass(frozen=True)
class HeaderCell:
    """One nonempty cell in the selected header row."""

    row: int
    column: int
    column_letter: str
    coordinate: str
    value: str


@dataclass(frozen=True)
class FormulaExample:
    """One formula retained to make derived columns reviewable."""

    coordinate: str
    expression: str


@dataclass(frozen=True)
class IdentifierCandidate:
    """A column that may identify samples, with uniqueness diagnostics."""

    column: int
    column_letter: str
    header: str | None
    reasons: tuple[IdentifierReason, ...]
    nonempty_values: int
    distinct_values: int
    duplicate_values: int
    formula_cells: int
    example_formula: str | None


@dataclass(frozen=True)
class SheetInventory:
    """Deterministic structural summary of one worksheet."""

    position: int
    title: str
    state: str
    declared_dimension: str
    max_row: int
    max_column: int
    observed_dimension: str | None
    nonempty_cells: int
    hidden_rows: int
    hidden_columns: int
    freeze_panes: str | None
    auto_filter: str | None
    merged_ranges: tuple[str, ...]
    formula_count: int
    formula_cached_values: int
    formula_missing_cached_values: int
    formula_examples: tuple[FormulaExample, ...]
    explicit_missing_markers: dict[str, int]
    header_blocks: tuple[HeaderBlock, ...]
    headers: tuple[HeaderCell, ...]
    identifier_candidates: tuple[IdentifierCandidate, ...]


@dataclass(frozen=True)
class WorkbookInventory:
    """Versioned inventory tied to the exact source workbook bytes."""

    schema_version: int
    source_filename: str
    sha256: str
    size_bytes: int
    sheet_count: int
    sheets: tuple[SheetInventory, ...]


def inspect_workbook(source: str | Path) -> WorkbookInventory:
    """Inspect an XLSX workbook without persisting or recalculating it."""
    path = Path(source)
    if path.suffix.casefold() != '.xlsx':
        message = f'Workbook must use the .xlsx extension: {path.name!r}.'
        raise WorkbookInspectionError(message)
    try:
        digest = _sha256(path)
        size_bytes = path.stat().st_size
        workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
        cached_workbook = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError) as error:
        message = f'Cannot inspect workbook {path.name!r}: {error}'
        raise WorkbookInspectionError(message) from error

    try:
        sheets = tuple(
            _inspect_sheet(position, sheet, cached_workbook[sheet.title])
            for position, sheet in enumerate(workbook.worksheets, start=1)
        )
    finally:
        workbook.close()
        cached_workbook.close()

    return WorkbookInventory(
        schema_version=1,
        source_filename=path.name,
        sha256=digest,
        size_bytes=size_bytes,
        sheet_count=len(sheets),
        sheets=sheets,
    )


def write_workbook_inventory(inventory: WorkbookInventory, destination: str | Path) -> None:
    """Write an inventory as stable UTF-8 JSON with no machine-specific source path."""
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = json.dumps(asdict(inventory), ensure_ascii=False, indent=2, sort_keys=True)
    path.write_text(f'{document}\n', encoding='utf-8')


def _inspect_sheet(position: int, sheet: Worksheet, cached_sheet: Worksheet) -> SheetInventory:
    nonempty_cells: list[Cell] = []
    formula_cells: list[Cell] = []
    missing_markers: Counter[str] = Counter()
    formula_columns: dict[int, list[Cell]] = defaultdict(list)
    header_blocks = _header_blocks(sheet)
    header_rows = frozenset(row for block in header_blocks for row in block.rows)

    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if value is None:
                continue
            nonempty_cells.append(cell)
            if cell.data_type == 'f':
                formula_cells.append(cell)
                formula_columns[cell.column].append(cell)
            elif isinstance(value, str) and cell.row not in header_rows:
                marker = value.strip().casefold()
                if marker in _EXPLICIT_MISSING_MARKERS:
                    missing_markers[marker] += 1

    headers = _headers(sheet, header_blocks)
    identifiers = _identifier_candidates(
        sheet,
        cached_sheet,
        header_rows,
        headers,
        formula_columns,
    )
    cached_formula_values = sum(
        cached_sheet[cell.coordinate].value is not None for cell in formula_cells
    )
    observed_dimension = _observed_dimension(nonempty_cells)
    freeze_panes = sheet.freeze_panes
    if isinstance(freeze_panes, Cell):
        freeze_panes = freeze_panes.coordinate

    return SheetInventory(
        position=position,
        title=sheet.title,
        state=sheet.sheet_state,
        declared_dimension=sheet.calculate_dimension(),
        max_row=sheet.max_row,
        max_column=sheet.max_column,
        observed_dimension=observed_dimension,
        nonempty_cells=len(nonempty_cells),
        hidden_rows=sum(1 for dimension in sheet.row_dimensions.values() if dimension.hidden),
        hidden_columns=sum(1 for dimension in sheet.column_dimensions.values() if dimension.hidden),
        freeze_panes=str(freeze_panes) if freeze_panes is not None else None,
        auto_filter=(str(sheet.auto_filter.ref) if sheet.auto_filter.ref is not None else None),
        merged_ranges=tuple(
            sorted(
                str(cell_range)
                for cell_range in cast('Iterable[object]', sheet.merged_cells.ranges)
            )
        ),
        formula_count=len(formula_cells),
        formula_cached_values=cached_formula_values,
        formula_missing_cached_values=len(formula_cells) - cached_formula_values,
        formula_examples=tuple(
            FormulaExample(cell.coordinate, str(cell.value))
            for cell in formula_cells[:_FORMULA_EXAMPLE_LIMIT]
        ),
        explicit_missing_markers=dict(sorted(missing_markers.items())),
        header_blocks=header_blocks,
        headers=headers,
        identifier_candidates=identifiers,
    )


@dataclass(frozen=True)
class _RowProfile:
    row: int
    nonempty_cells: int
    text_cells: int
    numeric_cells: int
    formula_cells: int
    columns: frozenset[int]


def _header_blocks(sheet: Worksheet) -> tuple[HeaderBlock, ...]:
    profiles = tuple(_row_profile(sheet, row) for row in range(1, sheet.max_row + 1))
    nonempty = tuple(profile for profile in profiles if profile.nonempty_cells)
    blocks: list[HeaderBlock] = []
    used_rows: set[int] = set()
    for position, profile in enumerate(nonempty[:-1]):
        following = nonempty[position + 1]
        if not _looks_like_header(profile) or not _looks_like_data(following, profile):
            continue
        rows = [profile]
        previous_row = profile.row - 1
        while previous_row >= 1:
            previous = profiles[previous_row - 1]
            if (
                previous.row in used_rows
                or not _looks_like_header(previous)
                or not previous.columns.intersection(rows[0].columns)
            ):
                break
            rows.insert(0, previous)
            previous_row -= 1
        row_numbers = tuple(item.row for item in rows)
        if any(row in used_rows for row in row_numbers):
            continue
        used_rows.update(row_numbers)
        blocks.append(
            HeaderBlock(
                rows=row_numbers,
                nonempty_cells=sum(item.nonempty_cells for item in rows),
                text_cells=sum(item.text_cells for item in rows),
            )
        )
    if blocks:
        return tuple(blocks)

    textual = [profile for profile in nonempty if profile.text_cells]
    if not textual:
        return ()
    best = max(
        textual,
        key=lambda profile: (
            profile.text_cells,
            profile.text_cells / profile.nonempty_cells,
            profile.nonempty_cells,
            -profile.row,
        ),
    )
    return (HeaderBlock((best.row,), best.nonempty_cells, best.text_cells),)


def _row_profile(sheet: Worksheet, row_number: int) -> _RowProfile:
    values = tuple(cell for cell in sheet[row_number] if cell.value is not None)
    text_cells = sum(isinstance(cell.value, str) and cell.data_type != 'f' for cell in values)
    formula_cells = sum(cell.data_type == 'f' for cell in values)
    numeric_cells = sum(
        not isinstance(cell.value, str) and cell.data_type != 'f' for cell in values
    )
    return _RowProfile(
        row=row_number,
        nonempty_cells=len(values),
        text_cells=text_cells,
        numeric_cells=numeric_cells,
        formula_cells=formula_cells,
        columns=frozenset(cell.column for cell in values),
    )


def _looks_like_header(profile: _RowProfile) -> bool:
    return (
        profile.text_cells >= _MIN_HEADER_TEXT_CELLS
        and profile.numeric_cells == 0
        and profile.formula_cells == 0
        and profile.text_cells == profile.nonempty_cells
    )


def _looks_like_data(following: _RowProfile, header: _RowProfile) -> bool:
    has_values = following.numeric_cells > 0 or following.formula_cells > 0
    fewer_labels = following.text_cells < header.text_cells
    same_columns = bool(following.columns.intersection(header.columns))
    return has_values and fewer_labels and same_columns


def _headers(sheet: Worksheet, blocks: tuple[HeaderBlock, ...]) -> tuple[HeaderCell, ...]:
    result: list[HeaderCell] = []
    for block in blocks:
        for row_number in block.rows:
            for cell in sheet[row_number]:
                if cell.value is None:
                    continue
                result.append(
                    HeaderCell(
                        row=row_number,
                        column=cell.column,
                        column_letter=get_column_letter(cell.column),
                        coordinate=cell.coordinate,
                        value=str(cell.value),
                    )
                )
    return tuple(result)


def _identifier_candidates(
    sheet: Worksheet,
    cached_sheet: Worksheet,
    header_rows: frozenset[int],
    headers: tuple[HeaderCell, ...],
    formula_columns: dict[int, list[Cell]],
) -> tuple[IdentifierCandidate, ...]:
    reasons: dict[int, set[IdentifierReason]] = defaultdict(set)
    headers_by_column: dict[int, list[str]] = defaultdict(list)
    for header in headers:
        headers_by_column[header.column].append(header.value)
        if _normalize_header(header.value) in _IDENTIFIER_HEADERS:
            reasons[header.column].add('header_match')
    for column, cells in formula_columns.items():
        if any(_is_composite_identifier_formula(str(cell.value)) for cell in cells):
            reasons[column].add('formula_composite')

    candidates: list[IdentifierCandidate] = []
    for column in sorted(reasons):
        values = [
            cached_sheet.cell(row=row, column=column).value
            for row in range(1, sheet.max_row + 1)
            if row not in header_rows
        ]
        populated = [value for value in values if value is not None and value != '']
        distinct = {_stable_scalar(value) for value in populated}
        formulas = formula_columns.get(column, [])
        candidates.append(
            IdentifierCandidate(
                column=column,
                column_letter=get_column_letter(column),
                header=next(iter(headers_by_column.get(column, [])), None),
                reasons=tuple(sorted(reasons[column])),
                nonempty_values=len(populated),
                distinct_values=len(distinct),
                duplicate_values=len(populated) - len(distinct),
                formula_cells=len(formulas),
                example_formula=str(formulas[0].value) if formulas else None,
            )
        )
    return tuple(candidates)


def _normalize_header(value: str) -> str:
    normalized = re.sub(r'[^\w]+', ' ', value.casefold(), flags=re.UNICODE)
    return normalized.replace('_', ' ').strip()


def _is_composite_identifier_formula(expression: str) -> bool:
    compact = expression.replace(' ', '')
    return '&' in compact and ('"_"' in compact or "'_'" in compact)


def _observed_dimension(cells: list[Cell]) -> str | None:
    if not cells:
        return None
    min_row = min(cast('int', cell.row) for cell in cells)
    max_row = max(cast('int', cell.row) for cell in cells)
    min_column = min(cast('int', cell.column) for cell in cells)
    max_column = max(cast('int', cell.column) for cell in cells)
    return f'{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{max_row}'


def _stable_scalar(value: object) -> JsonScalar:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()
