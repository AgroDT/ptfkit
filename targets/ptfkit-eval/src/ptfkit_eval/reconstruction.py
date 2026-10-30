"""Audit wide-to-long retention reconstruction without inferring scientific meaning."""

from __future__ import annotations

import hashlib
import json
import math
from collections import Counter, defaultdict
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import TYPE_CHECKING, TypeAlias, TypeVar, cast
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.exceptions import InvalidFileException


if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

    from .datasets import DatasetMapping, JsonScalar


_AxisValue: TypeAlias = Decimal | str
_Point: TypeAlias = tuple[_AxisValue, Decimal | None]
_Curve: TypeAlias = Counter[_Point]
_ParameterValue: TypeAlias = Decimal | str | None
_Parameter: TypeAlias = tuple[str, _ParameterValue]
_ParameterCurve: TypeAlias = Counter[_Parameter]
_CounterValue = TypeVar('_CounterValue')


class RetentionReconstructionError(ValueError):
    """Raised when reconstruction evidence cannot be audited under its contract."""


@dataclass(frozen=True)
class InferredRetentionCrosswalk:
    """One raw long-sheet sample group matched to exactly one wide curve."""

    raw_sample_id: str
    canonical_sample_id: str
    canonical_source: str
    canonical_sample_number: JsonScalar


@dataclass(frozen=True)
class RetentionReconstructionRuleAudit:
    """Count-only reconstruction evidence for one source-specific rule."""

    rule_id: str
    rule_status: str
    master_source: str
    long_source: str
    water_scale_status: str
    pressure_replacement_count: int
    proposed_pressure_replacement_count: int
    axis_override_count: int
    proposed_axis_override_count: int
    master_curve_count: int
    master_point_count: int
    wide_missing_water_count: int
    excluded_wide_water_count: int
    long_missing_water_count: int
    long_sample_group_count: int
    long_observation_count: int
    matched_point_count: int
    master_only_point_count: int
    long_only_point_count: int
    unique_row_match_count: int
    ambiguous_row_match_count: int
    unmatched_row_count: int
    resolved_group_count: int
    ambiguous_group_count: int
    unresolved_group_count: int
    master_parameter_count: int
    long_parameter_group_count: int
    long_parameter_count: int
    matched_parameter_count: int
    master_only_parameter_count: int
    long_only_parameter_count: int
    resolved_parameter_group_count: int
    ambiguous_parameter_group_count: int
    unresolved_parameter_group_count: int
    parameter_multiset_status: str
    parameter_identity_status: str
    parameter_reconstructable: bool | None
    confirmed_crosswalk_count: int
    matching_confirmed_crosswalk_count: int
    contradictory_confirmed_crosswalk_count: int
    point_multiset_status: str
    identity_status: str
    observation_reconstructable: bool | None
    reconstructable: bool
    normalization_ready: bool
    inferred_crosswalks: tuple[InferredRetentionCrosswalk, ...]
    blockers: tuple[str, ...]


@dataclass(frozen=True)
class RetentionReconstructionAudit:
    """Checksum-bound audit of all declared wide-to-long reconstruction rules."""

    schema_version: int
    source_filename: str
    sha256: str
    rule_count: int
    profile_source_count: int
    unaudited_source_count: int
    reconstructable_rule_count: int
    normalization_ready_rule_count: int
    unaudited_sources: tuple[str, ...]
    rules: tuple[RetentionReconstructionRuleAudit, ...]


def audit_retention_reconstruction_workbook(
    source: str | Path, mapping: DatasetMapping
) -> RetentionReconstructionAudit:
    """Compare declared wide curves with long observations and audit sample identity."""
    path = Path(source)
    digest = _sha256(path)
    _validate_source_identity(path, digest, mapping)
    rules = _sequence(mapping.document['retention_reconstruction_rules'])

    try:
        workbook = load_workbook(path, data_only=True, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError) as error:
        message = f'Cannot audit retention reconstruction in {path.name!r}: {error}'
        raise RetentionReconstructionError(message) from error

    try:
        audits = tuple(_audit_rule(workbook, _mapping(value), mapping) for value in rules)
    finally:
        workbook.close()

    profile_sources = {
        _string(_mapping(value)['source'])
        for value in _sequence(mapping.retention_profile['sources'])
    }
    audited_sources = {_string(_mapping(value)['long_source']) for value in rules}
    unaudited_sources = tuple(sorted(profile_sources - audited_sources))

    return RetentionReconstructionAudit(
        schema_version=1,
        source_filename=mapping.source_filename,
        sha256=digest,
        rule_count=len(audits),
        profile_source_count=len(profile_sources),
        unaudited_source_count=len(unaudited_sources),
        reconstructable_rule_count=sum(audit.reconstructable for audit in audits),
        normalization_ready_rule_count=sum(audit.normalization_ready for audit in audits),
        unaudited_sources=unaudited_sources,
        rules=audits,
    )


def write_retention_reconstruction_audit(
    audit: RetentionReconstructionAudit, destination: str | Path
) -> None:
    """Write count-only reconstruction evidence as stable UTF-8 JSON."""
    path = Path(destination)
    path.write_text(
        json.dumps(asdict(audit), ensure_ascii=False, indent=2) + '\n',
        encoding='utf-8',
    )


def _audit_rule(  # noqa: PLR0915
    workbook: Workbook, rule: Mapping[str, object], mapping: DatasetMapping
) -> RetentionReconstructionRuleAudit:
    rule_id = _string(rule['id'])
    master_sheet_name = _string(rule['master_worksheet'])
    long_sheet_name = _string(rule['long_worksheet'])
    for sheet_name in (master_sheet_name, long_sheet_name):
        if sheet_name not in workbook.sheetnames:
            message = f'Reconstruction rule {rule_id!r} references missing sheet {sheet_name!r}.'
            raise RetentionReconstructionError(message)
    master_sheet = workbook[master_sheet_name]
    long_sheet = workbook[long_sheet_name]

    places = _integer(rule['comparison_decimal_places'])
    replacements = _pressure_replacements(rule, places)
    curves, wide_missing_water_count, excluded_wide_water_count = _wide_curves(
        master_sheet, rule, replacements, places
    )
    long_groups, long_missing_water_count = _long_curves(long_sheet, rule, places)
    master_points = _sum_counters(curves.values())
    long_points = _sum_counters(long_groups.values())
    matched = master_points & long_points
    master_only = master_points - long_points
    long_only = long_points - master_points

    unique_rows, ambiguous_rows, unmatched_rows = _row_match_counts(long_points, curves)
    group_candidates = {
        raw_sample_id: tuple(
            sample_number for sample_number, curve in curves.items() if curve == long_curve
        )
        for raw_sample_id, long_curve in long_groups.items()
    }
    inferred = tuple(
        InferredRetentionCrosswalk(
            raw_sample_id=raw_sample_id,
            canonical_sample_id=f'{_string(rule["master_source"])}_{sample_number}',
            canonical_source=_string(rule['master_source']),
            canonical_sample_number=sample_number,
        )
        for raw_sample_id, candidates in sorted(group_candidates.items())
        if len(candidates) == 1
        for sample_number in candidates
    )
    resolved_count = sum(len(candidates) == 1 for candidates in group_candidates.values())
    ambiguous_count = sum(len(candidates) > 1 for candidates in group_candidates.values())
    unresolved_count = sum(not candidates for candidates in group_candidates.values())

    assigned_samples = [
        candidates[0] for candidates in group_candidates.values() if len(candidates) == 1
    ]
    observation_applicable = bool(_sequence(rule['blocks']))
    identity_exact = (
        resolved_count == len(long_groups)
        and len(long_groups) == len(curves)
        and len(set(assigned_samples)) == len(curves)
        and set(assigned_samples) == set(curves)
    )
    point_multiset_exact = not master_only and not long_only
    observation_reconstructable = (
        point_multiset_exact and identity_exact if observation_applicable else None
    )
    identity_status = (
        'exact_unique'
        if observation_applicable and identity_exact
        else 'unresolved'
        if observation_applicable and unresolved_count
        else 'ambiguous'
        if observation_applicable
        else 'not_applicable'
    )

    parameter_curves = _wide_parameter_curves(master_sheet, rule, places)
    long_parameter_groups = _long_parameter_curves(long_sheet, rule, places)
    parameter_applicable = rule['parameter_block'] is not None
    if not parameter_applicable and long_parameter_groups:
        message = (
            f'Reconstruction rule {rule_id!r} declares no parameter block but '
            'the long source contains VGN rows.'
        )
        raise RetentionReconstructionError(message)
    master_parameters = _sum_counters(parameter_curves.values())
    long_parameters = _sum_counters(long_parameter_groups.values())
    matched_parameters = master_parameters & long_parameters
    master_only_parameters = master_parameters - long_parameters
    long_only_parameters = long_parameters - master_parameters
    parameter_candidates = {
        raw_sample_id: tuple(
            sample_number
            for sample_number, curve in parameter_curves.items()
            if curve == long_curve
        )
        for raw_sample_id, long_curve in long_parameter_groups.items()
    }
    resolved_parameters = sum(len(candidates) == 1 for candidates in parameter_candidates.values())
    ambiguous_parameters = sum(len(candidates) > 1 for candidates in parameter_candidates.values())
    unresolved_parameters = sum(not candidates for candidates in parameter_candidates.values())
    assigned_parameter_samples = [
        candidates[0] for candidates in parameter_candidates.values() if len(candidates) == 1
    ]
    parameter_identity_exact = (
        resolved_parameters == len(long_parameter_groups)
        and len(long_parameter_groups) == len(parameter_curves)
        and len(set(assigned_parameter_samples)) == len(parameter_curves)
        and set(assigned_parameter_samples) == set(parameter_curves)
    )
    parameter_multiset_exact = not master_only_parameters and not long_only_parameters
    parameter_reconstructable = (
        parameter_multiset_exact and parameter_identity_exact if parameter_applicable else None
    )
    parameter_multiset_status = (
        'exact'
        if parameter_applicable and parameter_multiset_exact
        else 'mismatch'
        if parameter_applicable
        else 'not_applicable'
    )
    parameter_identity_status = (
        'exact_unique'
        if parameter_applicable and parameter_identity_exact
        else 'unresolved'
        if parameter_applicable and unresolved_parameters
        else 'ambiguous'
        if parameter_applicable
        else 'not_applicable'
    )
    reconstructable = (
        (observation_reconstructable is not False)
        and (parameter_reconstructable is not False)
        and (observation_applicable or parameter_applicable)
    )

    if not observation_applicable:
        inferred = tuple(
            InferredRetentionCrosswalk(
                raw_sample_id=raw_sample_id,
                canonical_sample_id=f'{_string(rule["master_source"])}_{candidates[0]}',
                canonical_source=_string(rule['master_source']),
                canonical_sample_number=candidates[0],
            )
            for raw_sample_id, candidates in sorted(parameter_candidates.items())
            if len(candidates) == 1
        )

    confirmed_crosswalks = _confirmed_crosswalks(mapping, rule)
    inferred_by_raw = {item.raw_sample_id: item for item in inferred}
    matching_crosswalks = sum(
        _crosswalk_matches(declared, inferred_by_raw.get(raw_sample_id))
        for raw_sample_id, declared in confirmed_crosswalks.items()
    )
    contradictory_crosswalks = sum(
        not _crosswalk_matches(declared, inferred_by_raw.get(raw_sample_id))
        for raw_sample_id, declared in confirmed_crosswalks.items()
    )
    transforms_confirmed = (
        rule['water_scale_status'] == 'confirmed'
        and all(
            _mapping(value)['status'] == 'confirmed'
            for value in _sequence(rule['pressure_replacements'])
        )
        and all(
            _mapping(override)['status'] == 'confirmed'
            for block_value in _sequence(rule['blocks'])
            for override in _mapping(_mapping(block_value).get('axis_overrides', {})).values()
        )
    )
    normalization_ready = (
        reconstructable
        and rule['status'] == 'confirmed'
        and transforms_confirmed
        and matching_crosswalks == len(set(long_groups) | set(long_parameter_groups))
        and len(confirmed_crosswalks) == len(set(long_groups) | set(long_parameter_groups))
    )

    return RetentionReconstructionRuleAudit(
        rule_id=rule_id,
        rule_status=_string(rule['status']),
        master_source=_string(rule['master_source']),
        long_source=_string(rule['long_source']),
        water_scale_status=_string(rule['water_scale_status']),
        pressure_replacement_count=len(replacements),
        proposed_pressure_replacement_count=sum(
            _mapping(value)['status'] == 'proposed'
            for value in _sequence(rule['pressure_replacements'])
        ),
        axis_override_count=sum(
            len(_mapping(_mapping(value).get('axis_overrides', {})))
            for value in _sequence(rule['blocks'])
        ),
        proposed_axis_override_count=sum(
            _mapping(override)['status'] == 'proposed'
            for block_value in _sequence(rule['blocks'])
            for override in _mapping(_mapping(block_value).get('axis_overrides', {})).values()
        ),
        master_curve_count=len(curves),
        master_point_count=sum(master_points.values()),
        wide_missing_water_count=wide_missing_water_count,
        excluded_wide_water_count=excluded_wide_water_count,
        long_missing_water_count=long_missing_water_count,
        long_sample_group_count=len(long_groups),
        long_observation_count=sum(long_points.values()),
        matched_point_count=sum(matched.values()),
        master_only_point_count=sum(master_only.values()),
        long_only_point_count=sum(long_only.values()),
        unique_row_match_count=unique_rows,
        ambiguous_row_match_count=ambiguous_rows,
        unmatched_row_count=unmatched_rows,
        resolved_group_count=resolved_count,
        ambiguous_group_count=ambiguous_count,
        unresolved_group_count=unresolved_count,
        master_parameter_count=sum(master_parameters.values()),
        long_parameter_group_count=len(long_parameter_groups),
        long_parameter_count=sum(long_parameters.values()),
        matched_parameter_count=sum(matched_parameters.values()),
        master_only_parameter_count=sum(master_only_parameters.values()),
        long_only_parameter_count=sum(long_only_parameters.values()),
        resolved_parameter_group_count=resolved_parameters,
        ambiguous_parameter_group_count=ambiguous_parameters,
        unresolved_parameter_group_count=unresolved_parameters,
        parameter_multiset_status=parameter_multiset_status,
        parameter_identity_status=parameter_identity_status,
        parameter_reconstructable=parameter_reconstructable,
        confirmed_crosswalk_count=len(confirmed_crosswalks),
        matching_confirmed_crosswalk_count=matching_crosswalks,
        contradictory_confirmed_crosswalk_count=contradictory_crosswalks,
        point_multiset_status=(
            'exact'
            if observation_applicable and point_multiset_exact
            else 'mismatch'
            if observation_applicable
            else 'not_applicable'
        ),
        identity_status=identity_status,
        observation_reconstructable=observation_reconstructable,
        reconstructable=reconstructable,
        normalization_ready=normalization_ready,
        inferred_crosswalks=inferred,
        blockers=tuple(_string(value) for value in _sequence(rule['blockers'])),
    )


def _wide_curves(
    sheet: Worksheet,
    rule: Mapping[str, object],
    replacements: Mapping[Decimal, Decimal],
    places: int,
) -> tuple[dict[JsonScalar, _Curve], int, int]:
    curves: dict[JsonScalar, _Curve] = {}
    missing_water_count = 0
    excluded_water_count = 0
    excluded_cells = {_string(value) for value in _sequence(rule.get('wide_excluded_cells', ()))}
    seen_excluded_cells: set[str] = set()
    missing_markers = {
        _quantize(_decimal(value, 'wide_missing_markers'), places)
        for value in _sequence(rule.get('wide_missing_markers', ()))
    }
    master_source = _string(rule['master_source'])
    for block_value in _sequence(rule['blocks']):
        block = _mapping(block_value)
        scale = _decimal(block.get('water_scale', rule['water_scale']), 'water_scale')
        axes = _wide_axes(sheet, rule, block, replacements, places)
        first_row = _integer(block['first_sample_row'])
        last_row = _integer(block['last_sample_row'])
        for row in range(first_row, last_row + 1):
            if sheet.cell(row, 3).value != master_source:
                message = (
                    f'{sheet.title}!C{row} does not equal declared master source {master_source!r}.'
                )
                raise RetentionReconstructionError(message)
            sample_number = _json_scalar(sheet.cell(row, 2).value, sheet.cell(row, 2).coordinate)
            if sample_number in curves:
                message = (
                    f'Duplicate master sample number {sample_number!r} in rule {rule["id"]!r}.'
                )
                raise RetentionReconstructionError(message)
            curve, missing_count, seen_excluded = _wide_curve(
                sheet,
                row,
                axes=axes,
                scale=scale,
                missing_markers=missing_markers,
                excluded_cells=excluded_cells,
                preserve_missing=rule['wide_missing_representation'] == 'explicit_long_null',
                places=places,
            )
            missing_water_count += missing_count
            excluded_water_count += len(seen_excluded)
            seen_excluded_cells.update(seen_excluded)
            curves[sample_number] = curve
    _validate_wide_counts(
        rule,
        missing_water_count,
        excluded_water_count,
        excluded_cells,
        seen_excluded_cells,
    )
    return curves, missing_water_count, excluded_water_count


def _wide_axes(
    sheet: Worksheet,
    rule: Mapping[str, object],
    block: Mapping[str, object],
    replacements: Mapping[Decimal, Decimal],
    places: int,
) -> dict[int, _AxisValue]:
    first_column = column_index_from_string(_string(block['first_column']))
    last_column = column_index_from_string(_string(block['last_column']))
    axis_row = _integer(block['axis_row'])
    axis_overrides = _mapping(block.get('axis_overrides', {}))
    block_columns = {get_column_letter(column) for column in range(first_column, last_column + 1)}
    unknown_overrides = set(axis_overrides) - block_columns
    if unknown_overrides:
        message = (
            f'Reconstruction rule {rule["id"]!r} declares axis overrides outside '
            f'the block: {sorted(unknown_overrides)!r}.'
        )
        raise RetentionReconstructionError(message)
    return {
        column: _replace_pressure(
            _axis_value(
                _mapping(axis_overrides[get_column_letter(column)])['long_value']
                if get_column_letter(column) in axis_overrides
                else sheet.cell(axis_row, column).value,
                sheet.cell(axis_row, column).coordinate,
                places,
            ),
            replacements,
            places,
        )
        for column in range(first_column, last_column + 1)
    }


def _wide_curve(
    sheet: Worksheet,
    row: int,
    *,
    axes: Mapping[int, _AxisValue],
    scale: Decimal,
    missing_markers: set[Decimal],
    excluded_cells: set[str],
    preserve_missing: bool,
    places: int,
) -> tuple[_Curve, int, set[str]]:
    curve: _Curve = Counter()
    missing_count = 0
    seen_excluded: set[str] = set()
    for column, pressure in axes.items():
        cell = sheet.cell(row, column)
        if cell.coordinate in excluded_cells:
            seen_excluded.add(cell.coordinate)
            continue
        raw_water = (
            None if cell.value is None else _quantize(_decimal(cell.value, cell.coordinate), places)
        )
        if raw_water is None or raw_water in missing_markers:
            missing_count += 1
            if preserve_missing:
                curve[(pressure, None)] += 1
            continue
        curve[(pressure, _quantize(raw_water * scale, places))] += 1
    return curve, missing_count, seen_excluded


def _validate_wide_counts(
    rule: Mapping[str, object],
    missing_count: int,
    excluded_count: int,
    excluded_cells: set[str],
    seen_excluded_cells: set[str],
) -> None:
    expected_missing = _integer(rule['expected_missing_water_count'])
    if missing_count != expected_missing:
        message = (
            f'Reconstruction rule {rule["id"]!r} expected {expected_missing} missing '
            f'wide water values but found {missing_count}.'
        )
        raise RetentionReconstructionError(message)
    expected_excluded = _integer(rule.get('expected_excluded_wide_water_count', 0))
    if excluded_count != expected_excluded or seen_excluded_cells != excluded_cells:
        message = (
            f'Reconstruction rule {rule["id"]!r} expected {expected_excluded} excluded '
            f'wide water cells but matched {excluded_count} of {len(excluded_cells)} '
            'declared cells.'
        )
        raise RetentionReconstructionError(message)


def _long_curves(
    sheet: Worksheet, rule: Mapping[str, object], places: int
) -> tuple[dict[str, _Curve], int]:
    groups: defaultdict[str, _Curve] = defaultdict(Counter)
    missing_water_count = 0
    long_source = _string(rule['long_source'])
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 2).value != long_source or sheet.cell(row, 4).value == 'VGN':
            continue
        raw_sample_id = sheet.cell(row, 1).value
        if not isinstance(raw_sample_id, str) or not raw_sample_id:
            message = f'{sheet.title}!A{row} must contain a non-empty raw sample identifier.'
            raise RetentionReconstructionError(message)
        pressure_cell = sheet.cell(row, 6)
        water_cell = sheet.cell(row, 7)
        pressure = _axis_value(pressure_cell.value, pressure_cell.coordinate, places)
        if water_cell.value is None:
            missing_water_count += 1
            water = None
        else:
            water = _quantize(_decimal(water_cell.value, water_cell.coordinate), places)
        groups[raw_sample_id][(pressure, water)] += 1
    expected_missing = _integer(rule['expected_missing_long_water_count'])
    if missing_water_count != expected_missing:
        message = (
            f'Reconstruction rule {rule["id"]!r} expected {expected_missing} missing '
            f'long water values but found {missing_water_count}.'
        )
        raise RetentionReconstructionError(message)
    return dict(groups), missing_water_count


def _wide_parameter_curves(
    sheet: Worksheet, rule: Mapping[str, object], places: int
) -> dict[JsonScalar, _ParameterCurve]:
    parameter_value = rule['parameter_block']
    if parameter_value is None:
        return {}
    parameter_block = _mapping(parameter_value)
    header_row = _integer(parameter_block['header_row'])
    columns = [
        column_index_from_string(_string(value)) for value in _sequence(parameter_block['columns'])
    ]
    names: dict[int, str] = {}
    for column in columns:
        cell = sheet.cell(header_row, column)
        if not isinstance(cell.value, str) or not cell.value.strip():
            message = f'{sheet.title}!{cell.coordinate} must name a fitted parameter.'
            raise RetentionReconstructionError(message)
        names[column] = cell.value.strip()

    curves: dict[JsonScalar, _ParameterCurve] = {}
    for rows_value in _sequence(parameter_block['sample_rows']):
        rows = _mapping(rows_value)
        first_row = _integer(rows['first'])
        last_row = _integer(rows['last'])
        for row in range(first_row, last_row + 1):
            sample_number = _json_scalar(sheet.cell(row, 2).value, sheet.cell(row, 2).coordinate)
            if sample_number in curves:
                message = (
                    f'Duplicate master sample number {sample_number!r} in rule {rule["id"]!r}.'
                )
                raise RetentionReconstructionError(message)
            curve: _ParameterCurve = Counter()
            for column, name in names.items():
                cell = sheet.cell(row, column)
                curve[(name, _parameter_value(cell.value, cell.coordinate, places))] += 1
            curves[sample_number] = curve
    return curves


def _long_parameter_curves(
    sheet: Worksheet, rule: Mapping[str, object], places: int
) -> dict[str, _ParameterCurve]:
    groups: defaultdict[str, _ParameterCurve] = defaultdict(Counter)
    long_source = _string(rule['long_source'])
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 2).value != long_source or sheet.cell(row, 4).value != 'VGN':
            continue
        raw_sample_id = sheet.cell(row, 1).value
        if not isinstance(raw_sample_id, str) or not raw_sample_id:
            message = f'{sheet.title}!A{row} must contain a non-empty raw sample identifier.'
            raise RetentionReconstructionError(message)
        name_cell = sheet.cell(row, 6)
        if not isinstance(name_cell.value, str) or not name_cell.value.strip():
            message = f'{sheet.title}!{name_cell.coordinate} must name a fitted parameter.'
            raise RetentionReconstructionError(message)
        value_cell = sheet.cell(row, 7)
        parameter = (
            name_cell.value.strip(),
            _parameter_value(value_cell.value, value_cell.coordinate, places),
        )
        groups[raw_sample_id][parameter] += 1
    return dict(groups)


def _pressure_replacements(rule: Mapping[str, object], places: int) -> dict[Decimal, Decimal]:
    return {
        _quantize(_decimal(item['master_value'], 'master_value'), places): _quantize(
            _decimal(item['long_value'], 'long_value'), places
        )
        for value in _sequence(rule['pressure_replacements'])
        for item in [_mapping(value)]
    }


def _replace_pressure(
    value: _AxisValue, replacements: Mapping[Decimal, Decimal], places: int
) -> _AxisValue:
    if isinstance(value, str):
        return value
    normalized = _quantize(value, places)
    return replacements.get(normalized, normalized)


def _row_match_counts(
    long_points: _Curve, curves: Mapping[JsonScalar, _Curve]
) -> tuple[int, int, int]:
    unique = 0
    ambiguous = 0
    unmatched = 0
    for point, count in long_points.items():
        candidate_count = sum(curve[point] > 0 for curve in curves.values())
        if candidate_count == 1:
            unique += count
        elif candidate_count > 1:
            ambiguous += count
        else:
            unmatched += count
    return unique, ambiguous, unmatched


def _confirmed_crosswalks(
    mapping: DatasetMapping, rule: Mapping[str, object]
) -> dict[str, Mapping[str, object]]:
    identity = _mapping(mapping.document['identity'])
    return {
        _string(crosswalk['raw_sample_id']): crosswalk
        for value in _sequence(identity['sample_crosswalks'])
        for crosswalk in [_mapping(value)]
        if crosswalk['worksheet'] == rule['long_worksheet']
        and crosswalk['canonical_source'] == rule['master_source']
        and crosswalk['status'] == 'confirmed'
    }


def _crosswalk_matches(
    declared: Mapping[str, object], inferred: InferredRetentionCrosswalk | None
) -> bool:
    return (
        inferred is not None
        and declared['canonical_sample_id'] == inferred.canonical_sample_id
        and declared['canonical_source'] == inferred.canonical_source
        and declared['canonical_sample_number'] == inferred.canonical_sample_number
    )


def _sum_counters(
    counters: Iterable[Counter[_CounterValue]],
) -> Counter[_CounterValue]:
    result: Counter[_CounterValue] = Counter()
    for counter in counters:
        result.update(counter)
    return result


def _validate_source_identity(path: Path, digest: str, mapping: DatasetMapping) -> None:
    if path.name != mapping.source_filename:
        message = (
            f'Workbook filename {path.name!r} does not match mapping source '
            f'{mapping.source_filename!r}.'
        )
        raise RetentionReconstructionError(message)
    if digest != mapping.source_sha256:
        message = (
            f'Workbook SHA-256 {digest} does not match mapping source {mapping.source_sha256}.'
        )
        raise RetentionReconstructionError(message)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open('rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                digest.update(chunk)
    except OSError as error:
        message = f'Cannot read workbook {path}: {error}'
        raise RetentionReconstructionError(message) from error
    return digest.hexdigest()


def _decimal(value: object, cell: str) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        message = f'{cell} must contain a finite numeric retention value.'
        raise RetentionReconstructionError(message)
    if isinstance(value, float) and not math.isfinite(value):
        message = f'{cell} must contain a finite numeric retention value.'
        raise RetentionReconstructionError(message)
    try:
        return Decimal(str(value))
    except InvalidOperation as error:
        message = f'{cell} cannot be represented as a decimal value.'
        raise RetentionReconstructionError(message) from error


def _axis_value(value: object, cell: str, places: int) -> _AxisValue:
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            message = f'{cell} must contain a non-empty retention axis value.'
            raise RetentionReconstructionError(message)
        try:
            return _quantize(Decimal(stripped.replace(',', '.')), places)
        except InvalidOperation:
            return stripped
    return _quantize(_decimal(value, cell), places)


def _parameter_value(value: object, cell: str, places: int) -> _ParameterValue:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        try:
            return _quantize(Decimal(stripped.replace(',', '.')), places)
        except InvalidOperation:
            return stripped
    return _quantize(_decimal(value, cell), places)


def _quantize(value: Decimal, places: int) -> Decimal:
    return value.quantize(Decimal(1).scaleb(-places))


def _json_scalar(value: object, cell: str) -> JsonScalar:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    message = f'{cell} contains unsupported sample number {type(value).__name__}.'
    raise RetentionReconstructionError(message)


def _mapping(value: object) -> Mapping[str, object]:
    if not isinstance(value, Mapping):
        message = f'Expected a mapping object, got {type(value).__name__}.'
        raise RetentionReconstructionError(message)
    return cast('Mapping[str, object]', value)


def _sequence(value: object) -> Sequence[object]:
    if isinstance(value, (str, bytes)) or not isinstance(value, Sequence):
        message = f'Expected a sequence, got {type(value).__name__}.'
        raise RetentionReconstructionError(message)
    return cast('Sequence[object]', value)


def _string(value: object) -> str:
    if not isinstance(value, str):
        message = f'Expected a string, got {type(value).__name__}.'
        raise RetentionReconstructionError(message)
    return value


def _integer(value: object) -> int:
    if isinstance(value, bool) or not isinstance(value, int):
        message = f'Expected an integer, got {type(value).__name__}.'
        raise RetentionReconstructionError(message)
    return value
