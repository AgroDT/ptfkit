"""Accounting, exact comparison support and diagnostic region regression contracts."""

import copy
import runpy
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook


def _helpers() -> dict:
    return runpy.run_path(
        str(Path(__file__).resolve().parents[3] / 'evaluation/experiments/meta_v3_3/summaries.py')
    )


def test_row_accounting_rejects_double_dispositions_and_unaccounted_rows():
    reconcile = _helpers()['reconcile_rows']
    document: dict = {
        'summary': {'total_data_rows': 3},
        'observations': [{'row': 2}],
        'fitted_parameters': [{'row': 3}],
        'exclusions': [{'row': 4, 'reason': 'missing_value'}],
    }
    assert [r['status'] for r in reconcile(document)] == [
        'normalized',
        'stored_parameter_not_observation',
        'excluded',
    ]
    duplicate = copy.deepcopy(document)
    duplicate['exclusions'][0]['row'] = 2
    with pytest.raises(ValueError, match='exactly one'):
        reconcile(duplicate)
    document['exclusions'] = []
    with pytest.raises(ValueError, match='exactly one'):
        reconcile(document)


def test_matrix_covers_all_functions_rows_and_targets_without_promoting_raw_readiness():
    build = _helpers()['eligibility_matrix']
    coverage = {
        'functions': [
            {'id': 'public', 'manifest_eligible': True, 'missing_measurement_paths': []},
            {
                'id': 'new_omitted_function',
                'manifest_eligible': False,
                'missing_measurement_paths': ['outputs.theta.measurement'],
            },
        ]
    }
    predictors = {
        'summary': {'candidate_sample_count': 2},
        'eligibility': [
            {
                'sample_key': key,
                'source': 'study',
                'sample_id': sample,
                'function': 'public',
                'eligible': True,
                'reasons': [],
            }
            for key, sample in [('sheet!2', 'a'), ('sheet!3', 'b')]
        ],
    }
    datasets = [
        {'target': {'quantity': quantity}, 'observations': observations}
        for quantity, observations in [
            ('volumetric_water_content', [{'source_id': 'study', 'sample_id': 'a'}]),
            ('gravimetric_water_content', []),
        ]
    ]
    executed = [
        {
            'source_id': 'study',
            'sample_id': 'a',
            'function': 'public',
            'target': 'volumetric_water_content',
        }
    ]
    rows = build(coverage, predictors, datasets, executed)
    assert len(rows) == 8
    assert {(r['master_key'], r['function'], r['target']) for r in rows} == {
        (key, function, quantity)
        for key in ('sheet!2', 'sheet!3')
        for function in ('public', 'new_omitted_function')
        for quantity in ('volumetric_water_content', 'gravimetric_water_content')
    }
    evaluated = [r for r in rows if r['status'] == 'eligible_and_evaluated']
    assert len(evaluated) == 1
    assert evaluated[0]['sample_id'] == 'a'
    omitted = next(r for r in rows if r['function'] == 'new_omitted_function')
    assert 'unresolved_measurement_contract' in omitted['reasons']
    assert omitted['details']['missing_measurement_paths'] == ['outputs.theta.measurement']
    # Numerically complete raw predictors cannot authorize another target or soil.
    raw = next(r for r in rows if r['function'] == 'public' and r['sample_id'] == 'b')
    assert raw['status'] == 'blocked'
    assert {'target_unavailable', 'source_target_protocol_review_required'} <= set(raw['reasons'])
    with pytest.raises(ValueError, match='Every executed'):
        build(coverage, predictors, datasets, [{**executed[0], 'sample_id': 'unknown'}])
    with pytest.raises(ValueError, match='complete public'):
        build(coverage, predictors, datasets, [{**executed[0], 'function': 'new_omitted_function'}])


def test_pressure_regions_keep_exact_boundaries_and_equal_sample_weights():
    diagnostics = _helpers()['region_metrics']
    points = [
        {
            'source_id': 'study',
            'sample_id': sample,
            'suction_kpa': pressure,
            'observed': 0.2,
            'predicted': predicted,
        }
        for sample, pressure, predicted in [
            ('a', 0, 0.2),
            ('a', 9.9, 0.4),
            ('b', 0, 0.6),
            ('a', 10, 0.3),
            ('a', 1499.9, 0.3),
            ('a', 1500, 0.0),
        ]
    ]
    result = diagnostics(points)
    wet, intermediate, dry = [r['summary'] for r in result['regions']]
    assert [r['point_count'] for r in (wet, intermediate, dry)] == [3, 2, 1]
    assert sum(r['point_count'] for r in (wet, intermediate, dry)) == len(points)
    # Sample a has mean error .1 and MSE .02; b has error .4 and MSE .16.
    assert wet['metrics']['mean_error'] == pytest.approx(0.25)
    assert wet['metrics']['root_mean_squared_error'] == pytest.approx(0.3)
    assert intermediate['metrics']['mean_error'] == pytest.approx(0.1)
    assert dry['metrics']['mean_error'] == pytest.approx(-0.2)
    assert result['selection_criterion'] is False
    assert diagnostics(points[:1])['regions'][1]['status'] == 'no_observations'
    with pytest.raises(ValueError, match='Duplicate'):
        diagnostics([*points, points[0]])


def test_scope_does_not_infer_row_roles_from_control_like_numbers():
    scan = runpy.run_path(
        str(Path(__file__).resolve().parents[3] / 'evaluation/experiments/meta_v3_3/scope.py')
    )['workbook_scope']
    book = Workbook()
    master = book.active
    master.title = 'master'
    master.append(['id', 'number', 'source'])
    master.append(['axis', 0.1, 'known'])
    master.append(['known_a', 1, 'known'])
    master.append(['unknown', 0.1, 'new_source'])
    long_sheet = book.create_sheet('long')
    long_sheet.append(['header'])
    long_sheet.append(['point'])
    mapping = {
        'worksheets': [
            {'id': 'samples_master', 'name': 'master', 'header_rows': [[1]]},
            {'id': 'retention', 'name': 'long', 'header_rows': [[1]]},
        ],
        'retention_reconstruction_rules': [
            {'master_worksheet': 'master', 'blocks': [{'axis_row': 2}]},
        ],
    }
    result = scan(
        book, mapping, [{'sheet_name': 'long', 'row': 2, 'status': 'normalized'}], {'master!3'}
    )
    assert result['populated_row_count'] == 6
    statuses = {r['key']: r['status'] for r in result['rows']}
    assert statuses['master!2'] == 'declared_retention_axis_not_sample'
    assert statuses['master!3'] == 'audited_retention_master_row'
    assert statuses['master!4'] == 'unresolved_row_role_and_measurement_contract'
    assert result['unreviewed_master_sources'][0]['sample_count'] is None
    assert result['unreviewed_master_sources'][0]['rows'] == [4]


def test_matrix_does_not_snap_nearby_pressure_to_a_fixed_output():
    build = _helpers()['eligibility_matrix']
    coverage = {
        'functions': [
            {
                'id': 'fixed',
                'manifest_eligible': True,
                'missing_measurement_paths': [],
            }
        ]
    }
    predictors = {
        'summary': {'candidate_sample_count': 1},
        'eligibility': [
            {
                'sample_key': 'sheet!2',
                'source': 'study',
                'sample_id': 'a',
                'function': 'fixed',
                'eligible': False,
                'reasons': ['missing_input'],
            }
        ],
    }
    dataset = {
        'target': {'quantity': 'volumetric_water_content'},
        'observations': [
            {
                'source_id': 'study',
                'sample_id': 'a',
                'suction_kpa': 33.000000001,
            }
        ],
    }
    support = {
        'fixed': {
            'curves': [],
            'fixed_outputs': [
                {
                    'output': 'water',
                    'measurement': {'quantity': 'volumetric_water_content'},
                    'suction_kpa': 33,
                }
            ],
        }
    }
    row = build(coverage, predictors, [dataset], [], support)[0]
    assert 'exact_pressure_target_unavailable' in row['reasons']
    dataset['observations'][0]['suction_kpa'] = 33
    row = build(coverage, predictors, [dataset], [], support)[0]
    assert 'exact_pressure_target_unavailable' not in row['reasons']
    assert row['status'] == 'blocked'
    assert row['scientific_status'] == 'unresolved'
    assert row['reasons'] == ['missing_input']
    predictors['eligibility'][0]['reasons'] = ['measurement_contract_mismatch']
    row = build(coverage, predictors, [dataset], [], support)[0]
    assert row['scientific_status'] == 'scientifically_incompatible'


def test_crosswalk_requires_unique_literal_ids_and_detects_permutations():
    crosswalk = runpy.run_path(
        str(
            Path(__file__).resolve().parents[3] / 'evaluation/experiments/meta_v3_3/source_audit.py'
        )
    )['exact_crosswalk']
    master = [
        {'id': 'a', 'key': 'master!2', 'values': [1, 2]},
        {'id': 'b', 'key': 'master!3', 'values': [3, 4]},
        {'id': 'unknown', 'key': 'master!4', 'values': [1, 2]},
    ]
    original = [
        {'id': 'b', 'key': 'wide!2', 'values': [3, 4]},
        {'id': 'a', 'key': 'wide!3', 'values': [2, 1]},
        {'id': 'other', 'key': 'wide!4', 'values': [1, 2]},
    ]
    a, b, unknown = crosswalk(master, original)
    assert a['original_key'] == 'wide!3'
    assert a['ordered_values_equal'] is False
    assert a['value_multiset_equal'] is True
    assert b['original_key'] == 'wide!2'
    assert b['ordered_values_equal'] is True
    assert unknown['original_key'] is None
    assert unknown['status'] == 'unresolved_identity'
    assert unknown['numeric_candidate_keys'] == ['wide!4']
    assert all(r['measurement_eligibility'] == 'unresolved' for r in (a, b, unknown))
    duplicate = crosswalk(master, [*original, original[1]])[0]
    assert duplicate['original_key'] is None
    duplicate_master = crosswalk([*master, master[0]], original)[0]
    assert duplicate_master['original_key'] is None


def test_blocker_impact_excludes_incompatibility_and_counts_joint_gaps_once():
    impact = _helpers()['blocker_impact']
    rows = [
        {
            'master_key': 'sheet!2',
            'function': 'a',
            'scientific_status': 'unresolved',
            'status': 'blocked',
            'reasons': ['density', 'identity'],
        },
        {
            'master_key': 'sheet!3',
            'function': 'a',
            'scientific_status': 'unresolved',
            'status': 'blocked',
            'reasons': ['density'],
        },
        {
            'master_key': 'sheet!4',
            'function': 'b',
            'scientific_status': 'scientifically_incompatible',
            'status': 'blocked',
            'reasons': ['density', 'measurement_contract_mismatch'],
        },
        {
            'master_key': 'sheet!5',
            'function': 'b',
            'scientific_status': 'eligible',
            'status': 'eligible_and_evaluated',
            'reasons': [],
        },
    ]
    result = impact([{'source_id': 'study', **row} for row in rows])
    assert result['coverage'] == {
        'currently_evaluated': 1,
        'scientifically_incompatible': 1,
        'unresolved_but_potentially_evaluable': 2,
        'structurally_impossible': 0,
    }
    density, identity = result['ranked_blockers']
    assert density['blocker'] == 'density'
    assert density['affected_rows'] == 2
    assert density['affected_functions'] == 1
    assert density['blocked_combinations'] == 2
    assert density['sole_blocker_combinations'] == 1
    assert density['co_blockers'] == {'identity': 1}
    assert identity['potential_gain_upper_bound'] == 1
    assert identity['sole_blocker_combinations'] == 0
    assert result['ranked_source_reviews'][0]['potential_combinations_upper_bound'] == 2
    assert result['ranked_source_reviews'][0]['unresolved_claim_count'] == 2
    assert result['ranked_source_reviews'][0]['unlock_efficiency_upper_bound'] == 1


def test_source_audit_separates_texture_columns_and_preserves_excel_dates():
    audit = runpy.run_path(
        str(
            Path(__file__).resolve().parents[3] / 'evaluation/experiments/meta_v3_3/source_audit.py'
        )
    )['priority_source_audit']
    book = Workbook()
    master = book.active
    master.title = 'master'
    original = book.create_sheet('original')
    for sheet, first in ((master, 24), (original, 36)):
        sheet.cell(1, 1, 'study_a')
        sheet.cell(1, 3, 'study')
        sheet.cell(1, first - 1, 'Q')
        sheet.cell(1, first, 0.3)
    original.cell(1, 12, datetime(2001, 1, 2))  # noqa: DTZ001 - Excel stores naive dates.
    original.cell(1, 96, 17.5)
    master.cell(1, 85, 17.5)
    review = {
        'additional_master_blocks': [
            {
                'raw_source': 'study',
                'workbook_schema': {
                    'original_sheet_index': 1,
                    'first_value_column': 36,
                    'last_value_column': 95,
                    'master_last_value_column': 84,
                    'water_vector_label': 'Q',
                },
                'required_claims': ['water_denominator_and_scale'],
            }
        ],
    }
    result = audit(book, review)['sources'][0]
    assert result['crosswalk'][0]['ordered_values_equal'] is True
    row = result['views'][1][0]
    assert row['values'] == [0.3]
    assert row['separate_predictor_cells'] == {'CR1': 17.5}
    assert row['reported_metadata_cells']['L1'] == {'excel_date_iso8601': '2001-01-02T00:00:00'}
    assert result['released_observation_count'] == 0
