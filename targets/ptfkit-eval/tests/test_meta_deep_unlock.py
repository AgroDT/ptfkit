"""Protect unit-aware coordinates, non-extrapolation and independent discovery."""

import importlib
import json
from pathlib import Path
from types import ModuleType

import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment

from ptfkit_eval.harmonization import EvidenceTier, EvidenceValue, reconstruct_water


@pytest.fixture
def audit(monkeypatch: pytest.MonkeyPatch) -> ModuleType:
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3]))
    return importlib.import_module('evaluation.experiments.meta_v3_3.deep_unlock')


@pytest.mark.parametrize(
    ('description', 'expected'),
    [
        ('Water at 30000 Pa.', 30),
        ('Water at -0.3 bar.', 30),
        ('Water at one-third bar.', 100 / 3),
        ('Water at 100 cm water head.', 9.80665),
        ('Water at pF 2.0.', 9.80665),
        ('Water at pF 0.0.', 0.0980665),
        ('Water at pF -1.0.', 0.00980665),
        ('Water at \u22121,500 kPa.', 1500),
        ('Water at 0 bar.', 0),
        ('Water at 0.3 bar or 1500 kPa.', None),
        ('Water at 10\u201333 kPa.', None),
        ('Field capacity.', None),
        ('Permanent wilting point.', None),
        ('Water at 20 cm soil depth.', None),
        ('Water at pF 1000.', None),
    ],
)
def test_pressure_coordinates_follow_physical_units(
    audit: ModuleType, description: str, expected: float | None
):
    value = audit.contracts.explicit_pressure(description)
    if expected is None:
        assert value is None
    else:
        assert value == pytest.approx(expected)


def test_scan_finds_unmapped_hidden_blocks_comments_and_formulas(audit: ModuleType, tmp_path: Path):
    workbook = Workbook()
    main = workbook.active
    main['B1'] = 'unrelated'
    extra = workbook.create_sheet('unmapped')
    extra.sheet_state = 'hidden'
    extra.merge_cells('B2:C2')
    extra['B2'] = 'Ksat'
    extra['B3'] = 'cm/day'
    extra['B4'] = 12
    extra['B7'] = 'Ksat'
    extra['B8'] = 'm/day'
    extra['B9'] = 0.4
    extra['D10'] = '=B4*2'
    extra['E11'].comment = Comment('particle density', 'reviewer')
    extra.row_dimensions[7].hidden = True
    extra.column_dimensions['B'].hidden = True
    path = tmp_path / 'workbook.xlsx'
    workbook.save(path)
    scan = audit.scan_workbook(path)
    sheet = scan['sheets'][1]
    assert sheet['state'] == 'hidden'
    assert sheet['merged_ranges'] == ['B2:C2']
    assert sheet['hidden_rows'] == [7]
    assert sheet['hidden_columns'] == ['B']
    assert sheet['formula_count'] == 1
    assert sheet['comment_count'] == 1
    locations = {c['location'] for c in scan['candidates']}
    assert {'unmapped!B2', 'unmapped!B7', 'unmapped!E11'} <= locations
    formula = next(c for c in scan['cells'] if c['cell'] == 'D10')
    assert formula['value'] == '=B4*2'
    assert formula['cached'] is None


def test_compound_requires_source_defined_pressure_pair(audit: ModuleType):
    field = {
        'name': 'water',
        'quantity': 'available_water_capacity',
        'unit': 'volume_fraction',
        'description': (
            'Difference between calculated volumetric water contents '
            'at -30 and -1500 kPa matric potential.'
        ),
    }
    function = {'id': 'synthetic', 'fields': [field], 'inputs': [], 'scalar_output': True}
    (target,) = audit.contracts.targets(function, {})
    assert target['kind'] == 'pressure_difference'
    assert target['component_pressures_kpa'] == [30, 1500]
    assert audit.contracts.pool_id(target) == 'volumetric_water_content:difference:30-1500kPa'
    function['documentation'] = {
        'warnings': ['Unclear whether outputs are gravimetric or volumetric.']
    }
    assert audit.contracts.targets(function, {}) == []
    function.pop('documentation')
    field['description'] = 'Available water capacity.'
    assert audit.contracts.targets(function, {}) == []


def test_new_bar_target_uses_existing_nonextrapolating_route(audit: ModuleType):
    # The reconstruction API receives the canonical pressure, never a nearest
    # grid coordinate. Its real interpolation kernel is exercised below.
    measured = [
        (10, EvidenceValue(0.4, EvidenceTier.STRICT, 'measured', ('A1',))),
        (100, EvidenceValue(0.2, EvidenceTier.STRICT, 'measured', ('A2',))),
    ]
    pressure = audit.contracts.explicit_pressure('Water content at 0.3 bar.')
    result = reconstruct_water(measured, pressure, water_quantity='volumetric_water_content')
    assert result is not None
    assert result.value == pytest.approx(0.3045757490560675)
    assert result.tier == EvidenceTier.RECONSTRUCTED
    assert result.sources == ('A1', 'A2')
    assert reconstruct_water(measured, 0, water_quantity='volumetric_water_content') is None
    assert reconstruct_water(measured, 101, water_quantity='volumetric_water_content') is None


def test_compound_keeps_both_parents_and_rejects_missing_component(audit: ModuleType):
    target = {
        'kind': 'pressure_difference',
        'adapter': 'explicit_pressure_difference_v1',
        'quantity': 'volumetric_water_content',
        'field': 'difference',
        'component_pressures_kpa': [30, 1500],
    }
    sample = {'source_id': 'source', 'sample_id': 'a'}
    dataset = {
        'target': {
            'quantity': 'volumetric_water_content',
            'unit': 'volume_fraction',
            'denominator_basis': 'bulk_volume',
            'material_basis': 'fine_earth',
        },
        'observations': [
            {
                **sample,
                'id': str(p),
                'suction_kpa': p,
                'value': v,
                'origin': 'measured',
                'review_status': 'confirmed',
                'quality_flags': [],
                'source_cells': [f'WRC!A{index}'],
            }
            for index, (p, v) in enumerate([(30, 0.4), (1500, 0.1)], 1)
        ],
    }
    context = {'fits': {}, 'old_registration': {}}
    function = {'fields': [{'name': 'difference'}]}
    result, reason = audit.campaign.reconstruct_target(
        context,
        function,
        target,
        dataset,
        sample,
        row=None,
        route='log_pressure_linear_interpolation_v1',
        disabled=set(),
    )
    assert reason is None
    assert result is not None
    assert result.value == pytest.approx(0.3)
    assert result.policy == 'explicit_pressure_difference_v1'
    assert [p.value for p in result.parents] == [0.4, 0.1]
    assert result.sources == ('WRC!A1', 'WRC!A2')
    target['component_pressures_kpa'] = [30, 1501]
    result, reason = audit.campaign.reconstruct_target(
        context,
        function,
        target,
        dataset,
        sample,
        row=None,
        route='log_pressure_linear_interpolation_v1',
        disabled=set(),
    )
    assert result is None
    assert reason == 'target_pressure_outside_measured_range'


def test_explicit_model_variant_uses_registered_evaluator_without_default_parameters(
    audit: ModuleType, monkeypatch: pytest.MonkeyPatch
):
    function = {
        'id': 'synthetic',
        'inputs': [],
        'scalar_output': False,
        'scope': {'models': {'h_theta': 'van Genuchten model with m = 1 - 1/n'}},
        'fields': [
            {
                'name': 'residual',
                'quantity': 'volumetric_water_content',
                'description': 'Residual water content.',
                'unit': 'volume_fraction',
            },
            {
                'name': 'saturated',
                'quantity': 'volumetric_water_content',
                'description': 'Saturated water content.',
                'unit': 'volume_fraction',
            },
            {
                'name': 'inverse_head',
                'quantity': 'van_genuchten_alpha',
                'unit': 'inverse_centimeter',
            },
            {'name': 'shape', 'quantity': 'van_genuchten_n', 'unit': 'dimensionless'},
        ],
    }
    signature = audit.contracts.explicit_vg_signature(function)
    assert signature is not None
    target = {
        'adapter': 'explicit_vg_mualem_signature_v1',
        'parameter_fields': signature,
        'pressure_kpa': 9.80665,
    }
    monkeypatch.setattr(
        audit.contracts,
        'native_outputs',
        lambda *_: {
            'residual': 0.1,
            'saturated': 0.5,
            'inverse_head': 0.01,
            'shape': 2,
        },
    )
    # At alpha*h=1 and n=2, the declared model has effective saturation 1/sqrt(2).
    result = audit.contracts.predicted(function, target, {}, audit.campaign.get_manifest())
    assert result == pytest.approx(0.382842712474619)
    function['scope']['models']['h_theta'] = 'van Genuchten model'
    assert audit.contracts.explicit_vg_signature(function) is None
    function['scope']['models']['h_theta'] = 'van Genuchten model with m = 1 - 1/n'
    function['fields'].pop(0)
    assert audit.contracts.explicit_vg_signature(function) is None


@pytest.mark.parametrize(
    ('name', 'description', 'expected'),
    [
        ('sand', 'Sand from 0.02 to 2 mm; percentage of fine earth.', (20, 2000)),
        ('silt', 'Silt from 0.002 to 0.02 mm; percentage of fine earth.', (2, 20)),
        ('clay_content', 'Clay smaller than 2 micrometres.', (0, 2)),
    ],
)
def test_source_particle_boundaries_override_default_texture_names(
    audit: ModuleType, name: str, description: str, expected: tuple[int, int]
):
    assert audit.campaign.inputs.interval({'name': name, 'description': description}) == expected


def test_pressure_support_stays_sample_specific_and_public_projection_removes_identities(
    audit: ModuleType,
):
    function = {
        'id': 'synthetic',
        'inputs': [],
        'scalar_output': True,
        'fields': [
            {
                'name': 'theta',
                'quantity': 'volumetric_water_content',
                'unit': 'volume_fraction',
                'description': 'Water content at 30 kPa.',
            },
        ],
    }
    samples = [{'source_id': 'private-source', 'sample_id': name} for name in ('a', 'b')]
    observations = []
    for sample, pressures in zip(samples, [(10, 100), (100, 1500)], strict=True):
        observations.extend(
            {
                **sample,
                'suction_kpa': pressure,
                'value': 0.2,
                'id': str(pressure),
                'origin': 'measured',
                'review_status': 'confirmed',
                'quality_flags': [],
                'source_cells': ['private-cell'],
            }
            for pressure in pressures
        )
    context = {
        'support': {},
        'datasets': [{'samples': samples, 'observations': observations, 'target': {}}],
    }
    rows = audit.support_audit(context, [function])
    assert len(rows) == 1
    assert [s['inside_support'] for s in rows[0]['samples']] == [True, False]
    assert [s['gap_kpa'] for s in rows[0]['samples']] == [0, 70]
    aggregate = audit.aggregate_report({'pressure_blockers': rows, 'predictors': []})
    assert aggregate['pressure_blockers'] == [
        {
            'function': 'synthetic',
            'pressure_kpa': 30,
            'assessed_samples': 2,
            'in_range_samples': 1,
            'observed_min_kpa': 10,
            'observed_max_kpa': 1500,
            'minimum_outside_gap_kpa': 70,
        }
    ]
    assert 'private-source' not in str(aggregate)
    assert '"sample_id":' not in json.dumps(aggregate)
    assert '"source_id":' not in json.dumps(aggregate)


def test_named_and_joint_source_bounds_reject_invalid_execution(
    audit: ModuleType, monkeypatch: pytest.MonkeyPatch
):
    inputs = audit.campaign.inputs
    assert inputs.within_named_domain('0.5 <= sand <= 96.4', {'sand': 0.5})
    assert not inputs.within_named_domain('0.5 <= sand <= 96.4', {'sand': 0})
    joint = '0 <= silt <= 100; 0.5 <= 100 - clay - silt <= 96.4'
    assert inputs.within_named_domain(joint, {'clay': 10, 'silt': 30})
    assert not inputs.within_named_domain(joint, {'clay': 10, 'silt': 90})
    function = {
        'id': 'synthetic',
        'inputs': [
            {'name': 'sand', 'unit': '%', 'domain': '0.5 <= sand <= 96.4'},
        ],
    }
    monkeypatch.setattr(
        inputs.previous,
        'texture_predictors',
        lambda *_: {
            'sand': EvidenceValue(0, EvidenceTier.STRICT, 'measured', ('A1',)),
        },
    )
    _, _, reasons = inputs.resolve_inputs(
        function,
        {'features': []},
        {'values': {}},
        'log_psd',
        {},
        representative=False,
        water_inputs={},
    )
    assert reasons == ['computational_domain:sand']
