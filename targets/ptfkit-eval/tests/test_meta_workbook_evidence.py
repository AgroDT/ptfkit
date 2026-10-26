"""Protect evidence gates against scientifically plausible but invalid substitutions."""

import importlib
from pathlib import Path
from types import ModuleType

import pytest
from openpyxl import Workbook


@pytest.fixture
def evidence(monkeypatch: pytest.MonkeyPatch) -> ModuleType:
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3]))
    return importlib.import_module('evaluation.experiments.meta_v3_3.workbook_evidence')


def test_conductivity_conversion_requires_reviewed_dimension(evidence: ModuleType):
    # SI definitions: 1 day = 86400 seconds, 1 cm = 0.01 m.
    assert evidence.conductivity_factor('meter_per_day') == pytest.approx(1 / 86400)
    assert evidence.conductivity_factor('centimeter_per_day') == pytest.approx(1 / 8640000)
    assert evidence.conductivity_factor('millimeter_per_minute') == pytest.approx(1 / 60000)
    assert evidence.conductivity_factor('centimeter_per_square_root_minute') is None
    assert evidence.conductivity_factor('cm/smin') is None
    assert evidence.conductivity_factor(None) is None


def test_unit_conversion_cannot_authorize_absorption_or_model_target(evidence: ModuleType):
    confirmed = {
        'identity_status': 'confirmed_identity',
        'quantity': 'measured_saturated_hydraulic_conductivity',
        'origin': 'measured',
        'method_status': 'confirmed',
        'material_status': 'confirmed',
        'unit': 'centimeter_per_day',
    }
    assert evidence.target_blockers(confirmed) == []
    for field, value, reason in (
        ('quantity', 'infiltration_rate', 'unconfirmed_quantity'),
        ('quantity', 'permeability', 'unconfirmed_quantity'),
        ('origin', 'model_derived', 'unconfirmed_origin'),
        ('method_status', 'unresolved', 'unconfirmed_method_status'),
        ('material_status', 'unresolved', 'unconfirmed_material_status'),
        ('identity_status', 'ambiguous_identity', 'unconfirmed_identity_status'),
        ('unit', None, 'unconfirmed_conductivity_unit'),
    ):
        assert reason in evidence.target_blockers({**confirmed, field: value})


def test_identity_uses_explicit_keys_and_rejects_conflicts(evidence: ModuleType):
    row = {
        'source_id': 'paper-A',
        'sample_id': 'core-8',
        'row': 100,
        'horizon': 'Ap',
        'depth': '0\u201420',
    }
    metadata = {'horizon': 'Ap', 'depth': '0-20 cm'}
    assert (
        evidence.identity_decision(('paper-A', 'core-8'), metadata, [row])['identity_status']
        == 'confirmed_identity'
    )
    # Matching depth/horizon and nearby row cannot repair a different sample key.
    assert (
        evidence.identity_decision(('paper-A', 'core-9'), metadata, [row])['identity_status']
        == 'unresolved_identity'
    )
    conflict = evidence.identity_decision(('paper-A', 'core-8'), {'depth': '20-40 cm'}, [row])
    assert conflict['identity_status'] == 'ambiguous_identity'
    assert conflict['metadata_conflicts'] == ['depth']
    assert (
        evidence.identity_decision(('paper-A', 'core-8'), metadata, [row, {**row, 'row': 200}])[
            'identity_status'
        ]
        == 'ambiguous_identity'
    )


def test_ksat_inventory_keeps_unknown_blocks_and_profile_missing(evidence: ModuleType):
    book = Workbook()
    for i in range(1, 7):
        book.create_sheet(str(i))
    master, thesis = book.worksheets[0], book.worksheets[6]
    for sheet, row in ((master, 7), (thesis, 10), (thesis, 11)):
        sheet[f'A{row}'] = 'explicit-core'
        sheet[f'C{row}'] = 'paper'
    thesis['AB9'] = 'm/day'
    thesis['AB10'] = 0.5
    thesis['K10'] = 'profile-1'
    thesis['AB11'] = 2
    review = {
        'ksat_blocks': [
            {
                'first': 10,
                'last': 10,
                'unit_cell': 'AB9',
                'unit_label': 'm/day',
                'unit': 'meter_per_day',
            }
        ]
    }
    rows = evidence.ksat_inventory(book, book, review, set())
    assert len(rows) == 2
    assert rows[0]['identity_status'] == 'confirmed_identity'
    assert rows[0]['canonical_sample'] is False
    assert rows[0]['target_blockers']
    assert rows[1]['profile'] is None
    assert 'unreviewed_ksat_block' in rows[1]['target_blockers']
    assert rows[1]['unit_conversion_to_m_per_s'] is None


def test_catalog_selection_uses_quantity_not_function_name(evidence: ModuleType):
    context = {
        'catalog': [
            {
                'id': 'unlisted_new_source',
                'inputs': [],
                'fields': [{'quantity': 'saturated_hydraulic_conductivity'}],
            },
            {
                'id': 'log_new_source',
                'inputs': [],
                'fields': [{'quantity': 'log_saturated_hydraulic_conductivity'}],
            },
            {
                'id': 'ahuja_name_but_wrong_quantity',
                'inputs': [],
                'fields': [{'quantity': 'unsaturated_hydraulic_conductivity'}],
            },
        ],
        'registration': {'psd_scenarios': ['reported_boundary_proxy']},
        'old_registration': {},
    }
    observation = {
        'master_rows': [],
        'identity_status': 'unresolved_identity',
        'source_id': 'p',
        'sample_id': 's',
        'profile': None,
        'cell': 'sheet!AB3',
        'target_blockers': ['unconfirmed_quantity'],
    }
    rows = evidence.ksat_catalog_audit(context, [observation], None)
    assert {r['function'] for r in rows} == {'unlisted_new_source', 'log_new_source'}
    assert all(not r['native_executed'] and r['target_blockers'] for r in rows)


def test_operational_targets_do_not_receive_default_pressure_endpoints(evidence: ModuleType):
    for description in ('Gravimetric field capacity.', 'Permanent wilting point.'):
        decision, blocker = evidence.unregistered_target_decision(
            {
                'fields': [{'quantity': 'gravimetric_water_content', 'description': description}],
                'inputs': [],
            }
        )
        assert decision == 'target_exists_related_but_operational_definition_unresolved'
        assert 'WRC water values exist' in blocker
    decision, blocker = evidence.unregistered_target_decision(
        {
            'fields': [{'quantity': 'osmotic_potential', 'description': 'Osmotic pressure.'}],
            'inputs': [],
        }
    )
    assert decision == 'no_compatible_target_measurement_identified'
    assert 'matric WRC pressure is not osmotic pressure' in blocker


def test_predictor_review_discards_title_hits_without_discarding_header(evidence: ModuleType):
    book = Workbook()
    book.active.title = 'master'
    book.active['E5'] = 'publication text containing pH'
    book.active['T1'] = 'pH'
    review = {'reviewed_header_columns': {'ph': ['T', 'W']}}
    assert evidence.reviewed_locations(book, ['master!E5', 'master!T1'], 'ph', review) == [
        'master!T1'
    ]
