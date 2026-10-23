import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from ptfkit_eval import (
    WorkbookInspectionError,
    inspect_workbook,
    write_workbook_inventory,
)


def _fixture_workbook(path: Path) -> None:
    workbook = Workbook()
    samples = workbook.active
    samples.title = 'Samples'
    samples.merge_cells('A1:D1')
    samples['A1'] = 'Soil observations'
    samples.append(['source', 'sample_number', 'sample_id', 'water_content'])
    samples.append(['paper_a', 1, '=A3&"_"&B3', '-'])
    samples.append(['paper_a', 2, '=A4&"_"&B4', 'N/A'])
    samples.freeze_panes = 'A3'
    samples.auto_filter.ref = 'A2:D4'
    samples.row_dimensions[4].hidden = True
    samples.column_dimensions['D'].hidden = True

    notes = workbook.create_sheet('Notes')
    notes.sheet_state = 'hidden'
    notes.append(['key', 'value'])
    notes.append(['owner', 'team'])
    workbook.save(path)
    workbook.close()


def test_inventory_reports_structure_without_modifying_the_workbook(tmp_path: Path):
    workbook_path = tmp_path / 'fixture.xlsx'
    _fixture_workbook(workbook_path)
    original = workbook_path.read_bytes()
    expected_digest = hashlib.sha256(original).hexdigest()

    inventory = inspect_workbook(workbook_path)

    assert workbook_path.read_bytes() == original
    assert inventory.source_filename == 'fixture.xlsx'
    assert inventory.sha256 == expected_digest
    assert inventory.sheet_count == 2

    samples = inventory.sheets[0]
    assert samples.title == 'Samples'
    assert samples.declared_dimension == 'A1:D4'
    assert samples.observed_dimension == 'A1:D4'
    assert samples.merged_ranges == ('A1:D1',)
    assert samples.freeze_panes == 'A3'
    assert samples.auto_filter == 'A2:D4'
    assert samples.hidden_rows == 1
    assert samples.hidden_columns == 1
    assert samples.formula_count == 2
    assert samples.formula_cached_values == 0
    assert samples.formula_missing_cached_values == 2
    assert [formula.coordinate for formula in samples.formula_examples] == ['C3', 'C4']
    assert samples.explicit_missing_markers == {'-': 1, 'n/a': 1}

    assert samples.header_blocks[0].rows == (2,)
    assert [header.value for header in samples.headers] == [
        'source',
        'sample_number',
        'sample_id',
        'water_content',
    ]
    assert [candidate.column_letter for candidate in samples.identifier_candidates] == ['B', 'C']
    sample_id = samples.identifier_candidates[1]
    assert sample_id.reasons == ('formula_composite', 'header_match')
    assert sample_id.formula_cells == 2
    assert sample_id.example_formula == '=A3&"_"&B3'

    assert inventory.sheets[1].state == 'hidden'


def test_detects_repeated_and_multiline_header_blocks(tmp_path: Path):
    workbook_path = tmp_path / 'blocks.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['coarse', 'medium', 'fine'])
    sheet.append([10, 20, 70])
    sheet.append(['coarse', 'medium', 'fine'])
    sheet.append([15, 25, 60])
    sheet.append(['particle size', 'particle size', 'particle size'])
    sheet.append(['> 1 mm', '1-0.1 mm', '< 0.1 mm'])
    sheet.append([5, 35, 60])
    workbook.save(workbook_path)
    workbook.close()

    inventory = inspect_workbook(workbook_path)

    assert [block.rows for block in inventory.sheets[0].header_blocks] == [
        (1,),
        (3,),
        (5, 6),
    ]


def test_inventory_json_is_stable_and_excludes_the_machine_specific_source_path(
    tmp_path: Path,
):
    source_directory = tmp_path / 'private-source-directory'
    source_directory.mkdir()
    workbook_path = source_directory / 'fixture.xlsx'
    _fixture_workbook(workbook_path)
    inventory = inspect_workbook(workbook_path)
    first = tmp_path / 'first.json'
    second = tmp_path / 'second.json'

    write_workbook_inventory(inventory, first)
    write_workbook_inventory(inventory, second)

    assert first.read_bytes() == second.read_bytes()
    document = json.loads(first.read_text(encoding='utf-8'))
    assert document['source_filename'] == 'fixture.xlsx'
    assert str(source_directory) not in first.read_text(encoding='utf-8')


def test_rejects_non_xlsx_sources_before_parsing(tmp_path: Path):
    source = tmp_path / 'fixture.csv'
    source.write_text('a,b\n1,2\n', encoding='utf-8')

    with pytest.raises(WorkbookInspectionError, match=r'\.xlsx extension'):
        inspect_workbook(source)
