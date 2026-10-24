import copy
import hashlib
import json
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from ptfkit_eval import (
    DatasetMapping,
    RetentionNormalizationError,
    normalize_retention_workbook,
    write_retention_normalization_summary,
)


HEADERS = [
    'sample_id',
    'source',
    'sample_number',
    'Формат данных влажности почвы',
    'Формат данных давления влаги',
    'Переменная/Давление влаги',
    'Значение/Содержание влаги',
]


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'ОГХ'
    sheet.append(HEADERS)
    sheet.append(['a_1', 'a', 1, 'W', 'pF', 2, 25])
    sheet.append(['a_1', 'a', 1, 'W', 'pF', 1, 0.25])
    sheet.append(['b_2', 'b', 2, 'Q', 'kPa', -33, 0.5])
    sheet.append(['c_3', 'c', 3, 'W', 'pF', 2, 25])
    sheet.append(['p_4', 'p', 4, 'VGN', None, 'Qr', 0.05])
    sheet.append(['a_5', 'a', 5, 'W', 'pF', 2, 125])
    sheet.append(['x_6', 'x', 6, 'X', 'pF', 2, 20])
    sheet.append(['missing', None, 7, 'W', 'pF', 2, 20])
    sheet.append(['stale_cached_id', 'a', 8, 'W', 'pF', 2, 25])
    workbook.save(path)
    workbook.close()


def _rule(
    *,
    source: str,
    water_encoding: str,
    pressure_encoding: str,
    status: str,
    water_quantity: str,
    water_unit: str | None,
    canonical_water_unit: str,
    pressure_representation: str,
    pressure_convention: str | None,
    blockers: list[str],
) -> dict[str, object]:
    return {
        'source': source,
        'observations': [
            {
                'water_encoding': water_encoding,
                'pressure_encoding': pressure_encoding,
                'status': status,
                'water_quantity': water_quantity,
                'water_unit': water_unit,
                'canonical_water_unit': canonical_water_unit,
                'pressure_representation': pressure_representation,
                'pressure_convention': pressure_convention,
                'canonical_pressure_unit': 'kilopascal',
                'blockers': blockers,
                'evidence': 'Test evidence.',
            }
        ],
    }


def _mapping(path: Path) -> DatasetMapping:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    rules = [
        _rule(
            source='a',
            water_encoding='W',
            pressure_encoding='pF',
            status='confirmed',
            water_quantity='gravimetric_water_content',
            water_unit='mass_percent',
            canonical_water_unit='mass_fraction',
            pressure_representation='pf_log10_centimeter_water_head',
            pressure_convention='magnitude',
            blockers=[],
        ),
        _rule(
            source='b',
            water_encoding='Q',
            pressure_encoding='kPa',
            status='confirmed',
            water_quantity='volumetric_water_content',
            water_unit='volume_fraction',
            canonical_water_unit='volume_fraction',
            pressure_representation='kilopascal',
            pressure_convention='negative',
            blockers=[],
        ),
        _rule(
            source='c',
            water_encoding='W',
            pressure_encoding='pF',
            status='blocked',
            water_quantity='gravimetric_water_content',
            water_unit=None,
            canonical_water_unit='mass_fraction',
            pressure_representation='pf_log10_centimeter_water_head',
            pressure_convention=None,
            blockers=['water_scale_by_source'],
        ),
    ]
    profile = {
        'sheet_name': 'ОГХ',
        'total_data_rows': 9,
        'observation_count': 6,
        'fitted_parameter_count': 1,
        'unclassified_count': 1,
        'missing_source_count': 1,
    }
    return DatasetMapping(
        dataset_id='test',
        status='draft_blocked',
        source_filename=path.name,
        source_sha256=digest,
        document={
            'identity': {
                'source_aliases': [
                    {
                        'worksheet': 'ОГХ',
                        'raw': 'a',
                        'canonical': 'canonical_a',
                        'status': 'confirmed',
                        'evidence': 'Test evidence.',
                    }
                ],
                'sample_crosswalks': [
                    {
                        'worksheet': 'ОГХ',
                        'raw_sample_id': 'a_1',
                        'canonical_sample_id': 'canonical_a_9',
                        'canonical_source': 'canonical_a',
                        'canonical_sample_number': 9,
                        'evidence_cells': ['Данные!X2:Y2'],
                        'status': 'confirmed',
                        'evidence': 'Test evidence.',
                    }
                ],
            },
            'retention_source_rules': rules,
        },
        inventory={},
        retention_profile=profile,
    )


def test_normalizes_only_confirmed_rules_and_balances_every_row(tmp_path: Path):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)

    result = normalize_retention_workbook(path, _mapping(path))

    assert result.summary.total_data_rows == 9
    assert result.summary.input_observation_count == 6
    assert result.summary.normalized_observation_count == 4
    assert result.summary.excluded_observation_count == 2
    assert result.summary.fitted_parameter_count == 1
    assert result.summary.unclassified_exclusion_count == 1
    assert result.summary.missing_source_exclusion_count == 1

    first, second, third, fourth = result.observations
    assert first.pressure_kilopascal == pytest.approx(9.80665)
    assert first.sample_id == 'canonical_a_9'
    assert first.source == 'canonical_a'
    assert first.sample_number == 9
    assert first.pressure_quantity == 'matric_potential'
    assert first.pressure_unit == 'kilopascal'
    assert first.pressure_convention == 'magnitude'
    assert first.water_content == pytest.approx(0.25)
    assert first.water_unit == 'mass_fraction'
    assert first.water_basis == 'dry_soil_mass'
    assert first.source_cells == (
        *(f'ОГХ!{column}2' for column in 'ABCDEFG'),
        'Данные!X2:Y2',
    )
    assert second.water_content == pytest.approx(0.0025)
    assert third.pressure_kilopascal == pytest.approx(33)
    assert third.water_content == pytest.approx(0.5)
    assert third.water_unit == 'volume_fraction'
    assert third.water_basis == 'bulk_volume'
    assert fourth.water_content == pytest.approx(1.25)
    assert fourth.water_basis == 'dry_soil_mass'

    assert [
        (item.model_family, item.parameter_name, item.value) for item in result.fitted_parameters
    ] == [('van_genuchten', 'Qr', 0.05)]
    assert [(item.row, item.reason) for item in result.exclusions] == [
        (5, 'source_rule_blocked'),
        (8, 'unclassified_water_encoding'),
        (9, 'missing_source'),
        (10, 'sample_identity_mismatch'),
    ]
    assert [(item.label, item.count) for item in result.summary.blocker_counts] == [
        ('retention_reconstruction', 1),
        ('water_scale_by_source', 1),
    ]


def test_summary_is_stable_and_does_not_redistribute_measurements(tmp_path: Path):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    result = normalize_retention_workbook(path, _mapping(path))
    first = tmp_path / 'first.json'
    second = tmp_path / 'second.json'

    write_retention_normalization_summary(result, first)
    write_retention_normalization_summary(result, second)

    assert first.read_bytes() == second.read_bytes()
    document = json.loads(first.read_text(encoding='utf-8'))
    assert document['normalized_observation_count'] == 4
    assert '9.80665' not in first.read_text(encoding='utf-8')
    assert '0.05' not in first.read_text(encoding='utf-8')


@pytest.mark.parametrize(
    ('row', 'water', 'expected'),
    [(2, 125, 1.25), (2, -1, None), (4, 1.25, None), (4, -0.01, None)],
)
def test_water_bounds_distinguish_dry_mass_ratio_from_bulk_volume_fraction(
    tmp_path: Path, row: int, water: float, expected: float | None
):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    workbook = load_workbook(path)
    workbook.worksheets[0].cell(row, 7, water)
    workbook.save(path)
    workbook.close()

    result = normalize_retention_workbook(path, _mapping(path))

    if expected is None:
        excluded = next(item for item in result.exclusions if item.row == row)
        assert excluded.reason == 'water_content_out_of_range'
        assert not any(item.row == row for item in result.observations)
    else:
        observation = next(item for item in result.observations if item.row == row)
        # 1.25 g water / 1 g dry soil is 125%; it is not a volume fraction.
        assert observation.water_content == pytest.approx(expected)
        assert observation.water_basis == 'dry_soil_mass'
        assert not any(item.row == row for item in result.exclusions)


def test_rejects_workbook_bytes_not_bound_to_the_mapping(tmp_path: Path):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    mapping = _mapping(path)
    changed_path = tmp_path / 'changed.xlsx'
    changed_path.write_bytes(path.read_bytes() + b'changed')
    changed_mapping = DatasetMapping(
        dataset_id=mapping.dataset_id,
        status=mapping.status,
        source_filename=changed_path.name,
        source_sha256=mapping.source_sha256,
        document=mapping.document,
        inventory=mapping.inventory,
        retention_profile=mapping.retention_profile,
    )

    with pytest.raises(RetentionNormalizationError, match='checksum'):
        normalize_retention_workbook(changed_path, changed_mapping)


@pytest.mark.parametrize(
    ('representation', 'encoding', 'axis', 'convention', 'expected'),
    [
        ('log10_kilopascal', 'pF', -1, 'magnitude', 0.1),
        ('log10_kilopascal', 'pF', 0, 'magnitude', 1),
        ('log10_kilopascal', 'pF', 3, 'magnitude', 1000),
        ('log10_kilopascal', 'pF', 2, 'negative', None),
        ('log10_kilopascal', 'pF', 400, 'magnitude', None),
        ('kilopascal', 'kPa', 0, 'magnitude', 0),
        ('megapascal', 'mPa', 0.001, 'magnitude', 1),
        ('megapascal', 'mPa', 140, 'magnitude', 140000),
        ('megapascal', 'mPa', -0.01, 'negative', 10),
        ('megapascal', 'mPa', -0.01, 'magnitude', None),
        ('megapascal', 'mPa', 1e308, 'magnitude', None),
    ],
)
def test_reviewed_pressure_representation_overrides_raw_label(
    tmp_path: Path,
    *,
    representation: str,
    encoding: str,
    axis: float,
    convention: str,
    expected: float | None,
):
    path = tmp_path / 'log-pressure.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '\u041e\u0413\u0425'
    sheet.append(HEADERS)
    sheet.append(['b_2', 'b', 2, 'Q', encoding, axis, 0.5])
    workbook.save(path)
    workbook.close()
    mapping = _mapping(path)
    document = dict(mapping.document)
    document['retention_source_rules'] = [
        _rule(
            source='b',
            water_encoding='Q',
            pressure_encoding=encoding,
            status='confirmed',
            water_quantity='volumetric_water_content',
            water_unit='volume_fraction',
            canonical_water_unit='volume_fraction',
            pressure_representation=representation,
            pressure_convention=convention,
            blockers=[],
        )
    ]
    mapping = replace(
        mapping,
        document=document,
        retention_profile={
            'sheet_name': '\u041e\u0413\u0425',
            'total_data_rows': 1,
            'observation_count': 1,
            'fitted_parameter_count': 0,
            'unclassified_count': 0,
            'missing_source_count': 0,
        },
    )

    result = normalize_retention_workbook(path, mapping)

    if expected is None:
        assert result.observations == ()
        assert [item.reason for item in result.exclusions] == ['pressure_convention_mismatch']
        return
    assert len(result.observations) == 1
    observation = result.observations[0]
    assert observation.pressure_encoding == encoding
    assert observation.raw_pressure == axis
    assert observation.pressure_kilopascal == pytest.approx(expected)
    assert observation.water_content == 0.5
    assert observation.water_quantity == 'volumetric_water_content'
    assert observation.water_unit == 'volume_fraction'
    assert observation.water_basis == 'bulk_volume'


def _row_review_mapping(path: Path) -> DatasetMapping:
    workbook = load_workbook(path)
    workbook['\u041e\u0413\u0425']['G10'] = 21
    master = workbook.create_sheet('master')
    master['X1'], master['Y1'] = 2, 1
    for row, number, column, water in [(2, 9, 24, 25), (3, 10, 25, 0.25), (4, 11, 24, 21)]:
        master.cell(row, 1, f'canonical_a_{number}')
        master.cell(row, 2, number)
        master.cell(row, 3, 'canonical_a')
        master.cell(row, column, water)
    workbook.save(path)
    workbook.close()
    mapping = _mapping(path)
    document = dict(mapping.document)
    document['retention_reconstruction_rules'] = [
        {
            'long_source': 'a',
            'master_source': 'canonical_a',
            'master_worksheet': 'master',
            'water_scale_status': 'confirmed',
            'water_scale': 1,
            'comparison_decimal_places': 12,
            'blocks': [
                {
                    'first_sample_row': 2,
                    'last_sample_row': 4,
                    'axis_row': 1,
                    'first_column': 'X',
                    'last_column': 'Y',
                }
            ],
        }
    ]
    rules = document['retention_source_rules']
    assert isinstance(rules, list)
    rules[0]['observations'][0]['requires_row_review'] = True
    document['retention_row_reviews'] = [
        {
            'source': 'a',
            'worksheet': '\u041e\u0413\u0425',
            'evidence': 'Independent row identities from the original wide records.',
            'decisions': [
                {
                    'row': row,
                    'canonical_sample_id': f'canonical_a_{number}',
                    'canonical_source': 'canonical_a',
                    'canonical_sample_number': number,
                    'evidence_cells': [
                        f'master!{column}1',
                        f'master!{column}{master_row}',
                        f'master!A{master_row}:C{master_row}',
                    ],
                }
                for row, number, master_row, column in [
                    (2, 9, 2, 'X'),
                    (3, 10, 3, 'Y'),
                    (10, 11, 4, 'X'),
                ]
            ]
            + [
                {
                    'row': 7,
                    'exclusion_reason': 'ambiguous_sample_identity',
                    'evidence_cells': ['master!X5', 'master!X6'],
                }
            ],
        }
    ]
    return replace(mapping, document=document)


def test_row_reviews_separate_shared_ids_and_repair_stale_cached_ids(tmp_path: Path):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    result = normalize_retention_workbook(path, _row_review_mapping(path))
    assert [(p.row, p.sample_id, p.sample_number) for p in result.observations] == [
        (2, 'canonical_a_9', 9),
        (3, 'canonical_a_10', 10),
        (4, 'b_2', 2),
        (10, 'canonical_a_11', 11),
    ]
    assert result.observations[0].water_content == pytest.approx(0.25)
    assert result.observations[1].water_content == pytest.approx(0.0025)
    assert result.observations[-1].pressure_kilopascal == pytest.approx(9.80665)
    assert 'master!A4:C4' in result.observations[-1].source_cells
    excluded = next(p for p in result.exclusions if p.row == 7)
    assert excluded.reason == 'ambiguous_sample_identity'
    assert excluded.raw_water_content == 125
    assert excluded.source_cells[-2:] == ('master!X5', 'master!X6')
    assert result.summary.normalized_observation_count == 4
    assert result.summary.excluded_observation_count == 2


@pytest.mark.parametrize('defect', ['missing', 'extra', 'duplicate', 'worksheet', 'identity'])
def test_row_reviews_fail_closed_on_incomplete_or_inconsistent_evidence(
    tmp_path: Path, defect: str
):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    mapping = _row_review_mapping(path)
    document = copy.deepcopy(dict(mapping.document))
    policies = document['retention_row_reviews']
    assert isinstance(policies, list)
    policy = policies[0]
    decisions = policy['decisions']
    if defect == 'missing':
        decisions.pop(0)
    elif defect == 'extra':
        decisions.append(dict(decisions[0], row=6))
    elif defect == 'duplicate':
        decisions.append(dict(decisions[0]))
    elif defect == 'worksheet':
        policy['worksheet'] = 'other'
    else:
        decisions[0]['canonical_sample_id'] = 'wrong'
    with pytest.raises(RetentionNormalizationError):
        normalize_retention_workbook(path, replace(mapping, document=document))


def test_row_review_identity_does_not_override_blocked_physical_semantics(tmp_path: Path):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    mapping = _mapping(path)
    document = dict(mapping.document)
    document['retention_row_reviews'] = [
        {
            'source': 'c',
            'worksheet': '\u041e\u0413\u0425',
            'evidence': 'Identity only.',
            'decisions': [
                {
                    'row': 5,
                    'canonical_sample_id': 'canonical_c_1',
                    'canonical_source': 'canonical_c',
                    'canonical_sample_number': 1,
                    'evidence_cells': ['master!A1:C1'],
                }
            ],
        }
    ]
    result = normalize_retention_workbook(path, replace(mapping, document=document))
    assert all(p.row != 5 for p in result.observations)
    excluded = next(p for p in result.exclusions if p.row == 5)
    assert excluded.reason == 'source_rule_blocked'
    assert excluded.blockers == ('water_scale_by_source',)


def test_required_row_review_cannot_be_removed_to_restore_cached_id_fallback(tmp_path: Path):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    mapping = _row_review_mapping(path)
    document = dict(mapping.document)
    del document['retention_row_reviews']
    with pytest.raises(RetentionNormalizationError, match='requires exhaustive'):
        normalize_retention_workbook(path, replace(mapping, document=document))


@pytest.mark.parametrize(
    'defect',
    ['wrong_identity', 'ambiguous_master', 'wrong_axis', 'unreviewed_scale', 'wrong_provenance'],
)
def test_row_release_rechecks_master_point_evidence(tmp_path: Path, defect: str):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    mapping = _row_review_mapping(path)
    document = copy.deepcopy(dict(mapping.document))
    if defect == 'wrong_identity':
        policies = document['retention_row_reviews']
        assert isinstance(policies, list)
        policies[0]['decisions'][0].update(
            canonical_sample_id='canonical_a_10', canonical_sample_number=10
        )
    elif defect == 'wrong_provenance':
        policies = document['retention_row_reviews']
        assert isinstance(policies, list)
        policies[0]['decisions'][0]['evidence_cells'] = ['master!A99:C99']
    elif defect == 'unreviewed_scale':
        rules = document['retention_reconstruction_rules']
        assert isinstance(rules, list)
        rules[0]['water_scale_status'] = 'proposed'
    else:
        workbook = load_workbook(path)
        if defect == 'ambiguous_master':
            workbook['master']['X4'] = 25
            policies = document['retention_row_reviews']
            assert isinstance(policies, list)
            policies[0]['decisions'][2] = {
                'row': 10,
                'exclusion_reason': 'ambiguous_sample_identity',
                'evidence_cells': ['master!X4'],
            }
        else:
            workbook['master']['X1'] = 3
        workbook.save(path)
        workbook.close()
        mapping = replace(mapping, source_sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    with pytest.raises(
        RetentionNormalizationError,
        match=r'master point|master identity|confirmed water scaling|provenance',
    ):
        normalize_retention_workbook(path, replace(mapping, document=document))


def test_two_long_rows_cannot_release_the_same_master_point_twice(tmp_path: Path):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    mapping = _row_review_mapping(path)
    workbook = load_workbook(path)
    workbook['\u041e\u0413\u0425'].append(['a_99', 'a', 99, 'W', 'pF', 2, 25])
    workbook.save(path)
    workbook.close()
    document = copy.deepcopy(dict(mapping.document))
    policies = document['retention_row_reviews']
    assert isinstance(policies, list)
    policies[0]['decisions'].append(dict(policies[0]['decisions'][0], row=11))
    profile = dict(mapping.retention_profile, total_data_rows=10, observation_count=7)
    mapping = replace(
        mapping,
        document=document,
        retention_profile=profile,
        source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
    )
    with pytest.raises(RetentionNormalizationError, match='unused exact master identity'):
        normalize_retention_workbook(path, mapping)


def test_missing_original_axis_cannot_be_released_but_does_not_block_other_reviewed_rows(
    tmp_path: Path,
):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    mapping = _row_review_mapping(path)
    workbook = load_workbook(path)
    workbook['master']['Y1'] = None
    workbook.save(path)
    workbook.close()
    mapping = replace(mapping, source_sha256=hashlib.sha256(path.read_bytes()).hexdigest())

    with pytest.raises(RetentionNormalizationError, match='master point'):
        normalize_retention_workbook(path, mapping)

    document = copy.deepcopy(dict(mapping.document))
    policies = document['retention_row_reviews']
    assert isinstance(policies, list)
    policies[0]['decisions'][1] = {
        'row': 3,
        'exclusion_reason': 'unreviewed_pressure_value',
        'evidence_cells': ['master!Y1', 'master!Y3'],
    }
    result = normalize_retention_workbook(path, replace(mapping, document=document))
    assert [item.row for item in result.observations] == [2, 4, 10]
    excluded = next(item for item in result.exclusions if item.row == 3)
    assert excluded.reason == 'unreviewed_pressure_value'
    assert excluded.source_cells[-2:] == ('master!Y1', 'master!Y3')


def test_reviewed_table_statistic_is_excluded_before_pressure_interpretation(tmp_path: Path):
    path = tmp_path / 'retention.xlsx'
    _workbook(path)
    mapping = _row_review_mapping(path)
    document = copy.deepcopy(dict(mapping.document))
    policies = document['retention_row_reviews']
    assert isinstance(policies, list)
    policies[0]['decisions'][0] = {
        'row': 2,
        'exclusion_reason': 'non_observation_table_value',
        'evidence_cells': ['master!X1', 'master!X2'],
    }

    result = normalize_retention_workbook(path, replace(mapping, document=document))

    assert [item.row for item in result.observations] == [3, 4, 10]
    excluded = next(item for item in result.exclusions if item.row == 2)
    assert excluded.reason == 'non_observation_table_value'
    assert excluded.raw_pressure == 2
    assert excluded.raw_water_content == 25
    assert excluded.source_cells[-2:] == ('master!X1', 'master!X2')
