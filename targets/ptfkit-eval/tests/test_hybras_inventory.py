import hashlib
import runpy
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from ptfkit_eval import load_retention_dataset


@pytest.mark.parametrize('orphan', [False, True])
def test_hybras_inventory_distinguishes_derived_saturation_and_rejects_orphan(
    tmp_path: Path, *, orphan: bool
):
    audit = runpy.run_path(
        str(Path(__file__).resolve().parents[3] / 'evaluation/experiments/hybras/audit.py')
    )['audit_archive']
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'HYBRAS'
    for _ in range(3):
        sheet.append(['header'])
    sheet.append([1, None, None, None, None, 7])
    sheet.append([2, None, None, None, None, 7])
    raw = workbook.create_sheet('RAWRET_lines')
    raw.append(['code', 'WR_ID', 'flag', 'head', 'theta', 'comments1'])
    for pressure in [0, 10, 100]:
        raw.append([1, '&1', 'l', pressure, 0.4, 'theta(h=0cm)=porosidade total=derived'])
    raw.append([99 if orphan else 2, '&2', 'l', 0, 0.5, 'measured saturation'])
    stream = BytesIO()
    workbook.save(stream)
    archive_path = tmp_path / 'source.zip'
    with ZipFile(archive_path, 'w') as archive:
        archive.writestr('HYBRAS_COMPLETE_V1.xlsx', stream.getvalue())
    if orphan:
        with pytest.raises(ValueError, match='no horizon metadata'):
            audit(archive_path)
        return
    result = audit(archive_path)
    assert result['horizons'] == 2
    assert result['distinct_profile_ids'] == 1
    assert result['raw_table_points'] == 4
    assert result['zero_head_points'] == 2
    assert result['zero_head_derived_from_porosity'] == 1
    assert result['positive_head_points'] == 2


def test_hybras_import_preserves_reported_heads_and_provenance(monkeypatch: pytest.MonkeyPatch):
    namespace = runpy.run_path(
        str(Path(__file__).resolve().parents[3] / 'evaluation/experiments/hybras/build.py')
    )
    build = namespace['build_dataset']
    with pytest.raises(ValueError, match='archive checksum'):
        build(b'unreviewed archive')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'HYBRAS'
    for _ in range(3):
        sheet.append(['header'])
    for code in [1, 2]:
        row = [None] * 143
        row[0], row[5], row[36] = code, 7, 'Independent fixture study'
        row[124], row[127], row[129] = 'pressure plate', 'secagem', 'indeformada'
        sheet.append(row)
    raw = workbook.create_sheet('RAWRET_lines')
    raw.append(['code', 'WR_ID', 'flag', 'head', 'theta', 'comments1', 'comments2'])
    raw.append([1, '&1', 'l', 0, 0.5, 'porosidade total=derived'])
    raw.append([1, '&1', 'l', 337, 0.3, 'porosidade total=derived', 'digitized; 1atm=1033'])
    raw.append([1, '&1', 'l', 1000, 0.4, 'porosidade total=derived', 'nonmonotone'])
    raw.append([2, '&2', 'l', 0, 0.4, 'measured saturation'])
    raw.append([2, '&2', 'l', 15300, 0.1, 'measured saturation', 'replicate mean; 1atm=1020'])
    stream = BytesIO()
    workbook.save(stream)
    archive_stream = BytesIO()
    with ZipFile(archive_stream, 'w') as archive:
        archive.writestr('HYBRAS_COMPLETE_V1.xlsx', stream.getvalue())
    archive_bytes = archive_stream.getvalue()
    monkeypatch.setitem(
        build.__globals__, 'ARCHIVE_SHA256', hashlib.sha256(archive_bytes).hexdigest()
    )
    monkeypatch.setitem(
        build.__globals__, 'WORKBOOK_SHA256', hashlib.sha256(stream.getvalue()).hexdigest()
    )
    document, provenance = build(archive_bytes)
    assert [p['value'] for p in document['observations']] == [0.3, 0.4, 0.1]
    # 337 cm is 33.0484105 kPa, not the approximate wide-table label 33 kPa.
    assert document['observations'][0]['suction_kpa'] == pytest.approx(33.0484105, abs=1e-12)
    assert document['observations'][0]['suction_kpa'] != 33
    assert document['observations'][0]['source_cells'] == ['RAWRET_lines!A3:H3']
    assert {s['profile_id'] for s in document['samples']} == {'hybras-profile-7'}
    assert len({s['source_id'] for s in document['samples']}) == 1
    assert all(not s['features'] for s in document['samples'])
    assert provenance['observations'][0]['comments'][1] == 'digitized; 1atm=1033'
    assert provenance['observations'][2]['comments'][1] == 'replicate mean; 1atm=1020'
    assert [p['reason'] for p in provenance['omitted_zero_head']] == [
        'derived_porosity',
        'saturation_deferred',
    ]
    summary = namespace['summarize_import'](document, provenance)
    assert summary['curves_with_physical_diagnostics'] == 1
    assert summary['diagnostics'] == {'increasing_water_with_suction': 1}
    assert summary['deferred_zero_head'] == {'derived_porosity': 1, 'saturation_deferred': 1}
    assert summary['pressure_conversion_blocked_observations'] == 1
    subset = namespace['confirmed_subset'](document)
    assert [p['value'] for p in subset['observations']] == [0.3, 0.4]
    assert len(subset['samples']) == 1
    assert document['observations'][2]['review_status'] == 'blocked'
    assert document['observations'][2]['quality_flags'] == ['pressure_conversion_review_required']
    assert [p['value'] for p in document['observations']] == [0.3, 0.4, 0.1]
    _assert_fit_reuse_requires_unchanged_observations(subset, document)


def _assert_fit_reuse_requires_unchanged_observations(subset: dict, document: dict) -> None:
    review = runpy.run_path(
        str(Path(__file__).resolve().parents[3] / 'evaluation/experiments/hybras/review_fits.py')
    )['review_fits']
    source_id = subset['samples'][0]['source_id']
    parent: dict = {
        'experiment': 'measured_retention_reference_fits_v1',
        'dataset': document,
        'dataset_sha256': load_retention_dataset(document).sha256,
        'results': [
            {
                'source_id': source_id,
                'sample_id': 'hybras-horizon-1',
                'fit': {
                    'plan': {'family_id': family},
                    'status': 'fitted',
                    'reason': None,
                    'observations': [
                        {
                            'reference': 'hybras-raw-row-3',
                            'suction_kpa': 337 * 0.0980665,
                            'water_content': 0.3,
                        },
                        {
                            'reference': 'hybras-raw-row-4',
                            'suction_kpa': 98.0665,
                            'water_content': 0.4,
                        },
                    ],
                },
            }
            for family in [
                'van_genuchten_mualem_v1',
                'van_genuchten_m1_v1',
                'brooks_corey_v1',
                'clapp1978_gradual_air_entry_v1',
            ]
        ],
    }
    reviewed = review(parent, subset)
    assert reviewed['fits'] == 4
    assert reviewed['samples'] == 1
    changed = deepcopy(parent)
    changed['results'][0]['fit']['observations'][0]['water_content'] = 0.31
    with pytest.raises(ValueError, match='changed observation series'):
        review(changed, subset)
    duplicate = deepcopy(parent)
    duplicate['results'].append(duplicate['results'][0])
    with pytest.raises(ValueError, match='Duplicate fit'):
        review(duplicate, subset)
    _assert_ordered_comparison_requires_complete_cohort(parent, subset)


def _assert_ordered_comparison_requires_complete_cohort(parent: dict, subset: dict) -> None:
    compare = runpy.run_path(
        str(
            Path(__file__).resolve().parents[3] / 'evaluation/experiments/hybras/compare_ordered.py'
        )
    )['compare']
    for record in parent['results']:
        record['fit']['plan']['target'] = subset['target']
        record['fit']['plan']['parameters'] = [
            {'name': 'theta_r', 'lower': 0, 'upper': 0.2},
            {'name': 'theta_s', 'lower': 0.3, 'upper': 0.8},
        ]
        record['fit']['attempts'] = [{'converged': True, 'sum_squared_error': 1}]
    records = deepcopy(
        [
            record
            for record in parent['results']
            if record['fit']['plan']['family_id'] in {'brooks_corey_v1', 'van_genuchten_mualem_v1'}
        ]
    )
    for record in records:
        record['fit']['plan']['parameters'][0]['upper'] = 0.533
        record['fit']['plan']['parameters'][1]['lower'] = 0.1
        record['fit']['attempts'][0]['sum_squared_error'] = 0.5
        record['fit']['parameters'] = [
            {'name': 'theta_r', 'parameter_interval': {'upper': 0.3}},
            {'name': 'theta_s', 'parameter_interval': {'lower': 0.4}},
        ]
        record['quality'] = {'diagnostics': []}
        record['uncertainty_support'] = {
            'parameter_intervals_outside_bounds': [],
            'mean_intervals_outside_physical_bounds': [],
            'prediction_intervals_outside_physical_bounds': [],
        }
    summary = compare(parent, records, subset)
    assert summary['records'] == 2
    assert all(f['objective_improved_gt_1e_10'] == 1 for f in summary['families'].values())
    with pytest.raises(ValueError, match='complete reviewed'):
        compare(parent, records[:1], subset)
    records[0]['fit']['observations'][0]['water_content'] = 0.31
    with pytest.raises(ValueError, match='canonical dataset'):
        compare(parent, records, subset)
