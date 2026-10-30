"""Protect explicit original-pressure recovery from unsupported coordinate snapping."""

import copy
import hashlib
import math
import runpy
from pathlib import Path

import openpyxl
import pytest

from ptfkit_eval.documents import canonical_json


METHOD = (
    'mesa de tens\u00e3o na tens\u00e3o de 6 kPa e panela de press\u00e3o nas tens\u00f5es '
    'de 10, 33, 100 e 1500 kPa.'
)
COMMENT = (
    'O valor da press\u00e3o foi dado originalmente em kPa. '
    'Foi utilizado o fator 10.20 para convers\u00e3o em cmh2o (1kPa=10.20 cmh2o). '
)


def _namespace() -> dict:
    root = Path(__file__).resolve().parents[3]
    return runpy.run_path(str(root / 'evaluation/experiments/hybras/restore_pressure_protocols.py'))


@pytest.mark.parametrize(('head', 'expected'), [(61.2, 6), (337, 33), (15300, 1500)])
def test_pressure_recovery_uses_declared_kpa_control_regime(head: float, expected: float):
    assert _namespace()['controlled_pressure'](METHOD, COMMENT, head) == expected


@pytest.mark.parametrize(
    ('method', 'comment', 'head'),
    [
        (METHOD, COMMENT, math.nextafter(337, math.inf)),
        (METHOD, COMMENT, 330),
        (METHOD, COMMENT, 336.6),
        ('pressure plate from 10 to 1500 kPa', COMMENT, 337),
        (METHOD, 'Originally in atm; 1atm=1020 cmh2o.', 337),
        (METHOD, 'Originally in cm.', 337),
        (None, COMMENT, 337),
    ],
)
def test_pressure_recovery_does_not_guess_nearby_or_unreviewed_coordinates(
    method: str | None, comment: str, head: float
):
    assert _namespace()['controlled_pressure'](method, comment, head) is None


def _fixture() -> tuple[dict, openpyxl.Workbook]:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'HYBRAS'
    sheet.append([])
    sheet.append([])
    headers = [None] * 125
    headers[80] = 337
    headers[94] = 15300
    sheet.append(headers)
    row = [None] * 125
    row[0], row[5], row[36], row[80], row[94], row[124] = 1, 8, 'study', 0.4, 0.2, METHOD
    sheet.append(row)
    raw = workbook.create_sheet('RAWRET_lines')
    raw.append(['code', 'WR_ID', 'flag', 'head', 'water', 'comment1', 'comment2', 'comment3'])
    raw.append([1, 1, 'l', 337, 0.4, None, COMMENT])
    raw.append([1, 1, 'l', 15300, 0.2, None, COMMENT])
    source = 'hybras-study-' + hashlib.sha256(canonical_json(['study', None])).hexdigest()[:16]
    document = {
        'samples': [
            {
                'source_id': source,
                'sample_id': 'hybras-horizon-1',
                'profile_id': 'hybras-profile-8',
                'features': [],
            }
        ],
        'observations': [
            {
                'id': f'hybras-raw-row-{number}',
                'source_id': source,
                'sample_id': 'hybras-horizon-1',
                'suction_kpa': head * 0.0980665,
                'value': water,
                'source_cells': [f'RAWRET_lines!A{number}:H{number}'],
            }
            for number, head, water in [(2, 337, 0.4), (3, 15300, 0.2)]
        ],
        'target': {'material_basis': 'unspecified'},
    }
    return document, workbook


def test_coordinate_overlay_preserves_targets_and_requires_wide_long_agreement():
    document, workbook = _fixture()
    before = copy.deepcopy(document)
    restore = _namespace()['restore']
    subset, review, summary = restore(document, workbook)
    assert document == before
    assert [p['suction_kpa'] for p in subset['observations']] == [33, 1500]
    for original, updated in zip(before['observations'], subset['observations'], strict=True):
        assert {**updated, 'suction_kpa': original['suction_kpa']} == original
    assert subset['samples'] == before['samples']
    assert subset['target'] == before['target']
    assert review['lineage'][0]['parent_suction_kpa'] == 33.0484105
    assert review['lineage'][0]['wide_water_source_cell'] == 'HYBRAS!CC4'
    assert summary['native_eligibility'] == 'not_granted'
    workbook['HYBRAS'].cell(row=4, column=81, value=0.41)
    subset, _, summary = restore(document, workbook)
    assert len(subset['observations']) == 1
    assert summary['exclusions'] == {'wide_long_value_or_coordinate_mismatch': 1}


def test_coordinate_overlay_rejects_tampered_canonical_water():
    document, workbook = _fixture()
    document['observations'][0]['value'] = 0.41
    with pytest.raises(ValueError, match='differs from the reviewed source row'):
        _namespace()['restore'](document, workbook)
