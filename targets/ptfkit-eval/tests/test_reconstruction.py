import hashlib
import json
from pathlib import Path
from typing import cast

import pytest
from openpyxl import Workbook, load_workbook

from ptfkit_eval import (
    DatasetMapping,
    RetentionReconstructionError,
    audit_retention_reconstruction_workbook,
    write_retention_reconstruction_audit,
)


def _workbook(path: Path, *, mixed_groups: bool = False) -> None:
    workbook = Workbook()
    master = workbook.active
    master.title = 'Данные'
    master['X2'] = 0.001
    master['Y2'] = 1
    master['B3'] = 1
    master['C3'] = 'master'
    master['X3'] = 0.1
    master['Y3'] = 0.2
    master['B4'] = 2
    master['C4'] = 'master'
    master['X4'] = 0.3 if mixed_groups else 0.1
    master['Y4'] = 0.4
    master['AK2'] = 'Qr'
    master['AL2'] = 'n'
    master['AK3'] = '0,01'
    master['AL3'] = 1.2
    master['AK4'] = '-'
    master['AL4'] = 1.3

    long = workbook.create_sheet('ОГХ')
    long.append(['sample_id', 'source', 'sample_number', 'water', 'pressure', 'axis', 'value'])
    if mixed_groups:
        long.append(['raw_1', 'long', 1, 'Q', 'pF', -1, 10])
        long.append(['raw_1', 'long', 1, 'Q', 'pF', 1, 40])
        long.append(['raw_2', 'long', 2, 'Q', 'pF', -1, 30])
        long.append(['raw_2', 'long', 2, 'Q', 'pF', 1, 20])
    else:
        long.append(['raw_1', 'long', 1, 'Q', 'pF', -1, 10])
        long.append(['raw_1', 'long', 1, 'Q', 'pF', 1, 40])
        long.append(['raw_2', 'long', 2, 'Q', 'pF', -1, 10])
        long.append(['raw_2', 'long', 2, 'Q', 'pF', 1, 20])
    long.append(['raw_1', 'long', 1, 'VGN', None, 'Qr', '-'])
    long.append(['raw_1', 'long', 1, 'VGN', None, 'n', 1.3])
    long.append(['raw_2', 'long', 2, 'VGN', None, 'Qr', 0.01])
    long.append(['raw_2', 'long', 2, 'VGN', None, 'n', 1.2])
    workbook.save(path)
    workbook.close()


def _mapping(
    path: Path, *, blocked: bool = False, expected_missing_water_count: int = 0
) -> DatasetMapping:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    crosswalks = []
    if not blocked:
        crosswalks = [
            {
                'worksheet': 'ОГХ',
                'raw_sample_id': 'raw_1',
                'canonical_sample_id': 'master_2',
                'canonical_source': 'master',
                'canonical_sample_number': 2,
                'status': 'confirmed',
            },
            {
                'worksheet': 'ОГХ',
                'raw_sample_id': 'raw_2',
                'canonical_sample_id': 'master_1',
                'canonical_source': 'master',
                'canonical_sample_number': 1,
                'status': 'confirmed',
            },
        ]
    return DatasetMapping(
        dataset_id='test',
        status='draft_blocked',
        source_filename=path.name,
        source_sha256=digest,
        document={
            'identity': {'sample_crosswalks': crosswalks},
            'retention_reconstruction_rules': [
                {
                    'id': 'test_source',
                    'status': 'blocked' if blocked else 'confirmed',
                    'master_worksheet': 'Данные',
                    'long_worksheet': 'ОГХ',
                    'master_source': 'master',
                    'long_source': 'long',
                    'water_scale': 100,
                    'water_scale_status': 'proposed' if blocked else 'confirmed',
                    'comparison_decimal_places': 12,
                    'expected_missing_water_count': expected_missing_water_count,
                    'expected_missing_long_water_count': 0,
                    'wide_missing_representation': 'omitted_from_long',
                    'pressure_replacements': [
                        {
                            'master_value': 0.001,
                            'long_value': -1,
                            'status': 'proposed' if blocked else 'confirmed',
                        }
                    ],
                    'blocks': [
                        {
                            'axis_row': 2,
                            'first_sample_row': 3,
                            'last_sample_row': 4,
                            'first_column': 'X',
                            'last_column': 'Y',
                        }
                    ],
                    'parameter_block': {
                        'header_row': 2,
                        'columns': ['AK', 'AL'],
                        'sample_rows': [{'first': 3, 'last': 4}],
                    },
                    'blockers': ['retention_reconstruction'] if blocked else [],
                }
            ],
        },
        inventory={},
        retention_profile={'sources': ({'source': 'long'},)},
    )


def _reconstruction_rule(mapping: DatasetMapping) -> dict[str, object]:
    rules = cast('list[dict[str, object]]', mapping.document['retention_reconstruction_rules'])
    return rules[0]


def test_reconstructs_complete_curves_and_requires_confirmed_crosswalks(tmp_path: Path):
    path = tmp_path / 'curves.xlsx'
    _workbook(path)

    result = audit_retention_reconstruction_workbook(path, _mapping(path))

    assert result.rule_count == 1
    assert result.profile_source_count == 1
    assert result.unaudited_source_count == 0
    assert result.unaudited_sources == ()
    assert result.reconstructable_rule_count == 1
    assert result.normalization_ready_rule_count == 1
    audit = result.rules[0]
    assert audit.master_curve_count == 2
    assert audit.master_point_count == 4
    assert audit.long_observation_count == 4
    assert audit.long_missing_water_count == 0
    assert audit.matched_point_count == 4
    assert audit.master_only_point_count == 0
    assert audit.long_only_point_count == 0
    assert audit.unique_row_match_count == 2
    assert audit.ambiguous_row_match_count == 2
    assert audit.unmatched_row_count == 0
    assert audit.master_parameter_count == 4
    assert audit.long_parameter_count == 4
    assert audit.matched_parameter_count == 4
    assert audit.parameter_multiset_status == 'exact'
    assert audit.parameter_identity_status == 'exact_unique'
    assert audit.parameter_reconstructable is True
    assert audit.identity_status == 'exact_unique'
    assert audit.observation_reconstructable is True
    assert audit.reconstructable is True
    assert audit.normalization_ready is True
    assert [
        (item.raw_sample_id, item.canonical_sample_number) for item in audit.inferred_crosswalks
    ] == [('raw_1', 2), ('raw_2', 1)]


def test_exact_point_multiset_does_not_accept_mixed_sample_groups(tmp_path: Path):
    path = tmp_path / 'mixed.xlsx'
    _workbook(path, mixed_groups=True)

    result = audit_retention_reconstruction_workbook(path, _mapping(path, blocked=True))

    audit = result.rules[0]
    assert audit.point_multiset_status == 'exact'
    assert audit.matched_point_count == 4
    assert audit.resolved_group_count == 0
    assert audit.unresolved_group_count == 2
    assert audit.identity_status == 'unresolved'
    assert audit.parameter_multiset_status == 'exact'
    assert audit.parameter_identity_status == 'exact_unique'
    assert audit.reconstructable is False
    assert audit.normalization_ready is False
    assert audit.inferred_crosswalks == ()


def test_rejects_workbook_bytes_not_bound_to_the_mapping(tmp_path: Path):
    path = tmp_path / 'curves.xlsx'
    _workbook(path)
    mapping = _mapping(path)
    path.write_bytes(path.read_bytes() + b'changed')

    with pytest.raises(RetentionReconstructionError, match='SHA-256'):
        audit_retention_reconstruction_workbook(path, mapping)


def test_rejects_unreviewed_wide_missing_value_drift(tmp_path: Path):
    path = tmp_path / 'curves.xlsx'
    _workbook(path)

    with pytest.raises(RetentionReconstructionError, match=r'expected 1 missing.*found 0'):
        audit_retention_reconstruction_workbook(
            path,
            _mapping(path, expected_missing_water_count=1),
        )


def test_reconstructs_categorical_axes_without_numeric_coercion(tmp_path: Path):
    path = tmp_path / 'categorical.xlsx'
    _workbook(path)
    workbook = load_workbook(path)
    master = workbook['Данные']
    long = workbook['ОГХ']
    master['X2'] = 'MG'
    master['Y2'] = 'PWP'
    for row in range(2, 6):
        long.cell(row, 6).value = 'MG' if row % 2 == 0 else 'PWP'
    workbook.save(path)
    workbook.close()

    result = audit_retention_reconstruction_workbook(path, _mapping(path))

    audit = result.rules[0]
    assert audit.point_multiset_status == 'exact'
    assert audit.identity_status == 'exact_unique'
    assert audit.reconstructable is True


def test_wide_block_can_override_the_source_default_water_scale(tmp_path: Path):
    path = tmp_path / 'block-scale.xlsx'
    _workbook(path)
    workbook = load_workbook(path)
    long = workbook['ОГХ']
    for row in range(2, 6):
        long.cell(row, 7).value /= 100
    workbook.save(path)
    workbook.close()
    mapping = _mapping(path)
    rule = _reconstruction_rule(mapping)
    blocks = cast('list[dict[str, object]]', rule['blocks'])
    blocks[0]['water_scale'] = 1

    audit = audit_retention_reconstruction_workbook(path, mapping).rules[0]

    assert audit.point_multiset_status == 'exact'
    assert audit.reconstructable is True


def test_wide_block_can_supply_a_reviewable_missing_axis_value(tmp_path: Path):
    path = tmp_path / 'axis-override.xlsx'
    _workbook(path)
    workbook = load_workbook(path)
    workbook['Данные']['X2'] = None
    workbook.save(path)
    workbook.close()
    mapping = _mapping(path)
    rule = _reconstruction_rule(mapping)
    blocks = cast('list[dict[str, object]]', rule['blocks'])
    blocks[0]['axis_overrides'] = {
        'X': {'long_value': -1, 'status': 'confirmed', 'evidence': 'Test fixture.'}
    }

    audit = audit_retention_reconstruction_workbook(path, mapping).rules[0]

    assert audit.point_multiset_status == 'exact'
    assert audit.reconstructable is True


def test_explicit_null_is_matched_by_sample_and_axis(tmp_path: Path):
    path = tmp_path / 'explicit-null.xlsx'
    _workbook(path)
    workbook = load_workbook(path)
    master = workbook['Данные']
    long = workbook['ОГХ']
    master['Y4'] = None
    long['G3'] = None
    workbook.save(path)
    workbook.close()
    mapping = _mapping(path, expected_missing_water_count=1)

    with pytest.raises(RetentionReconstructionError, match=r'expected 0 missing long.*found 1'):
        audit_retention_reconstruction_workbook(path, mapping)

    rule = _reconstruction_rule(mapping)
    rule['expected_missing_long_water_count'] = 1
    rule['wide_missing_representation'] = 'explicit_long_null'

    result = audit_retention_reconstruction_workbook(path, mapping)

    audit = result.rules[0]
    assert audit.master_point_count == 4
    assert audit.long_observation_count == 4
    assert audit.wide_missing_water_count == 1
    assert audit.long_missing_water_count == 1
    assert audit.point_multiset_status == 'exact'
    assert audit.reconstructable is True

    workbook = load_workbook(path)
    workbook['ОГХ']['F3'] = 2
    workbook.save(path)
    workbook.close()
    changed_mapping = _mapping(path, expected_missing_water_count=1)
    changed_rule = _reconstruction_rule(changed_mapping)
    changed_rule['expected_missing_long_water_count'] = 1
    changed_rule['wide_missing_representation'] = 'explicit_long_null'

    changed = audit_retention_reconstruction_workbook(path, changed_mapping).rules[0]
    assert changed.point_multiset_status == 'mismatch'
    assert changed.reconstructable is False


def test_declared_numeric_missing_marker_can_match_an_explicit_long_null(tmp_path: Path):
    path = tmp_path / 'numeric-missing-marker.xlsx'
    _workbook(path)
    workbook = load_workbook(path)
    workbook['Данные']['Y4'] = 0
    workbook['ОГХ']['G3'] = None
    workbook.save(path)
    workbook.close()
    mapping = _mapping(path, expected_missing_water_count=1)
    rule = _reconstruction_rule(mapping)
    rule['expected_missing_long_water_count'] = 1
    rule['wide_missing_representation'] = 'explicit_long_null'
    rule['wide_missing_markers'] = [0]

    audit = audit_retention_reconstruction_workbook(path, mapping).rules[0]

    assert audit.wide_missing_water_count == 1
    assert audit.long_missing_water_count == 1
    assert audit.point_multiset_status == 'exact'


def test_declared_populated_wide_exclusion_is_not_relabelled_as_missing(tmp_path: Path):
    path = tmp_path / 'wide-exclusion.xlsx'
    _workbook(path)
    workbook = load_workbook(path)
    workbook['ОГХ'].delete_rows(3)
    workbook.save(path)
    workbook.close()
    mapping = _mapping(path)
    rule = _reconstruction_rule(mapping)
    rule['expected_excluded_wide_water_count'] = 1
    rule['wide_excluded_cells'] = ['Y4']

    audit = audit_retention_reconstruction_workbook(path, mapping).rules[0]

    assert audit.excluded_wide_water_count == 1
    assert audit.wide_missing_water_count == 0
    assert audit.point_multiset_status == 'exact'
    assert audit.reconstructable is True


def test_parameter_only_source_is_audited_without_observation_blocks(tmp_path: Path):
    path = tmp_path / 'parameters-only.xlsx'
    _workbook(path)
    workbook = load_workbook(path)
    workbook['ОГХ'].delete_rows(2, 4)
    workbook.save(path)
    workbook.close()
    mapping = _mapping(path)
    rule = _reconstruction_rule(mapping)
    rule['blocks'] = []

    result = audit_retention_reconstruction_workbook(path, mapping)

    audit = result.rules[0]
    assert audit.master_curve_count == 0
    assert audit.long_observation_count == 0
    assert audit.point_multiset_status == 'not_applicable'
    assert audit.identity_status == 'not_applicable'
    assert audit.observation_reconstructable is None
    assert audit.matched_parameter_count == 4
    assert audit.parameter_identity_status == 'exact_unique'
    assert audit.parameter_reconstructable is True
    assert audit.reconstructable is True
    assert audit.normalization_ready is True


def test_writes_count_only_reconstruction_evidence(tmp_path: Path):
    path = tmp_path / 'curves.xlsx'
    destination = tmp_path / 'audit.json'
    _workbook(path)
    audit = audit_retention_reconstruction_workbook(path, _mapping(path))

    write_retention_reconstruction_audit(audit, destination)

    payload = json.loads(destination.read_text(encoding='utf-8'))
    assert payload['sha256'] == audit.sha256
    assert payload['rules'][0]['matched_point_count'] == 4
    assert payload['rules'][0]['inferred_crosswalks'][0] == {
        'raw_sample_id': 'raw_1',
        'canonical_sample_id': 'master_2',
        'canonical_source': 'master',
        'canonical_sample_number': 2,
    }
    assert 'points' not in payload['rules'][0]
