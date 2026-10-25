"""Audit literal cross-sheet identities without granting scientific eligibility."""

# ruff: noqa: INP001 - Standalone experiment helpers.

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date
from typing import TYPE_CHECKING

from openpyxl.utils import get_column_letter


if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook


def exact_crosswalk(master: list[dict], original: list[dict]) -> list[dict]:
    """Require unique literal IDs; numerical matches are diagnostics only."""
    indexed = defaultdict(list)
    for record in original:
        indexed[record['id']].append(record)
    counts = Counter(record['id'] for record in master)
    result = []
    for record in master:
        candidates = indexed[record['id']]
        unique = counts[record['id']] == 1 and len(candidates) == 1
        counterpart = candidates[0] if unique else None
        values = record['values']
        result.append(
            {
                'master_key': record['key'],
                'literal_id': record['id'],
                'status': 'exact_literal_id' if unique else 'unresolved_identity',
                'original_key': counterpart['key'] if counterpart else None,
                'ordered_values_equal': counterpart['values'] == values if counterpart else None,
                'value_multiset_equal': Counter(counterpart['values']) == Counter(values)
                if counterpart
                else None,
                'numeric_candidate_keys': [r['key'] for r in original if r['values'] == values],
                'measurement_eligibility': 'unresolved',
            }
        )
    return result


def priority_source_audit(book: Workbook, followup: dict) -> dict:  # noqa: C901
    """Retain row roles, crosswalks, raw predictor cells and vector reuse evidence."""
    master = book.worksheets[0]
    records = []
    for review in followup['additional_master_blocks']:
        if 'workbook_schema' not in review:
            continue
        source = review['raw_source']
        schema = review['workbook_schema']
        original = book.worksheets[schema['original_sheet_index']]
        start = schema['first_value_column']
        views: list[list[dict]] = []
        for sheet, first in ((master, 24), (original, start)):
            rows: list[dict] = []
            for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if values[2] != source:
                    continue
                last = schema.get(
                    'last_value_column' if sheet is original else 'master_last_value_column'
                )
                tail = list(values[first - 1 : last])
                while tail and tail[-1] is None:
                    tail.pop()
                label = values[first - 2]
                rows.append(
                    {
                        'key': f'{sheet.title}!{number}',
                        'id': values[0],
                        'reported_vector_label': label,
                        'role': 'reported_head_vector'
                        if label == 'm water h'
                        else 'coordinate_vector_unit_unresolved'
                        if sheet is original and number in schema.get('control_rows', [])
                        else 'reported_water_vector'
                        if label is not None and label == schema['water_vector_label']
                        else schema['noncontrol_role']
                        if sheet is original and 'noncontrol_role' in schema
                        else 'unresolved_axis_or_water_vector',
                        'values': tail,
                        'populated_value_count': sum(v is not None for v in tail),
                        'repeated_numeric_coordinate_count': sum(
                            n - 1
                            for v, n in Counter(tail).items()
                            if isinstance(v, (int, float)) and n > 1
                        ),
                        'reported_metadata_cells': {
                            f'{get_column_letter(column)}{number}': {
                                'excel_date_iso8601': v.isoformat()
                            }
                            if isinstance(v, date)
                            else v
                            for column, v in enumerate(values[: first - 2], 1)
                            if v is not None
                        },
                        'separate_predictor_cells': {
                            f'{get_column_letter(column)}{number}': v
                            for column, v in enumerate(values, 1)
                            if last is not None and column > last and v is not None
                        },
                    }
                )
            views.append(rows)
        crosswalk = exact_crosswalk(*views)
        duplicates = defaultdict(list)
        water_multisets = defaultdict(list)
        for row in views[0]:
            if row['values']:
                duplicates[tuple(row['values'])].append(row['key'])
        for row in views[1]:
            if row['role'] == 'reported_water_vector' and row['values']:
                water_multisets[frozenset(Counter(row['values']).items())].append(row['key'])
        records.append(
            {
                'raw_source': source,
                'claim_status': 'exactly_derived',
                'master_row_count': len(views[0]),
                'original_row_count': len(views[1]),
                'crosswalk': crosswalk,
                'views': views,
                'duplicate_master_vectors': [keys for keys in duplicates.values() if len(keys) > 1],
                'unique_nonempty_vector_count': len(duplicates),
                'nonempty_reported_water_vector_count': sum(
                    r['role'] == 'reported_water_vector' and bool(r['values']) for r in views[1]
                ),
                'reused_water_value_multisets': [
                    keys for keys in water_multisets.values() if len(keys) > 1
                ],
                'independent_profile_count': None,
                'released_observation_count': 0,
                'remaining_claims': review['required_claims'],
                'column_header_evidence': {
                    f'{get_column_letter(column)}{number}': value
                    for number, values in enumerate(
                        original.iter_rows(min_row=1, max_row=3, values_only=True), 1
                    )
                    for column, value in enumerate(values, 1)
                    if isinstance(value, str)
                },
                'review': review,
            }
        )
    return {
        'schema_version': 1,
        'policy': 'literal_ids_not_row_order_or_numerical_similarity',
        'sources': records,
    }
