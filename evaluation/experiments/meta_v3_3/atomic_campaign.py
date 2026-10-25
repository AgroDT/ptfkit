"""Review atomic coverage opportunities against exact released measurement targets."""

# ruff: noqa: INP001 - Standalone source-review campaign.

from __future__ import annotations

import argparse
import hashlib
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from ptfkit_eval.documents import canonical_json


_WATER = {'volumetric_water_content', 'gravimetric_water_content'}


def targets(function: dict, support: dict) -> list[dict]:  # noqa: C901, PLR0912 - Source-specific review decisions.
    """Resolve reviewed target groups, without interpreting arbitrary pressure names.

    Legacy point contracts below are discovery contracts from the reviewed YAML
    descriptions, not public API measurement migrations or execution permission.
    Field FC and available-moisture differences are explicitly not curve points.
    """
    fid, source = function['id'], function['source']
    result: list[dict] = [
        {
            'quantity': p['measurement']['quantity'],
            'pressure_kpa': p['suction_kpa'],
            'kind': 'fixed_pressure',
        }
        for p in support.get(fid, {}).get('fixed_outputs', [])
        if p['suction_kpa'] is not None
    ]
    if source == 'beniaich2023':
        result.append({'quantity': 'gravimetric_saturation', 'kind': 'operational'})
    legacy_pressures: list[int] = []
    quantity = 'volumetric_water_content'
    if source in ('rawls1982', 'rawls2003'):
        legacy_pressures = [int(o['name'].rsplit('_', 1)[1]) for o in function['outputs']]
    elif source == 'chakraborty2011':
        quantity, legacy_pressures = 'gravimetric_water_content', [33, 100, 500, 1500]
    elif source == 'myeni2021':
        quantity = 'gravimetric_water_content'
        legacy_pressures = [33 if fid.endswith('_fc') else 1500]
    elif source == 'dharumarajan2019' and not fid.endswith('_infiltration'):
        quantity, legacy_pressures = 'gravimetric_water_content', [33, 1500]
    elif source == 'tiwary2014' and fid.endswith('_bsr'):
        quantity, legacy_pressures = 'gravimetric_water_content', [33, 100, 1500]
    elif source == 'oosterveld1980' and fid.endswith('_wilting_point'):
        quantity, legacy_pressures = 'gravimetric_water_content', [1500]
    elif source == 'wang2012':
        legacy_pressures = [33]
    result.extend(
        {'quantity': quantity, 'pressure_kpa': p, 'kind': 'fixed_pressure'}
        for p in legacy_pressures
    )
    if source == 'pidgeon1972':
        if '_awc' in fid or '_eawc' in fid:
            result.append({'quantity': 'available_water_capacity', 'kind': 'operational'})
        else:
            result.append(
                {
                    'quantity': 'volumetric_field_capacity'
                    if '_fc_vol_' in fid
                    else 'gravimetric_permanent_wilting_point'
                    if '_pwp' in fid
                    else 'gravimetric_field_capacity',
                    'kind': 'operational',
                }
            )
    elif source == 'oosterveld1980':
        if fid.endswith('_retention'):
            result.append(
                {
                    'quantity': 'gravimetric_water_content',
                    'kind': 'curve',
                    'pressure_range_kpa': [10, 1500],
                }
            )
        elif not fid.endswith('_wilting_point'):
            result.append(
                {
                    'quantity': 'gravimetric_available_moisture'
                    if fid.endswith('_available_water')
                    else 'field_capacity_tension'
                    if fid.endswith('_field_capacity_tension')
                    else 'gravimetric_field_capacity',
                    'kind': 'operational',
                }
            )
    elif not legacy_pressures and source != 'pidgeon1972':
        curves = support.get(fid, {}).get('curves', [])
        quantities = {o['quantity'] for o in function['outputs']}
        if curves or (
            quantities & _WATER
            and any(
                o['quantity'] in {'van_genuchten_alpha', 'hutson_cass_b', 'varallyai_m'}
                for o in function['outputs']
            )
        ):
            result.append({'quantity': 'volumetric_water_content', 'kind': 'curve'})
        elif quantities & _WATER and not result:
            result.append(
                {'quantity': 'retention_output_contract_unresolved', 'kind': 'operational'}
            )
    result.extend(
        {'quantity': q, 'kind': 'operational'}
        for q in {o['quantity'] for o in function['outputs']}
        & {
            'saturated_hydraulic_conductivity',
            'infiltration_rate',
            'dry_bulk_density',
            'available_water_capacity',
            'unsaturated_hydraulic_conductivity',
            'osmotic_potential',
        }
    )
    if not result:
        result.append({'quantity': 'other_output_contract_unresolved', 'kind': 'operational'})
    return list({json.dumps(r, sort_keys=True): r for r in result}.values())


def pool_id(target: dict) -> str:
    """Keep water denominator, operational target and exact pressure distinct."""
    return f'{target["kind"]}:{target["quantity"]}' + (
        f':{target["pressure_kpa"]:g}kPa' if target['kind'] == 'fixed_pressure' else ''
    )


def matching_observations(target: dict, observations: list[dict], quantity: str) -> list[dict]:
    """Filter measured coordinates exactly; never round, interpolate or convert water basis."""
    if target['quantity'] != quantity or target['kind'] == 'operational':
        return []
    if target['kind'] == 'fixed_pressure':
        return [o for o in observations if o['suction_kpa'] == target['pressure_kpa']]
    lower, upper = target.get('pressure_range_kpa', [0, float('inf')])
    return [o for o in observations if lower <= o['suction_kpa'] <= upper]


def assess(
    coverage: dict, support: dict, datasets: list[dict], function_audit: dict, raw_rows: dict
) -> dict:
    """Recompute target readiness separately from predictor eligibility for every function."""
    prior = {f['function']: f for f in function_audit['functions']}
    pools: dict[str, dict] = defaultdict(
        lambda: {
            'functions': set(),
            'sources': set(),
            'scored': set(),
            'samples': set(),
            'unreviewed_raw_sources': set(),
        }
    )
    records: list[dict] = []
    for function in coverage['functions']:
        fid = function['id']
        bindings = prior[fid]['predictor_definitions']
        source_records = []
        declared = targets(function, support)
        for target in declared:
            pool = pools[pool_id(target)]
            pool['functions'].add(fid)
            if target['quantity'] == 'dry_bulk_density':
                pool['unreviewed_raw_sources'].update(
                    source for (source, _), row in raw_rows.items() if row.get('Q') is not None
                )
            for dataset in datasets:
                points = matching_observations(
                    target, dataset['observations'], dataset['target']['quantity']
                )
                sources = sorted({o['source_id'] for o in points})
                pool['sources'].update(sources)
                pool['samples'].update((o['source_id'], o['sample_id']) for o in points)
                for source in sources:
                    samples = {o['sample_id'] for o in points if o['source_id'] == source}
                    complete = sum(
                        all(
                            (i['name'] == 'tension' and function['source'] == 'oosterveld1980')
                            or (
                                i['candidate_column'] is not None
                                and row.get(i['candidate_column']) is not None
                            )
                            for i in bindings
                        )
                        for (raw_source, sample), row in raw_rows.items()
                        if raw_source == source and sample in samples
                    )
                    source_records.append(
                        {
                            'source_id': source,
                            'target_pool': pool_id(target),
                            'measured_target_samples': len(samples),
                            'complete_candidate_predictor_rows': complete,
                            'predictor_status': 'unresolved_measurement_contract'
                            if complete
                            else 'missing_predictor_values',
                            'target_status': (
                                'measured_coordinates_available_protocol_review_required'
                            ),
                        }
                    )
            if prior[fid]['real_data_scored_samples'] and target['kind'] == 'curve':
                pool['scored'].add(fid)
        scored = bool(prior[fid]['real_data_scored_samples'])
        records.append(
            {
                'function': fid,
                'source': function['source'],
                'targets': declared,
                'predictor_status': 'confirmed_for_scored_contract'
                if scored
                else 'reviewed_unresolved'
                if any(r['complete_candidate_predictor_rows'] for r in source_records)
                else 'missing_predictors_for_available_targets'
                if source_records
                else 'not_prioritized_until_target_available',
                'target_status': 'scored'
                if scored
                else 'measured_coordinates_available_protocol_review_required'
                if source_records
                else 'no_compatible_target_in_current_release',
                'real_data_scored': scored,
                'source_opportunities': source_records,
                'execution_authorized': scored,
            }
        )
    return {
        'functions': records,
        'target_pools': [
            {
                'target_pool': name,
                'catalog_ptfs': len(p['functions']),
                'catalog_functions': sorted(p['functions']),
                'potential_measured_sources': len(p['sources'] | p['unreviewed_raw_sources']),
                'unreviewed_raw_sources': sorted(p['unreviewed_raw_sources']),
                'reviewed_coordinate_sources': sorted(p['sources']),
                'reviewed_sources': len(p['sources']),
                'potential_measured_samples': len(p['samples']),
                'scored_ptfs': len(p['scored']),
                'main_target_blocker': 'predictor_and_source_protocol_review'
                if p['sources']
                else 'no_exact_or_operational_measured_target_in_current_release',
            }
            for name, p in sorted(pools.items())
        ],
        'counts': {
            'implemented': len(records),
            'scored_atomic': sum(r['real_data_scored'] for r in records),
            'new_scored_atomic': 0,
            'functions_with_measured_target_coordinates': sum(
                bool(r['source_opportunities']) for r in records
            ),
        },
        'scope': 'current_released_canonical_data_not_all_unreviewed_workbook_cells',
    }


def render(report: dict) -> str:
    """Expose a target-first work queue and the limits of the campaign evidence."""
    lines = [
        '# Atomic coverage source campaign',
        '',
        (
            'No new numeric predictor binding is authorized. Atomic coverage remains '
            '2/146. Ensemble algorithms and statistical evidence gates are unchanged.'
        ),
        '',
        (
            'Coordinate availability establishes water quantity and exact pressure, not complete '
            'PTF observation-method/material compatibility. '
            'Operational FC/AWC targets are separate.'
        ),
        '',
        '## Target availability',
        '',
        (
            '| Target pool | Catalog PTFs | Potential measured sources | '
            'Reviewed coordinate sources | Scored PTFs | Main target blocker |'
        ),
        '| --- | ---: | ---: | ---: | ---: | --- |',
    ]
    lines.extend(
        f'| {p["target_pool"]} | {p["catalog_ptfs"]} | '
        f'{p["potential_measured_sources"]} | {p["reviewed_sources"]} | '
        f'{p["scored_ptfs"]} | {p["main_target_blocker"]} |'
        for p in report['target_pools']
    )
    lines += [
        '',
        '## Function decisions',
        '',
        '| Function | Predictor status | Target status | Scored |',
        '| --- | --- | --- | --- |',
    ]
    lines += [
        f'| {r["function"]} | {r["predictor_status"]} | {r["target_status"]} | '
        f'{r["real_data_scored"]} |'
        for r in report['functions']
    ]
    lines += ['', '## Scientific group reviews', '']
    for review in report['review']['family_reviews']:
        lines += [
            f'### {review["id"]}',
            '',
            review['review'],
            '',
            'Evidence: ' + '; '.join(review['evidence']) + '.',
            '',
        ]
    lines += ['## Source findings', '']
    for name, source in report['review']['sources'].items():
        lines += [f'### {name}', '', f'Reference: {source["reference"]}', '']
        lines += [
            f'- **{key}**: {value}'
            for key, value in source.items()
            if key not in {'reference', 'local_evidence', 'reviewed_locations'}
        ]
        lines += ['']
    lines += [
        '## Next scientific work',
        '',
        (
            'First obtain numeric PSD, carbon/organic matter and density measurements matched '
            'to the 18 Badenko horizons: their measured theta33 targets already exist. '
            'This could support four Gunarathna theta33 functions, two Rawls2003 theta33 '
            'functions and Saxton2006, subject to their separate measurement protocols. '
            'These seven are a conditional function roster, not seven approved evaluations.'
        ),
        '',
        (
            'For a continuous WRC expansion, independently establish mass-basis PSD and '
            'organic-matter definitions for the Panina horizons. The current laser fractions '
            'and carbon-times-1.724 values are insufficient. No source-only gain is guaranteed.'
        ),
        '',
        (
            'The high-yield groups above have a documented blocker on the current release. '
            'This does not prove fewer than ten functions could ever be scored from all '
            'unclassified workbook blocks. Full-source identity and method evidence remains '
            'unresolved, and scientific review is not replaced by the old discovery ceilings.'
        ),
    ]
    return '\n'.join(lines) + '\n'


def run(benchmark: Path, workbook: Path, output: Path) -> dict:
    """Use immutable benchmark inputs and retain hashes of all reviewed evidence."""
    inventory = json.loads((benchmark / 'artifact-hashes.json').read_bytes())
    for name, digest in inventory.items():
        if hashlib.sha256((benchmark / name).read_bytes()).hexdigest() != digest:
            message = f'Benchmark artifact checksum mismatch: {name}'
            raise ValueError(message)
    summary = json.loads((benchmark / 'summary.json').read_bytes())
    if hashlib.sha256(workbook.read_bytes()).hexdigest() != summary['workbook_sha256']:
        message = 'Workbook checksum differs from the reviewed benchmark.'
        raise ValueError(message)
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        raw_rows = {
            (r[2], r[0]): {chr(65 + i): v for i, v in enumerate(r)}
            for r in book.worksheets[0].iter_rows(min_row=2, max_col=23, values_only=True)
            if isinstance(r[0], str)
        }
    finally:
        book.close()
    report = assess(
        json.loads((benchmark / 'implemented-functions.json').read_bytes()),
        json.loads((benchmark / 'catalog-support.json').read_bytes()),
        [
            json.loads((benchmark / f'canonical-{kind}.json').read_bytes())
            for kind in ('volume', 'mass')
        ],
        json.loads((benchmark / 'function-audit.json').read_bytes()),
        raw_rows,
    )
    root = Path(__file__).resolve().parents[3]
    review = Path(__file__).with_name('atomic-campaign-review.json')
    report['review'] = json.loads(review.read_bytes())
    report['input_sha256'] = {
        'workbook': summary['workbook_sha256'],
        'review': hashlib.sha256(review.read_bytes()).hexdigest(),
        'benchmark_inventory': hashlib.sha256(
            (benchmark / 'artifact-hashes.json').read_bytes()
        ).hexdigest(),
        'implementation': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
    }
    report['evidence_sha256'] = {
        s['local_evidence']: hashlib.sha256((root / s['local_evidence']).read_bytes()).hexdigest()
        for s in report['review']['sources'].values()
        if 'local_evidence' in s
    }
    for source in report['review']['sources'].values():
        if (
            'local_evidence' in source
            and report['evidence_sha256'][source['local_evidence']] != source['sha256']
        ):
            message = 'Scientific review evidence checksum mismatch.'
            raise ValueError(message)
    report['baseline'] = {
        'scored_atomic': summary['function_audit']['counts']['real_data_scored'],
        'evaluated_combinations': summary['blocker_impact']['coverage']['currently_evaluated'],
        'reviewed_scoring_sources': summary['class_comparison']['source_count'],
        'reviewed_scoring_profiles': summary['class_comparison']['profile_count'],
    }
    report['scoring_result'] = {
        **report['baseline'],
        'new_functions': [],
        'new_target_pools': [],
        'new_sources': [],
        'new_profiles': [],
        'new_maximal_common_cohorts': [],
        'ensemble_diagnostic': 'unchanged_completed_clapp_cosby_experiment',
    }
    output.mkdir(parents=True, exist_ok=False)
    refreshed = json.loads((benchmark / 'function-audit.json').read_bytes())
    decisions = {r['function']: r for r in report['functions']}
    for row in refreshed['functions']:
        row.update(
            {
                key: decisions[row['function']][key]
                for key in ('predictor_status', 'target_status', 'targets', 'source_opportunities')
            }
        )
    refreshed['previous_source_impacts_are_discovery_only'] = True
    refreshed['source_impact_review'] = (
        'atomic-campaign-review.json withdraws Chestnova/Panina actionable ceilings; '
        'no additional authorized atomic contract'
    )
    (output / 'function-audit.json').write_bytes(canonical_json(refreshed))
    (output / 'summary.json').write_bytes(canonical_json(report))
    (output / 'report.md').write_text(render(report), encoding='utf-8', newline='\n')
    (output / 'artifact-hashes.json').write_bytes(
        canonical_json(
            {
                p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                for p in output.iterdir()
                if p.is_file()
            }
        )
    )
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--benchmark', type=Path, required=True)
    parser.add_argument('--workbook', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    run(args.benchmark, args.workbook, args.output)
