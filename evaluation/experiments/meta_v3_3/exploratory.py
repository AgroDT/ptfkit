"""Run preregistered approximate META scenarios separately from the main benchmark."""

# ruff: noqa: INP001, SLF001 - Share helpers within the experiment without changing the main API.

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import cast

import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from ptfkit_eval import evaluate_candidate
from ptfkit_eval.artifacts import artifact_value
from ptfkit_eval.documents import document_hash, plain

from evaluation.experiments.catalog_coverage.audit import audit
from evaluation.experiments.meta_v3_3 import harmonized as main_benchmark
from evaluation.experiments.meta_v3_3.atomic_campaign import pool_id
from evaluation.experiments.meta_v3_3.exploratory_inputs import (
    approximate_inputs,
    approximate_target,
)
from ptfkit import get_manifest


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify_artifacts(directory: Path) -> dict:
    """Reject incomplete, modified or path-escaping main benchmark artifacts."""
    hashes = json.loads((directory / 'artifact-hashes.json').read_bytes())
    for name, expected in hashes.items():
        path = (directory / name).resolve()
        if not path.is_relative_to(directory.resolve()) or _hash(path) != expected:
            message = f'Benchmark artifact integrity failure: {name}'
            raise ValueError(message)
    return {**hashes, 'artifact-hashes.json': _hash(directory / 'artifact-hashes.json')}


def _clay_upper(row: int, mapping: dict) -> float:
    upper = 1.0
    for binding in mapping['predictor_bindings']:
        if binding['input'] != 'clay' or binding.get('column') != 'N':
            continue
        for override in binding.get('measurement_overrides', []):
            if any(span['first'] <= row <= span['last'] for span in override['sample_rows']):
                upper = override['measurement']['particle_size']['upper']
    return float(upper)


def raw_rows(workbook: Path, predictor_audit: dict, mapping: dict) -> dict:
    """Join through the existing audited master keys; never guess sample/source aliases."""
    indexed: dict[tuple[str, str], set[str]] = {}
    for item in predictor_audit['eligibility']:
        indexed.setdefault((item['source'], item['sample_id']), set()).add(item['sample_key'])
    sheet_name = next(s['name'] for s in mapping['worksheets'] if s['id'] == 'samples_master')
    with workbook.open('rb') as stream:
        book = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
        try:
            rows = {
                number: tuple(values)
                for number, values in enumerate(
                    book[sheet_name].iter_rows(min_row=2, max_col=22, values_only=True), 2
                )
            }
        finally:
            book.close()
    result = {}
    for identity, keys in sorted(indexed.items()):
        if len(keys) != 1:
            continue
        key = next(iter(keys))
        sheet, _, row_number = key.rpartition('!')
        if sheet != sheet_name or not row_number.isdigit() or int(row_number) not in rows:
            continue
        number = int(row_number)
        values = rows[number]
        if (values[2], values[0]) != identity:
            message = 'Raw workbook identity differs from the audited master identity.'
            raise ValueError(message)
        result[identity] = {
            'worksheet': sheet_name,
            'row': number,
            'clay_upper_um': _clay_upper(number, mapping),
            'values': {get_column_letter(i): value for i, value in enumerate(values, 1)},
        }
    return result


def evaluate_scenario(  # noqa: C901, PLR0912, PLR0915 - Independent sample/target decisions.
    coverage: dict,
    support: dict,
    datasets: list[dict],
    rows: dict,
    manifest: dict,
    *,
    scenario: str,
    registration: dict,
) -> tuple[list[dict], list[dict]]:
    """Execute only supported native contracts under explicit approximate data assumptions."""
    public = {f['id']: f for f in manifest['functions']}
    matrix, predictions = [], []
    for function in coverage['functions']:
        fid = function['id']
        for target in main_benchmark.catalog_targets(function, support):
            candidate, candidate_reasons = main_benchmark._candidate(
                function, target, support, manifest
            )
            for dataset in datasets:
                for sample in dataset['samples']:
                    raw = rows.get((sample['source_id'], sample['sample_id']))
                    values, evidence, reasons = {}, {}, list(candidate_reasons)
                    if fid not in public:
                        reasons.append('missing_public_measurement_contract')
                    else:
                        values, evidence, _ = main_benchmark._inputs(public[fid], sample, manifest)
                        if raw is not None:
                            values, evidence, blockers = approximate_inputs(
                                public[fid],
                                raw,
                                scenario,
                                registration,
                                manifest,
                                reviewed_values=values,
                                reviewed_evidence=evidence,
                            )
                            reasons.extend(blockers)
                        elif len(values) != len(public[fid]['inputs']):
                            reasons.append('no_exact_raw_predictor_identity')
                    # Reconstruct only measured WRC coordinates in their original water basis.
                    points, target_reasons = main_benchmark._target_points(
                        dataset, sample, {**target, 'quantity': dataset['target']['quantity']}
                    )
                    reasons.extend(target_reasons)
                    if target['quantity'] not in (
                        'volumetric_water_content',
                        'gravimetric_water_content',
                    ):
                        reasons.append('unsupported_target_quantity')
                    if not points:
                        reasons.append('no_supported_target_observations')
                    record: dict = {
                        'function': fid,
                        'scenario': scenario,
                        'source_id': sample['source_id'],
                        'sample_id': sample['sample_id'],
                        'profile_id': sample.get('profile_id'),
                        'dataset_id': dataset['dataset_id'],
                        'target_pool': pool_id(target),
                        'view': 'unsupported' if reasons else 'exploratory_candidate',
                        'reasons': sorted(set(reasons)),
                        'predictors': evidence,
                        'executed_points': 0,
                        'failed_points': 0,
                        'target_evidence': [],
                    }
                    matrix.append(record)
                    if reasons or candidate is None:
                        continue
                    required = cast('dict', artifact_value(candidate.target))
                    if dataset['target']['quantity'] == required['quantity'] and any(
                        dataset['target'][key] != required[key]
                        for key in ('unit', 'denominator_basis')
                    ):
                        record['view'] = 'unsupported'
                        record['reasons'].append('unsupported_water_representation')
                        continue
                    for pressure, observation in points:
                        try:
                            converted = approximate_target(
                                observation, dataset['target'], required, raw, registration
                            )
                        except ValueError:
                            record['view'] = 'unsupported'
                            record['reasons'].append('invalid_density_converted_target')
                            break
                        if converted is None:
                            record['view'] = 'unsupported'
                            record['reasons'].append('missing_sample_density_for_target_conversion')
                            break
                        target_evidence = converted.document()
                        record['target_evidence'].append(target_evidence)
                        tiers = {v['tier'] for v in [target_evidence, *evidence.values()]}
                        if 'approximate' not in tiers:
                            record['view'] = 'main_reference'
                            continue
                        record['view'] = 'exploratory'
                        try:
                            result = evaluate_candidate(
                                candidate, [pressure], values, manifest=manifest
                            )
                        except (ValueError, ArithmeticError) as error:
                            record['failed_points'] += 1
                            record.setdefault('execution_errors', []).append(
                                type(error).__name__ + ': ' + str(error)
                            )
                            continue
                        record['executed_points'] += 1
                        predictions.append(
                            {
                                **{
                                    key: record[key]
                                    for key in (
                                        'function',
                                        'scenario',
                                        'source_id',
                                        'sample_id',
                                        'profile_id',
                                        'dataset_id',
                                        'target_pool',
                                    )
                                },
                                'view': 'exploratory',
                                'suction_kpa': pressure,
                                'observed': converted.value,
                                'predicted': result.points[0].value,
                                'water_quantity': required['quantity'],
                                'target_contract': required,
                                'target_evidence': target_evidence,
                                'predictor_evidence': evidence,
                                'recommendation_eligible': False,
                            }
                        )
    return matrix, predictions


def summarize(
    coverage: dict, matrix: list[dict], predictions: list[dict], registration: dict
) -> dict:
    """Report scenarios without producing a winner or changing the main leaderboard."""
    functions = []
    for function in coverage['functions']:
        fid = function['id']
        for scenario in registration['scenarios']:
            points = [p for p in predictions if p['function'] == fid and p['scenario'] == scenario]
            rows = [r for r in matrix if r['function'] == fid and r['scenario'] == scenario]
            functions.append(
                {
                    'function': fid,
                    'scenario': scenario,
                    'scored_samples': len({(p['source_id'], p['sample_id']) for p in points}),
                    'scored_points': len(points),
                    'failed_points': sum(r['failed_points'] for r in rows),
                    'metrics_by_pool': {
                        pool: main_benchmark._metrics(
                            [p for p in points if p['target_pool'] == pool]
                        )
                        for pool in sorted({p['target_pool'] for p in points})
                    },
                    'blockers': dict(
                        sorted(Counter(reason for r in rows for reason in r['reasons']).items())
                    ),
                }
            )
    scenarios = {
        scenario: len({p['function'] for p in predictions if p['scenario'] == scenario})
        for scenario in registration['scenarios']
    }
    pools = []
    for scenario, pool in sorted({(p['scenario'], p['target_pool']) for p in predictions}):
        points = [p for p in predictions if p['scenario'] == scenario and p['target_pool'] == pool]
        pools.append(
            {
                'scenario': scenario,
                'target_pool': pool,
                'scored_functions': len({p['function'] for p in points}),
                'samples': len({(p['source_id'], p['sample_id']) for p in points}),
                'profiles': len(
                    {(p['source_id'], p['profile_id']) for p in points if p['profile_id']}
                ),
            }
        )
    return {
        'schema_version': 1,
        'protocol': registration['protocol'],
        'view': 'exploratory',
        'implemented_functions': len(coverage['functions']),
        'scored_functions_by_scenario': scenarios,
        'matrix_counts': dict(sorted(Counter(r['view'] for r in matrix).items())),
        'functions': functions,
        'target_pools': pools,
        'paired_sensitivity': paired_sensitivity(predictions, registration),
        'recommendation': None,
        'main_leaderboard_eligible': False,
        'limitations': [
            (
                'Approximate predictors and target basis assumptions '
                'do not establish validation eligibility.'
            ),
            (
                'Legacy functions without public measurement contracts '
                'or registered target adapters remain unsupported.'
            ),
            (
                'Observed profiles are not inferred from sample rows; '
                'no learned ensemble or final recommendation is produced.'
            ),
        ],
    }


def paired_sensitivity(predictions: list[dict], registration: dict) -> list[dict]:
    """Compare the same model and exact common coordinates between preregistered scenarios."""
    left, right = registration['scenarios']
    records = []
    for fid, pool in sorted({(p['function'], p['target_pool']) for p in predictions}):
        by_scenario = {}
        for scenario in (left, right):
            by_scenario[scenario] = {
                (p['source_id'], p['sample_id'], p['suction_kpa']): p
                for p in predictions
                if p['function'] == fid and p['target_pool'] == pool and p['scenario'] == scenario
            }
        keys = sorted(by_scenario[left].keys() & by_scenario[right].keys())
        if not keys:
            continue
        metrics = {
            scenario: cast(
                'dict', main_benchmark._metrics([by_scenario[scenario][key] for key in keys])
            )
            for scenario in (left, right)
        }
        records.append(
            {
                'function': fid,
                'target_pool': pool,
                'common_points': len(keys),
                'metrics': metrics,
                'difference_right_minus_left': {
                    key: metrics[right][key] - metrics[left][key] for key in metrics[left]
                },
                'scenarios': [left, right],
                'recommendation': None,
            }
        )
    return records


def render(report: dict) -> str:
    lines = [
        '# Exploratory sensitivity benchmark',
        '',
        (
            'Approximate assumptions only. '
            'Excluded from the main leaderboard and final recommendations.'
        ),
        '',
        (
            f'Catalog: {report["implemented_functions"]} functions. '
            f'Scored by scenario: {report["scored_functions_by_scenario"]}.'
        ),
        '',
        '| PTF | Scenario | Samples | Points | Failed points |',
        '|---|---|---:|---:|---:|',
    ]
    lines.extend(
        '| '
        + ' | '.join(
            str(row[key])
            for key in ('function', 'scenario', 'scored_samples', 'scored_points', 'failed_points')
        )
        + ' |'
        for row in report['functions']
    )
    lines.extend(
        [
            '',
            'Errors by target pool and paired common-coordinate sensitivity are in summary.json.',
            '',
            *report['limitations'],
            '',
        ]
    )
    return '\n'.join(lines)


def run(  # noqa: PLR0915 - One hash-verified output transaction.
    reference: Path, main: Path, workbook: Path, output: Path
) -> dict:
    """Snapshot assumptions before scoring and verify that both main input runs remain unchanged."""
    root = Path(__file__).resolve().parents[3]
    if any(output.resolve().is_relative_to(path.resolve()) for path in (reference, main)):
        message = 'Exploratory outputs must be outside the strict and harmonized directories.'
        raise ValueError(message)
    before = {'strict': verify_artifacts(reference), 'harmonized': verify_artifacts(main)}
    main_summary = json.loads((main / 'summary.json').read_bytes())
    if any(
        _hash(reference / name) != expected
        for name, expected in main_summary['input_sha256'].items()
    ):
        message = 'The harmonized run does not match the supplied strict benchmark.'
        raise ValueError(message)
    summary = json.loads((reference / 'summary.json').read_bytes())
    if _hash(workbook) != summary['workbook_sha256']:
        message = 'Exploratory predictors must come from the exact strict benchmark workbook.'
        raise ValueError(message)
    manifest = cast('dict', plain(get_manifest()))
    generated = json.loads((root / 'targets/ptfkit-py/src/ptfkit/_manifest.json').read_bytes())
    if document_hash(manifest) != document_hash(generated):
        message = 'Installed and generated PTF manifests must agree.'
        raise ValueError(message)
    coverage = audit(root)
    support = json.loads((reference / 'catalog-support.json').read_bytes())
    datasets = [
        json.loads((reference / f'canonical-{name}.json').read_bytes())
        for name in ('volume', 'mass')
    ]
    reviewed = json.loads(
        (reference / 'native-class-comparison/artifacts/dataset.json').read_bytes()
    )
    reviewed_ids = {(s['source_id'], s['sample_id']) for s in reviewed['samples']}
    for dataset in datasets:
        dataset['samples'] = [
            s for s in dataset['samples'] if (s['source_id'], s['sample_id']) not in reviewed_ids
        ]
        dataset['observations'] = [
            o
            for o in dataset['observations']
            if (o['source_id'], o['sample_id']) not in reviewed_ids
        ]
    datasets.append(reviewed)
    mapping_path = root / 'evaluation/datasets/meta_v3_3.yaml'
    mapping = yaml.safe_load(mapping_path.read_bytes())
    predictors = json.loads((reference / 'predictor-audit.json').read_bytes())
    rows = raw_rows(workbook, predictors, mapping)
    registration_path = Path(__file__).with_name('exploratory-policies.json')
    registration = json.loads(registration_path.read_bytes())
    output.mkdir(parents=True, exist_ok=False)
    main_benchmark._write(output, 'policies.json', registration)
    main_benchmark._write(output, 'core-policies.json', main_benchmark.policy_registry())
    matrix, predictions = [], []
    for scenario in registration['scenarios']:
        scenario_matrix, scenario_predictions = evaluate_scenario(
            coverage,
            support,
            datasets,
            rows,
            manifest,
            scenario=scenario,
            registration=registration,
        )
        matrix.extend(scenario_matrix)
        predictions.extend(scenario_predictions)
    report = summarize(coverage, matrix, predictions, registration)
    report['input_sha256'] = {
        **before,
        'workbook': _hash(workbook),
        'mapping': _hash(mapping_path),
        'policy_registration': _hash(registration_path),
        'manifest': document_hash(manifest),
    }
    report['implementation_sha256'] = {
        p.relative_to(root).as_posix(): _hash(p)
        for p in [
            Path(__file__),
            Path(__file__).with_name('exploratory_inputs.py'),
            Path(__file__).with_name('harmonized.py'),
        ]
    }
    if before != {'strict': verify_artifacts(reference), 'harmonized': verify_artifacts(main)}:
        message = 'Main benchmark artifacts changed during exploratory evaluation.'
        raise ValueError(message)
    report['main_artifacts_unchanged'] = True
    main_benchmark._write(output, 'exploratory-matrix.json', matrix)
    main_benchmark._write(output, 'exploratory-predictions.json', predictions)
    main_benchmark._write(output, 'summary.json', report)
    (output / 'report.md').write_text(render(report), encoding='utf-8')
    main_benchmark._write(
        output, 'artifact-hashes.json', {p.name: _hash(p) for p in sorted(output.iterdir())}
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--benchmark', type=Path, required=True)
    parser.add_argument('--harmonized', type=Path, required=True)
    parser.add_argument('--workbook', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    run(args.benchmark, args.harmonized, args.workbook, args.output)


if __name__ == '__main__':
    main()
