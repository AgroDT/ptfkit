"""Export predictor-defined testing subsets without selecting on target values."""

# ruff: noqa: INP001, E501, PLR0915, PLR2004 - Standalone workbook audit and report tables.

from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import load_workbook

from evaluation.experiments.meta_v3_3 import campaign_inputs
from evaluation.experiments.meta_v3_3 import expanded_participation as participation
from evaluation.experiments.meta_v3_3 import expanded_predictors as predictors
from evaluation.experiments.meta_v3_3 import workbook_semantic as semantic
from evaluation.experiments.meta_v3_3.pointwise_wrc import identity, sample_key


if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook


ROOT = Path(__file__).resolve().parents[3]
PUBLIC = Path(__file__).with_name('final_results')
OUTPUT = ROOT / 'local/meta-predictor-subsets-v1_0'
BASES = ('volumetric_water_content', 'gravimetric_water_content', 'ksat')
POOLS = {
    'texture_class': ('soil_texture',),
    'texture_class_bd': ('soil_texture', 'bulk_density'),
    'clay': ('clay',),
    'sand': ('sand',),
    'bulk_density': ('bulk_density',),
    'organic_carbon': ('organic_carbon',),
    'clay_bd': ('clay', 'bulk_density'),
    'clay_oc': ('clay', 'organic_carbon'),
    'clay_bd_oc': ('clay', 'bulk_density', 'organic_carbon'),
    'sand_clay': ('sand', 'clay'),
    'sand_clay_bd': ('sand', 'clay', 'bulk_density'),
    'sand_clay_oc': ('sand', 'clay', 'organic_carbon'),
    'sand_clay_bd_oc': ('sand', 'clay', 'bulk_density', 'organic_carbon'),
    'texture': ('sand', 'silt', 'clay'),
    'texture_bd': ('sand', 'silt', 'clay', 'bulk_density'),
    'texture_oc': ('sand', 'silt', 'clay', 'organic_carbon'),
    'texture_om': ('sand', 'silt', 'clay', 'organic_matter'),
    'texture_bd_oc': ('sand', 'silt', 'clay', 'bulk_density', 'organic_carbon'),
    'texture_bd_om': ('sand', 'silt', 'clay', 'bulk_density', 'organic_matter'),
    'texture_bd_oc_ph': ('sand', 'silt', 'clay', 'bulk_density', 'organic_carbon', 'ph'),
    'texture_bd_oc_pd': (
        'sand',
        'silt',
        'clay',
        'bulk_density',
        'organic_carbon',
        'particle_density',
    ),
    'texture_bd_oc_porosity': (
        'sand',
        'silt',
        'clay',
        'bulk_density',
        'organic_carbon',
        'porosity',
    ),
}
PROBE = {
    'id': 'predictor_availability_probe',
    'inputs': [
        {'name': name, 'unit': unit, 'domain': None}
        for name, unit in (
            ('sand', '%'),
            ('silt', '%'),
            ('clay', '%'),
            ('bulk_density', 'g/cm^3'),
            ('organic_carbon', '%'),
            ('organic_matter', '%'),
            ('particle_density', 'g/cm^3'),
            ('porosity', 'm^3/m^3'),
            ('ph', 'dimensionless'),
            ('depth', 'cm'),
        )
    ],
}
PROBE['inputs'].append(
    {
        'name': 'soil_texture',
        'type': 'enum',
        '_ref_type': 'UsdaTextureClass',
        'unit': None,
        'domain': None,
    }
)


def read(path: Path):  # noqa: ANN201 - Scientific JSON documents have mixed roots.
    return json.loads(path.read_bytes())


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8'
    )


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(
            {
                k: json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict, tuple)) else v
                for k, v in row.items()
            }
            for row in rows
        )


def finite(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def reported_features(row: dict, book: Workbook, formulas: Workbook) -> tuple[dict, dict]:
    """Read actual cells, before PSD closure, organic conversion or density derivation."""
    features, generic = {}, {}
    descriptors = (
        ('M', 'physical_clay_lt_10_um', None, '%', 0, 100),
        ('N', f'clay_lt_{row["clay_upper_um"]}_um', 'clay', '%', 0, 100),
        ('O', 'silt_1_50_um', 'silt', '%', 0, 100),
        ('P', 'sand_50_2000_um', 'sand', '%', 0, 100),
        ('Q', 'bulk_density', 'bulk_density', 'g/cm^3', 0, 10),
        ('R', 'organic_carbon', 'organic_carbon', '%', 0, 100),
        ('S', 'organic_matter', 'organic_matter', '%', 0, 100),
        ('T', 'ph', 'ph', '1', 0, 14),
        ('U', 'porosity_reported_scale', 'porosity', 'workbook_fraction_or_percent', 0, 100),
        ('V', 'particle_density', 'particle_density', 'g/cm^3', 0, 10),
    )
    sheet = book[row['worksheet']]
    formula_sheet = formulas[row['worksheet']]
    for alias, name, canonical, unit, low, high in descriptors:
        address = f'{row["columns"][alias]}{row["row"]}'
        cell, formula = sheet[address], formula_sheet[address]
        value = cell.value
        if formula.data_type == 'f' or not finite(value) or not low <= value <= high:
            continue
        if canonical in ('bulk_density', 'particle_density', 'porosity') and value == 0:
            continue
        features[name] = {
            'value': value,
            'unit': unit,
            'source_cells': [f'{sheet.title}!{address}'],
            'status': 'reported_numeric_workbook_cell',
            'laboratory_method': 'not_validated_here',
        }
        if canonical:
            generic[canonical] = name
    bins = row.get('psd_bins')
    if bins:
        for (low, high), value, address in zip(
            bins['intervals_um'], bins['original_values'], bins['cells'], strict=True
        ):
            sheet_name, cell_address = address.rsplit('!', 1)
            if (
                finite(value)
                and 0 <= value <= 100
                and formulas[sheet_name][cell_address].data_type != 'f'
            ):
                features[f'psd_{low}_{high}_um'] = {
                    'value': value,
                    'unit': '%',
                    'source_cells': [address],
                    'status': 'reported_numeric_workbook_cell',
                    'laboratory_method': 'not_validated_here',
                }
    return features, generic


def partition(records: list[dict]) -> dict[str, list[dict]]:
    """Exact signatures retain reported versus derived predictors and PSD cutoffs."""
    groups = defaultdict(list)
    for record in records:
        signature = json.dumps(
            {
                'reported': sorted(record['reported']),
                'e5_usable': sorted(record['e5_usable']),
                'row_status': record['row_status'],
            },
            sort_keys=True,
        )
        group = 'predictors_' + hashlib.sha256(signature.encode()).hexdigest()[:12]
        record['subset_id'] = group
        groups[group].append(record)
    return dict(sorted(groups.items()))


def pool_members(records: list[dict], required: tuple[str, ...], policy: str) -> list[dict]:
    """Require a complete input set on one predictor record, never a row union."""
    field = 'reported_canonical' if policy == 'reported' else 'e5_usable'
    return [
        r
        for r in records
        if r['row_status'] == 'unique_predictor_record' and set(required) <= set(r[field])
    ]


def counts(records: list[dict], basis: str) -> dict:
    points = [p for r in records for p in r['observations'] if p['basis'] == basis]
    return {
        'basis': basis,
        'n_samples': len({sample_key(p) for p in points}),
        'n_predictor_records': sum(
            any(p['basis'] == basis for p in r['observations']) for r in records
        ),
        'n_observations': len(points),
        'n_sources': len({p['source_id'] for p in points}),
    }


def export_dataset(path: Path, records: list[dict], basis: str, selection: dict) -> None:
    samples = [
        {
            k: r[k]
            for k in (
                'record_id',
                'source_id',
                'sample_id',
                'predictor_record',
                'row_status',
                'reported',
                'reported_canonical',
                'e5_usable',
                'predictor_evidence',
                'subset_id',
            )
        }
        for r in records
        if any(p['basis'] == basis for p in r['observations'])
    ]
    observations = [
        {**p, 'predictor_record_id': r['record_id']}
        for r in records
        for p in r['observations']
        if p['basis'] == basis
    ]
    write_json(
        path,
        {
            'format': 'ptfkit_predictor_testing_subset_v1',
            'selection': selection,
            'basis': basis,
            'counts': counts(records, basis),
            'samples': samples,
            'observations': observations,
            'note': 'Predictor availability only. Model domains, physical compatibility, pressure support and native evaluation still require checks. Split by original source/sample identity.',
        },
    )


def eligibility(record: dict, row: dict, sample: dict, context: dict) -> list[dict]:
    result = []
    for function in context['catalog']:
        targets = semantic.target_list(function, context)
        axis = {t['axis_input'] for t in targets if 'axis_input' in t}
        required = [b['name'] for b in function['inputs'] if b['name'] not in axis]
        conditioned = any(
            b.get('measurement', {}).get('quantity')
            in (
                'volumetric_water_content',
                'gravimetric_water_content',
            )
            for b in function['inputs']
            if b['name'] not in axis
        )
        _, _, reasons = predictors.resolve(function, sample, row, context, axis)
        result.append(
            {
                'record_id': record['record_id'],
                'subset_id': record['subset_id'],
                'source_id': record['source_id'],
                'sample_id': record['sample_id'],
                'predictor_record': record['predictor_record'],
                'function': function['id'],
                'required_inputs_excluding_axis': required,
                'reasons': reasons,
                'input_status': 'water_conditioned'
                if conditioned
                else 'blocked'
                if reasons
                else 'input_checks_passed',
                'native_evaluation_performed': False,
            }
        )
    return result


def main() -> None:  # noqa: C901, PLR0912 - Export and conservation checks.
    context, observations, rows, _ = participation.build_context()
    workbook = ROOT / 'local/meta_v3.3.xlsx'
    fingerprint = hashlib.sha256(workbook.read_bytes()).hexdigest()
    benchmark = read(ROOT / 'local/meta-final-v1_0-pointwise/experimental-wrc-ledger.json')
    if fingerprint != benchmark['workbook_sha256']:
        msg = 'Workbook differs from the frozen experimental ledger.'
        raise ValueError(msg)
    ledger = benchmark['observations'] + [p for p in observations if p['basis'] == 'ksat']
    scan_index = {identity(p): p for p in observations}
    if len(scan_index) != len(observations):
        msg = 'Rescan has duplicate experimental coordinates.'
        raise ValueError(msg)
    samples = {
        (s['source_id'], s['sample_id']): s for d in context['datasets'] for s in d['samples']
    }
    book = load_workbook(workbook, data_only=True, keep_links=False)
    formulas = load_workbook(workbook, data_only=False, keep_links=False)
    records, row_index = {}, {}
    try:
        for point in ledger:
            scanned = scan_index[identity(point)]
            if not math.isclose(point['observed'], scanned['observed'], rel_tol=1e-12):
                msg = 'Target value changed since the frozen ledger.'
                raise ValueError(msg)
            row = participation.matched_row(scanned, rows)
            row_label = f'{row["worksheet"]}!{row["row"]}' if row else 'unresolved'
            key = (*sample_key(point), row_label)
            if key not in records:
                record_id = hashlib.sha256(json.dumps(key, ensure_ascii=True).encode()).hexdigest()[
                    :16
                ]
                reported, canonical = reported_features(row, book, formulas) if row else ({}, {})
                sample = samples.get(sample_key(point), {'features': []})
                reported_class = campaign_inputs.reported_class(sample)
                if reported_class is not None:
                    reported['usda_texture_class'] = reported_class.document()
                    canonical['soil_texture'] = 'usda_texture_class'
                usable, proofs, _ = (
                    predictors.resolve(PROBE, sample, row, context) if row else ({}, {}, [])
                )
                record = {
                    'record_id': record_id,
                    'source_id': point['source_id'],
                    'sample_id': point['sample_id'],
                    'predictor_record': row_label,
                    'row_status': 'unique_predictor_record'
                    if row
                    else 'missing_or_ambiguous_predictor_record',
                    'reported': reported,
                    'reported_canonical': canonical,
                    'e5_usable': usable,
                    'predictor_evidence': proofs,
                    'observations': [],
                }
                records[key] = record
                if row:
                    row_index[record_id] = row
            records[key]['observations'].append(point)
    finally:
        book.close()
        formulas.close()
    records = list(records.values())
    groups = partition(records)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    write_json(
        OUTPUT / 'predictor-records.json', {'workbook_sha256': fingerprint, 'records': records}
    )
    memberships = [
        {
            'record_id': r['record_id'],
            'subset_id': r['subset_id'],
            'source_id': r['source_id'],
            'sample_id': r['sample_id'],
            'predictor_record': r['predictor_record'],
            'row_status': r['row_status'],
            'reported_predictors': sorted(r['reported']),
            'e5_usable_predictors': sorted(r['e5_usable']),
            **counts([r], b),
        }
        for r in records
        for b in BASES
        if counts([r], b)['n_observations']
    ]
    write_csv(OUTPUT / 'sample_membership.csv', memberships)
    summaries, pools = [], []
    for group, members in groups.items():
        for basis in BASES:
            summary = counts(members, basis)
            if not summary['n_observations']:
                continue
            summaries.append(
                {
                    'subset_id': group,
                    'reported_predictors': sorted(members[0]['reported']),
                    'e5_usable_predictors': sorted(members[0]['e5_usable']),
                    'row_status': members[0]['row_status'],
                    **summary,
                }
            )
            export_dataset(
                OUTPUT / 'exact' / group / f'{basis}.json',
                members,
                basis,
                {'exact_predictor_signature': group},
            )
    for policy in ('reported', 'e5_usable'):
        for name, required in POOLS.items():
            members = pool_members(records, required, policy)
            for basis in BASES:
                summary = counts(members, basis)
                pools.append(
                    {
                        'pool_id': name,
                        'availability_policy': policy,
                        'required_predictors': required,
                        **summary,
                    }
                )
                if summary['n_observations']:
                    export_dataset(
                        OUTPUT / 'pools' / policy / name / f'{basis}.json',
                        members,
                        basis,
                        {
                            'required_predictors': required,
                            'availability_policy': policy,
                            'overlapping_pools': True,
                        },
                    )
    expected_datasets = {
        (OUTPUT / 'exact' / r['subset_id'] / f'{r["basis"]}.json').resolve() for r in summaries
    } | {
        (
            OUTPUT / 'pools' / r['availability_policy'] / r['pool_id'] / f'{r["basis"]}.json'
        ).resolve()
        for r in pools
        if r['n_observations']
    }
    for folder in ('exact', 'pools'):
        for file in (OUTPUT / folder).rglob('*.json'):
            resolved = file.resolve()
            if resolved not in expected_datasets:
                if (
                    not resolved.is_relative_to(OUTPUT.resolve())
                    or read(file).get('format') != 'ptfkit_predictor_testing_subset_v1'
                ):
                    msg = 'Unexpected file in the generated subset directories.'
                    raise ValueError(msg)
                file.unlink()
    decisions = []
    for r in records:
        if r['record_id'] in row_index:
            decisions.extend(
                eligibility(
                    r,
                    row_index[r['record_id']],
                    samples.get((r['source_id'], r['sample_id']), {'features': []}),
                    context,
                )
            )
    write_csv(OUTPUT / 'function_input_eligibility.csv', decisions)
    write_csv(OUTPUT / 'predictor_subset_summary.csv', summaries)
    write_csv(OUTPUT / 'predictor_pool_summary.csv', pools)
    write_csv(PUBLIC / 'predictor_subset_summary.csv', summaries)
    write_csv(PUBLIC / 'predictor_pool_summary.csv', pools)
    lines = [
        '# ptfkit predictor-defined testing subsets',
        '',
        'Every experimental target is assigned once to a predictor record using exact source/sample identity and the existing reviewed record matcher. No cross-row union, imputation or target-based selection is used to construct the reported predictor sets. Missing or ambiguous rows remain an explicit subset. Water retention and Ksat are separate targets.',
        '',
        f'Exact predictor signatures: {len(groups)}. Predictor records: {len(records)}.',
        '',
        '| Basis | Samples | Predictor records | Observations | Sources |',
        '|---|---:|---:|---:|---:|',
        *[
            f'| {b} | {counts(records, b)["n_samples"]} | {counts(records, b)["n_predictor_records"]} | {counts(records, b)["n_observations"]} | {counts(records, b)["n_sources"]} |'
            for b in BASES
        ],
        '',
        '## Available testing pools',
        '',
        'Pools require the named predictors and include records with additional predictors. They overlap. Exact signature subsets are disjoint at the predictor-record level; a physical sample appearing in several target views must remain in the same validation fold.',
        '',
        '`reported` uses actual numeric, non-formula workbook cells. `e5_usable` additionally uses existing registered same-record transformations/proxies, with full evidence saved. Derived OC/OM, PSD closure, boundary proxies and density/porosity conversions are not relabelled as measurements. Availability is not proof of a compatible laboratory method or applicability of a specific PTF.',
        '',
        'Texture cutoffs are retained in exact signatures (including clay <1 versus <2 micrometres and original PSD bins). Generic texture pools mean availability of the reported components, not verified USDA fractions. Pressure and measured target water contents are not predictors. Specific formulas still require domain, pressure-support, water-basis and native-output checks.',
        '',
        'Generic E5 texture availability excludes filling missing fractions with class-average values. Source-specific reviewed class-midpoint inputs, such as the Cosby route, remain model-specific and are assessed in function_input_eligibility.csv. A reported or derived texture class is its own predictor, not three measured fractions.',
        '',
        '| Pool | Policy | Basis | Samples | Points |',
        '|---|---|---|---:|---:|',
        *[
            f'| {r["pool_id"]} | {r["availability_policy"]} | {r["basis"]} | {r["n_samples"]} | {r["n_observations"]} |'
            for r in pools
        ],
        '',
        '## Files and testing',
        '',
        '- `predictor_subset_summary.csv`: aggregate exact signatures, without private identities.',
        '- `predictor_pool_summary.csv`: aggregate counts for overlapping testing pools.',
        '- `local/meta-predictor-subsets-v1_0/sample_membership.csv`: exact membership and predictor record.',
        '- `local/meta-predictor-subsets-v1_0/predictor-records.json`: reported cells, E5 values, evidence and targets.',
        '- `local/meta-predictor-subsets-v1_0/exact/`: one complete dataset per exact signature and target basis.',
        '- `local/meta-predictor-subsets-v1_0/pools/`: complete datasets for reported and E5 testing pools.',
        '- `local/meta-predictor-subsets-v1_0/function_input_eligibility.csv`: function-specific input/domain checks. Native predictions were not evaluated by this grouping step.',
        '',
        'Split by source, or by profile/sample when appropriate, before model selection. Keep every pressure of the same physical sample in the same fold, including water-basis views and memberships in overlapping pools. These are availability subsets, not statistically independent train/test splits.',
        '',
        '## Two observations versus two samples',
        '',
        'The current primary WRC table has no model scored on only two observations. Its smallest rows are Chakraborty2011 equations 2, 3 and 5: 8 experimental points in 2 samples. Those are different counts. Two samples are weak validation evidence and must be reported with their 0.75% point coverage (8/1069) and 2.60% sample coverage (2/77); a small error on that subset is not evidence for general superiority.',
        '',
        'Secondary fixed-pressure functions may define only two native pressure outputs, such as 33 and 1500 kPa. That is the number of model outputs per sample, not the number of independent experimental observations in the primary benchmark. Reconstructed secondary targets remain explicitly separate.',
        '',
        'Reproduce: `python -B -m evaluation.experiments.meta_v3_3.predictor_subsets` in the prepared ptfkit-eval environment.',
        '',
    ]
    report = '\n'.join(lines)
    (PUBLIC / 'PREDICTOR-SUBSETS.md').write_text(report, encoding='utf-8')
    (OUTPUT / 'README.md').write_text(report, encoding='utf-8')
    assigned = [identity(p) for r in records for p in r['observations']]
    if len(assigned) != len(set(assigned)) or set(assigned) != {identity(p) for p in ledger}:
        msg = 'Subsets do not partition the entire experimental target ledger.'
        raise ValueError(msg)
    write_json(
        OUTPUT / 'validation.json',
        {
            'workbook_sha256': fingerprint,
            'exact_signatures': len(groups),
            'target_assignments': len(assigned),
            'duplicates': 0,
            'counts': [counts(records, b) for b in BASES],
        },
    )
    hashes = read(PUBLIC / 'artifact-hashes.json')
    manifest = {
        p.relative_to(OUTPUT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(OUTPUT.rglob('*'))
        if p.is_file() and p.name != 'artifact-hashes.json'
    }
    write_json(OUTPUT / 'artifact-hashes.json', manifest)
    for name in (
        'PREDICTOR-SUBSETS.md',
        'predictor_subset_summary.csv',
        'predictor_pool_summary.csv',
    ):
        hashes[name] = hashlib.sha256((PUBLIC / name).read_bytes()).hexdigest()
    write_json(PUBLIC / 'artifact-hashes.json', hashes)
    print(  # noqa: T201 - Command result.
        f'Exported {len(groups)} exact signatures, {len(records)} records and {len(assigned)} experimental targets to {OUTPUT}',
        flush=True,
    )


if __name__ == '__main__':
    main()
