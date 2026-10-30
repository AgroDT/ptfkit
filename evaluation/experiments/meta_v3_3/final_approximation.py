"""Full-ledger, basis-separated approximation benchmark; no PTF execution."""

# ruff: noqa: INP001, E501, PLR2004, C901, PLR0912, PLR0915, PERF401 - Explicit scientific protocol and report tables.

from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
from collections import Counter, defaultdict
from dataclasses import asdict
from importlib.util import find_spec
from itertools import pairwise
from pathlib import Path


ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
for package, location in (
    ('ptfkit', 'targets/ptfkit-py/src'),
    ('ptfkit_eval', 'targets/ptfkit-eval/src'),
):
    if find_spec(package) is None:
        sys.path.append(str(ROOT / location))

import numpy as np  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from ptfkit_eval import load_dataset_mapping, normalize_retention_workbook  # noqa: E402
from ptfkit_eval.families import evaluate_retention_family, retention_families  # noqa: E402
from scipy.optimize import least_squares  # noqa: E402
from scipy.stats import t as student_t  # noqa: E402

from evaluation.experiments.meta_v3_3 import prepare_selection_report as selection  # noqa: E402
from evaluation.experiments.meta_v3_3.selection_audit import additional_wrc, inventory, thesis_wrc  # noqa: E402


FAMILIES = tuple(retention_families())
BASES = ('volumetric_water_content', 'gravimetric_water_content')
SEED = 20261027
POLICY = {
    'id': 'full_ledger_approximation_v1',
    'seed': SEED,
    'minimum_observations': 'n >= free curve parameters + 1; at least that many distinct pressures',
    'objective': 'unweighted squared water-content residuals within each curve',
    'water_scale_upper': {'volumetric_water_content': 1.0, 'gravimetric_water_content': 5.0},
    'water_upper_explanation': 'Volume fraction <=1 is physical. The mass-basis upper bound 5 kg/kg is a numerical box, not a universal physical maximum; hits are rejected strictly.',
    'pressure_parameter_log10_bounds': [-12.0, 12.0],
    'shape_bounds': {
        'van_genuchten_mualem_v1': [1.0001, 20.0],
        'van_genuchten_m1_v1': [0.05, 20.0],
        'brooks_corey_v1': [0.01, 10.0],
        'clapp1978_gradual_air_entry_v1': [0.05, 11.4999],
    },
    'restarts': 'Twelve deterministic starts: six pressure quantiles crossed with two shape starts; BC additionally every pressure-partition midpoint.',
    'strict': 'Converged; finite and physical; n>p; no critical box hit; full normalized-Jacobian rank; condition <=1e10; nondegenerate fitted range.',
    'relaxed': 'Converged; finite and physical; n>p; nonconstant predictions. Boundary and rank warnings retained; descriptive only.',
    'critical_bounds': 'All optimizer coordinates except a zero residual-water fraction; zero residual water is an admissible nested model.',
    'boundary_tolerance': 1e-6,
    'rank_relative_tolerance': 1e-10,
    'maximum_condition': 1e10,
    'information_criteria': 'Gaussian ML; K=p+1 includes estimated residual variance; additive Gaussian constants omitted. AICc undefined for n<=K+1.',
    'comparisons': 'Same observed points; primary descriptive rankings on samples with relaxed success for all four families. Strict common-cohort sensitivity reported separately.',
    'error_distribution_quantile_levels': [0, 0.05, 0.25, 0.5, 0.75, 0.95, 1],
    'pressure_bins_kpa': [0, 0.1, 1, 10, 100, 1500, 10000, 100000, 1000000, 10000000],
    'regions_kpa': {'wet': [0, 10], 'intermediate': [10, 1500], 'dry': [1500, None]},
    'bootstrap': '1000 whole-sample cluster draws per bin; equal sample weights, within-sample points retained together; no source-independence claim.',
    'intervals': 'Empirical residual quantiles describe fitted-point errors, not validated prospective prediction intervals. Bootstrap bands quantify uncertainty of summaries; parameter intervals are separate conditional local-linear intervals.',
}


def json_write(path: Path, value: object) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=True, indent=2, sort_keys=True, allow_nan=False) + '\n',
        encoding='utf-8',
    )


def reconstruct(workbook: Path) -> tuple[list[dict], dict]:
    mapping = load_dataset_mapping(ROOT / 'evaluation/datasets/meta_v3_3.yaml')
    normalized = normalize_retention_workbook(workbook, mapping)
    review = json.loads(
        (ROOT / 'evaluation/datasets/meta_v3_3.clapp-source-review.json').read_bytes()
    )
    profiles = {
        (r['source_id'], r['sample_id']): r['profile_id'] for r in review['selected_samples']
    }
    points = []
    datasets = []
    for basis in BASES:
        selected = [o for o in normalized.observations if o.water_quantity == basis]
        datasets.append(
            {
                'samples': [
                    {'source_id': source, 'sample_id': sample}
                    for source, sample in sorted({(o.source, o.sample_id) for o in selected})
                ]
            }
        )
        for o in selected:
            points.append(
                {
                    'id': f'normalized:{o.row}',
                    'source_id': o.source,
                    'sample_id': o.sample_id,
                    'profile_id': profiles.get((o.source, o.sample_id)),
                    'horizon': None,
                    'pressure_kpa': o.pressure_kilopascal,
                    'observed': o.water_content,
                    'basis': basis,
                    'source_cells': list(o.source_cells),
                    'origin': 'reviewed_normalized_workbook_observation',
                    'raw_pressure': o.raw_pressure,
                    'raw_water': o.raw_water_content,
                    'raw_water_encoding': o.water_encoding,
                    'raw_pressure_encoding': o.pressure_encoding,
                }
            )
    context = {'root': ROOT, 'datasets': datasets}
    book = load_workbook(workbook, data_only=True, keep_links=False)
    try:
        master = book.worksheets[0]
        horizons = {
            (master[f'C{number}'].value, master[f'A{number}'].value): master[f'I{number}'].value
            for number in range(2, master.max_row + 1)
        }
        for point in points:
            point['horizon'] = horizons.get((point['source_id'], point['sample_id']))
        extra, decisions = additional_wrc(context, book)
        thesis, thesis_decisions = thesis_wrc(context, book)
    finally:
        book.close()
    points.extend(extra + thesis)
    frozen = json.loads(selection.SOURCE.read_bytes())
    actual = inventory(points)
    for key in frozen['dataset']['wrc']:
        if actual[key] != frozen['dataset']['wrc'][key]:
            message = f'Reconstructed ledger differs from the selection checkpoint: {key}: {actual[key]!r}'
            raise ValueError(message)
    unique = {(p['source_id'], p['sample_id'], p['basis'], p['pressure_kpa']) for p in points}
    if len(unique) != len(points):
        message = 'Repeated sample/pressure records need independent replicate evidence.'
        raise ValueError(message)
    audit = {
        'inventory': actual,
        'normalization': asdict(normalized.summary),
        'normalization_exclusions': [asdict(o) for o in normalized.exclusions],
        'additional_decisions': decisions,
        'thesis_decisions': thesis_decisions,
        'fitted_workbook_parameters_excluded': len(normalized.fitted_parameters),
    }
    return points, audit


def decode(family: str, x: np.ndarray, scale: float) -> dict:
    clapp = family == 'clapp1978_gradual_air_entry_v1'
    parameters = {'theta_s': float(x[0] * scale)}
    if not clapp:
        parameters['theta_r'] = float(x[0] * x[1] * scale)
    pressure_index = 1 if clapp else 2
    pressure_name = (
        'alpha'
        if family.startswith('van_genuchten')
        else 'saturation_suction'
        if clapp
        else 'air_entry'
    )
    shape_name = 'n' if family.startswith('van_genuchten') else 'b' if clapp else 'lambda'
    parameters[pressure_name] = float(10 ** x[pressure_index])
    parameters[shape_name] = float(x[pressure_index + 1])
    return parameters


def predict(family: str, parameters: dict, pressures: np.ndarray, scale: float) -> np.ndarray:
    # Homogeneous amplitudes permit an explicit mass-basis extension of the same shape.
    normalized = {k: v / scale if k in ('theta_s', 'theta_r') else v for k, v in parameters.items()}
    return scale * evaluate_retention_family(family, normalized, pressures)


def fit_curve(points: list[dict], family: str) -> dict:
    h = np.array([p['pressure_kpa'] for p in points])
    y = np.array([p['observed'] for p in points])
    basis = points[0]['basis']
    scale = POLICY['water_scale_upper'][basis]
    clapp = family == 'clapp1978_gradual_air_entry_v1'
    p = 3 if clapp else 4
    result = {
        'basis': basis,
        'family': family,
        'source_id': points[0]['source_id'],
        'sample_id': points[0]['sample_id'],
        'n': len(points),
        'free_parameters': p,
        'pressure_min_kpa': float(min(h)),
        'pressure_max_kpa': float(max(h)),
        'status': 'ineligible',
        'strict_accepted': False,
        'relaxed_accepted': False,
        'converged': False,
        'reasons': [],
        'parameters': None,
        'boundary_parameters': [],
        'physically_invalid_parameters': [],
        'rmse': None,
        'mae': None,
        'bias': None,
        'r2': None,
        'aicc': None,
        'bic': None,
        'jacobian_rank': None,
        'condition_number': None,
        'parameter_intervals': None,
    }
    if len(points) < p + 1 or len(np.unique(h)) < p + 1:
        result['reasons'] = ['insufficient_observations_or_distinct_pressures']
        return result
    if np.any(y < 0) or np.any(y >= scale) or not np.all(np.isfinite(y)):
        result['reasons'] = ['observations_outside_declared_water_box']
        return result
    shape_low, shape_high = POLICY['shape_bounds'][family]
    lower = np.array([1e-9, -12, shape_low] if clapp else [1e-9, 0, -12, shape_low])
    upper = np.array([1 - 1e-9, 12, shape_high] if clapp else [1 - 1e-9, 1 - 1e-9, 12, shape_high])
    positive = h[h > 0]
    pressure_starts = list(np.quantile(np.log10(positive), [0, 0.2, 0.4, 0.6, 0.8, 1]))
    if family == 'brooks_corey_v1':
        logs = np.unique(np.log10(positive))
        pressure_starts.extend((logs[:-1] + logs[1:]) / 2)
    shapes = [1.5, 4] if family.startswith('van_genuchten') else [2, 8] if clapp else [0.2, 1]
    starts = []
    for pressure in pressure_starts:
        head = -pressure if family.startswith('van_genuchten') else pressure
        for shape in shapes:
            saturation = min(0.99, max(1e-5, float(max(y)) * 1.1 / scale))
            residual_fraction = min(0.9, float(min(y) / max(y)) * 0.5) if max(y) > 0 else 0.1
            start = (
                [saturation, head, shape] if clapp else [saturation, residual_fraction, head, shape]
            )
            starts.append(np.clip(start, lower + 1e-8, upper - 1e-8))
    solutions = []
    for start in starts:
        optimized = least_squares(
            lambda x: predict(family, decode(family, x, scale), h, scale) - y,
            start,
            bounds=(lower, upper),
            method='trf',
            x_scale='jac',
            max_nfev=1500,
            ftol=1e-10,
            xtol=1e-10,
            gtol=1e-10,
        )
        if optimized.success and np.all(np.isfinite(optimized.fun)):
            solutions.append(optimized)
    result['attempted_starts'] = len(starts)
    result['converged_starts'] = len(solutions)
    if not solutions:
        result.update(status='failed', reasons=['optimizer_not_converged'])
        return result
    best = min(solutions, key=lambda s: float(s.fun @ s.fun))
    parameters = decode(family, best.x, scale)
    errors = best.fun
    fitted = errors + y
    sse = float(errors @ errors)
    sst = float(np.sum((y - np.mean(y)) ** 2))
    names = (
        ['saturation_fraction', 'log10_saturation_suction', 'b']
        if clapp
        else ['saturation_fraction', 'residual_fraction', 'log10_pressure_parameter', 'shape']
    )
    hits = [
        name + (':lower' if abs(best.x[i] - lower[i]) <= 1e-6 * (upper[i] - lower[i]) else ':upper')
        for i, name in enumerate(names)
        if min(best.x[i] - lower[i], upper[i] - best.x[i]) <= 1e-6 * (upper[i] - lower[i])
    ]
    critical = [name for name in hits if name != 'residual_fraction:lower']
    norms = np.linalg.norm(best.jac, axis=0)
    jac = best.jac / np.where(norms > 0, norms, 1)
    singular = np.linalg.svd(jac, compute_uv=False)
    rank = int(np.sum(singular > singular[0] * 1e-10)) if singular[0] > 0 else 0
    condition = float(singular[0] / singular[-1]) if singular[-1] > 0 else None
    if condition is not None and not math.isfinite(condition):
        condition = None
    admissible = bool(
        np.all(np.isfinite(fitted))
        and np.all(fitted >= 0)
        and np.all(fitted <= scale)
        and parameters['theta_s'] > parameters.get('theta_r', 0)
    )
    nondegenerate = bool(np.ptp(fitted) > 1e-8 * max(1, float(np.ptp(y))))
    relaxed = admissible and nondegenerate
    reasons = []
    if not admissible:
        reasons.append('physically_invalid_prediction_or_parameters')
    if not nondegenerate:
        reasons.append('degenerate_constant_curve')
    if critical:
        reasons.append('critical_boundary')
    if rank < p or condition is None or condition > 1e10:
        reasons.append('nonidentifiable_or_ill_conditioned')
    k, n = p + 1, len(y)
    likelihood = n * math.log(max(sse / n, 1e-300))
    aic = likelihood + 2 * k
    result.update(
        status='fitted',
        converged=True,
        strict_accepted=bool(relaxed and not reasons),
        relaxed_accepted=relaxed,
        reasons=reasons,
        parameters=parameters,
        boundary_parameters=hits,
        physically_invalid_parameters=[] if admissible else ['water_or_shape_contract'],
        rmse=math.sqrt(sse / n),
        mae=float(np.mean(abs(errors))),
        bias=float(np.mean(errors)),
        r2=1 - sse / sst if sst > 0 else None,
        aicc=aic + 2 * k * (k + 1) / (n - k - 1) if n > k + 1 else None,
        bic=likelihood + k * math.log(n),
        jacobian_rank=rank,
        condition_number=condition,
        optimizer_message=best.message,
        optimality=float(best.optimality),
        points=[
            {**point, 'predicted': float(value), 'residual': float(error)}
            for point, value, error in zip(points, fitted, errors, strict=True)
        ],
    )
    if result['strict_accepted'] and not hits:
        covariance = sse / (n - p) * np.linalg.pinv(best.jac.T @ best.jac, rcond=1e-14)
        derivatives = []
        for i in range(p):
            step = 1e-6 * max(1, abs(best.x[i]))
            dx = np.zeros(p)
            dx[i] = step
            plus, minus = decode(family, best.x + dx, scale), decode(family, best.x - dx, scale)
            derivatives.append([(plus[name] - minus[name]) / (2 * step) for name in parameters])
        transform = np.asarray(derivatives).T
        physical_cov = transform @ covariance @ transform.T
        se = np.sqrt(np.maximum(0, np.diag(physical_cov)))
        critical_t = float(student_t.ppf(0.975, n - p))
        result['parameter_intervals'] = {
            name: {
                'estimate': value,
                'standard_error': float(error),
                'lower': value - critical_t * float(error),
                'upper': value + critical_t * float(error),
            }
            for (name, value), error in zip(parameters.items(), se, strict=True)
        }
    return result


def weighted_quantiles(
    values: np.ndarray, weights: np.ndarray, quantiles: list[float]
) -> list[float]:
    order = np.argsort(values, kind='stable')
    ordered, mass = values[order], weights[order]
    cumulative = np.cumsum(mass) / np.sum(mass)
    return [
        float(ordered[min(len(ordered) - 1, int(np.searchsorted(cumulative, q, side='left')))])
        for q in quantiles
    ]


def residual_summary(groups: list[np.ndarray], *, bootstrap: int = 0, seed: int = SEED) -> dict:
    if not groups:
        return {'samples': 0, 'observations': 0}
    values = np.concatenate(groups)
    weights = np.concatenate([np.full(len(g), 1 / (len(groups) * len(g))) for g in groups])
    q = weighted_quantiles(values, weights, [0.025, 0.5, 0.975])
    ae = weighted_quantiles(abs(values), weights, [0.5, 0.95])
    per_sample = np.array([[np.mean(g), np.mean(abs(g)), np.mean(g**2)] for g in groups])
    means = np.mean(per_sample, axis=0)
    result = {
        'samples': len(groups),
        'observations': len(values),
        'bias': float(means[0]),
        'mae': float(means[1]),
        'rmse': math.sqrt(float(means[2])),
        'median_residual': q[1],
        'empirical_error_lower_95': q[0],
        'empirical_error_upper_95': q[2],
        'median_absolute_error': ae[0],
        'p95_absolute_error': ae[1],
    }
    if bootstrap and len(groups) >= 2:
        rng = np.random.default_rng(seed)
        indices = rng.integers(len(groups), size=(bootstrap, len(groups)))
        boot = np.mean(per_sample[indices], axis=1)
        medians = []
        for draw in indices:
            counts = np.bincount(draw, minlength=len(groups))
            w = np.concatenate(
                [
                    np.full(len(g), count / (len(groups) * len(g)))
                    for count, g in zip(counts, groups, strict=True)
                ]
            )
            medians.append(weighted_quantiles(values, w, [0.5])[0])
        result['bootstrap_95'] = {
            name: list(map(float, np.quantile(array, [0.025, 0.975])))
            for name, array in [
                ('bias', boot[:, 0]),
                ('mae', boot[:, 1]),
                ('rmse', np.sqrt(boot[:, 2])),
                ('median_residual', medians),
            ]
        }
        result['bootstrap_draws'] = bootstrap
    else:
        result['bootstrap_95'] = None
    return result


def support(records: list[dict]) -> dict:
    pressures = [np.array([p['pressure_kpa'] for p in r['points']]) for r in records]
    if not pressures:
        return {'min_kpa': None, 'max_kpa': None, 'p05_kpa': None, 'p95_kpa': None}
    values = np.concatenate(pressures)
    weights = np.concatenate([np.full(len(g), 1 / (len(pressures) * len(g))) for g in pressures])
    low, high = weighted_quantiles(values, weights, [0.05, 0.95])
    return {
        'min_kpa': float(min(values)),
        'max_kpa': float(max(values)),
        'p05_kpa': low,
        'p95_kpa': high,
    }


def aggregate(records: list[dict]) -> dict:
    output = {}
    for basis in BASES:
        selected = [r for r in records if r['basis'] == basis]
        indices = {
            family: {(r['source_id'], r['sample_id']): r for r in selected if r['family'] == family}
            for family in FAMILIES
        }
        common = set.intersection(
            *[{k for k, r in index.items() if r['relaxed_accepted']} for index in indices.values()]
        )
        strict_common = set.intersection(
            *[{k for k, r in index.items() if r['strict_accepted']} for index in indices.values()]
        )
        rows = []
        details = {}
        for family in FAMILIES:
            all_rows = list(indices[family].values())
            eligible = [r for r in all_rows if r['status'] != 'ineligible']
            relaxed = [r for r in eligible if r['relaxed_accepted']]
            strict = [r for r in eligible if r['strict_accepted']]
            paired = [indices[family][k] for k in sorted(common)]

            def quantiles(name: str, items: list[dict]) -> list | None:
                values = [r[name] for r in items if r[name] is not None]
                return (
                    list(map(float, np.quantile(values, [0, 0.05, 0.25, 0.5, 0.75, 0.95, 1])))
                    if values
                    else None
                )

            delta = {}
            for metric in ('aicc', 'bic'):
                valid_keys = [
                    k
                    for k in sorted(common)
                    if all(index[k][metric] is not None for index in indices.values())
                ]
                differences = [
                    indices[family][k][metric] - min(index[k][metric] for index in indices.values())
                    for k in valid_keys
                ]
                delta[metric] = float(np.median(differences)) if differences else None
                delta[metric + '_samples'] = len(valid_keys)
            row = {
                'basis': basis,
                'family': family,
                'total_samples': len(all_rows),
                'eligible_samples': len(eligible),
                'strict_successes': len(strict),
                'relaxed_successes': len(relaxed),
                'strict_success_rate': len(strict) / len(eligible) if eligible else None,
                'relaxed_success_rate': len(relaxed) / len(eligible) if eligible else None,
                'paired_relaxed_samples': len(common),
                'paired_strict_samples': len(strict_common),
                'median_rmse': float(np.median([r['rmse'] for r in paired])) if paired else None,
                'median_mae': float(np.median([r['mae'] for r in paired])) if paired else None,
                'strict_own_median_rmse': float(np.median([r['rmse'] for r in strict]))
                if strict
                else None,
                'strict_common_median_rmse': float(
                    np.median([indices[family][k]['rmse'] for k in strict_common])
                )
                if strict_common
                else None,
                'median_delta_aicc': delta['aicc'],
                'aicc_paired_samples': delta['aicc_samples'],
                'median_delta_bic': delta['bic'],
                'bic_paired_samples': delta['bic_samples'],
                'boundary_fraction': sum(bool(r['boundary_parameters']) for r in eligible)
                / len(eligible)
                if eligible
                else None,
                'ill_conditioned_fraction': sum(
                    'nonidentifiable_or_ill_conditioned' in r['reasons'] for r in eligible
                )
                / len(eligible)
                if eligible
                else None,
                **{'strict_support_' + k: v for k, v in support(strict).items()},
                **{'relaxed_support_' + k: v for k, v in support(relaxed).items()},
            }
            rows.append(row)
            bins = []
            edges = POLICY['pressure_bins_kpa']
            for low, high in pairwise(edges):
                groups = [
                    np.array(
                        [p['residual'] for p in r['points'] if low <= p['pressure_kpa'] < high]
                    )
                    for r in paired
                ]
                groups = [g for g in groups if len(g)]
                bins.append(
                    {
                        'pressure_min_kpa': low,
                        'pressure_max_kpa': high,
                        **residual_summary(groups, bootstrap=1000),
                    }
                )
            regions = {}
            for name, (low, high) in POLICY['regions_kpa'].items():
                groups = [
                    np.array(
                        [
                            p['residual']
                            for p in r['points']
                            if low <= p['pressure_kpa'] < (math.inf if high is None else high)
                        ]
                    )
                    for r in paired
                ]
                regions[name] = residual_summary([g for g in groups if len(g)], bootstrap=1000)
            details[family] = {
                'error_quantiles': {
                    name: quantiles(name, paired) for name in ('rmse', 'mae', 'bias')
                },
                'rmse_distribution': [r['rmse'] for r in paired],
                'bins': bins,
                'regions': regions,
                'rejection_counts': dict(
                    Counter(reason for r in eligible for reason in r['reasons'])
                ),
            }
        ranking = sorted(
            [r for r in rows if r['median_rmse'] is not None], key=lambda r: r['median_rmse']
        )
        comparisons = []
        if ranking:
            best = ranking[0]['family']
            for other in FAMILIES:
                if other == best or not common:
                    continue
                differences = np.array(
                    [indices[best][k]['rmse'] - indices[other][k]['rmse'] for k in sorted(common)]
                )
                draws = np.random.default_rng(SEED).integers(len(common), size=(2000, len(common)))
                comparisons.append(
                    {
                        'candidate': best,
                        'comparator': other,
                        'samples': len(common),
                        'median_paired_rmse_difference': float(np.median(differences)),
                        'bootstrap_median_difference_95': list(
                            map(
                                float,
                                np.quantile(np.median(differences[draws], axis=1), [0.025, 0.975]),
                            )
                        ),
                    }
                )
        regional_comparisons = {}
        for region, (low, high) in POLICY['regions_kpa'].items():
            sample_mae = {
                family: {
                    key: float(
                        np.mean(
                            [
                                abs(point['residual'])
                                for point in indices[family][key]['points']
                                if low
                                <= point['pressure_kpa']
                                < (math.inf if high is None else high)
                            ]
                        )
                    )
                    for key in sorted(common)
                    if any(
                        low <= point['pressure_kpa'] < (math.inf if high is None else high)
                        for point in indices[family][key]['points']
                    )
                }
                for family in FAMILIES
            }
            region_keys = sorted(set.intersection(*[set(index) for index in sample_mae.values()]))
            best_region = (
                min(FAMILIES, key=lambda f: np.mean([sample_mae[f][key] for key in region_keys]))
                if region_keys
                else None
            )
            comparisons_region = []
            if best_region:
                draws = np.random.default_rng(SEED).integers(
                    len(region_keys), size=(2000, len(region_keys))
                )
                for other in FAMILIES:
                    if other == best_region:
                        continue
                    diff = np.array(
                        [
                            sample_mae[best_region][key] - sample_mae[other][key]
                            for key in region_keys
                        ]
                    )
                    comparisons_region.append(
                        {
                            'candidate': best_region,
                            'comparator': other,
                            'mean_paired_mae_difference': float(np.mean(diff)),
                            'bootstrap_mean_difference_95': list(
                                map(
                                    float, np.quantile(np.mean(diff[draws], axis=1), [0.025, 0.975])
                                )
                            ),
                            'samples': len(region_keys),
                        }
                    )
            regional_comparisons[region] = comparisons_region
        output[basis] = {
            'summary': rows,
            'families': details,
            'descriptive_lowest_median_rmse': ranking[0]['family'] if ranking else None,
            'paired_comparisons': comparisons,
            'region_paired_comparisons': regional_comparisons,
            'region_lowest_mae': {
                region: min(
                    (f for f in FAMILIES if details[f]['regions'][region]['samples']),
                    key=lambda f: details[f]['regions'][region]['mae'],
                    default=None,
                )
                for region in POLICY['regions_kpa']
            },
        }
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=Path, default=ROOT / 'local/meta_v3.3.xlsx')
    parser.add_argument('--private-output', type=Path, default=ROOT / 'local/meta-final-benchmark')
    parser.add_argument('--aggregate-only', action='store_true')
    args = parser.parse_args()
    args.private_output.mkdir(parents=True, exist_ok=True)
    fit_path = args.private_output / 'approximation-fits.json'
    if args.aggregate_only:
        artifact = json.loads(fit_path.read_bytes())
        points, audit = reconstruct(args.workbook)
        by_id = {point['id']: point for point in points}
        for record in artifact['results']:
            for point in record.get('points', []):
                original = by_id[point['id']]
                if any(
                    point[key] != original[key]
                    for key in ('source_id', 'sample_id', 'pressure_kpa', 'observed', 'basis')
                ):
                    message = 'Stored fit observations differ from the reconstructed ledger.'
                    raise ValueError(message)
                point.update(original)
        artifact['ledger_inventory'] = audit['inventory']
        json_write(args.private_output / 'observations.json', points)
        json_write(args.private_output / 'ledger-audit.json', audit)
    else:
        points, audit = reconstruct(args.workbook)
        json_write(args.private_output / 'observations.json', points)
        json_write(args.private_output / 'ledger-audit.json', audit)
        groups = defaultdict(list)
        for point in points:
            groups[point['basis'], point['source_id'], point['sample_id']].append(point)
        records = []
        for number, (key, group) in enumerate(sorted(groups.items()), 1):
            for family in FAMILIES:
                records.append(fit_curve(sorted(group, key=lambda p: p['pressure_kpa']), family))
            print(f'Completed curve {number}/{len(groups)} ({key[0]})', flush=True)  # noqa: T201
            json_write(
                fit_path,
                {
                    'policy': POLICY,
                    'workbook_sha256': hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
                    'ledger_inventory': audit['inventory'],
                    'results': records,
                },
            )
        artifact = json.loads(fit_path.read_bytes())
    identities = [
        (r['basis'], r['source_id'], r['sample_id'], r['family']) for r in artifact['results']
    ]
    expected = {
        (p['basis'], p['source_id'], p['sample_id'], family) for p in points for family in FAMILIES
    }
    if len(set(identities)) != len(identities) or set(identities) != expected:
        message = 'Every ledger sample and registered family must have exactly one fit decision.'
        raise ValueError(message)
    for record in artifact['results']:
        record['parameter_units'] = {
            name: ('m3/m3' if record['basis'] == BASES[0] else 'kg/kg')
            if name in ('theta_s', 'theta_r')
            else '1/kPa'
            if name == 'alpha'
            else 'kPa'
            if name in ('air_entry', 'saturation_suction')
            else 'dimensionless'
            for name in (record['parameters'] or {})
        }
        record['parameter_interval_physical_violations'] = []
        for name, interval in (record['parameter_intervals'] or {}).items():
            low, high = interval['lower'], interval['upper']
            valid = low >= 0
            if name == 'theta_s':
                valid = low > 0 and high <= POLICY['water_scale_upper'][record['basis']]
            elif name == 'theta_r':
                valid = low >= 0 and high < record['parameters']['theta_s']
            elif name == 'n' and record['family'] == 'van_genuchten_mualem_v1':
                valid = low > 1
            elif name == 'b':
                valid = low > 0 and high < 11.5
            else:
                valid = low > 0
            if not valid:
                record['parameter_interval_physical_violations'].append(name)
    json_write(fit_path, artifact)
    aggregates = aggregate(artifact['results'])
    for basis in BASES:
        for row in aggregates[basis]['summary']:
            family = row['family']
            detail = aggregates[basis]['families'][family]
            populated = {
                name: value for name, value in detail['regions'].items() if value['samples']
            }
            best_region = (
                min(populated, key=lambda name: populated[name]['mae']) if populated else None
            )
            row['lowest_observed_mae_region'] = best_region
            row['lowest_error_region_samples'] = (
                populated[best_region]['samples'] if best_region else 0
            )
            row['lowest_error_region_median_absolute_error'] = (
                populated[best_region]['median_absolute_error'] if best_region else None
            )
            row['lowest_error_region_p95_absolute_error'] = (
                populated[best_region]['p95_absolute_error'] if best_region else None
            )
            examples = [
                r
                for r in artifact['results']
                if r['basis'] == basis and r['family'] == family and r['parameter_intervals']
            ]
            detail['parameter_uncertainty'] = {
                'available_fits': len(examples),
                'fits_with_interval_physical_violations': sum(
                    bool(r['parameter_interval_physical_violations']) for r in examples
                ),
                'qualification': 'conditional_local_linear_student_t; no validated population coverage',
            }
            if examples:
                typical = min(
                    examples, key=lambda r: abs(r['rmse'] - row['strict_own_median_rmse'])
                )
                json_write(
                    args.private_output / f'parameter-uncertainty-example-{basis}-{family}.json',
                    typical,
                )
    json_write(args.private_output / 'approximation-aggregates.json', aggregates)
    selection.OUTPUT.mkdir(exist_ok=True)
    selection.write_csv(
        selection.OUTPUT / 'wrc_approximation_summary.csv',
        [r for basis in BASES for r in aggregates[basis]['summary']],
    )
    public = {
        'status': 'calculations_complete_presentation_pending',
        'policy': POLICY,
        'source_path': selection.SOURCE.relative_to(ROOT).as_posix(),
        'source_sha256': hashlib.sha256(selection.SOURCE.read_bytes()).hexdigest(),
        'selection': json.loads(selection.SOURCE.read_bytes()),
        'approximation': {
            'workbook_sha256': artifact['workbook_sha256'],
            'ledger_inventory': artifact['ledger_inventory'],
            'aggregates': aggregates,
            'private_fit_artifact_sha256': hashlib.sha256(fit_path.read_bytes()).hexdigest(),
        },
    }
    json_write(selection.OUTPUT / 'benchmark_aggregates.json', public)
    print(  # noqa: T201 - Command result.
        json.dumps(
            {
                basis: {'preferred': a['descriptive_lowest_median_rmse'], 'summary': a['summary']}
                for basis, a in aggregates.items()
            },
            indent=2,
        )
    )


if __name__ == '__main__':
    main()
