"""E5 workbook semantic exploration, isolated from frozen strict and E0-E4 runs."""

# ruff: noqa: INP001, SLF001, E501, C901, PLR0912, PLR0915, PLR0917, PLR2004, ISC004 - Workbook coordinates and physical branches.

from __future__ import annotations

import argparse
import json
import math
import re
from collections import Counter, defaultdict
from itertools import pairwise
from pathlib import Path

from openpyxl import load_workbook
from ptfkit_eval.documents import canonical_json
from ptfkit_eval.execution import convert_value
from ptfkit_eval.harmonization import EvidenceTier, EvidenceValue

from evaluation.experiments.meta_v3_3 import campaign
from evaluation.experiments.meta_v3_3 import campaign_contracts as contracts
from evaluation.experiments.meta_v3_3 import campaign_inputs as inputs
from evaluation.experiments.meta_v3_3 import campaign_targets as targets
from evaluation.experiments.meta_v3_3 import exploratory_inputs as previous
from evaluation.experiments.meta_v3_3.source_domains import within_named_domain
from evaluation.experiments.meta_v3_3.workbook_evidence import (
    conductivity_factor,
    normalized_metadata,
    numeric,
)


PROTOCOL = 'workbook_semantic_exploratory_v1'
ROUTE = 'log_pressure_linear_interpolation_v1'
MECHANISMS = (
    'permissive Ksat',
    'plasticity targets',
    'pH',
    'particle density / porosity',
    'FC/PWP/AWC approximations',
    'inverse WRC',
    'PSD aliases',
    'other',
)
# Header positions are specific to the pinned workbook; the registry preserves literal labels.
MASTER = {
    'M': 'M',
    'N': 'N',
    'O': 'O',
    'P': 'P',
    'Q': 'Q',
    'R': 'R',
    'S': 'S',
    'T': 'T',
    'U': 'U',
    'V': 'V',
    'J': 'J',
    'I': 'I',
}
THESIS = dict(
    zip(MASTER, ('P', 'Q', 'R', 'S', 'T', 'U', 'V', 'W', 'X', 'Y', 'M', 'L'), strict=True)
)
ALIASES = {
    'k_sat': 'saturated_hydraulic_conductivity',
    'Ksat': 'saturated_hydraulic_conductivity',
    'hydraulic conductivity': 'saturated_hydraulic_conductivity',
    'water permeability': 'saturated_hydraulic_conductivity',
    'OC': 'organic_carbon',
    'SOC': 'organic_carbon',
    'OM': 'organic_matter',
    'particle density': 'particle_density',
    'total porosity': 'total_porosity',
    'FC': 'theta33',
    'PWP': 'theta1500',
    'plastic limit': 'lower_plastic_limit',
    'liquid limit': 'upper_plastic_limit',
}


def proof(value: float, mechanism: str, cells: tuple[str, ...], **details: object) -> EvidenceValue:
    return EvidenceValue(
        value,
        EvidenceTier.APPROXIMATE,
        PROTOCOL,
        cells,
        (('mechanism', mechanism), *details.items()),
    )


def cell_value(
    sheet: object, address: str, mechanism: str, **details: object
) -> EvidenceValue | None:
    value = sheet[address].value
    if not numeric(value):
        return None
    return proof(
        float(value), mechanism, (f'{sheet.title}!{address}',), original_value=value, **details
    )


def workbook_rows(book: object, formulas: object | None = None) -> list[dict]:
    """Keep rows independent; cross-sheet joins require unique keys and matching metadata."""
    result = []
    psd_index = defaultdict(list)
    for cells in book.worksheets[2].iter_rows(min_row=2, values_only=True):
        psd_index[cells[1], cells[0]].append(cells)
    for sheet_index, columns in ((0, MASTER), (3, MASTER), (4, MASTER), (6, THESIS)):
        sheet = book.worksheets[sheet_index]
        clay_upper = 1
        source = None
        for number in range(2, sheet.max_row + 1):
            key = (sheet[f'C{number}'].value, sheet[f'A{number}'].value)
            if not all(isinstance(v, str) and v.strip() for v in key):
                continue
            if key[0] != source:
                clay_upper, source = 1, key[0]
            clay_label = sheet[f'{columns["N"]}{number}'].value
            if isinstance(clay_label, str) and '< 2' in clay_label:
                clay_upper = 2
            row = {
                'source_id': key[0],
                'sample_id': key[1],
                'worksheet': sheet.title,
                'row': number,
                'values': {c: sheet[f'{v}{number}'].value for c, v in columns.items()},
                'columns': columns,
                'formulas': {},
                'clay_upper_um': clay_upper,
                'sheet_index': sheet_index,
                'properties': {},
            }
            row['value_cells'] = {c: (f'{sheet.title}!{v}{number}',) for c, v in columns.items()}
            if sheet_index == 6:
                bins = [sheet[f'{c}{number}'].value for c in ('CR', 'CS', 'CT', 'CU', 'CV', 'CW')]
                bin_cells = [
                    f'{sheet.title}!{c}{number}' for c in ('CR', 'CS', 'CT', 'CU', 'CV', 'CW')
                ]
            elif len(psd_index[key]) == 1:
                psd_row = psd_index[key][0]
                bins = list(psd_row[3:9])
                psd_number = next(
                    r[0].row
                    for r in book.worksheets[2].iter_rows(min_row=2)
                    if r[0].value == key[1] and r[1].value == key[0]
                )
                bin_cells = [f'{book.worksheets[2].title}!{c}{psd_number}' for c in 'DEFGHI']
            else:
                bins, bin_cells = [], []
            if bins:
                original_bins = [v if numeric(v) else None for v in bins]
                bins = [v if numeric(v) and 0 <= v <= 100 else None for v in bins]
                bin_sources = [(c,) for c in bin_cells]
                missing_bin = None
                known = [v for v in bins if numeric(v) and 0 <= v <= 100]
                if len(known) == 5 and 99 <= sum(known) <= 100:
                    missing_bin = bins.index(None)
                    bins[missing_bin] = 100 - sum(known)
                    bin_sources[missing_bin] = tuple(
                        c for i, c in enumerate(bin_cells) if i != missing_bin
                    )
                for column, indices in (
                    ('P', (0, 1)),
                    ('O', (2, 3, 4)),
                    ('N', (5,)),
                    ('M', (3, 4, 5)),
                ):
                    if not numeric(row['values'][column]) and all(
                        numeric(bins[i]) for i in indices
                    ):
                        row['values'][column] = sum(bins[i] for i in indices)
                        row['value_cells'][column] = tuple(
                            sorted({c for i in indices for c in bin_sources[i]})
                        )
                        if column == 'N':
                            row['clay_upper_um'] = 1
                row['psd_bins'] = {
                    'values': bins,
                    'cells': bin_cells,
                    'sources': bin_sources,
                    'original_values': original_bins,
                    'missing_bin_derived': missing_bin,
                    'closure_formula': 'missing percent = 100 - sum(other five observed bins)'
                    if missing_bin is not None
                    else None,
                    'intervals_um': [[250, 1000], [50, 250], [10, 50], [5, 10], [1, 5], [0, 1]],
                }
            if formulas is not None:
                row['formulas'] = {
                    c: formulas[sheet.title][f'{v}{number}'].value
                    for c, v in columns.items()
                    if formulas[sheet.title][f'{v}{number}'].data_type == 'f'
                }
            for name, column, mechanism in (
                ('ph', 'T', 'pH'),
                ('particle_density', 'V', 'particle density / porosity'),
                ('porosity', 'U', 'particle density / porosity'),
            ):
                item = cell_value(
                    sheet,
                    f'{columns[column]}{number}',
                    mechanism,
                    header=str(sheet[f'{columns[column]}1'].value),
                    measurement_method='unknown',
                )
                if item is None:
                    continue
                value = float(item.value)
                if name == 'ph' and 0 <= value <= 14:
                    item = proof(
                        value,
                        mechanism,
                        item.sources,
                        ph_method='unknown',
                        method_contract_approximate=True,
                        original_value=value,
                    )
                elif name == 'particle_density' and 0 < value < 10:
                    item = proof(
                        value,
                        mechanism,
                        item.sources,
                        original_value=value,
                        interpreted_unit='gram_per_cubic_centimeter',
                        measurement_method='unknown',
                    )
                elif name == 'porosity' and 0 < value <= 100:
                    item = proof(
                        value if value <= 1 else value / 100,
                        mechanism,
                        item.sources,
                        original_value=value,
                        interpreted_scale='fraction' if value <= 1 else 'percent',
                    )
                else:
                    continue
                row['properties'][name] = item
            if sheet_index == 6:
                for name, column in (('lower_plastic_limit', 'AF'), ('upper_plastic_limit', 'AG')):
                    item = cell_value(
                        sheet,
                        f'{column}{number}',
                        'plasticity targets',
                        original_unit_label='not_recorded',
                        interpreted_unit='mass_percent',
                        interpretation_status='workbook_semantic_assumption',
                        measurement_method='unknown',
                    )
                    if item is not None and 0 <= item.value <= 100:
                        row['properties'][name] = item
                lower = row['properties'].get('lower_plastic_limit')
                upper = row['properties'].get('upper_plastic_limit')
                if lower and upper and upper.value >= lower.value:
                    row['properties']['plasticity_index'] = proof(
                        upper.value - lower.value,
                        'plasticity targets',
                        (*lower.sources, *upper.sources),
                        formula='upper - lower',
                        parents=[lower.document(), upper.document()],
                    )
                unit = (
                    'meter_per_day'
                    if 33 <= number <= 40
                    else 'centimeter_per_minute'
                    if 42 <= number <= 71
                    else 'centimeter_per_day'
                    if number == 87
                    else None
                )
                if unit:
                    item = cell_value(
                        sheet,
                        f'AB{number}',
                        'permissive Ksat',
                        original_unit_label='cm/smin'
                        if 42 <= number <= 71
                        else 'm/day'
                        if number <= 40
                        else 'cm/day',
                        literal_workbook_unit_label=sheet[
                            'AB41' if 42 <= number <= 71 else 'AB32' if number <= 40 else 'AB86'
                        ].value,
                        interpreted_unit='cm/min' if 42 <= number <= 71 else unit,
                        interpretation_status='workbook_semantic_assumption',
                        measurement_method='unknown_or_not_recorded',
                        canonical_unit='meter_per_second',
                    )
                    if item and item.value >= 0:
                        row['properties']['ksat'] = proof(
                            item.value * conductivity_factor(unit),
                            'permissive Ksat',
                            item.sources,
                            **{k: v for k, v in item.details if k != 'mechanism'},
                            conversion_factor=conductivity_factor(unit),
                        )
                # Source-local porosity block, not a global interpretation of the WRC columns.
                if 42 <= number <= 71 and numeric(sheet[f'AK{number}'].value):
                    item = cell_value(sheet, f'AK{number}', 'particle density / porosity')
                    if item and 0 < item.value <= 1 and 'porosity' not in row['properties']:
                        row['properties']['porosity'] = proof(
                            item.value,
                            'particle density / porosity',
                            item.sources,
                            original_value=item.value,
                            interpreted_scale='fraction',
                            header_cell='AK41',
                        )
            result.append(row)
    return result


def compatible(left: dict, right: dict) -> bool:
    return (left['source_id'], left['sample_id']) == (
        right['source_id'],
        right['sample_id'],
    ) and all(
        left['values'].get(c) is None
        or right['values'].get(c) is None
        or normalized_metadata(left['values'][c]) == normalized_metadata(right['values'][c])
        for c in ('I', 'J')
    )


def inverse_wrc(points: list[tuple[float, float]], theta: float) -> float | None:
    """Invert a monotone measured curve in log pressure, with no extrapolation."""
    ordered = sorted(points)
    if (
        len({p for p, _ in ordered}) != len(ordered)
        or any(not (numeric(p) and numeric(w)) or p < 0 for p, w in ordered)
        or any(a[1] < b[1] for a, b in pairwise(ordered))
    ):
        return None
    exact = [p for p, w in ordered if w == theta]
    if len(exact) == 1:
        return exact[0]
    if exact:
        return None  # A plateau has no unique inverse.
    for (p0, w0), (p1, w1) in pairwise(ordered):
        if p0 > 0 and w1 < theta < w0:
            return math.exp(math.log(p0) + (theta - w0) / (w1 - w0) * math.log(p1 / p0))
    return None


def semantic_inputs(
    function: dict, sample: dict, row: dict, context: dict, water: dict, axis: set[str]
) -> tuple[dict, dict, list[str]]:
    values, proofs, _ = inputs.resolve_inputs(
        function,
        sample,
        row,
        'reported_boundary_proxy',
        context['old_registration'],
        representative=False,
        water_inputs=water,
        axis_inputs=axis,
    )

    def remap(document: dict) -> dict:
        document = dict(document)
        sources = []
        for source in document['sources']:
            column = next(
                (c for c in row['columns'] if source == f'{row["worksheet"]}!{c}{row["row"]}'), None
            )
            sources.extend(row['value_cells'][column] if column else (source,))
        document['sources'] = sorted(set(sources))
        details = dict(document.get('details', {}))
        if details.get('reported_column') in row['columns']:
            alias = details['reported_column']
            details.update(normalized_column_alias=alias, reported_column=row['columns'][alias])
        document['details'] = details
        document['parents'] = [remap(p) for p in document['parents']]
        return document

    proofs = {name: remap(p) for name, p in proofs.items()}
    properties = row['properties']
    texture = previous.texture_predictors(
        row, 'reported_boundary_proxy', context['old_registration']
    )
    for binding in function['inputs']:
        name = binding['name']
        if name in axis:
            continue
        if function['id'] == 'calc_ptf_cosby1984_univariate' and name in ('sand', 'silt', 'clay'):
            # The resolver requires the complete reviewed Table 2 midpoint. A measured
            # fraction must never replace one component of that categorical coordinate.
            continue
        item, source_unit = None, None
        if name in (
            'ph',
            'ph_h2o',
            'ph_kcl',
            'ph_cacl2',
            'particle_density',
            'porosity',
            'total_porosity',
        ):
            item = properties.get(
                'porosity' if name == 'total_porosity' else 'ph' if name.startswith('ph') else name
            )
            source_unit = (
                'dimensionless'
                if name.startswith('ph')
                else 'gram_per_cubic_centimeter'
                if name == 'particle_density'
                else 'volume_fraction'
            )
            if item is None and name in ('porosity', 'total_porosity'):
                bulk, particle = previous.raw_number(row, 'Q'), properties.get('particle_density')
                if bulk and particle and 'Q' not in row['formulas'] and 0 < bulk < particle.value:
                    item = proof(
                        1 - bulk / particle.value,
                        'particle density / porosity',
                        (
                            f'{row["worksheet"]}!{row["columns"]["Q"]}{row["row"]}',
                            *particle.sources,
                        ),
                        formula='1 - bulk_density / particle_density',
                        measurement_method='unknown',
                    )
        elif name in ('clay', 'clay_content', 'sand', 'sand_content', 'silt', 'silt_content'):
            component = name.removesuffix('_content')
            item = texture.get(component)
            if item is not None:
                mapped = remap(item.document())
                item = proof(
                    float(item.value), 'PSD aliases', tuple(mapped['sources']), parents=[mapped]
                )
            if item is None:
                column = {'clay': 'N', 'silt': 'O', 'sand': 'P'}[component]
                value = previous.raw_number(row, column)
                if value is None and component == 'clay':
                    column, value = 'M', previous.raw_number(row, 'M')
                if value is None and component == 'sand':
                    physical_clay = previous.raw_number(row, 'M')
                    if physical_clay is not None and 0 <= physical_clay <= 100:
                        item = proof(
                            100 - physical_clay,
                            'PSD aliases',
                            row['value_cells']['M'],
                            formula='100 - physical_clay_percent',
                            reported_interval_um=[10, 2000],
                            requested_interval_um=inputs.interval(binding),
                            cutoff_mismatch_approximation=True,
                        )
                if value is not None and 0 <= value <= 100:
                    item = proof(
                        value,
                        'PSD aliases',
                        row['value_cells'][column],
                        reported_interval_um={
                            'M': [0, 10],
                            'N': [0, row['clay_upper_um']],
                            'O': [1, 50],
                            'P': [50, 2000],
                        }[column],
                        requested_interval_um=inputs.interval(binding),
                        cutoff_mismatch_approximation=True,
                    )
            if item:
                item = proof(
                    item.value,
                    'PSD aliases',
                    item.sources,
                    reported_clay_upper_um=row['clay_upper_um'],
                    requested_interval_um=inputs.interval(binding),
                    cutoff_mismatch_approximation=True,
                    parents=[item.document()],
                    same_sample_psd_bins=row.get('psd_bins'),
                )
            source_unit = 'mass_percent'
        elif name in ('fine_sand', 'coarse_sand', 'very_fine_sand') and row.get('psd_bins'):
            bounds = (
                inputs.interval(binding)
                or {
                    'fine_sand': (50, 250),
                    'coarse_sand': (250, 1000),
                    'very_fine_sand': (50, 100),
                }[name]
            )
            bins = row['psd_bins']
            value, sources, complete = 0, [], True
            for number, (low, high) in enumerate(bins['intervals_um']):
                left, right = max(low, bounds[0]), min(high, bounds[1])
                if left >= right:
                    continue
                raw = bins['values'][number]
                if not numeric(raw) or low == 0:
                    complete = False
                    break
                weight = math.log(right / left) / math.log(high / low)
                value += raw * weight
                sources.extend(bins['sources'][number])
            if complete and sources:
                item = proof(
                    value,
                    'PSD aliases',
                    tuple(sources),
                    requested_interval_um=bounds,
                    original_intervals_um=bins['intervals_um'],
                    formula='sum of observed bins, log-diameter interpolation for partial bins',
                    cutoff_mismatch_approximation=True,
                )
            source_unit = 'mass_percent'
        elif name in ('depth', 'mean_depth'):
            raw = str(row['values'].get('J', '')).strip().casefold()
            numbers = re.fullmatch(
                r'\(?\s*(\d+(?:\.\d+)?)\s*(?:[-\u2013\u2014]\s*(\d+(?:\.\d+)?))?\s*\)?\s*(?:cm|\u0441\u043c)?',
                raw,
            )
            if numbers:
                lower, upper = float(numbers[1]), float(numbers[2] or numbers[1])
                if 0 <= lower <= upper:
                    item = proof(
                        (lower + upper) / 2,
                        'other',
                        row['value_cells']['J'],
                        original_value=raw,
                        interpreted_unit='centimeter',
                        formula='interval midpoint or explicit depth',
                    )
            source_unit = 'centimeter'
        elif name == 'bulk_volume' and binding.get('unit') == 'cm^3/g':
            bulk = inputs.density(row, context['old_registration'])
            if bulk and bulk.value > 0:
                item = proof(
                    1 / float(bulk.value),
                    'other',
                    row['value_cells']['Q'],
                    formula='1 / bulk_density_g_cm3',
                    interpreted_unit='cubic_centimeter_per_gram',
                    original_bulk_density=float(bulk.value),
                    measurement_method='unknown',
                    clod_preparation_contract_approximate=True,
                )
                source_unit = 'cubic_centimeter_per_gram'
        if item is not None:
            unit = contracts.canonical_input_unit(binding)
            if name == 'bulk_volume' and binding.get('unit') == 'cm^3/g':
                unit = 'cubic_centimeter_per_gram'
            try:
                values[name] = convert_value(float(item.value), source_unit, unit)
                proofs[name] = item.document()
            except ValueError:
                values.pop(name, None)
    reasons = [
        'missing_predictor:' + b['name']
        for b in function['inputs']
        if b['name'] not in values and b['name'] not in axis
    ]
    for binding in function['inputs']:
        name = binding['name']
        if name in values and not within_named_domain(
            binding.get('domain') or '', {**values, 'value': values[name]}
        ):
            reasons.append('computational_domain:' + name)
        if name in proofs:
            proofs[name].update(
                native_value=values.get(name),
                native_unit=binding.get('unit'),
                canonical_input_unit=contracts.canonical_input_unit(binding),
            )
    return values, proofs, reasons


def target_list(function: dict, context: dict) -> list[dict]:
    result = contracts.targets(function, context['support'])
    for field in function['fields']:
        name, quantity = field['name'], field.get('quantity')
        if quantity == 'volumetric_water_content' and any(
            text in field.get('description', '').lower()
            for text in ('represented by soil porosity', 'taken as total porosity')
        ):
            result.append(
                {
                    'kind': 'porosity_parameter',
                    'field': name,
                    'quantity': quantity,
                    'unit': field['unit'],
                    'adapter': 'native_spec_fixed_pressure_v1',
                    'mechanism': 'particle density / porosity',
                }
            )
        if quantity in ('saturated_hydraulic_conductivity', 'log_saturated_hydraulic_conductivity'):
            # Standard deviations and separate capillary components are not whole-soil Ksat.
            if name.startswith('sd_') or any(
                w in field.get('description', '').lower()
                for w in ('noncapillary', 'capillary hydraulic', 'matching-point')
            ):
                continue
            result.append(
                {
                    'kind': 'ksat',
                    'field': name,
                    'quantity': quantity,
                    'unit': field['unit'],
                    'mechanism': 'permissive Ksat',
                }
            )
        elif name in ('lower_plastic_limit', 'upper_plastic_limit', 'plasticity_index'):
            result.append(
                {
                    'kind': 'plasticity',
                    'field': name,
                    'quantity': quantity,
                    'unit': field['unit'],
                    'mechanism': 'plasticity targets',
                }
            )
        elif not any(t.get('field') == name for t in result):
            description = field.get('description', '').lower()
            is_awc = quantity == 'available_water_capacity' or name in (
                'plant_available_water',
                'available_water',
            )
            is_fc = name in ('field_capacity', 'theta_fc', 'fc') or 'field capacity' in description
            is_pwp = (
                name in ('wilting_point', 'permanent_wilting_point', 'pwp')
                or 'wilting point' in description
            )
            if (
                (is_awc or is_fc or is_pwp)
                and 'extended' not in description
                and 'extended' not in name
            ):
                # Explicit alternatives require their own parsed contract; never overwrite them.
                scope = function.get('scope', {}).get('prediction_target', '')
                if (
                    contracts.explicit_pressure(description + ' ' + scope) is not None
                    or 'cylinder' in description + scope.lower()
                ):
                    continue
                result.append(
                    {
                        'kind': 'operational',
                        'field': name,
                        'quantity': quantity
                        if quantity in contracts.WATER
                        else 'volumetric_water_content',
                        'unit': field['unit'],
                        'pressures': [33, 1500] if is_awc else [33] if is_fc else [1500],
                        'mechanism': 'FC/PWP/AWC approximations',
                    }
                )
    if any(i['name'] == 'theta' for i in function['inputs']):
        result.extend(
            {
                'kind': 'inverse',
                'field': f['name'],
                'quantity': 'matric_potential',
                'unit': f['unit'],
                'mechanism': 'inverse WRC',
            }
            for f in function['fields']
            if f.get('quantity') == 'matric_potential'
        )
    return result


def observe(  # noqa: PLR0911 - Separate physical target dispositions.
    target: dict, function: dict, dataset: dict | None, sample: dict, row: dict, context: dict
) -> tuple[EvidenceValue | None, str | None]:
    kind = target['kind']
    if kind in ('ksat', 'plasticity'):
        return row['properties'].get('ksat' if kind == 'ksat' else target['field']), None
    if kind == 'direct_property':
        value = inputs.density(row, context['old_registration'])
        return (
            proof(
                float(value.value),
                'other',
                row['value_cells']['Q'],
                original_unit='gram_per_cubic_centimeter',
                measurement_method='unknown',
            )
            if value
            else None
        ), None
    if kind == 'porosity_parameter':
        return row['properties'].get('porosity'), None
    if dataset is None:
        return None, 'same_sample_wrc_absent'
    if kind == 'operational':
        parents = []
        for pressure in target['pressures']:
            value, reason = campaign.reconstruct_target(
                context,
                function,
                {
                    'kind': 'fixed_pressure',
                    'quantity': target['quantity'],
                    'field': target['field'],
                    'pressure_kpa': pressure,
                    'adapter': 'native_spec_fixed_pressure_v1',
                },
                dataset,
                sample,
                row=row,
                route=ROUTE,
                disabled=set(),
            )
            if value is None:
                return None, reason
            parents.append(value)
        value = parents[0].value - (parents[1].value if len(parents) == 2 else 0)
        return proof(
            value,
            target['mechanism'],
            tuple(sorted({s for p in parents for s in p.sources})),
            operational_target_approximation=True,
            pressures_kpa=target['pressures'],
            parents=[p.document() for p in parents],
        ), None
    return campaign.reconstruct_target(
        context, function, target, dataset, sample, row=row, route=ROUTE, disabled=set()
    )


def prediction(function: dict, target: dict, values: dict, context: dict) -> float:
    if (
        target.get('unit') == 'cubic_centimeter_per_gram'
        and target.get('quantity') == 'gravimetric_water_content'
    ):
        value = contracts.native_outputs(function, values, target)[target['field']]
        return campaign.finite_prediction(
            value * context['old_registration']['water_density_g_cm3']
        )
    if target['kind'] not in ('ksat', 'plasticity', 'operational', 'inverse'):
        return campaign.finite_prediction(
            contracts.predicted(function, target, values, context['manifest'])
        )
    value = contracts.native_outputs(function, values, {})[target['field']]
    unit = target['unit']
    if target['kind'] == 'ksat':
        if unit.startswith('log10_'):
            value, unit = 10**value, unit.removeprefix('log10_')
        factor = conductivity_factor(unit)
        if factor is None and unit == 'inch_per_hour':
            factor = 0.0254 / 3600
        if factor is None:
            message = 'Unresolved conductivity logarithm or reference unit.'
            raise ValueError(message)
        value *= factor
    elif target['kind'] == 'plasticity':
        value = convert_value(value, unit, 'mass_percent')
    elif target['kind'] == 'operational':
        value = (
            value / 1000
            if unit == 'millimeter_per_meter'
            else convert_value(
                value,
                unit,
                'mass_fraction'
                if target['quantity'] == 'gravimetric_water_content'
                else 'volume_fraction',
            )
        )
    elif target['kind'] == 'inverse':
        value = convert_value(value, unit, 'kilopascal')
    return campaign.finite_prediction(value)


def measured_targets(candidates: list[dict], pairs: list[tuple], key: tuple) -> list[tuple]:
    """Use every observed WRC pressure for continuous outputs, retaining scalar contracts."""
    result = []
    curves = {
        'existing_model_adapter_v1',
        'explicit_vg_mualem_signature_v1',
        'hutson_cass_dry_v1',
        'logistic_pressure_scale_v1',
    }
    for target in candidates:
        if pairs and ('axis_input' in target or target.get('adapter') in curves):
            for dataset, sample in pairs:
                pressures = {
                    target['pressure_kpa'],
                    *(p for p, _ in targets.measured_series(dataset, sample)),
                }
                result.extend(
                    ({**target, 'pressure_kpa': pressure}, [(dataset, sample)])
                    for pressure in sorted(pressures)
                )
        else:
            result.append(
                (
                    target,
                    pairs or [(None, {'source_id': key[0], 'sample_id': key[1], 'features': []})],
                )
            )
    # Five registered curve anchors can lead to the same observed coordinate.
    unique = {}
    for target, selected in result:
        dataset, sample = selected[0]
        identity = (
            json.dumps(target, sort_keys=True),
            dataset['target']['quantity'] if dataset else None,
            sample['source_id'],
            sample['sample_id'],
        )
        unique.setdefault(identity, (target, selected))
    return list(unique.values())


def evaluate(context: dict, rows: list[dict]) -> tuple[list[dict], list[dict]]:
    decisions, points = [], []
    measured = defaultdict(list)
    for dataset in context['datasets']:
        for sample in dataset['samples']:
            measured[sample['source_id'], sample['sample_id']].append((dataset, sample))
    masters = defaultdict(list)
    for row in rows:
        if row['sheet_index'] == 0:
            masters[row['source_id'], row['sample_id']].append(row)
    for function in context['catalog']:
        candidates = target_list(function, context)
        if not candidates:
            decisions.append(
                {
                    'function': function['id'],
                    'status': 'blocked',
                    'reasons': ['physical_target_unavailable_or_definition_unresolved'],
                }
            )
        for row in rows:
            key = (row['source_id'], row['sample_id'])
            joins = masters[key]
            matched = len(joins) == 1 and compatible(row, joins[0])
            pairs = measured.get(key, []) if matched else []
            for target, selected_pairs in measured_targets(candidates, pairs, key):
                for dataset, sample in selected_pairs:
                    record = {
                        'function': function['id'],
                        'source_id': key[0],
                        'sample_id': key[1],
                        'workbook_row': f'{row["worksheet"]}!{row["row"]}',
                        'target': target,
                        'status': 'blocked',
                        'reasons': [],
                    }
                    decisions.append(record)
                    if target['kind'] == 'inverse':
                        if (
                            dataset is None
                            or dataset['target']['quantity'] != 'volumetric_water_content'
                        ):
                            record['reasons'] = ['same_sample_volumetric_wrc_absent']
                            continue
                        curve = targets.measured_series(dataset, sample)
                        # Evaluate interior water coordinates supplied by actual observations.
                        coordinates = [(p, v) for p, v in curve if p > 0]
                        coordinates.extend(
                            (
                                None,
                                proof(
                                    (float(a.value) + float(b.value)) / 2,
                                    'inverse WRC',
                                    (*a.sources, *b.sources),
                                    formula='midpoint of adjacent measured water contents',
                                    parents=[a.document(), b.document()],
                                ),
                            )
                            for (p0, a), (p1, b) in pairwise(
                                sorted(curve, key=lambda item: item[0])
                            )
                            if 0 < p0 < p1 and a.value > b.value
                        )
                    else:
                        observed, reason = observe(target, function, dataset, sample, row, context)
                        if observed is None:
                            record['reasons'] = [reason or 'target_property_absent']
                            continue
                        coordinates = [(None, observed)]
                    axis = {target['axis_input']} if 'axis_input' in target else set()
                    if target['kind'] == 'inverse':
                        axis.add('theta')
                    water = (
                        campaign.water_predictors(
                            function,
                            dataset,
                            sample,
                            row,
                            context['old_registration'],
                            enabled=True,
                        )
                        if dataset
                        else {}
                    )
                    values, proofs, reasons = semantic_inputs(
                        function, sample, row, context, water, axis
                    )
                    record['reasons'] = reasons
                    if reasons:
                        continue
                    if target['kind'] == 'inverse':
                        record['reasons'] = ['invalid_or_ambiguous_inverse_wrc']
                    for pressure, coordinate_observation in coordinates:
                        observed = coordinate_observation
                        selected, selected_proofs = dict(values), dict(proofs)
                        if target['kind'] == 'inverse':
                            theta = float(observed.value)
                            inverse = inverse_wrc([(p, float(v.value)) for p, v in curve], theta)
                            if inverse is None:
                                continue
                            selected['theta'] = theta
                            selected_proofs['theta'] = observed.document()
                            observed = proof(
                                inverse,
                                'inverse WRC',
                                observed.sources,
                                theta=theta,
                                interpolation_space='log_pressure',
                                extrapolation=False,
                            )
                        try:
                            predicted = prediction(function, target, selected, context)
                        except (ValueError, TypeError, ArithmeticError) as error:
                            if record['status'] != 'executed':
                                record['reasons'] = ['native_domain_or_adapter_failure']
                                record['execution_error'] = str(error)
                            continue
                        mechanisms = {target['mechanism']} if target.get('mechanism') else set()
                        mechanisms.update(
                            v['details']['mechanism']
                            for v in selected_proofs.values()
                            if v.get('details', {}).get('mechanism')
                        )
                        mechanisms = mechanisms or {'other'}
                        record['status'], record['reasons'] = 'executed', []
                        record.pop('execution_error', None)
                        points.append(
                            {
                                **record,
                                'observed': observed.value,
                                'predicted': predicted,
                                'pressure_kpa': pressure
                                if target['kind'] == 'inverse'
                                else target.get('pressure_kpa'),
                                'theta_coordinate': selected.get('theta')
                                if target['kind'] == 'inverse'
                                else None,
                                'scenario': {
                                    0: 'master',
                                    3: 'podzolic_view',
                                    4: 'chernozem_view',
                                    6: 'thesis',
                                }[row['sheet_index']],
                                'source_specification': function['specification'],
                                'source_status': function['status'],
                                'prediction_transform': {
                                    'formula': 'water_volume_per_dry_mass * water_density',
                                    'water_density_g_cm3': context['old_registration'][
                                        'water_density_g_cm3'
                                    ],
                                }
                                if target.get('unit') == 'cubic_centimeter_per_gram'
                                else None,
                                'predictor_evidence': selected_proofs,
                                'target_evidence': observed.document(),
                                'mechanisms': sorted(mechanisms),
                                'minimum_tier': 'E5',
                                'recommendation_eligible': False,
                            }
                        )
    return decisions, points


def render(summary: dict) -> str:
    lines = [
        '# E5 workbook semantic exploratory benchmark',
        '',
        'Coverage-first descriptive exploration; no strict or harmonized recommendation.',
        '',
        '| Tier | Executable PTF |',
        '|---|---:|',
    ]
    lines.extend(
        f'| {r["evidence_ceiling"]} | {r["functions_executable"]} |'
        for r in summary['assumption_ladder']
    )
    lines += [
        '',
        f'Catalog: {summary["catalog_count"]}; functions with fresh E5 executions: {summary["e5_executable"]}; '
        f'new beyond E4: {len(summary["new_functions"])}.',
        '',
        f'Executions: {summary["executions"]}; distinct target cells used: {summary["target_cells_used"]}; '
        f'distinct predictor cells used: {summary["predictor_cells_used"]}.',
        f'Original WRC observations consumed as targets: {summary["consumption"]["wrc_target_observations"]}; '
        f'as predictors: {summary["consumption"]["wrc_predictor_observations"]}; '
        f'unique samples scored: {summary["consumption"]["samples"]}.',
        '',
        '| Mechanism | New PTF (overlapping) |',
        '|---|---:|',
    ]
    lines.extend(f'| {k} | {v} |' for k, v in summary['attribution'].items())
    lines += [
        '',
        'Attribution counts functions with a successful execution using the mechanism; groups overlap.',
        'The cumulative E5 ceiling inherits frozen E4 coverage. Inherited E4 class means remain historical E4 evidence;',
        'E5 executions never invent missing particle density, arbitrary mean predictors or WRC extrapolations.',
        '',
        'Ksat is normalized to m/s. Log outputs require explicit base and velocity reference.',
        'Continuous WRC outputs use every observed pressure of the same sample, alongside the registered anchors.',
        'Fixed scalar outputs retain their source-defined coordinates.',
        'Plastic limits use the AF/AG labels and an explicit percentage interpretation.',
        'Unknown laboratory methods are recorded as assumptions. Different physical quantities remain separate.',
        '',
        'Every function has a scored result or a concrete blocker in the private matrix.',
        'Private rows, values, identities, predictions and evidence remain in the local output directory.',
        '',
        '## Functions without a fresh E5 execution',
        '',
        'Historical E4 coverage remains inherited even when a function has no fresh E5 execution.',
        '',
    ]
    lines.extend(
        f'- `{f["function"]}`: {", ".join(f["reasons"])}.' for f in summary['blocked_functions']
    )
    lines += ['', '## Uncertain physical definitions', '']
    for decision in summary['needs_user_decision']:
        lines.append(
            f'- `{decision["function"]}`: {decision["workbook_field"]}; '
            + '; '.join(
                f'{i["interpretation"]}: potential +{i["expected_coverage_gain"]} PTF'
                for i in decision['interpretations']
            )
            + '.'
        )
    lines += [
        '',
        '## Per-pool descriptive errors',
        '',
        '| Function | Target | Scenario | Samples | Points | Unit | RMSE |',
        '|---|---|---|---:|---:|---|---:|',
    ]
    lines.extend(
        f'| `{m["function"]}` | {m["target"]} | {m["scenario"]} | {m["samples"]} | {m["points"]} | {m["unit"]} | {m["rmse"]:.6g} |'
        for m in summary['metrics']
    )
    return '\n'.join(lines) + '\n'


def pool_metrics(points: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for p in points:
        target = p['target']
        kind = target['kind']
        unit = (
            'meter_per_second'
            if kind == 'ksat'
            else 'mass_percent'
            if kind == 'plasticity'
            else 'kilopascal'
            if kind == 'inverse'
            else 'gram_per_cubic_centimeter'
            if kind == 'direct_property'
            else 'mass_fraction'
            if target['quantity'] == 'gravimetric_water_content'
            else 'volume_fraction'
        )
        key = (p['function'], target.get('field', target.get('adapter')), p['scenario'], unit, kind)
        groups[key].append(p)
    result = []
    for (function, field, scenario, unit, kind), selected in sorted(groups.items()):
        samples = defaultdict(list)
        for p in selected:
            samples[p['source_id'], p['sample_id']].append(p['predicted'] - p['observed'])
        count = len(samples)
        result.append(
            {
                'function': function,
                'target': field,
                'target_kind': kind,
                'scenario': scenario,
                'unit': unit,
                'samples': count,
                'points': len(selected),
                'rmse': math.sqrt(
                    math.fsum(
                        math.fsum(e * e for e in errors) / len(errors)
                        for errors in samples.values()
                    )
                    / count
                ),
                'mae': math.fsum(
                    math.fsum(abs(e) for e in errors) / len(errors) for errors in samples.values()
                )
                / count,
                'bias': math.fsum(math.fsum(errors) / len(errors) for errors in samples.values())
                / count,
            }
        )
    return result


def consumption(points: list[dict]) -> dict:
    def observation_ids(value: object) -> set[str]:
        if isinstance(value, dict):
            found = {value['observation_id']} if 'observation_id' in value else set()
            return found | {
                identity for item in value.values() for identity in observation_ids(item)
            }
        if isinstance(value, (list, tuple)):
            return {identity for item in value for identity in observation_ids(item)}
        return set()

    target_ids = {identity for p in points for identity in observation_ids(p['target_evidence'])}
    predictor_ids = {
        identity for p in points for identity in observation_ids(p['predictor_evidence'])
    }
    return {
        'wrc_target_observations': len(target_ids),
        'wrc_predictor_observations': len(predictor_ids),
        'wrc_observations_union': len(target_ids | predictor_ids),
        'samples': len({(p['source_id'], p['sample_id']) for p in points}),
        'function_sample_pairs': len(
            {(p['function'], p['source_id'], p['sample_id']) for p in points}
        ),
    }


def physical_decisions(context: dict, rows: list[dict]) -> list[dict]:
    """Quantify alternate water bases without admitting either into scored E5 results."""
    result = []
    for function in context['catalog']:
        if not contracts.unknown_water_basis(function):
            continue
        interpretations = []
        for quantity, unit in (
            ('gravimetric_water_content', 'mass_percent'),
            ('volumetric_water_content', 'volume_percent'),
        ):
            hypothetical = {
                **function,
                'documentation': {**function.get('documentation', {}), 'warnings': []},
                'fields': [{**f, 'quantity': quantity, 'unit': unit} for f in function['fields']],
            }
            _, points = evaluate({**context, 'catalog': [hypothetical]}, rows)
            interpretations.append(
                {
                    'interpretation': quantity,
                    'expected_coverage_gain': int(bool(points)),
                    'feasible_executions': len(points),
                }
            )
        result.append(
            {
                'function': function['id'],
                'workbook_field': 'WRC water-content basis versus unspecified source output percentage',
                'interpretations': interpretations,
                'status': 'needs_user_decision',
                'gains_overlap_between_interpretations': True,
            }
        )
    return result


def run(args: argparse.Namespace) -> dict:
    frozen = {
        str(p): campaign.prior.verify_artifacts(p)
        for p in (args.baseline, args.benchmark, args.harmonized)
    }
    if any(
        args.output.resolve().is_relative_to(p.resolve())
        for p in (args.baseline, args.benchmark, args.harmonized)
    ):
        message = 'E5 output must be separate from frozen inputs.'
        raise ValueError(message)
    context = campaign.prepare(args.benchmark, args.harmonized, args.workbook)
    baseline = json.loads((args.baseline / 'summary.json').read_bytes())
    book = load_workbook(args.workbook, data_only=True, keep_links=False)
    formulas = load_workbook(args.workbook, data_only=False, keep_links=False)
    try:
        rows = workbook_rows(book, formulas)
        headers = {
            s.title: {cell.coordinate: str(cell.value) for cell in s[1] if cell.value is not None}
            for s in (book.worksheets[0], book.worksheets[6])
        }
    finally:
        book.close()
        formulas.close()
    records, points = evaluate(context, rows)
    executed = {p['function'] for p in points}
    inherited = set(baseline['stages'][-1]['functions'])
    new = executed - inherited
    by_function = defaultdict(list)
    for r in records:
        by_function[r['function']].append(r)
    if set(by_function) != {f['id'] for f in context['catalog']}:
        message = 'Every catalog function must be assessed.'
        raise ValueError(message)
    blockers = [
        {'function': fid, 'reasons': sorted({reason for r in decisions for reason in r['reasons']})}
        for fid, decisions in sorted(by_function.items())
        if fid not in executed
    ]
    target_cells = {s for p in points for s in p['target_evidence']['sources']}
    predictor_cells = {
        s for p in points for v in p['predictor_evidence'].values() for s in v['sources']
    }
    summary = {
        'protocol': PROTOCOL,
        'catalog_count': len(context['catalog']),
        'e5_executable': len(executed),
        'functions': sorted(executed),
        'new_functions': sorted(new),
        'executions': len(points),
        'consumption': consumption(points),
        'cumulative_executable': len(executed | inherited),
        'target_cells_used': len(target_cells),
        'predictor_cells_used': len(predictor_cells),
        'assumption_ladder': [
            *baseline['assumption_ladder'],
            {'evidence_ceiling': 'E5', 'functions_executable': len(executed | inherited)},
        ],
        'attribution': {
            m: len({p['function'] for p in points if p['function'] in new and m in p['mechanisms']})
            for m in MECHANISMS
        },
        'target_usage': dict(Counter(p['target']['kind'] for p in points)),
        'target_cells_by_kind': {
            kind: len(
                {
                    cell
                    for p in points
                    if p['target']['kind'] == kind
                    for cell in p['target_evidence']['sources']
                }
            )
            for kind in sorted({p['target']['kind'] for p in points})
        },
        'conductivity_catalog': sorted(
            f['id']
            for f in context['catalog']
            if any(t['kind'] == 'ksat' for t in target_list(f, context))
        ),
        'blocked_functions': blockers,
        'needs_user_decision': physical_decisions(context, rows),
        'metrics': pool_metrics(points),
        'predictor_usage': dict(Counter(name for p in points for name in p['predictor_evidence'])),
        'candidate_property_counts': dict(Counter(name for r in rows for name in r['properties'])),
        'recommendation': None,
        'input_sha256': {'workbook': campaign.prior._hash(args.workbook), **frozen},
        'specification_sha256': {
            p.relative_to(context['root']).as_posix(): campaign.prior._hash(p)
            for p in sorted((context['root'] / 'specs').rglob('*.yaml'))
        },
        'manifest_sha256': campaign.prior._hash(
            context['root'] / 'targets/ptfkit-py/src/ptfkit/_manifest.json'
        ),
        'native_extension_sha256': campaign.prior._hash(Path(campaign.native_extension.__file__)),
        'implementation_sha256': {
            p.name: campaign.prior._hash(p) for p in sorted(Path(__file__).parent.glob('*.py'))
        },
        'frozen_artifacts_unchanged': True,
    }
    if frozen != {
        str(p): campaign.prior.verify_artifacts(p)
        for p in (args.baseline, args.benchmark, args.harmonized)
    }:
        message = 'Frozen artifacts changed during E5 execution.'
        raise ValueError(message)
    args.output.mkdir(parents=True, exist_ok=False)
    for name, value in (
        ('summary', summary),
        ('matrix', records),
        ('predictions', points),
        (
            'workbook-properties',
            [
                {
                    'source_id': r['source_id'],
                    'sample_id': r['sample_id'],
                    'workbook_row': f'{r["worksheet"]}!{r["row"]}',
                    'properties': {name: p.document() for name, p in r['properties'].items()},
                }
                for r in rows
                if r['properties']
            ],
        ),
        (
            'policy',
            {
                'protocol': PROTOCOL,
                'aliases': ALIASES,
                'headers': headers,
                'hard_rules': [
                    'same physical quantity',
                    'same sample/horizon',
                    'no invented values',
                ],
            },
        ),
    ):
        (args.output / (name + '.json')).write_bytes(canonical_json(value))
    (args.output / 'report.md').write_text(render(summary), encoding='utf-8')
    (args.output / 'artifact-hashes.json').write_bytes(
        canonical_json({p.name: campaign.prior._hash(p) for p in sorted(args.output.iterdir())})
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ('baseline', 'benchmark', 'harmonized', 'workbook', 'output'):
        parser.add_argument('--' + name, type=Path, required=True)
    summary = run(parser.parse_args())
    print(  # noqa: T201 - Command-line result.
        json.dumps(
            {
                k: summary[k]
                for k in (
                    'catalog_count',
                    'e5_executable',
                    'cumulative_executable',
                    'executions',
                    'attribution',
                )
            },
            indent=2,
        )
    )


if __name__ == '__main__':
    main()
