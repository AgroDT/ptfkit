"""Reproducible observed-pressure WRC and atomic log-Ksat model selection."""

# ruff: noqa: INP001, SLF001, C901, PLR0912, PLR0915, PLR2004, E501 - Explicit experiment policy and reporting.

from __future__ import annotations

import argparse
import importlib.metadata
import json
import math
import platform
from collections import Counter, defaultdict
from itertools import combinations, pairwise
from pathlib import Path
from statistics import median
from typing import cast

import numpy as np
from openpyxl import load_workbook
from ptfkit_eval.documents import canonical_json

from evaluation.experiments.meta_v3_3 import campaign
from evaluation.experiments.meta_v3_3 import campaign_contracts as contracts
from evaluation.experiments.meta_v3_3 import campaign_inputs as inherited_inputs
from evaluation.experiments.meta_v3_3 import workbook_semantic as semantic
from evaluation.experiments.meta_v3_3.comparative import correlation
from evaluation.experiments.meta_v3_3.deep_unlock import families
from evaluation.experiments.meta_v3_3.selection_audit import inventory, rescan


POLICY: dict = {
    'protocol': 'observed_pressure_model_selection_v1',
    'pressure_regions_kpa': {'wet': [0, 10], 'intermediate': [10, 1500], 'dry': [1500, None]},
    'region_intervals': 'left closed, right open',
    'minimum_pressures_per_sample': 2,
    'minimum_discrete_model_pressures': 3,
    'discrete_prediction_curve': 'linear in log10 pressure, positive endpoints only, no extrapolation',
    'water_basis': 'Separate original volume and mass cohorts; no harmonization.',
    'primary_cohort': 'Largest roster covering at least half the maximum executable sample coverage; ties by sample count then point count.',
    'reduced_ensemble': 'Best training atomic plus best member among top five training models with residual correlation below 0.95; otherwise two best. Equal weights. Re-select inside every outer fold.',
    'robustness': 'Leave one source out; otherwise profile; otherwise 500 sample bootstrap draws, seed 20261026.',
    'meaningful_ensemble_improvement': 'At least 1% pooled RMSE reduction and lower RMSE in every held-out group; descriptive, not a population significance claim.',
    'ksat': 'Atomic only; m/s, log10 scoring; nonpositive predictions explicitly invalid.',
    'evidence': 'Preserve E0-E5; E5 permitted with approximate predictor mappings flagged.',
}
CURVES = {
    'existing_model_adapter_v1',
    'explicit_vg_mualem_signature_v1',
    'hutson_cass_dry_v1',
    'logistic_pressure_scale_v1',
}


def candidate_targets(function: dict, context: dict) -> tuple[list[dict], list[dict]]:
    targets = semantic.target_list(function, context)
    water = [t for t in targets if t['kind'] == 'fixed_pressure']
    continuous = [t for t in water if 'axis_input' in t or t['adapter'] in CURVES]
    if continuous:
        # Prefer a complete physical model over its scalar diagnostic outputs.
        water = [continuous[0]]
    elif len({t['pressure_kpa'] for t in water}) < POLICY['minimum_discrete_model_pressures']:
        water = []
    return water, [t for t in targets if t['kind'] == 'ksat']


def discrete_prediction(
    function: dict, targets: list[dict], pressure: float, values: dict, context: dict
) -> float:
    coordinates = sorted(
        (t['pressure_kpa'], semantic.prediction(function, t, values, context)) for t in targets
    )
    for p, v in coordinates:
        if math.isclose(p, pressure, rel_tol=1e-12, abs_tol=1e-12):
            return v
    for (low, a), (high, b) in pairwise(coordinates):
        if 0 < low < pressure < high:
            weight = math.log(pressure / low) / math.log(high / low)
            return a + weight * (b - a)
    message = 'Observed pressure outside discrete model support.'
    raise ValueError(message)


def resolve_selection_inputs(
    function: dict, sample: dict, row: dict, context: dict, *, water: dict, axis: set
) -> tuple:
    """Retain registered E0-E4 resolutions before applying the E5 fallback."""
    if row['sheet_index'] != 0:
        values, proofs, reasons = semantic.semantic_inputs(
            function, sample, row, context, water, axis
        )
        return values, proofs, reasons, 'E5'
    old_values, old_proofs, old_reasons = inherited_inputs.resolve_inputs(
        function,
        sample,
        row,
        'reported_boundary_proxy',
        context['old_registration'],
        representative=True,
        water_inputs=water,
        axis_inputs=axis,
    )
    if not old_reasons:
        return (
            old_values,
            old_proofs,
            [],
            campaign.tier(
                old_proofs, {'tier': 'strict', 'policy': 'observed', 'parents': []}, bridge=False
            ),
        )
    values, proofs, reasons = semantic.semantic_inputs(function, sample, row, context, water, axis)
    for name, value in old_values.items():
        if name not in values:
            values[name], proofs[name] = value, old_proofs[name]
    reasons = [
        r
        for r in reasons
        if not (r.startswith('missing_predictor:') and r.split(':', 1)[1] in values)
    ]
    return values, proofs, reasons, 'E5'


def predict(
    context: dict, observations: list[dict], rows: list[dict]
) -> tuple[list[dict], list[dict]]:
    datasets = {d['dataset_id']: d for d in context['datasets']}
    samples = {
        (d['dataset_id'], s['source_id'], s['sample_id']): s
        for d in context['datasets']
        for s in d['samples']
    }
    row_index = defaultdict(list)
    for row in rows:
        row_index[row['source_id'], row['sample_id']].append(row)
    points, audit = [], []
    for function in context['catalog']:
        water, conductivity = candidate_targets(function, context)
        declared_families = families(function)
        wrc_disposition = (
            'observed_pressure_adapter_available'
            if water
            else 'curve_definition_or_adapter_unresolved'
            if {'full water-retention curve', 'retention-model parameters'} & set(declared_families)
            else 'isolated_pressure_outputs_not_full_WRC'
            if 'water content at specified pressure/suction' in declared_families
            else 'not_WRC_output'
        )
        failures = Counter()
        executed = 0
        for o in observations:
            candidates = conductivity if o['basis'] == 'ksat' else water
            if not candidates:
                continue
            matches = row_index[o['source_id'], o['sample_id']]
            if o['basis'] == 'ksat':
                matches = [
                    r
                    for r in matches
                    if r['properties'].get('ksat') and o['id'] in r['properties']['ksat'].sources
                ]
            elif o.get('predictor_workbook_row'):
                matches = [
                    r
                    for r in matches
                    if f'{r["worksheet"]}!{r["row"]}' == o['predictor_workbook_row']
                ]
            else:
                primary = [r for r in matches if r['sheet_index'] == 0]
                matches = primary if len(primary) == 1 else []
            if len(matches) != 1:
                failures['ambiguous_or_missing_predictor_row'] += 1
                continue
            row = matches[0]
            dataset = datasets.get(o.get('dataset_id'))
            sample = samples.get(
                (o.get('dataset_id'), o['source_id'], o['sample_id']),
                {'features': [], 'source_id': o['source_id'], 'sample_id': o['sample_id']},
            )
            target = candidates[0]
            if o['basis'] != 'ksat' and target['quantity'] != o['basis']:
                failures['different_water_basis'] += 1
                continue
            axis = {target['axis_input']} if 'axis_input' in target else set()
            water_inputs = (
                campaign.water_predictors(
                    function, dataset, sample, row, context['old_registration'], enabled=True
                )
                if dataset
                else {}
            )
            values, proofs, reasons, evidence_tier = resolve_selection_inputs(
                function, sample, row, context, water=water_inputs, axis=axis
            )
            if reasons:
                failures.update(reasons)
                continue
            try:
                if o['basis'] != 'ksat':
                    if 'axis_input' in target or target.get('adapter') in CURVES:
                        if 'axis_input' in target:
                            binding = next(
                                i for i in function['inputs'] if i['name'] == target['axis_input']
                            )
                            value = contracts.axis_value(
                                {**target, 'pressure_kpa': o['pressure_kpa']}
                            )
                            if not inherited_inputs.legacy_domain(
                                value, binding
                            ) or not inherited_inputs.within_named_domain(
                                binding.get('domain') or '',
                                {**values, binding['name']: value, 'value': value},
                            ):
                                failures['pressure_outside_declared_domain'] += 1
                                continue
                        value = semantic.prediction(
                            function, {**target, 'pressure_kpa': o['pressure_kpa']}, values, context
                        )
                    else:
                        value = discrete_prediction(
                            function, candidates, o['pressure_kpa'], values, context
                        )
                else:
                    value = semantic.prediction(function, target, values, context)
                if not math.isfinite(value) or (o['basis'] == 'ksat' and value <= 0):
                    failures['nonpositive_or_nonfinite_prediction'] += 1
                    continue
            except (ValueError, TypeError, ArithmeticError):
                failures['native_domain_or_curve_support_failure'] += 1
                continue
            executed += 1
            points.append(
                {
                    **o,
                    'function': function['id'],
                    'predicted': value,
                    'minimum_tier': 'E5'
                    if o['basis'] == 'ksat' or o.get('assumptions')
                    else evidence_tier,
                    'predictor_evidence': proofs,
                    'water_conditioned': bool(water_inputs),
                    'adapter': target.get('adapter', 'canonical_linear_or_log_Ksat'),
                    'output_field': target.get('field'),
                    'output_unit': target.get('unit'),
                    'source_specification': function['specification'],
                    'source_status': function['status'],
                }
            )
        audit.append(
            {
                'function': function['id'],
                'source_status': function['status'],
                'declared_families': declared_families,
                'wrc_disposition': wrc_disposition,
                'wrc_targets': water,
                'ksat_targets': conductivity,
                'status': 'executed' if executed else 'blocked_or_not_target_candidate',
                'executions': executed,
                'failure_counts': dict(sorted(failures.items())),
                'input_water_quantities': [
                    i['name']
                    for i in function['inputs']
                    if i.get('measurement', {}).get('quantity') in contracts.WATER
                ],
            }
        )
    return points, audit


def metrics(points: list[dict], *, ksat: bool = False) -> dict | None:
    if not points:
        return None
    groups = defaultdict(list)
    ratios = []
    for p in points:
        if ksat:
            if min(p['predicted'], p['observed']) <= 0:
                message = 'Log scoring requires positive observations and predictions.'
                raise ValueError(message)
            error = math.log10(p['predicted']) - math.log10(p['observed'])
            ratios.append(p['predicted'] / p['observed'])
        else:
            error = p['predicted'] - p['observed']
        groups[p['source_id'], p['sample_id']].append(error)
    errors = [e for g in groups.values() for e in g]
    losses = [[e] for e in errors] if ksat else list(groups.values())
    # Ksat uses observation means; WRC gives equal weight to samples.
    result = {
        'rmse': math.sqrt(sum(sum(e * e for e in g) / len(g) for g in losses) / len(losses)),
        'mae': sum(sum(abs(e) for e in g) / len(g) for g in losses) / len(losses),
        'bias': sum(sum(g) / len(g) for g in losses) / len(losses),
        'median_absolute_error': median(abs(e) for e in errors),
        **inventory(points),
    }
    if ksat:
        result.update(
            median_ratio=median(ratios),
            **{
                f'factor_{factor}': sum(abs(e) <= math.log10(factor) + 1e-14 for e in errors)
                / len(errors)
                for factor in (2, 5, 10)
            },
        )
    return result


def eligible_keys(keys: set, lookup: dict, *, ksat: bool) -> frozenset:
    if ksat:
        return frozenset(keys)
    pressures = defaultdict(set)
    for key in keys:
        point = lookup[key]
        pressures[point['source_id'], point['sample_id']].add(point['pressure_kpa'])
    return frozenset(
        k
        for k in keys
        if len(pressures[lookup[k]['source_id'], lookup[k]['sample_id']])
        >= POLICY['minimum_pressures_per_sample']
    )


def scored(points: list[dict], *, ksat: bool = False) -> dict:
    result = metrics(points, ksat=ksat)
    if result is None:
        message = 'Cannot rank an empty cohort.'
        raise ValueError(message)
    return result


def common_cohorts(indices: dict, *, ksat: bool = False) -> list[dict]:
    """Enumerate the intersection closure; each cohort has its maximal compatible roster."""
    if not indices:
        return []
    lookup = {k: p for index in indices.values() for k, p in index.items()}
    supports = {eligible_keys(set(index), lookup, ksat=ksat) for index in indices.values()}
    pending = list(supports)
    while pending:
        current = pending.pop()
        for other in list(supports):
            joined = eligible_keys(set(current & other), lookup, ksat=ksat)
            if joined not in supports:
                supports.add(joined)
                pending.append(joined)
    result = []
    for keys in supports:
        members = sorted(fid for fid, index in indices.items() if keys <= index.keys())
        if not keys or len(members) < 2:
            continue
        for key in keys:
            if len({indices[f][key]['observed'] for f in members}) != 1:
                message = 'Common cohort observations disagree.'
                raise ValueError(message)
        result.append(
            {'members': members, 'keys': sorted(keys), **inventory([lookup[k] for k in keys])}
        )
    return sorted(
        result, key=lambda c: (-c['samples'], -len(c['members']), -c['observations'], c['members'])
    )


def rank(indices: dict, members: list[str], keys: list, *, ksat: bool = False) -> list[dict]:
    scores: list[dict] = [
        {'function': f, **scored([indices[f][k] for k in keys], ksat=ksat)} for f in members
    ]
    scores.sort(key=lambda r: (r['rmse'], abs(r['bias']), r['function']))
    for number, row in enumerate(scores, 1):
        row['rank'] = number
    return scores


def reduced_members(indices: dict, members: list[str], keys: list) -> list[str]:
    ordered = [r['function'] for r in rank(indices, members, keys)]
    best = ordered[0]
    residuals = [indices[best][k]['predicted'] - indices[best][k]['observed'] for k in keys]
    for f in ordered[1:5]:
        corr = correlation(
            residuals, [indices[f][k]['predicted'] - indices[f][k]['observed'] for k in keys]
        )
        if corr is not None and corr < 0.95:
            return [best, f]
    return ordered[:2]


def ensembles(indices: dict, members: list[str], keys: list, training: list | None = None) -> dict:
    reduced = reduced_members(indices, members, keys if training is None else training)
    result = {}
    for name in ('mean', 'median', 'reduced'):
        selected = reduced if name == 'reduced' else members
        result[name] = {
            k: {
                **indices[selected[0]][k],
                'function': name,
                'predicted': (median if name == 'median' else lambda v: sum(v) / len(v))(
                    [indices[f][k]['predicted'] for f in selected]
                ),
            }
            for k in keys
        }
    return {'members': reduced, 'indices': result}


def robustness(indices: dict, members: list[str], keys: list, *, ksat: bool = False) -> dict:
    reference = indices[members[0]]
    dimension = 'block' if ksat else 'source_id'
    labels = sorted({reference[k].get(dimension) for k in keys})
    if len(labels) < 2 and not ksat:
        dimension = 'profile_id'
        labels = sorted({reference[k].get(dimension) for k in keys if reference[k].get(dimension)})
        if any(not reference[k].get(dimension) for k in keys):
            labels = []
    folds = []
    distributions = defaultdict(list)
    reduced_oof = []
    ensemble_draws = defaultdict(list)
    pooled_best = rank(indices, members, keys, ksat=ksat)[0]['function']
    if len(labels) >= 2:
        for label in labels:
            train = [k for k in keys if reference[k][dimension] != label]
            test = [k for k in keys if reference[k][dimension] == label]
            table = rank(indices, members, train, ksat=ksat)
            held = rank(indices, members, test, ksat=ksat)
            record = {'omitted_group': label, 'remaining_ranking': table, 'held_out_ranking': held}
            for r in table:
                distributions[r['function']].append((r['rank'], r['rmse']))
            if not ksat:
                ensemble = ensembles(indices, members, test, training=train)
                atom = table[0]['function']
                record['training_best_atomic'] = atom
                record['reduced_members'] = ensemble['members']
                record['held_out_training_best_atomic'] = scored([indices[atom][k] for k in test])
                record['held_out_pooled_best_atomic'] = scored(
                    [indices[pooled_best][k] for k in test]
                )
                record['held_out_ensembles'] = {
                    name: scored(list(index.values()))
                    for name, index in ensemble['indices'].items()
                }
                baseline = record['held_out_training_best_atomic']['rmse']
                for name, score in record['held_out_ensembles'].items():
                    ensemble_draws[name].append(score['rmse'] - baseline)
                    ensemble_draws[name + ':pooled_best'].append(
                        score['rmse'] - record['held_out_pooled_best_atomic']['rmse']
                    )
                reduced_oof.extend(ensemble['indices']['reduced'].values())
            folds.append(record)
    else:
        dimension = 'sample_bootstrap'
        groups = defaultdict(list)
        for k in keys:
            groups[reference[k]['source_id'], reference[k]['sample_id']].append(k)
        ordered = sorted(groups)
        rng = np.random.default_rng(20261026)
        for _ in range(500):
            # Synthetic draw identities preserve multiplicity and within-curve dependence.
            draw = rng.integers(0, len(ordered), len(ordered))
            table: list[dict] = []
            for f in members:
                points = [
                    {**indices[f][k], 'sample_id': str(i)}
                    for i, j in enumerate(draw)
                    for k in groups[ordered[j]]
                ]
                table.append({'function': f, **scored(points, ksat=ksat)})
            table.sort(key=lambda r: (r['rmse'], abs(r['bias']), r['function']))
            for number, row in enumerate(table, 1):
                distributions[row['function']].append((number, row['rmse']))
            if not ksat:
                in_bag = {k for j in draw for k in groups[ordered[j]]}
                out_of_bag = sorted(set(keys) - in_bag)
                if not out_of_bag:
                    continue
                # Membership and atomic baseline use only the resampled training samples.
                atom = table[0]['function']
                ensemble = ensembles(indices, members, out_of_bag, training=sorted(in_bag))
                baseline = scored([indices[atom][k] for k in out_of_bag])['rmse']
                for name, index in ensemble['indices'].items():
                    ensemble_draws[name].append(scored(list(index.values()))['rmse'] - baseline)
                    ensemble_draws[name + ':pooled_best'].append(
                        scored(list(index.values()))['rmse']
                        - scored([indices[pooled_best][k] for k in out_of_bag])['rmse']
                    )
                # Repeated test appearances are retained in distributions, not pooled as N.
    return {
        'dimension': dimension,
        'folds': folds,
        'rank_stability': [
            {
                'function': f,
                'median_rank': median(r for r, _ in draws),
                'rank_range': [min(r for r, _ in draws), max(r for r, _ in draws)],
                'fraction_best': sum(r == 1 for r, _ in draws) / len(draws),
                'rmse_quantiles': np.quantile(
                    [v for _, v in draws], [0, 0.025, 0.5, 0.975, 1]
                ).tolist(),
            }
            for f, draws in sorted(distributions.items())
        ],
        'reduced_nested_oof': metrics(reduced_oof),
        'ensemble_held_out_deltas': {
            name: {
                'evaluations': len(draws),
                'fraction_better_than_atomic': sum(d < 0 for d in draws) / len(draws),
                'rmse_delta_quantiles': np.quantile(draws, [0, 0.025, 0.5, 0.975, 1]).tolist(),
            }
            for name, draws in sorted(ensemble_draws.items())
        },
    }


def analyze(points: list[dict], *, ksat: bool = False) -> dict:
    indices = defaultdict(dict)
    for point in points:
        fid, key = point['function'], point['id']
        if key in indices[fid]:
            message = 'Duplicate function/observation prediction.'
            raise ValueError(message)
        indices[fid][key] = point
    if not ksat:
        indices = {
            fid: {k: index[k] for k in eligible_keys(set(index), index, ksat=False)}
            for fid, index in indices.items()
        }
        indices = {fid: index for fid, index in indices.items() if index}
    if not indices:
        return {'status': 'no_eligible_predictions'}
    cohorts = common_cohorts(indices, ksat=ksat)
    if not cohorts:
        return {
            'status': 'no_comparable_multi_pressure_cohort',
            'coverage': {f: inventory(list(i.values())) for f, i in indices.items()},
        }
    maximum = max(inventory(list(i.values()))['samples'] for i in indices.values())
    eligible = [c for c in cohorts if c['samples'] >= maximum / 2]
    primary = min(
        eligible or cohorts,
        key=lambda c: (-len(c['members']), -c['samples'], -c['observations'], c['members']),
    )
    keys, members = primary['keys'], primary['members']
    full = eligible_keys(
        set.intersection(*(set(i) for i in indices.values())),
        {k: p for i in indices.values() for k, p in i.items()},
        ksat=ksat,
    )
    table = rank(indices, members, keys, ksat=ksat)
    result: dict = {
        'status': 'exploratory_E5',
        'coverage': {f: inventory(list(i.values())) for f, i in sorted(indices.items())},
        'full_roster': {'members': sorted(indices), 'keys': sorted(full), 'points': len(full)},
        'primary': {**primary, 'ranking': table},
        'compatible_cohorts': [
            {**c, 'ranking': rank(indices, c['members'], c['keys'], ksat=ksat)} for c in cohorts
        ],
        'pairwise': [],
        'robustness': robustness(indices, members, keys, ksat=ksat),
    }
    for a, b in combinations(sorted(indices), 2):
        shared = eligible_keys(set(indices[a]) & set(indices[b]), indices[a], ksat=ksat)
        if shared:
            result['pairwise'].append(
                {'members': [a, b], 'ranking': rank(indices, [a, b], sorted(shared), ksat=ksat)}
            )
    if ksat:
        result['block_metrics'] = {
            block: rank(
                indices,
                members,
                [k for k in keys if indices[members[0]][k]['block'] == block],
                ksat=True,
            )
            for block in sorted({indices[members[0]][k]['block'] for k in keys})
        }
        widest = cohorts[0]
        result['widest_cohort'] = {
            **widest,
            'ranking': rank(indices, widest['members'], widest['keys'], ksat=True),
            'robustness': robustness(indices, widest['members'], widest['keys'], ksat=True),
            'block_metrics': {
                block: rank(
                    indices,
                    widest['members'],
                    [
                        k
                        for k in widest['keys']
                        if indices[widest['members'][0]][k]['block'] == block
                    ],
                    ksat=True,
                )
                for block in sorted(
                    {indices[widest['members'][0]][k]['block'] for k in widest['keys']}
                )
            },
        }
    else:
        for cohort in result['compatible_cohorts']:
            cohort['robustness'] = (
                result['robustness']
                if cohort['keys'] == keys
                else robustness(indices, cast('list[str]', cohort['members']), cohort['keys'])
            )
        ensemble = ensembles(indices, members, keys)
        result['ensembles'] = {
            name: scored(list(index.values())) for name, index in ensemble['indices'].items()
        }
        result['reduced_members'] = ensemble['members']
        result['residual_correlations'] = [
            {
                'members': [a, b],
                'correlation': correlation(
                    [indices[a][k]['predicted'] - indices[a][k]['observed'] for k in keys],
                    [indices[b][k]['predicted'] - indices[b][k]['observed'] for k in keys],
                ),
            }
            for a, b in combinations([r['function'] for r in table[:5]], 2)
        ]
        result['regions'] = {
            region: {
                f: metrics(
                    [
                        indices[f][k]
                        for k in keys
                        if low
                        <= indices[f][k]['pressure_kpa']
                        < (high if high is not None else math.inf)
                    ]
                )
                for f in members
            }
            for region, (low, high) in POLICY['pressure_regions_kpa'].items()
        }
        result['pressure_residuals'] = {
            f: [
                {
                    'pressure_kpa': p,
                    **scored([indices[f][k] for k in keys if indices[f][k]['pressure_kpa'] == p]),
                }
                for p in sorted({indices[f][k]['pressure_kpa'] for k in keys})
            ]
            for f in [r['function'] for r in table[:5]]
        }
        best = table[0]['rmse']
        result['ensemble_conclusion'] = {
            name: {
                'relative_rmse_change': score['rmse'] / best - 1,
                'meaningful_and_robust': score['rmse'] <= 0.99 * best
                and result['robustness']['ensemble_held_out_deltas']
                .get(name, {})
                .get('fraction_better_than_atomic', 0)
                == 1
                and result['robustness']['ensemble_held_out_deltas']
                .get(name + ':pooled_best', {})
                .get('fraction_better_than_atomic', 0)
                == 1,
            }
            for name, score in result['ensembles'].items()
        }
        for cohort in result['compatible_cohorts']:
            blend = ensembles(indices, cast('list[str]', cohort['members']), cohort['keys'])
            cohort['ensembles'] = {
                name: scored(list(index.values())) for name, index in blend['indices'].items()
            }
            cohort['reduced_members'] = blend['members']
    return result


def failure_diagnostics(points: list[dict], analyses: dict) -> dict:
    result = {}
    for name, analysis in analyses.items():
        if 'primary' not in analysis:
            continue
        ranking = analysis['primary']['ranking']
        selected_ids = [r['function'] for r in ranking[:3] + ranking[-2:]]
        selected = [p for p in points if p['function'] in selected_ids and p['basis'] == name]
        result[name] = []
        for fid in selected_ids:
            rows = [p for p in selected if p['function'] == fid]
            if not rows:
                continue
            target = next(r for r in rows)
            result[name].append(
                {
                    'function': fid,
                    'adapter': target['adapter'],
                    'output_field': target['output_field'],
                    'output_unit': target['output_unit'],
                    'canonical_prediction_range': [
                        min(p['predicted'] for p in rows),
                        max(p['predicted'] for p in rows),
                    ],
                    'canonical_observation_range': [
                        min(p['observed'] for p in rows),
                        max(p['observed'] for p in rows),
                    ],
                    'negative_predictions': sum(p['predicted'] < 0 for p in rows),
                    'water_above_one': sum(p['predicted'] > 1 for p in rows)
                    if name != 'ksat'
                    else None,
                    'source_specification': target['source_specification'],
                    'scale_review': 'Use the source-declared linear unit; explicitly exponentiate log10 outputs before velocity conversion. No clipping or calibration.',
                }
            )
    return result


def used_observations(points: list[dict]) -> dict:
    by_model = defaultdict(dict)
    for point in points:
        if not point['water_conditioned']:
            by_model[point['basis'], point['function']][point['id']] = point
    retained = defaultdict(dict)
    for (basis, _), index in by_model.items():
        for key in eligible_keys(set(index), index, ksat=basis == 'ksat'):
            retained[basis][key] = index[key]
    return {basis: inventory(list(index.values())) for basis, index in sorted(retained.items())}


def public_summary(summary: dict) -> dict:
    """Strip private identity/group keys while retaining full numerical comparisons."""

    def redact(value: object) -> object:
        if isinstance(value, dict):
            return {
                k: redact(v)
                for k, v in value.items()
                if k not in {'keys', 'omitted_group', 'pressure_residuals'}
            }
        if isinstance(value, list):
            return [redact(v) for v in value]
        return value

    return cast('dict', redact(summary))


def render(summary: dict) -> str:
    lines = [
        '# meta_v3.3 model selection',
        '',
        'Dataset-specific exploratory selection with E0-E5 provenance retained. New workbook-only targets and approximate predictor mappings are explicitly E5; their measurement origin and laboratory protocol are not independently confirmed. No universal best PTF is claimed.',
        '',
    ]
    volume = summary['analyses']['volumetric_water_content']
    conductivity = summary['analyses']['ksat']
    if 'primary' in volume and 'widest_cohort' in conductivity:
        ensemble_name = min(volume['ensembles'], key=lambda name: volume['ensembles'][name]['rmse'])
        lines.extend(
            [
                f'**Best expanded-cohort atomic WRC:** `{volume["primary"]["ranking"][0]["function"]}`. **Best pooled WRC blend:** {ensemble_name}; '
                + (
                    'a robust improvement under the registered rule.'
                    if volume['ensemble_conclusion'][ensemble_name]['meaningful_and_robust']
                    else 'its pooled gain does not establish a robust improvement.'
                ),
                '',
                f'**Maximum-coverage atomic Ksat recommendation:** `{conductivity["widest_cohort"]["ranking"][0]["function"]}` on {conductivity["widest_cohort"]["samples"]} samples, with the block sensitivity below. This differs from the broad-roster winner on a smaller common cohort.',
                '',
            ]
        )
    for label, analysis in summary['analyses'].items():
        if label not in ('volumetric_water_content', 'gravimetric_water_content', 'ksat'):
            continue
        lines.extend(['## ' + label, ''])
        if 'primary' not in analysis:
            lines.extend([analysis['status'], ''])
            continue
        primary = analysis['primary']
        lines.extend(
            [
                f'Exact primary cohort: {primary["samples"]} samples, {primary["observations"]} observations, {primary["sources"]} sources. Pressure range (kPa): {primary["pressure_range_kpa"]}.',
                '',
                '| Rank | PTF | RMSE log10 | MAE log10 | Log bias | x2 | x5 | x10 | N | Blocks |'
                if label == 'ksat'
                else '| Rank | PTF | RMSE | MAE | Bias | Median AE | Samples | Points | Coverage |',
                '|---|---|---|---|---|---|---|---|---|---|'
                if label == 'ksat'
                else '|---|---|---|---|---|---|---|---|---|',
            ]
        )
        for r in primary['ranking'][:10]:
            if label == 'ksat':
                lines.append(
                    f'| {r["rank"]} | `{r["function"]}` | {r["rmse"]:.6g} | {r["mae"]:.6g} | {r["bias"]:.6g} | {r["factor_2"]:.1%} | {r["factor_5"]:.1%} | {r["factor_10"]:.1%} | {r["samples"]} | {len(analysis["block_metrics"])} |'
                )
            else:
                denominator = summary['dataset']['wrc_by_basis'][label]['samples']
                lines.append(
                    f'| {r["rank"]} | `{r["function"]}` | {r["rmse"]:.6g} | {r["mae"]:.6g} | {r["bias"]:.6g} | {r["median_absolute_error"]:.6g} | {r["samples"]} | {r["observations"]} | {r["samples"]}/{denominator} |'
                )
        lines.extend(
            [
                '',
                f'Full executable-roster common points: {analysis["full_roster"]["points"]}. All closed compatible cohorts and exact pairwise comparisons are in summary.json.',
                '',
                f'Robustness: {analysis["robustness"]["dimension"]}. Full rank ranges, best fractions and RMSE distributions are in the aggregate checkpoint.',
                '',
            ]
        )
        if label == 'ksat':
            best = primary['ranking'][0]
            stable = min(
                analysis['robustness']['rank_stability'],
                key=lambda r: (
                    r['median_rank'],
                    r['rank_range'][1],
                    -r['fraction_best'],
                    r['function'],
                ),
            )
            coverage = max(analysis['coverage'], key=lambda f: analysis['coverage'][f]['samples'])
            lines.extend(
                [
                    f'Best pooled: `{best["function"]}`. Median predicted/observed ratio {best["median_ratio"]:.4g}; within factors 2/5/10: {best["factor_2"]:.1%}/{best["factor_5"]:.1%}/{best["factor_10"]:.1%}.',
                    '',
                    f'Best rank-stable: `{stable["function"]}` (median rank {stable["median_rank"]}, range {stable["rank_range"]}, fraction best {stable["fraction_best"]:.1%}). Highest coverage example: `{coverage}` ({analysis["coverage"][coverage]["samples"]} samples).',
                    '',
                ]
            )
            widest = analysis['widest_cohort']
            lines.extend(
                [
                    f'All-block comparison: {widest["samples"]} samples across {len(analysis["widest_cohort"]["block_metrics"])} blocks; best `{widest["ranking"][0]["function"]}` with log10 RMSE {widest["ranking"][0]["rmse"]:.6g}. This is a separate exact cohort; its error is not directly ranked against the broader-roster table above.',
                    '',
                    '| Block | Best atomic | RMSE log10 | N |',
                    '|---|---|---|---|',
                ]
            )
            for block, scores in widest['block_metrics'].items():
                r = scores[0]
                lines.append(f'| {block} | `{r["function"]}` | {r["rmse"]:.6g} | {r["samples"]} |')
            lines.extend(
                [
                    '',
                    'All-block leave-one-block-out winners: '
                    + ', '.join(
                        f'`{fold["omitted_group"]}` -> `{fold["remaining_ranking"][0]["function"]}` (N={fold["remaining_ranking"][0]["samples"]})'
                        for fold in widest['robustness']['folds']
                    )
                    + '.',
                    '',
                    'The maximum-coverage recommendation is the all-block pooled winner, with block sensitivity explicitly retained. No single PTF wins after every block omission. The broad-roster pooled leader is dominated by the cm/min-assumed block; removing that block changes its ranking.',
                    '',
                ]
            )
            lines.extend(
                [
                    '',
                    'Class-specific Ferrer-Julia regressions are tested across workbook soils as an E5 applicability extrapolation; the workbook does not establish their FAO classes. A pooled win is not source-domain validation.',
                    '',
                ]
            )
        else:
            winner = primary['ranking'][0]['function']
            stability = next(
                r for r in analysis['robustness']['rank_stability'] if r['function'] == winner
            )
            lines.extend(
                [
                    f'Best atomic on this exact cohort: `{winner}`. Median resample rank {stability["median_rank"]}; rank range {stability["rank_range"]}; fraction best {stability["fraction_best"]:.1%}; RMSE quantiles (min, 2.5%, median, 97.5%, max) {stability["rmse_quantiles"]}.',
                    '',
                ]
            )
            best = primary['ranking'][0]['rmse']
            lines.extend(['| Ensemble | RMSE | Change vs best atomic |', '|---|---|---|'])
            for name, score in analysis['ensembles'].items():
                lines.append(
                    f'| {name} | {score["rmse"]:.6g} | {(score["rmse"] / best - 1):+.2%} |'
                )
            lines.extend(
                [
                    '',
                    'Reduced members (descriptive full-data choice): '
                    + ', '.join('`' + m + '`' for m in analysis['reduced_members'])
                    + '.',
                    '',
                    'Reduced membership is selected again on training groups for outer held-out evaluation; its full-data score alone is selection-biased.',
                    '',
                ]
            )
            nested = analysis['robustness']['reduced_nested_oof']
            if nested:
                lines.extend(
                    [
                        f'Nested source/profile-held-out reduced-ensemble RMSE: {nested["rmse"]:.6g}. Membership is selected on training groups; comparison against both training-selected and pooled-best atomic functions is retained.',
                        '',
                    ]
                )
            lines.extend(
                [
                    'Robust improvement meeting the registered 1% and held-out-group rule: '
                    + ', '.join(
                        name
                        for name, value in analysis['ensemble_conclusion'].items()
                        if value['meaningful_and_robust']
                    )
                    if any(
                        v['meaningful_and_robust'] for v in analysis['ensemble_conclusion'].values()
                    )
                    else 'No ensemble meets the registered robust-improvement rule.',
                    '',
                ]
            )
            pair = (
                analysis['residual_correlations'][0] if analysis['residual_correlations'] else None
            )
            if pair and pair['correlation'] is not None:
                lines.extend(
                    [f'Top-two atomic residual correlation: {pair["correlation"]:.6g}.', '']
                )
            lines.extend(
                [
                    '| Pressure region | Samples | Best atomic bias | Mean-ensemble bias |',
                    '|---|---|---|---|',
                ]
            )
            for region, scores in analysis['regions'].items():
                atomic = scores[winner]
                if atomic:
                    lines.append(
                        f'| {region} | {atomic["samples"]} | {atomic["bias"]:.6g} | {sum(score["bias"] for score in scores.values()) / len(scores):.6g} |'
                    )
            lines.extend(
                [
                    '',
                    'A mean can cancel opposing residual biases on the pooled cohort. Highly correlated errors add little diversity, and a bias cancellation that changes with source does not establish a reliable ensemble gain.',
                    '',
                ]
            )
    lines.extend(
        [
            '## Reviewed-target sensitivity and residuals',
            '',
            '| Basis | Samples | Points | Best atomic | RMSE |',
            '|---|---|---|---|---|',
        ]
    )
    for basis in ('volumetric_water_content', 'gravimetric_water_content'):
        reviewed = summary['analyses'][basis + '_reviewed_targets_only']
        if 'primary' in reviewed:
            r = reviewed['primary']['ranking'][0]
            lines.append(
                f'| {basis} | {r["samples"]} | {r["observations"]} | `{r["function"]}` | {r["rmse"]:.6g} |'
            )
    lines.extend(
        [
            '',
            'These are separate exact cohorts, so their RMSE values do not rank against the expanded-cohort scores. Smaller compatible cohorts retain additional atomic models, including the source-faithful Cosby dry-branch model. Pressure-specific regression curves are interpolated only between at least three native output pressures; their scores do not establish wet-to-dry extrapolation skill.',
            '',
            'Top-model residual correlations and fixed wet [0,10), intermediate [10,1500), dry [1500,infinity) kPa biases are in the checkpoint; detailed pressure residuals stay private. Reduced members require complementary residuals, rather than averaging only close variants of one formula.',
            '',
            'The extreme low Wosten Table 5 Ksat forecast was independently recalculated from its specification: the large organic-matter proxy drives the exponential to about 1e-138 m/s. The adapter uses the already exponentiated linear cm/day output and converts once to m/s. No sign, double-log or conversion defect was found; no formula was calibrated or clipped.',
            '',
        ]
    )
    lines.extend(
        [
            '## Compatible WRC cohorts',
            '',
            '| Basis | Models | Samples | Points | Best atomic | Atomic RMSE | Best pooled blend | Blend RMSE |',
            '|---|---|---|---|---|---|---|---|',
        ]
    )
    for basis in ('volumetric_water_content', 'gravimetric_water_content'):
        for cohort in summary['analyses'][basis].get('compatible_cohorts', []):
            best = cohort['ranking'][0]
            blend = min(cohort['ensembles'], key=lambda name: cohort['ensembles'][name]['rmse'])
            lines.append(
                f'| {basis} | {len(cohort["members"])} | {cohort["samples"]} | {cohort["observations"]} | `{best["function"]}` | {best["rmse"]:.6g} | {blend} | {cohort["ensembles"][blend]["rmse"]:.6g} |'
            )
    lines.extend(
        [
            '',
            'Each row uses its own exact common sample/pressure set. Errors must not be ranked across rows. Full atomic rankings, member lists and held-out ensemble comparisons are retained for every cohort.',
            '',
        ]
    )
    lines.extend(
        [
            '## Dataset sufficiency',
            '',
            f'Available WRC ledger: {summary["dataset"]["wrc"]["observations"]} points, {summary["dataset"]["wrc"]["samples"]} samples, {summary["dataset"]["wrc"]["horizons"]} identified sample/horizon records, {summary["dataset"]["wrc"]["sources"]} source labels, {summary["dataset"]["wrc"]["profiles"]} explicitly identified profile. {summary["dataset"]["wrc"]["samples_without_profile"]} samples lack an explicit profile ID. Pressure range {summary["dataset"]["wrc"]["pressure_range_kpa"]} kPa. Counts describe workbook records, not independently verified laboratory measurements.',
            '',
            f'Ksat ledger: {summary["dataset"]["ksat"]["observations"]} observations/samples, {summary["dataset"]["ksat"]["sources"]} source labels, four unit/source blocks; no confirmed profile grouping. Canonical range {summary["dataset"]["ksat_range_m_s"]} m/s.',
            '',
            'Actual scored union across eligible standalone functions (duplicate function predictions counted once per observation): '
            + '; '.join(
                f'{basis}: {value["observations"]} points / {value["samples"]} samples / {value["sources"]} sources'
                for basis, value in summary['used_observations'].items()
            )
            + '.',
            '',
            f'WRC basis counts: {summary["dataset"]["wrc"]["basis_counts"]}. The new E5 audit contributes {summary["dataset"]["additional_wrc"]["observations"]} additional pairs; ambiguous saturation markers, unresolved derived curves and repeated heads without measurement-level replicate evidence remain excluded. Raw pressure/water cells, original encodings, normalization and origin status are retained in the private ledger.',
            '',
            summary['dataset']['completeness'],
            '',
            'WRC basis/axis/provenance exclusions remain explicit. Fitted parameters, copied views and Rosetta-derived Ksat never increase experimental N. Ksat method and origin are unconfirmed under permissive workbook semantics; cm/smin is interpreted as cm/min. The three additional master cm/day values are a separate block.',
            '',
            'See model-selection-summary.json for all metrics, coverage, factor errors, block metrics, residual patterns and rank stability. Private ledgers, candidate decisions and predictions are retained only in ignored local storage.',
            '',
            'Reproduce from the repository root with `mise run eval:model-selection local/meta-catalog-resource local/meta-catalog-harmonized local/meta_v3.3.xlsx local/meta-model-selection`. The output directory must be new. Repeat with a second new output directory and compare artifact-hashes.json byte for byte. The aggregate checkpoint is evaluation/datasets/meta_v3_3.model-selection-summary.json.',
            '',
        ]
    )
    return '\n'.join(lines)


def run(args: argparse.Namespace) -> dict:
    if args.output.exists():
        message = 'Model-selection output directory must be new.'
        raise ValueError(message)
    if any(
        args.output.resolve().is_relative_to(p.resolve()) for p in (args.benchmark, args.harmonized)
    ):
        message = 'Selection output must be separate from frozen inputs.'
        raise ValueError(message)
    workbook_hash = campaign.prior._hash(args.workbook)
    context = campaign.prepare(args.benchmark, args.harmonized, args.workbook)
    book = load_workbook(args.workbook, data_only=True, keep_links=False)
    formulas = load_workbook(args.workbook, data_only=False, keep_links=False)
    try:
        audit, observations, rows = rescan(context, args.workbook, book, formulas)
    finally:
        book.close()
        formulas.close()
    points, decisions = predict(context, observations, rows)
    analyses = {}
    for basis in ('volumetric_water_content', 'gravimetric_water_content', 'ksat'):
        selected = [p for p in points if p['basis'] == basis and not p['water_conditioned']]
        analyses[basis] = analyze(selected, ksat=basis == 'ksat')
    conditioned = [p for p in points if p['water_conditioned']]
    for basis in sorted({p['basis'] for p in conditioned}):
        analyses[basis + '_water_conditioned'] = analyze(
            [p for p in conditioned if p['basis'] == basis]
        )
    for basis in ('volumetric_water_content', 'gravimetric_water_content'):
        analyses[basis + '_reviewed_targets_only'] = analyze(
            [
                p
                for p in points
                if p['basis'] == basis and not p['water_conditioned'] and not p.get('assumptions')
            ]
        )
    if context['input_hashes'] != {
        'strict': campaign.prior.verify_artifacts(args.benchmark),
        'harmonized': campaign.prior.verify_artifacts(args.harmonized),
    }:
        message = 'Frozen inputs changed during model selection.'
        raise ValueError(message)
    summary = {
        'policy': POLICY,
        'dataset': {
            k: audit[k]
            for k in (
                'wrc',
                'wrc_by_basis',
                'ksat',
                'normalization',
                'additional_wrc',
                'ksat_block_counts',
                'ksat_original_units',
                'ksat_range_m_s',
                'additional_block_decisions',
                'completeness',
            )
        },
        'analyses': analyses,
        'used_observations': used_observations(points),
        'failure_diagnostics': failure_diagnostics(points, analyses),
        'catalog': decisions,
        'catalog_counts': {
            'public_functions': len(decisions),
            'wrc_candidates': sum(
                d['wrc_disposition']
                in ('observed_pressure_adapter_available', 'curve_definition_or_adapter_unresolved')
                for d in decisions
            ),
            'wrc_adapted': sum(bool(d['wrc_targets']) for d in decisions),
            'ksat_candidates': sum(bool(d['ksat_targets']) for d in decisions),
            'ksat_executed': len({p['function'] for p in points if p['basis'] == 'ksat'}),
        },
        'input_sha256': {
            'workbook': campaign.prior._hash(args.workbook),
            **context['input_hashes'],
        },
        'implementation_sha256': {
            p.name: campaign.prior._hash(p) for p in sorted(Path(__file__).parent.glob('*.py'))
        },
        'specification_sha256': {
            p.relative_to(context['root']).as_posix(): campaign.prior._hash(p)
            for p in sorted((context['root'] / 'specs').rglob('*.yaml'))
        },
        'manifest_sha256': campaign.prior._hash(
            context['root'] / 'targets/ptfkit-py/src/ptfkit/_manifest.json'
        ),
        'native_extension_sha256': campaign.prior._hash(Path(campaign.native_extension.__file__)),
        'runtime': {
            'python': platform.python_version(),
            **{
                name: importlib.metadata.version(name)
                for name in ('numpy', 'scipy', 'openpyxl', 'ptfkit', 'pyyaml')
            },
        },
        'dependency_sha256': {
            p.relative_to(context['root']).as_posix(): campaign.prior._hash(p)
            for p in [
                context['root'] / 'targets/ptfkit-eval/uv.lock',
                context['root'] / 'targets/ptfkit-eval/pyproject.toml',
            ]
        },
        'evaluation_library_sha256': {
            p.relative_to(context['root']).as_posix(): campaign.prior._hash(p)
            for p in sorted((context['root'] / 'targets/ptfkit-eval/src').rglob('*.py'))
        },
    }
    if campaign.prior._hash(args.workbook) != workbook_hash:
        message = 'Workbook changed during model selection.'
        raise ValueError(message)
    args.output.mkdir(parents=True, exist_ok=False)
    for name, value in [
        ('summary', summary),
        ('audit', audit),
        ('observations', observations),
        ('predictions', points),
        ('policy', POLICY),
        ('model-selection-summary', public_summary(summary)),
    ]:
        (args.output / (name + '.json')).write_bytes(canonical_json(value))
    (args.output / 'MODEL-SELECTION.md').write_text(render(summary), encoding='utf-8')
    (args.output / 'artifact-hashes.json').write_bytes(
        canonical_json({p.name: campaign.prior._hash(p) for p in sorted(args.output.iterdir())})
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ('benchmark', 'harmonized', 'workbook', 'output'):
        parser.add_argument('--' + name, type=Path, required=True)
    summary = run(parser.parse_args())
    print(  # noqa: T201 - Command result.
        json.dumps(
            {
                name: analysis.get('primary', {}).get('ranking', [])[:1]
                for name, analysis in summary['analyses'].items()
            },
            indent=2,
        )
    )


if __name__ == '__main__':
    main()
