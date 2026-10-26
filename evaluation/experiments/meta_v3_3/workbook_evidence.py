"""Fail-closed Ksat and predictor evidence review, separate from frozen campaigns."""

# ruff: noqa: INP001, SLF001, E501, ISC004, PERF401 - Standalone experiment and report tables.

from __future__ import annotations

import argparse
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from ptfkit_eval.documents import canonical_json

from evaluation.experiments.meta_v3_3 import campaign, deep_unlock, exploratory
from evaluation.experiments.meta_v3_3 import campaign_contracts as contracts
from evaluation.experiments.meta_v3_3 import campaign_inputs as inputs


KSAT_QUANTITIES = {'saturated_hydraulic_conductivity', 'log_saturated_hydraulic_conductivity'}
THESIS_INDEX = 6
LENGTH_METERS = {'meter': 1, 'centimeter': 0.01, 'millimeter': 0.001}
TIME_SECONDS = {'day': 86400, 'hour': 3600, 'minute': 60, 'second': 1}


def conductivity_factor(unit: str | None) -> float | None:
    """Convert an explicitly reviewed length/time unit to m/s; never parse labels."""
    length, _, time = (unit or '').partition('_per_')
    if length not in LENGTH_METERS or time not in TIME_SECONDS:
        return None
    return LENGTH_METERS[length] / TIME_SECONDS[time]


def target_blockers(observation: dict) -> list[str]:
    """Require more than a numeric value with convertible units to authorize Ksat."""
    reasons = list(observation.get('blockers', []))
    for key, required in (
        ('identity_status', 'confirmed_identity'),
        ('quantity', 'measured_saturated_hydraulic_conductivity'),
        ('origin', 'measured'),
        ('material_status', 'confirmed'),
        ('method_status', 'confirmed'),
    ):
        if observation.get(key) != required:
            reasons.append('unconfirmed_' + key)
    if conductivity_factor(observation.get('unit')) is None:
        reasons.append('unconfirmed_conductivity_unit')
    return sorted(set(reasons))


def normalized_metadata(value: object) -> str:
    """Normalize typography only, without filling missing units or profile labels."""
    text = str(value or '').casefold().strip()
    text = re.sub(r'[\u00ad\u2010-\u2014\u2212]', '-', text)
    return re.sub(r'\s+', '', text).removesuffix('cm').removesuffix('\u0441\u043c')


def identity_decision(key: tuple, metadata: dict, candidates: list[dict]) -> dict:
    """Require unique explicit source/sample keys and reject contradictory metadata."""
    matches = [r for r in candidates if (r['source_id'], r['sample_id']) == key]
    status = 'unresolved_identity'
    conflicts = []
    if len(matches) > 1:
        status = 'ambiguous_identity'
    elif matches:
        conflicts = [
            name
            for name in ('horizon', 'depth')
            if metadata.get(name) is not None
            and matches[0].get(name) is not None
            and normalized_metadata(metadata[name]) != normalized_metadata(matches[0][name])
        ]
        status = 'ambiguous_identity' if conflicts else 'confirmed_identity'
    return {
        'identity_status': status,
        'master_rows': [r['row'] for r in matches],
        'identity_evidence': ['explicit_source_id', 'explicit_cached_sample_id'],
        'metadata_conflicts': conflicts,
        'matching_metadata': [
            name
            for name in ('horizon', 'depth')
            if len(matches) == 1
            and metadata.get(name) is not None
            and matches[0].get(name) is not None
            and name not in conflicts
        ],
    }


def numeric(value: object) -> bool:
    return isinstance(value, (float, int)) and not isinstance(value, bool) and math.isfinite(value)


def reviewed_locations(book: Workbook, locations: list[str], alias: str, review: dict) -> list[str]:
    """Remove lexical hits inside publication titles, soil names and site descriptions."""
    columns = review.get('reviewed_header_columns', {}).get(alias)
    if columns is None:
        return locations
    accepted = []
    for location in locations:
        sheet_name, address = location.rsplit('!', 1)
        index = book.sheetnames.index(sheet_name)
        expected = columns[1 if index == THESIS_INDEX else 0].split(',')
        if index in (0, 3, 4, 5, 6) and book[sheet_name][address].column_letter in expected:
            accepted.append(location)
    return accepted


def ksat_inventory(
    book: Workbook, formulas: Workbook, review: dict, canonical: set[tuple]
) -> list[dict]:
    """Account for every numeric Ksat cell, including unregistered future blocks."""
    main, thesis = book.worksheets[0], book.worksheets[6]
    masters = [
        {
            'source_id': main[f'C{r}'].value,
            'sample_id': main[f'A{r}'].value,
            'row': r,
            'horizon': main[f'I{r}'].value,
            'depth': main[f'J{r}'].value,
        }
        for r in range(2, main.max_row + 1)
    ]
    result = []
    for r in range(2, thesis.max_row + 1):
        if not numeric(thesis[f'AB{r}'].value):
            continue
        blocks = [b for b in review['ksat_blocks'] if b['first'] <= r <= b['last']]
        block = blocks[0] if len(blocks) == 1 else {}
        if block and thesis[block['unit_cell']].value != block['unit_label']:
            message = 'Reviewed Ksat unit header changed.'
            raise ValueError(message)
        key = (thesis[f'C{r}'].value, thesis[f'A{r}'].value)
        metadata = {
            name: thesis[f'{column}{r}'].value
            for name, column in (
                ('profile', 'K'),
                ('horizon', 'L'),
                ('depth', 'M'),
                ('site', 'H'),
            )
        }
        decision = identity_decision(key, metadata, masters)
        result.append(
            {
                **block,
                **decision,
                **metadata,
                'source_id': key[0],
                'sample_id': key[1],
                'cell': f'{thesis.title}!AB{r}',
                'row': r,
                'value': thesis[f'AB{r}'].value,
                'formula': formulas.worksheets[6][f'AB{r}'].data_type == 'f',
                'canonical_sample': key in canonical,
                'canonical_identity_status': 'confirmed_identity'
                if key in canonical
                else 'unresolved_identity',
                'canonical_blocker': None
                if key in canonical
                else 'no_retained_canonical_wrc_sample_with_exact_key',
                'material_status': 'unresolved',
                'method_status': 'unresolved',
                'evidence_tier': 'candidate_not_released',
                'unit_conversion_to_m_per_s': conductivity_factor(block.get('unit')),
                'blockers': block.get('blockers', ['unreviewed_ksat_block']),
                'source_wrc_method': thesis[f'E{r}'].value,
            }
        )
    for observation in result:
        observation['target_blockers'] = target_blockers(observation)
    return result


def predictor_inventory(
    book: Workbook, scan: dict, old: dict, review: dict, canonical: set[tuple]
) -> list[dict]:
    """Review every prior predictor candidate, retaining source-specific numeric coverage."""
    result = []
    for predictor in old['predictors']:
        name = predictor['predictor']
        alias = deep_unlock.CANONICAL.get(name, name)
        locations = reviewed_locations(book, predictor['candidate_locations'], alias, review)
        cells = {}
        for candidate in scan['candidates']:
            if candidate['location'] not in locations:
                continue
            sheet_name, address = candidate['location'].rsplit('!', 1)
            sheet = book[sheet_name]
            header = sheet[address]
            for row in range(header.row + 1, sheet.max_row + 1):
                value = sheet.cell(row, header.column).value
                if numeric(value):
                    cells[(sheet_name, row, header.column)] = {
                        'source_id': sheet[f'C{row}'].value,
                        'sample_id': sheet[f'A{row}'].value,
                        'cell': f'{sheet_name}!{sheet.cell(row, header.column).coordinate}',
                    }
        identities = {(c['source_id'], c['sample_id']) for c in cells.values()}
        groups = Counter(str(c['source_id']) for c in cells.values())
        note = review['predictor_decisions'].get(name) or review['predictor_decisions'].get(alias)
        if not note:
            note = (
                'Existing sample-specific exploratory evidence retained; no new contract inferred.'
                if predictor['existing_evidence_function_count']
                else 'No reviewed numeric measurement contract; names and labels do not establish evidence.'
            )
        result.append(
            {
                'predictor': name,
                'candidate_locations': locations,
                'rejected_lexical_locations': sorted(
                    set(predictor['candidate_locations']) - set(locations)
                ),
                'numeric_candidate_cells': len(cells),
                'distinct_explicit_keys': len(identities),
                'canonical_overlap_keys': len(identities & canonical),
                'source_numeric_cell_counts': dict(sorted(groups.items())),
                'requirements': predictor['requirements'],
                'functions_potentially_affected': sorted(
                    {r['function'] for r in predictor['requirements']}
                ),
                'previous_status': predictor['scientific_status'],
                'review_decision': note,
                'newly_released_measurements': 0,
                'unit_method_material_status': 'See source-specific review; no new full contract released',
            }
        )
    return result


def ksat_catalog_audit(
    context: dict,
    observations: list[dict],
    book: Workbook | None,
    formulas: Workbook | None = None,
) -> list[dict]:
    """Discover outputs by registered quantities, then assess every candidate row."""
    rows = []
    for function in context['catalog']:
        outputs = [f for f in function['fields'] if f.get('quantity') in KSAT_QUANTITIES]
        if not outputs:
            continue
        for observation in observations:
            master = observation['master_rows']
            raw = None
            if (
                book is not None
                and len(master) == 1
                and observation['identity_status'] == 'confirmed_identity'
            ):
                r = master[0]
                sheet = book.worksheets[0]
                # This row is diagnostic only. No target is released, and no native call is made.
                raw = {
                    'worksheet': sheet.title,
                    'row': r,
                    'clay_upper_um': exploratory._clay_upper(r, context['mapping']),
                    'values': {
                        sheet.cell(r, c).column_letter: sheet.cell(r, c).value for c in range(1, 23)
                    },
                }
                if formulas is not None:
                    raw['formulas'] = {
                        formulas.worksheets[0].cell(r, c).column_letter: formulas.worksheets[0]
                        .cell(r, c)
                        .value
                        for c in range(1, 23)
                        if formulas.worksheets[0].cell(r, c).data_type == 'f'
                    }
            sample = {
                'source_id': observation['source_id'],
                'sample_id': observation['sample_id'],
                'profile_id': observation['profile'],
                'features': [],
            }
            for scenario in context['registration']['psd_scenarios']:
                values, proofs, reasons = inputs.resolve_inputs(
                    function,
                    sample,
                    raw,
                    scenario,
                    context['old_registration'],
                    representative=False,
                    water_inputs={},
                )
                rows.append(
                    {
                        'function': function['id'],
                        'outputs': outputs,
                        'cell': observation['cell'],
                        'scenario': scenario,
                        'identity_status': observation['identity_status'],
                        'target_blockers': observation['target_blockers'],
                        'predictor_blockers': sorted(set(reasons)),
                        'resolved_predictors': sorted(values),
                        'predictor_evidence': proofs,
                        'source_domain_status': 'partial_predictor_bounds_only; full domain not established',
                        'native_executed': False,
                    }
                )
    return rows


def unregistered_target_decision(function: dict) -> tuple[str, str]:  # noqa: PLR0911 - Distinct scientific contracts.
    """Distinguish operational endpoint gaps, inverse targets and unrelated quantities."""
    fields = function['fields']
    quantities = {f.get('quantity') for f in fields}
    description = ' '.join(f.get('description', '') for f in fields).casefold()
    if 'osmotic_potential' in quantities:
        return (
            'no_compatible_target_measurement_identified',
            'No reviewed osmotic potential or salinity measurements; matric WRC pressure is not osmotic pressure',
        )
    if 'available_water_capacity' in quantities or (
        'minus' in description and 'wilting' in description
    ):
        return (
            'target_derivable_but_endpoint_contract_unresolved',
            'WRC differences exist, but source FC/PWP or extended-AWC endpoints are not registered; no default 33/1500-kPa substitution',
        )
    if 'cylinder-method' in description:
        return (
            'target_exists_related_but_protocol_unresolved',
            'WRC pressure exists; cylinder-method field-capacity state is not identified in workbook rows',
        )
    if 'field capacity' in description or 'field-capacity' in description:
        return (
            'target_exists_related_but_operational_definition_unresolved',
            'WRC water values exist; source operational field-capacity criterion is not bound to a measured workbook state',
        )
    if 'wilting' in description:
        return (
            'target_exists_related_but_operational_definition_unresolved',
            'WRC water values exist; source wilting protocol/pressure endpoint is not specified in resolved output contract',
        )
    if 'matric_potential' in quantities and any(i['name'] == 'theta' for i in function['inputs']):
        return (
            'target_derivable_but_inverse_contract_unregistered',
            'Observed WRC pairs could supply inverse pressure targets; no reviewed inverse adapter or joint same-sample anchor/support policy. No new pressure route introduced in this audit',
        )
    if 'campbell_b' in quantities:
        return (
            'target_derivable_but_model_fit_contract_unresolved',
            'No reviewed Campbell parameter fit/nominal air-entry target; source-fixed density and soil-group semantics do not identify workbook samples',
        )
    return (
        'target_contract_unresolved',
        'Resolved output quantities: '
        + ', '.join(sorted(str(q) for q in quantities))
        + '; no compatible registered observation or physical conversion contract',
    )


def summarize_functions(
    context: dict, old: dict, records: list[dict], points: list[dict], ks: list[dict]
) -> list[dict]:
    by_function, ks_by_function = defaultdict(list), defaultdict(list)
    for record in records:
        by_function[record['function']].append(record)
    for record in ks:
        ks_by_function[record['function']].append(record)
    executed = {p['function'] for p in points} | context['strict_ids']
    prior_by_id = {f['function']: f for f in old['functions']}
    result = []
    for function in context['catalog']:
        fid = function['id']
        previous = prior_by_id[fid]
        current = by_function[fid]
        bindings = contracts.targets(function, context['support'])
        reasons = sorted({reason for r in current for reason in r['reasons']})
        input_sets = [
            tuple(sorted(x for x in r['reasons'] if 'predictor' in x or 'domain' in x))
            for r in current
            if r.get('sample_id')
        ]
        best = sorted(set(input_sets), key=lambda x: (len(x), x))[:3]
        ks_rows = ks_by_function[fid]
        family = set(previous['families'])
        if fid in executed:
            outcome, blocker = 'executable', 'none'
        elif ks_rows:
            outcome, blocker = (
                'ambiguous_evidence',
                'Ksat candidates exist; quantity/method/origin/material not confirmed',
            )
        elif 'plasticity properties' in family:
            outcome, blocker = (
                'ambiguous_evidence',
                'AF/AG candidates exist; source table, measurement protocol and water basis unresolved',
            )
        elif bindings:
            outcome = (
                'predictors_satisfied_but_target_support_blocked'
                if () in input_sets
                else 'sample_domain_blocked'
                if any(
                    s and all(x.startswith('computational_domain:') for x in s) for s in input_sets
                )
                else 'predictor_contract_blocked'
            )
            blocker = '; '.join(reasons)
        elif family <= {'moisture equivalent', 'unsaturated conductivity', 'infiltration'}:
            outcome, blocker = (
                'no_compatible_target_measurement_identified',
                'No reviewed direct measurement; WRC alone cannot establish this property',
            )
        else:
            outcome, blocker = unregistered_target_decision(function)
        result.append(
            {
                'function': fid,
                'previous_category': previous['decision'],
                'outcome': outcome,
                'blocker': blocker,
                'remaining_reasons': reasons,
                'closest_sample_predictor_blocker_sets': best,
                'per_sample_records': len(input_sets),
                'registered_targets': bindings,
                'target_outputs': function['fields'],
                'required_predictors': function['inputs'],
                'candidate_target_locations': previous['candidate_target_locations'],
                'ksat_candidate_rows_assessed': len(ks_rows),
                'ksat_predictor_reasons': sorted(
                    {reason for r in ks_rows for reason in r['predictor_blockers']}
                ),
                'ksat_complete_predictor_rows': sum(not r['predictor_blockers'] for r in ks_rows),
            }
        )
    return result


def render(report: dict) -> str:
    lines = [
        '# ptfkit workbook evidence audit',
        '',
        f'Confirmed executable coverage: **{report["baseline_executable"]} -> {report["executable"]} / {report["catalog_count"]}**.',
        '',
        'All three Ksat blocks were reviewed separately. No Ksat target or new predictor was released. '
        'A source-level method description does not establish the origin of an individual workbook value. '
        'No compatible Ksat pool exists yet; no Ksat error metric or pooled scientific conclusion is reported.',
        '',
        'Strict predictions are recomputed and byte-identical to the frozen campaign. '
        'The full catalog was reevaluated at the final policy stage in both registered PSD scenarios. '
        'Existing approximate assumptions remain approximate; this audit does not upgrade their evidence.',
        '',
        '## Ksat blocks',
        '',
        '| Range | Numeric rows | Unit decision | Master identity | Canonical WRC overlap | Method / origin decision |',
        '| --- | ---: | --- | --- | ---: | --- |',
    ]
    for block in report['ksat_blocks']:
        lines.append(
            f'| {block["range"]} | {block["numeric_rows"]} | {block["unit"] or "unresolved: literal cm/smin"} | {block["identity_counts"]} | {block["canonical_overlap"]} | {block["reported_related_method"]}; exact measurement/model origin unresolved |'
        )
    lines += [
        '',
        'Workbook/master identity and retained canonical WRC membership are separate. '
        'Absence from the WRC cohort would not prevent a new Ksat-only cohort if its target contract were confirmed. '
        'The private identity map preserves explicit source/sample IDs, profile, site, horizon, depth, '
        'metadata conflicts, exact cell and master row. Missing profile labels are never forward-filled.',
        '',
        f'All **{report["ksat_function_count"]}** functions with registered saturated or log-saturated conductivity outputs '
        f'were selected automatically, producing {report["ksat_assessments"]} row/scenario assessments. '
        'The historical 55 count referred only to a blocked subset of the catalog. '
        'Conductivity input requirements and computable bounds are assessed even when targets fail; '
        'no native call or full application-domain claim is made for an unconfirmed target. '
        'Log outputs additionally require an explicit logarithm/reference-unit contract before scoring.',
        '',
        '## Predictor review',
        '',
        '| Predictor | Numeric candidate cells / canonical keys | Candidate locations | Review decision | Affected functions |',
        '| --- | --- | --- | --- | ---: |',
    ]
    for p in report['predictors']:
        if p['candidate_locations'] or p['predictor'] in {
            'nitrogen',
            'particle_density',
            'total_porosity',
            'ph',
        }:
            lines.append(
                f'| `{p["predictor"]}` | {p["numeric_candidate_cells"]} / {p["canonical_overlap_keys"]} | {", ".join(p["candidate_locations"])} | {p["review_decision"]} | {len(p["functions_potentially_affected"])} |'
            )
    lines += [
        '',
        'Counts are discovery counts, not approved measurements; duplicate workbook views may repeat a key. '
        'The private scan contains all 8 sheets, formulas and cached values; per-source counts and exact function '
        'input contracts are retained in the machine-readable checkpoint.',
        '',
        '## Corrections to the historical deep-unlock snapshot',
        '',
        report['corrections'],
        '',
        '## Unlock attribution',
        '',
        '| Mechanism | New functions |',
        '| --- | ---: |',
    ]
    lines.extend(f'| {key} | {value} |' for key, value in report['attribution'].items())
    lines += [
        '',
        '## Reassessment of prior categories',
        '',
        '| Previous category | Functions | Current outcomes |',
        '| --- | ---: | --- |',
    ]
    for category, counts in report['category_reassessment'].items():
        lines.append(f'| {category} | {sum(counts.values())} | {counts} |')
    lines += [
        '',
        '## Function decisions',
        '',
        'All catalog functions are listed below. The private matrix keeps per-sample support decisions. '
        'Closest-sample predictor sets are alternatives from actual sample records, never a union across samples. '
        'Target routes are not equivalent to confirmed observations for every sample.',
        '',
        '| Function | Previous category | Outcome / concrete blocker | Closest sample predictor gaps; Ksat diagnostic gaps |',
        '| --- | --- | --- | --- |',
    ]
    for f in report['functions']:
        gaps = f['closest_sample_predictor_blocker_sets'] or f['ksat_predictor_reasons']
        lines.append(
            f'| `{f["function"]}` | {f["previous_category"]} | {f["outcome"]}: {f["blocker"]} | {gaps} |'
        )
    lines += ['', '## Source evidence', '']
    for name, reference in report['review']['references'].items():
        lines.append(
            f'- [{name}]({reference["url"]}): {reference["locator"]}. {reference["decision"]}'
        )
    lines += [
        '',
        'Primary documents are retained privately with SHA256 fingerprints. Unavailable full theses and failed '
        'downloads were not treated as evidence. Completing the unresolved contracts requires an exact original '
        'table/figure-to-row crosswalk, preparation/material and measurement protocol, and origin of each value. '
        'The ambiguous unit additionally needs an authoritative correction. These are evidence limitations, '
        'not assertions that the properties do not exist.',
        '',
    ]
    return '\n'.join(lines)


def run(args: argparse.Namespace) -> dict:  # noqa: C901, PLR0915 - One immutable audit transaction.
    directories = (args.baseline, args.discovery, args.benchmark, args.harmonized)
    if any(args.output.resolve().is_relative_to(p.resolve()) for p in directories):
        message = 'Output must be separate from every frozen input.'
        raise ValueError(message)
    initial_hashes = {str(p): exploratory.verify_artifacts(p) for p in directories}
    review_path = Path(__file__).with_name('workbook-evidence-review.json')
    review = json.loads(review_path.read_bytes())
    if exploratory._hash(args.workbook) != review['workbook_sha256']:
        message = 'Workbook differs from reviewed evidence.'
        raise ValueError(message)
    old = json.loads((args.discovery / 'report.json').read_bytes())
    scan = json.loads((args.discovery / 'workbook-scan.json').read_bytes())
    context = campaign.prepare(args.benchmark, args.harmonized, args.workbook)
    context['mapping'] = yaml.safe_load(
        (context['root'] / 'evaluation/datasets/meta_v3_3.yaml').read_bytes()
    )
    canonical = {
        (s['source_id'], s['sample_id']) for d in context['datasets'] for s in d['samples']
    }
    book = load_workbook(args.workbook, data_only=True)
    formulas = load_workbook(args.workbook, data_only=False)
    observations = ksat_inventory(book, formulas, review, canonical)
    predictors = predictor_inventory(book, scan, old, review, canonical)
    ks = ksat_catalog_audit(context, observations, book, formulas)
    book.close()
    formulas.close()
    records, points = [], []
    for scenario in context['registration']['psd_scenarios']:
        r, p = campaign.evaluate(
            context,
            context['registration']['stages'][-1],
            scenario,
            context['registration']['primary_target_route'],
        )
        records.extend(r)
        points.extend(p)
    functions = summarize_functions(context, old, records, points, ks)
    before = json.loads((args.baseline / 'summary.json').read_bytes())
    before_ids = {f['function'] for f in before['functions'] if f['scored']} | context['strict_ids']
    after_ids = {p['function'] for p in points} | context['strict_ids']
    if not before_ids <= after_ids:
        message = 'Review lost frozen executable functions.'
        raise ValueError(message)
    strict = canonical_json(context['strict_points'])
    if strict != (args.baseline / 'strict-predictions.json').read_bytes():
        message = 'Recomputed strict predictions differ from baseline bytes.'
        raise ValueError(message)
    prediction_bytes = canonical_json(points)
    if prediction_bytes != (args.baseline / 'predictions.json').read_bytes():
        message = 'Changed exploratory predictions require a scored campaign before publication.'
        raise ValueError(message)
    blocks = []
    for block in review['ksat_blocks']:
        selected = [o for o in observations if block['first'] <= o['row'] <= block['last']]
        blocks.append(
            {
                **block,
                'range': f'{scan["sheets"][6]["sheet"]}!AB{block["first"]}:AB{block["last"]}',
                'numeric_rows': len(selected),
                'identity_counts': dict(Counter(o['identity_status'] for o in selected)),
                'canonical_overlap': sum(o['canonical_sample'] for o in selected),
            }
        )
    category_reassessment = {}
    for category in sorted({f['previous_category'] for f in functions}):
        category_reassessment[category] = dict(
            Counter(f['outcome'] for f in functions if f['previous_category'] == category)
        )
    report = {
        'protocol': review['protocol'],
        'catalog_count': len(functions),
        'baseline_executable': len(before_ids),
        'executable': len(after_ids),
        'new_functions': sorted(after_ids - before_ids),
        'functions': functions,
        'ksat_blocks': blocks,
        'ksat_observations': len(observations),
        'ksat_function_count': len({r['function'] for r in ks}),
        'ksat_assessments': len(ks),
        'released_ksat_observations': sum(not o['target_blockers'] for o in observations),
        'method_specific_ksat_pools': [],
        'predictors': predictors,
        'category_reassessment': category_reassessment,
        'review': review,
        'strict_byte_identical': True,
        'exploratory_predictions_byte_identical': True,
        'pressure_policies_changed': False,
        'attribution': dict.fromkeys(
            (
                'Ksat target recovery',
                'particle density / porosity',
                'pH',
                'nitrogen',
                'other confirmed predictors',
            ),
            0,
        ),
        'corrections': 'AE1 is nitrogen-adsorption specific surface, not nitrogen concentration; its column has no numeric observations. '
        'AF/AG contain plastic/liquid-limit candidates, retracting the historical blanket absence classification for 17 functions. '
        'AC is heat of wetting. Literal AB41 is cm/smin, not confirmed cm/min. Historical artifacts remain frozen.',
        'input_sha256': {
            str(p): exploratory._hash(p / 'artifact-hashes.json') for p in directories
        },
        'review_sha256': exploratory._hash(review_path),
        'implementation_sha256': {
            p.name: exploratory._hash(p) for p in sorted(Path(__file__).parent.glob('*.py'))
        },
        'specification_sha256': {
            p.relative_to(context['root']).as_posix(): exploratory._hash(p)
            for p in sorted((context['root'] / 'specs').rglob('*.yaml'))
        },
        'source_document_sha256': {
            v['local_file']: exploratory._hash(args.sources / v['local_file'])
            for v in review['references'].values()
        },
    }
    # No approval is silently ignored by an audit-only execution path.
    if report['released_ksat_observations'] or after_ids != before_ids:
        message = (
            'New approved evidence requires a scored campaign and attribution before publishing.'
        )
        raise ValueError(message)
    args.output.mkdir(parents=True, exist_ok=False)
    artifacts = {
        'summary': report,
        'ksat-identity-map': observations,
        'ksat-catalog-matrix': ks,
        'matrix': records,
        'predictions': points,
    }
    for name, value in artifacts.items():
        (args.output / (name + '.json')).write_bytes(canonical_json(value))
    (args.output / 'strict-predictions.json').write_bytes(strict)
    (args.output / 'report.md').write_text(render(report), encoding='utf-8')
    if initial_hashes != {str(p): exploratory.verify_artifacts(p) for p in directories}:
        message = 'Frozen inputs changed during review.'
        raise ValueError(message)
    (args.output / 'artifact-hashes.json').write_bytes(
        canonical_json({p.name: exploratory._hash(p) for p in sorted(args.output.iterdir())})
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    for name in (
        'baseline',
        'discovery',
        'benchmark',
        'harmonized',
        'workbook',
        'sources',
        'output',
    ):
        parser.add_argument('--' + name, type=Path, required=True)
    args = parser.parse_args()
    result = run(args)
    print(  # noqa: T201 - CLI summary.
        json.dumps(
            {
                k: result[k]
                for k in (
                    'baseline_executable',
                    'executable',
                    'ksat_function_count',
                    'ksat_observations',
                    'category_reassessment',
                )
            }
        )
    )


if __name__ == '__main__':
    main()
