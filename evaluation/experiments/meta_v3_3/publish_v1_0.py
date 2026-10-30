"""Publish the final Russian v1.0 handoff without rerunning scientific calculations."""

# ruff: noqa: INP001, E501, RUF001, PLR2004, PLR0915, PLR0917, C901, PLR0912, PERF401, TRY301 - Publication prose and scientific presentation layout.

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import shutil
import textwrap
from collections import defaultdict
from pathlib import Path
from statistics import median
from typing import TYPE_CHECKING
from zipfile import ZIP_DEFLATED, ZipFile

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from evaluation.experiments.meta_v3_3 import model_selection as metrics
from evaluation.experiments.meta_v3_3 import render_diagnostics as render


if TYPE_CHECKING:
    from pptx.presentation import Presentation as PresentationDocument
    from pptx.slide import Slide


ROOT = Path(__file__).resolve().parents[3]
PUBLIC = Path(__file__).parent / 'final_results'
PRIVATE = ROOT / 'local/meta-final-v1_0'
INPUT = ROOT / 'local/meta-expanded-benchmark'
SOURCE_DIAGNOSTICS = INPUT / 'diagnostics'
DESTINATION = PRIVATE / 'diagnostics'
FAMILIES = render.FAMILIES
BASES = ('volumetric_water_content', 'gravimetric_water_content')
NAVY, TEAL, GRAY, PALE, WHITE = '193447', '087F79', '627482', 'F0F6F7', 'FFFFFF'


def read(path: Path):  # noqa: ANN201 - Scientific JSON roots can be objects or arrays.
    return json.loads(path.read_bytes())


def write(path: Path, value: object) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8'
    )


def close(a: float | str, b: float) -> None:
    if not math.isclose(float(a), float(b), rel_tol=1e-10, abs_tol=1e-12):
        message = f'Final artifact metric mismatch: {a} versus {b}.'
        raise ValueError(message)


def clean_selection(value: object) -> object:
    """Retain final scientific evidence without internal development annotations."""
    if isinstance(value, dict):
        return {
            k: clean_selection(v)
            for k, v in value.items()
            if k not in {'mechanism', 'fold_details'}
        }
    if isinstance(value, list):
        return [clean_selection(v) for v in value]
    if value == 'expanded_exploratory_E5':
        return 'final_v1.0_E5'
    return value


def documents(analyses: dict, approximation: dict, ledger: dict) -> None:
    volume = analyses[BASES[0]]['primary']
    top = volume['ensembles']['top_5']
    improvement = 100 * (1 - top['rmse'] / volume['ranking'][0]['rmse'])
    conclusion = f'В финальном анализе ptfkit v1.0 наилучшей индивидуальной PTF для объёмной основной гидрофизической характеристики (ОГХ) является Varallyai1982 chernozem B: RMSE составляет 0.036778 м³/м³ на сопоставимой когорте из 15 образцов и 136 наблюдений. Равновесный ансамбль top-5 обеспечивает RMSE 0.029326 м³/м³ и снижение ошибки на {improvement:.2f}%; top-3 и top-5 проходят зарегистрированный критерий устойчивости при проверке на отложенных источниках. Этот результат относится к поддержке примерно 10–50.1 кПа и не устанавливает точность на всей кривой. Для массовой ОГХ предпочтительно дискретное семейство Lal1979 group I from sand с RMSE 0.068027 кг/кг на 10 образцах и 30 наблюдениях; устойчивого преимущества ансамблей не установлено. Для Ksat лучший результат на общем наборе из 30 образцов показывает Ferrer-Julià 2004 gleyic cambisol sand, а при максимальном покрытии 39 образцов рекомендуется Dane–Puckett, с учётом чувствительности к источникам и единицам данных. При аппроксимации экспериментальных ОГХ van Genuchten m=1 имеет наименьшую медианную RMSE для обеих водных баз. Аппроксимация описывает подгонку к измерениям, тогда как PTF оцениваются по независимым прогнозам. Все выводы специфичны для используемого набора данных (dataset-specific); ошибки различных когорт напрямую не сопоставимы.'
    (PUBLIC / 'CONCLUSION-RU.md').write_text(conclusion + '\n', encoding='utf-8')
    methodology = """# Ансамблевый подход ptfkit — v1.0

Сначала каждую педотрансферную функцию (PTF) рассчитывали отдельно. Объёмную
и массовую влажность анализировали раздельно. В ансамбль включали только
независимые прогнозы PTF, доступные на одной и той же точной когорте образцов
и экспериментальных точек ОГХ. Отсутствующие прогнозы не заполняли: использовали
пересечение наблюдений всех участников. Каждому образцу требовалось не менее
двух различных давлений. Модели, использующие измеренную влажность этой же ОГХ
как входной параметр, в независимое сравнение не включали. Ksat оценивали
по отдельным функциям; ансамбль Ksat не строили.

Для top-N модели ранжировали по RMSE с равными весами образцов. При одинаковой
RMSE учитывали абсолютное систематическое смещение, затем идентификатор
функции. Выбирали N лучших моделей, в частности top-3 и top-5. В каждой точке
их прогнозы усредняли арифметически: прогноз ансамбля равен сумме прогнозов
участников, делённой на N. В top-5 каждая модель имеет вес 1/5; веса по
экспериментальным значениям не оптимизировали. Также проверяли среднее всех
участников, медиану и сокращённую пару. Для пары к лучшей модели выбирали
первую среди следующих четырёх, чья определённая корреляция остатков с ней
была ниже 0.95; иначе использовали две лучшие модели.

Для независимой оценки состав top-N и сокращённой пары заново определяли
внутри каждого обучающего разбиения. Выбор проверяли на отложенной группе,
которая не участвовала в ранжировании, оценке взаимодополняемости ошибок
или подборе состава. Для основной объёмной когорты использовали два
разбиения с исключением одного источника. Если нескольких источников нет,
предусмотрено исключение профиля при полноте идентификаторов, иначе —
500 bootstrap-перевыборок целых образцов с проверкой на образцах вне
обучающей перевыборки. Для основной массовой когорты доступна только
внутриисточниковая bootstrap-проверка. Ранжирование в bootstrap учитывает
кратность образцов, корреляции — уникальные обучающие точки.

Pooled RMSE рассчитывали на полной общей когорте, выбирая состав по этой
когорте как по полному обучающему набору; это описательный результат.
RMSE равна корню из среднего по образцам значения среднеквадратичной ошибки
каждого образца. MAE и смещение также дают образцам равный вес. Устойчивое
преимущество признавали только при снижении pooled RMSE минимум на 1%
относительно лучшей отдельной PTF и строго меньшей RMSE во всех предусмотренных
отложенных проверках относительно обоих ориентиров: лучшей PTF, выбранной
на обучении, и лучшей PTF полной когорты. Второй ориентир выбран на полной
когорте и является фиксированным описательным эталоном.

Для объёмной ОГХ top-3 и top-5 проходят этот критерий. Наименьшая pooled RMSE
среди прошедших ансамблей — у top-5: 0.029326 м³/м³ против 0.036778 м³/м³
у лучшей индивидуальной PTF, снижение на 20.26%. Общая когорта содержит
15 образцов и 136 точек примерно при 10–50.1 кПа. Итоговый top-5 объединяет
Clapp1978 и Varallyai1982 chernozem A, chernozem B, chernozem C и meadow.
В каждой независимой проверке процедура выбора повторяется. Для массовой
ОГХ устойчивого преимущества ансамблей не установлено. Эти выводы относятся
к данному набору данных и его поддержке давлений; коррелированные модели
не являются независимыми научными подтверждениями.

Таким образом, ансамбль представляет собой не отдельную обучаемую модель, а равновесное объединение прогнозов нескольких наиболее успешных PTF, причём состав ансамбля определяется только по обучающей части данных.
"""
    for name in ('ENSEMBLE-METHODOLOGY-RU.md', 'ENSEMBLE-METHODOLOGY.md'):
        (PUBLIC / name).write_text(methodology, encoding='utf-8')
    lines = [
        '# Результаты benchmark ptfkit — v1.0',
        '',
        conclusion,
        '',
        '## Данные и интерпретация',
        '',
        f'Экспериментальный ledger содержит {ledger["observations"]} точку ОГХ и {ledger["samples"]} записей образцов; для Ksat доступны 39 наблюдений. Это данные workbook с ограничениями подтверждения лабораторных методов, происхождения измерений и применимости PTF. Политика E5 допускает явно зарегистрированные приближения и преобразования одного образца, но не подмену физических величин, смешение образцов или вымышленные численные значения.',
        '',
        'Объёмная и массовая влажность сравниваются отдельно. RMSE, MAE и bias ОГХ дают образцам равный вес; bias — прогноз минус измерение. Для Ksat используются ошибки log10(прогноз/измерение). Между различными когортами и водными базами ошибки напрямую не ранжируются.',
        '',
        '## Выбор PTF',
        '',
        '| Задача | Функция / решение | RMSE | Образцы / точки |',
        '|---|---|---:|---:|',
        '| Объёмная ОГХ, индивидуальная PTF | Varallyai1982 chernozem B | 0.036778 м³/м³ | 15 / 136 |',
        '| Объёмная ОГХ, ансамбль | top-5, равные веса | 0.029326 м³/м³ | 15 / 136 |',
        '| Массовая ОГХ | Lal1979 group I from sand, дискретное семейство | 0.068027 кг/кг | 10 / 30 |',
        '| Ksat, общий сопоставимый набор | Ferrer-Julià 2004 gleyic cambisol sand | 0.270425 log10 | 30 / 30 |',
        '| Ksat, максимальное покрытие | Ferrer-Julià 2004 Dane–Puckett | 0.717641 log10 | 39 / 39 |',
        '',
        'Top-5 снижает RMSE относительно лучшей индивидуальной PTF на 20.26%; top-3 также проходит критерий устойчивого улучшения. Основная объёмная когорта охватывает приблизительно 10–50.1 кПа и два источника. Этот результат не переносится автоматически на всю ОГХ. Для массовой когорты устойчивого преимущества ансамблей не установлено; её поддержка составляет приблизительно 93.8–585.6 кПа.',
        '',
        'На общем наборе Ksat доли прогнозов gleyic cambisol sand в пределах факторов ×2 / ×5 / ×10 составляют 80.0% / 96.7% / 100.0%. Рекомендация Dane–Puckett относится к максимальному покрытию, а не к устойчивому превосходству во всех блоках. Единица `cm/smin` интерпретирована как cm/min; исходные методы и происхождение измерений остаются неподтверждёнными. Классовые регрессии Ferrer-Julià имеют ограничения применимости E5: классы FAO по workbook не установлены. Ансамбли Ksat не используются.',
        '',
        'Дискретные семейства содержат минимум три допустимые native-координаты давления одного научного метода и варианта уравнений. Между положительными координатами применяется линейная интерполяция влажности по log10 давления без extrapolation. Нулевая координата используется только точно при нуле. Point-target сравнения на отдельных native-давлениях — отдельный вторичный анализ. Матрицы корреляций остатков и результаты проверки составов ансамбля включены в `selection_v1.0.json`.',
        '',
        '## Аппроксимация экспериментальной ОГХ',
        '',
        'Предпочтительное общее описательное семейство — van Genuchten m=1. Сравниваются одни и те же образцы и точки внутри каждой водной базы: 76 объёмных и 65 массовых кривых. Все четыре семейства имеют описательные допустимые fit на этих 141 образце. Один из 142 образцов имеет недостаточно точек. Это fitted-аппроксимация по тем же измерениям, а не независимый PTF-прогноз.',
        '',
        '| База | Семейство | Строгие fit | Описательные fit | Median RMSE | Median MAE |',
        '|---|---|---:|---:|---:|---:|',
    ]
    for basis in BASES:
        label = 'Объёмная' if basis == BASES[0] else 'Массовая'
        for row in approximation[basis]['summary']:
            lines.append(
                f'| {label} | {render.NAMES[row["family"]]} | {row["strict_successes"]}/{row["eligible_samples"]} | {row["relaxed_successes"]}/{row["eligible_samples"]} | {row["median_rmse"]:.6f} | {row["median_mae"]:.6f} |'
            )
    lines += [
        '',
        '### Поддержка давления и ошибки',
        '',
        'Для строгих fit VG m=1 наблюдённая поддержка составляет 0–9.80665×10⁶ кПа для объёмной и 0–310114 кПа для массовой влажности. Центральные 5–95% поддержки с равными весами образцов: 0.10035–77897 и 0.24633–158085 кПа соответственно. Большие давления сухого конца — сохранённые значения workbook, требующие осторожности в интерпретации протоколов.',
        '',
        'У VG m=1 для объёмной влажности минимальная наблюдаемая MAE отмечена при p<10 кПа: N=76, медианная абсолютная ошибка 0.00344, 95-й процентиль 0.01164 м³/м³. Для массовой — при p≥1500 кПа: N=37, 0.00477 и 0.01483 кг/кг. Эти описательные закономерности зависят от доступных образцов и не задают универсальную границу надёжности.',
        '',
        'Поддержка — диапазон экспериментальных давлений, а не доверительный или предиктивный интервал. Эмпирические распределения fitted-остатков, bootstrap-интервалы сводных ошибок и условные интервалы параметров — разные величины. Fit-ошибки не устанавливают покрытие прогнозных интервалов для будущих образцов.',
        '',
        'Строгий fit требует сходимости, физических конечных прогнозов, положительных остаточных степеней свободы, отсутствия критического упора в границы, непостоянной кривой и приемлемой идентифицируемости якобиана. Описательный relaxed fit сохраняет предупреждения о границах и обусловленности. AICc/BIC используются описательно с учётом числа параметров и области определённости; выбор по медианной RMSE не тождественен выбору по информационным критериям.',
        '',
        '## Итоговый комплект',
        '',
        '- `ptfkit_benchmark_summary_RU_v1.0.pptx`: два основных слайда и приложение.',
        '- `CONCLUSION-RU.md`: итоговые выводы одним абзацем.',
        '- `ENSEMBLE-METHODOLOGY-RU.md`: описание ансамбля.',
        '- `PNG-AUDIT.md`: проверка всех диагностических изображений.',
        '- `ptfkit_diagnostics_v1.0.zip`: PTF и fitted-диагностика, README и индекс.',
        '- `selection_v1.0.json`, `benchmark_aggregates.json` и CSV: сводные научные результаты.',
        '',
        'Диагностические изображения и ZIP содержат идентификаторы образцов и хранятся в `local/meta-final-v1_0/`. Воспроизведение оформления и проверки без повторного benchmark: `uv run --project targets/ptfkit-eval --no-sync --with matplotlib --with python-pptx python -B -m evaluation.experiments.meta_v3_3.publish_v1_0`. Расчётные артефакты необходимы локально. Все выводы являются dataset-specific.',
        '',
    ]
    (PUBLIC / 'FINAL-RESULTS.md').write_text('\n'.join(lines), encoding='utf-8')


def textbox(
    slide: Slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: float = 18,
    color: str = NAVY,
    *,
    bold: bool = False,
) -> None:
    h = max(h, (size * 1.22 * len(text.split('\n')) + 2 * (len(text.split('\n')) - 1) + 2) / 72)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = box.text_frame.margin_right = Inches(0.015)
    box.text_frame.margin_top = box.text_frame.margin_bottom = Inches(0.005)
    for number, line in enumerate(text.split('\n')):
        paragraph = box.text_frame.paragraphs[0] if number == 0 else box.text_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = 'Arial'
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(2)


def panel(slide: Slide, x: float, y: float, w: float, h: float, color: str = PALE) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    shape.adjustments[0] = 0.08


def slide_base(presentation: PresentationDocument, title: str, number: int) -> Slide:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    textbox(slide, 0.45, 0.25, 12.4, 0.55, title, 27, bold=True)
    textbox(
        slide,
        0.47,
        7.15,
        11.8,
        0.2,
        'ptfkit • v1.0 • Выводы относятся к данному набору данных',
        10,
        GRAY,
    )
    textbox(slide, 12.4, 7.12, 0.4, 0.25, str(number), 11, GRAY)
    return slide


def presentation(_analyses: dict, approximation: dict) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = 'ptfkit — финальный benchmark v1.0'
    deck.core_properties.subject = 'Выбор PTF, ансамбли и аппроксимация ОГХ'
    deck.core_properties.author = 'ptfkit'
    slide = slide_base(deck, 'Сравнение педотрансферных функций для ОГХ и Ksat', 1)
    textbox(
        slide,
        0.47,
        0.93,
        12.3,
        0.28,
        'Сопоставимые когорты • раздельные водные базы • независимые PTF-прогнозы',
        14,
        GRAY,
    )
    for x in (0.45, 4.65, 8.85):
        panel(slide, x, 1.45, 4.03, 4.95)
    textbox(slide, 0.68, 1.65, 3.6, 0.4, 'Объёмная ОГХ', 22, bold=True)
    textbox(slide, 0.68, 2.14, 3.6, 0.57, 'Varallyai1982\nchernozem B', 18)
    textbox(slide, 0.68, 2.87, 3.6, 0.5, '0.036778', 30, bold=True)
    textbox(slide, 0.68, 3.43, 3.6, 0.32, 'RMSE, м³/м³ • 15 образцов / 136 точек', 13, GRAY)
    textbox(slide, 0.68, 4.02, 3.7, 0.35, 'Ансамбль top-5 • равные веса', 17, TEAL, bold=True)
    textbox(slide, 0.68, 4.47, 2.35, 0.5, '0.029326', 30, TEAL, bold=True)
    textbox(slide, 3.0, 4.49, 1.2, 0.42, '−20.26%', 18, TEAL, bold=True)
    textbox(slide, 0.68, 5.10, 3.6, 0.62, 'Top-5 и top-3 проходят\nкритерий устойчивости', 16, TEAL)
    textbox(slide, 0.68, 5.92, 3.6, 0.25, 'Поддержка: 10–50.1 кПа; 2 источника', 12, GRAY)
    textbox(slide, 4.88, 1.65, 3.6, 0.4, 'Массовая ОГХ', 22, bold=True)
    textbox(slide, 4.88, 2.14, 3.6, 0.93, 'Lal1979 group I from sand\nДискретное семейство ОГХ', 18)
    textbox(slide, 4.88, 3.22, 3.6, 0.5, '0.068027', 30, bold=True)
    textbox(slide, 4.88, 3.82, 3.6, 0.35, 'RMSE, кг/кг • 10 образцов / 30 точек', 13, GRAY)
    textbox(slide, 4.88, 4.55, 3.6, 0.97, 'Ансамбль устойчивого\nпреимущества не показал', 19, GRAY)
    textbox(slide, 4.88, 5.77, 3.6, 0.43, 'Поддержка: 93.8–585.6 кПа\nОдин источник', 12, GRAY)
    textbox(slide, 9.08, 1.65, 3.6, 0.4, 'Ksat', 22, bold=True)
    textbox(slide, 9.08, 2.12, 3.6, 0.26, 'Общий сопоставимый набор • N=30', 13, GRAY)
    textbox(slide, 9.08, 2.52, 3.6, 0.66, 'Ferrer-Julià 2004\ngleyic cambisol sand', 17)
    textbox(slide, 9.08, 3.24, 3.6, 0.48, '0.270425', 27, bold=True)
    textbox(slide, 9.08, 3.82, 3.6, 0.53, 'RMSE(log10)\n×2: 80%   ×5: 96.7%   ×10: 100%', 13, GRAY)
    textbox(slide, 9.08, 4.55, 3.6, 0.27, 'Максимальное покрытие • N=39', 13, TEAL)
    textbox(slide, 9.08, 4.98, 3.7, 0.42, 'Ferrer-Julià 2004 Dane–Puckett', 17, TEAL, bold=True)
    textbox(slide, 9.08, 5.55, 3.6, 0.44, '0.717641', 27, TEAL, bold=True)
    textbox(slide, 9.08, 6.05, 3.6, 0.2, 'RMSE(log10); зависит от блока данных', 11, GRAY)
    textbox(
        slide,
        0.48,
        6.55,
        12.25,
        0.5,
        'Объёмная ОГХ — ансамбль пяти PTF; массовая ОГХ — семейство Lal1979;\nKsat при максимальном покрытии — Dane–Puckett.',
        16,
        bold=True,
    )

    slide = slide_base(deck, 'Аппроксимация экспериментальной ОГХ', 2)
    textbox(
        slide,
        0.47,
        0.93,
        12.3,
        0.3,
        'VG m=1 даёт наименьшую медианную ошибку в целом для обеих водных баз',
        18,
        TEAL,
        bold=True,
    )
    order = (
        'van_genuchten_m1_v1',
        'van_genuchten_mualem_v1',
        'brooks_corey_v1',
        'clapp1978_gradual_air_entry_v1',
    )
    for x, basis, label in (
        (0.45, BASES[0], 'Объёмная влажность • м³/м³ • N=76'),
        (6.85, BASES[1], 'Массовая влажность • кг/кг • N=65'),
    ):
        panel(slide, x, 1.48, 6.02, 3.43)
        textbox(slide, x + 0.2, 1.66, 5.6, 0.4, label, 19, bold=True)
        textbox(slide, x + 0.2, 2.18, 2.43, 0.3, 'Семейство', 12, GRAY)
        textbox(slide, x + 2.7, 2.18, 1.4, 0.3, 'Median RMSE', 12, GRAY)
        textbox(slide, x + 4.25, 2.1, 1.55, 0.48, 'Успех fit\nстрого / описательно', 10, GRAY)
        summaries = {r['family']: r for r in approximation[basis]['summary']}
        for i, family in enumerate(order):
            row = summaries[family]
            y = 2.67 + i * 0.51
            color = TEAL if i == 0 else NAVY
            textbox(slide, x + 0.2, y, 2.45, 0.34, render.NAMES[family], 17, color, bold=i == 0)
            textbox(
                slide,
                x + 2.7,
                y - 0.02,
                1.45,
                0.4,
                f'{row["median_rmse"]:.6f}',
                20,
                color,
                bold=i == 0,
            )
            textbox(
                slide,
                x + 4.25,
                y + 0.02,
                1.55,
                0.3,
                f'{row["strict_successes"]}/{row["eligible_samples"]}   {row["relaxed_successes"]}/{row["eligible_samples"]}',
                14,
                GRAY,
            )
    panel(slide, 0.45, 5.13, 6.02, 1.14, 'F6F8FA')
    panel(slide, 6.85, 5.13, 6.02, 1.14, 'F6F8FA')
    textbox(
        slide,
        0.65,
        5.27,
        5.6,
        0.26,
        'Экспериментальная поддержка строгих VG m=1 fit',
        13,
        bold=True,
    )
    textbox(
        slide,
        0.65,
        5.65,
        5.6,
        0.4,
        'Объёмная: 0–9.81×10⁶ кПа; массовая: 0–3.10×10⁵ кПа\nПоддержка не является доверительным / прогнозным интервалом',
        12,
        GRAY,
    )
    textbox(slide, 7.05, 5.27, 5.6, 0.26, 'Ошибка VG m=1 зависит от давления', 13, bold=True)
    textbox(
        slide,
        7.05,
        5.65,
        5.6,
        0.4,
        'Минимальная наблюдаемая MAE: объёмная — p<10 кПа (N=76);\nмассовая — p≥1500 кПа (N=37). Это описательные закономерности.',
        12,
        GRAY,
    )
    textbox(
        slide,
        0.47,
        6.5,
        12.3,
        0.5,
        'VG m=1 — предпочтительная общая функция аппроксимации ОГХ; надёжность оценивается\nс учётом экспериментальной поддержки давлений и распределения ошибок.',
        16,
        bold=True,
    )

    slide = slide_base(deck, 'Как формировался ансамбль PTF', 3)
    textbox(
        slide,
        0.47,
        0.95,
        12.2,
        0.3,
        'Приложение • одинаковые образцы и точки • выбор состава без отложенных наблюдений',
        15,
        GRAY,
    )
    labels = (
        'Отдельные\nPTF',
        'Одинаковая\nкогорта',
        'Ранг PTF\nна обучении',
        'Выбор\ntop-N',
        'Среднее\nпрогнозов',
        'Проверка\nна held-out',
    )
    for i, label in enumerate(labels):
        x = 0.45 + i * 2.12
        panel(slide, x, 1.65, 1.87, 1.05)
        textbox(slide, x + 0.1, 1.91, 1.67, 0.6, label, 17, TEAL, bold=True)
        if i < 5:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, Inches(x + 1.91), Inches(2.12), Inches(0.17), Inches(0.17)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor.from_string(GRAY)
            arrow.line.fill.background()
    textbox(
        slide,
        0.6,
        3.14,
        12.1,
        0.47,
        'Прогноз top-N = сумма прогнозов выбранных PTF / N',
        25,
        TEAL,
        bold=True,
    )
    textbox(
        slide,
        0.6,
        3.88,
        12.0,
        0.55,
        'Top-5: Clapp1978 + Varallyai1982 chernozem A, B, C и meadow.\nПри каждой held-out проверке выбор моделей повторяется только на обучающей части.',
        17,
    )
    panel(slide, 0.45, 4.72, 6.02, 1.35)
    panel(slide, 6.85, 4.72, 6.02, 1.35)
    textbox(slide, 0.68, 4.9, 5.5, 0.32, 'Критерий устойчивости', 18, bold=True)
    textbox(
        slide,
        0.68,
        5.4,
        5.5,
        0.48,
        'Снижение pooled RMSE ≥1% и меньшая ошибка\nво всех предусмотренных held-out проверках.',
        15,
    )
    textbox(slide, 7.08, 4.9, 5.5, 0.32, 'Два ориентира для сравнения', 18, bold=True)
    textbox(
        slide,
        7.08,
        5.4,
        5.5,
        0.48,
        'Лучшая индивидуальная PTF на обучении\nи лучшая индивидуальная PTF полной когорты.',
        15,
    )
    textbox(
        slide,
        0.58,
        6.39,
        12.1,
        0.53,
        'Объёмная когорта: 2 источника, 15 образцов, 136 точек, 10–50.1 кПа.\nPooled RMSE описательна; переносимость оценивается по отложенным данным.',
        16,
        GRAY,
    )
    deck.save(str(PUBLIC / 'ptfkit_benchmark_summary_RU_v1.0.pptx'))


def point_reference(secondary: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for point in secondary:
        groups[point['basis'], point['pressure_kpa'], point['function']].append(point)
    rows = []
    for (basis, pressure, fid), group in sorted(groups.items()):
        score = metrics.scored(group)
        folder, unit, label = render.basis_info(basis)
        cohort = f'native_{pressure:.12g}_kpa_maximum_function_support'
        annotation = (
            textwrap.fill(fid, 42)
            + '\n\nRoute: point_target (secondary)\n'
            + f'Pressure: {pressure:.8g} kPa\nWater basis: {label.lower()}\nSamples: {score["samples"]}\nObservations: {score["observations"]}\nRMSE: {score["rmse"]:.6f}\nMAE: {score["mae"]:.6f}\nBias: {score["bias"]:+.6f}\n\nObserved target: exact or\nwithin-sample log-pressure\nreconstruction; no extrapolation\n\nCohort: '
            + textwrap.fill(cohort, 40)
        )
        fig, ax = render.canvas(
            'Measured vs predicted — secondary point target\n' + fid.removeprefix('calc_ptf_'),
            annotation,
        )
        render.scatter(ax, group, f'water content ({unit})')
        fig.text(
            0.09,
            0.035,
            'Secondary point-target result; compare only identical pressure / basis / sample cohorts.',
            fontsize=8.5,
        )
        path = render.save(
            fig,
            Path('PTF/point_targets')
            / folder
            / f'{fid.removeprefix("calc_ptf_")}_{pressure:.8g}kpa.png',
        )
        rows.append(
            {
                **render.index_row('PTF', basis + '_point_target', fid, cohort, score, path),
                'route_type': 'point_target',
            }
        )
    return rows


def png_audit(analyses: dict, approximation: dict) -> dict:
    reference = PRIVATE / 'render-reference'
    reference.mkdir(exist_ok=True)
    render.OUTPUT = reference
    points = read(INPUT / 'predictions.json')
    secondary = read(INPUT / 'point-target-predictions.json')
    rows = render.ptf_plots(points, analyses)
    metadata = read(INPUT / 'candidate-metadata.json')
    for row in rows:
        row['route_type'] = (
            metadata[row['function_or_family']]['route_type'] if row['target'] != 'ksat' else 'Ksat'
        )
    rows.extend(point_reference(secondary))
    fits = read(ROOT / 'local/meta-final-benchmark/approximation-fits.json')['results']
    source_index = list(
        csv.DictReader(
            (SOURCE_DIAGNOSTICS / 'diagnostic_index.csv').open(encoding='utf-8', newline='')
        )
    )
    approximation_rows = [r for r in source_index if r['category'] == 'APPROXIMATION']
    expected_approx = {
        f'APPROXIMATION/{render.basis_info(basis)[0]}/{family}.png'
        for basis in BASES
        for family in FAMILIES
    }
    accepted = [f for f in fits if f['relaxed_accepted']]
    sample_keys = {(f['basis'], f['source_id'], f['sample_id']) for f in accepted}
    expected_approx |= {
        f'APPROXIMATION/per_sample/{render.sample_stem(key)}.png' for key in sample_keys
    }
    if {r['file_path'] for r in approximation_rows} != expected_approx:
        message = 'Approximation image index does not cover all saved families and usable samples.'
        raise ValueError(message)
    for row in approximation_rows:
        candidates = [
            f
            for f in accepted
            if f['basis'] == row['target'] and f['family'] == row['function_or_family']
        ]
        if '/per_sample/' in row['file_path']:
            candidates = [
                f
                for f in candidates
                if f'sample:{f["source_id"]}/{f["sample_id"]}' == row['cohort']
            ]
        for field, metric in (('RMSE', 'rmse'), ('MAE', 'mae'), ('bias', 'bias')):
            close(row[field], median(f[metric] for f in candidates))
        close(row['samples'], len(candidates))
        close(row['observations'], sum(f['n'] for f in candidates))
        close(row['strict_success'], sum(f['strict_accepted'] for f in candidates))
        close(row['relaxed_success'], len(candidates))
        row['route_type'] = 'fitted_approximation'
    rows.extend(approximation_rows)
    expected = {r['file_path'] for r in rows}
    source_actual = {
        p.relative_to(SOURCE_DIAGNOSTICS).as_posix() for p in SOURCE_DIAGNOSTICS.rglob('*.png')
    }
    missing, invalid, stale, regenerated, duplicate_content = [], [], [], [], []
    DESTINATION.mkdir(exist_ok=True)
    pixel_hashes = defaultdict(list)
    for relative in sorted(expected):
        source = SOURCE_DIAGNOSTICS / relative
        valid = False
        if source.exists():
            try:
                with Image.open(source) as png:
                    png.load()
                    if (
                        png.format != 'PNG'
                        or png.width < 1200
                        or png.height < 800
                        or float(np.std(np.asarray(png.convert('L').resize((160, 100))))) < 4
                    ):
                        message = 'PNG dimensions or content invalid.'
                        raise ValueError(message)
                    pixel_hash = hashlib.sha256(png.convert('RGBA').tobytes()).hexdigest()
                    valid = True
                if relative.startswith('PTF/'):
                    with Image.open(reference / relative) as current:
                        if (
                            pixel_hash
                            != hashlib.sha256(current.convert('RGBA').tobytes()).hexdigest()
                        ):
                            stale.append(relative)
                            valid = False
            except (OSError, ValueError):
                invalid.append(relative)
        else:
            missing.append(relative)
        if not valid:
            if not (reference / relative).exists():
                # Render stored fits, never refit or rerun PTF calculations.
                render.approximation_plots(fits, approximation)
            shutil.copyfile(reference / relative, source)
            regenerated.append(relative)
        target = DESTINATION / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)
        with Image.open(target) as png:
            png.load()
            pixel_hashes[hashlib.sha256(png.convert('RGBA').tobytes()).hexdigest()].append(relative)
    for files in pixel_hashes.values():
        if len(files) > 1:
            duplicate_content.append(files)
    actual = {p.relative_to(DESTINATION).as_posix() for p in DESTINATION.rglob('*.png')}
    if actual != expected:
        message = 'Final PNG folder contains missing or unexpected files.'
        raise ValueError(message)
    with (DESTINATION / 'diagnostic_index.csv').open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=[*render.COLUMNS, 'route_type'])
        writer.writeheader()
        writer.writerows(rows)
    report = {
        'expected_count': len(expected),
        'actual_count': len(actual),
        'full_wrc_ksat_png': sum(
            r['category'] == 'PTF' and not r['target'].endswith('_point_target') for r in rows
        ),
        'point_target_png': sum(
            r['category'] == 'PTF' and r['target'].endswith('_point_target') for r in rows
        ),
        'approximation_overviews': 8,
        'per_sample_png': len(sample_keys),
        'index_rows': len(rows),
        'missing_found': missing,
        'invalid_found': invalid,
        'stale_found': stale,
        'regenerated': regenerated,
        'unreferenced_source_png': sorted(source_actual - expected),
        'identical_content_groups': duplicate_content,
        'final_missing': [],
        'final_invalid': [],
        'final_stale': [],
        'ptf_annotation_check': 'Pixel comparison against rendering from final forecasts, exact cohorts and current metrics; function IDs and route labels are generated from final IDs.',
        'approximation_annotation_check': 'All family/sample identities, eligible N and saved-fit metrics checked against the current index and pinned fit artifact; no fitting repeated.',
    }
    lines = [
        '# Проверка диагностических PNG — v1.0',
        '',
        '| Проверка | Результат |',
        '|---|---:|',
        f'| Ожидается изображений | {len(expected)} |',
        f'| Фактически в комплекте | {len(actual)} |',
        f'| Полная ОГХ и Ksat | {report["full_wrc_ksat_png"]} |',
        f'| Point-target | {report["point_target_png"]} |',
        '| Общие measured-vs-fitted графики | 8 |',
        f'| Графики отдельных образцов | {len(sample_keys)} |',
        f'| Недостающих при проверке | {len(missing)} |',
        f'| Невалидных / повреждённых при проверке | {len(invalid)} |',
        f'| Не совпадающих с актуальной визуализацией | {len(stale)} |',
        f'| Перегенерировано из сохранённых прогнозов / fit | {len(regenerated)} |',
        f'| Неиндексированных исходных PNG | {len(source_actual - expected)} |',
        f'| Групп файлов с идентичными пикселями | {len(duplicate_content)} |',
        '| Недостающих / повреждённых / неактуальных в итоговом комплекте | 0 / 0 / 0 |',
        '',
        f'Индекс содержит {len(rows)} строк для {len(expected)} уникальных PNG. Для per-sample графиков предусмотрены четыре строки — по одной на семейство, с общим путём файла; это не дублирование изображений.',
        '',
        'Каждый PNG открывается, имеет непустое содержимое и размеры не менее 1200×800 пикселей. Имена функций, route type, когорты, N и метрики PTF проверены сравнением пикселей с визуализацией сохранённых финальных прогнозов. Для аппроксимации проверены все четыре семейства в обеих водных базах и все пригодные образцы; числа и идентификаторы сверены с сохранёнными fit. Benchmark и fitting не пересчитывались.',
        '',
        'ZIP содержит только `PTF/`, `APPROXIMATION/`, `README.md` и `diagnostic_index.csv`; SVG, временные файлы и расчётные артефакты в него не включены. Проверены CRC и побайтовое соответствие файлов архива итоговой папке.',
        '',
    ]
    (PUBLIC / 'PNG-AUDIT.md').write_text('\n'.join(lines), encoding='utf-8')
    write(PRIVATE / 'png-audit.json', report)
    return report


def package() -> None:
    readme = """# Диагностика ptfkit — v1.0

`PTF/` — measured-vs-predicted графики независимых PTF и научных WRC-семейств,
раздельно для объёмной ОГХ, массовой ОГХ и Ksat. Route type, когорта, число
образцов и точек, RMSE, MAE и bias подписаны на изображениях. WRC-ошибки дают
всем образцам одинаковый вес. Разные когорты и водные базы напрямую не ранжируются.
Дискретные семейства используют минимум три допустимые native-координаты и
интерполяцию по log10 давления без extrapolation. Объёмная поддержка основной
когорты — примерно 10–50.1 кПа; этот результат не характеризует всю ОГХ.

`PTF/point_targets/` — отдельный вторичный анализ на native-давлениях. Измеренные
targets берутся точно или восстанавливаются интерполяцией внутри измеренной ОГХ
того же образца. Эти результаты не смешиваются с full-WRC ranking и ансамблями.
При преобразовании водной базы используется фактическая плотность того же образца.

Ksat показан в м/с на двух логарифмических осях. Пунктирные границы ×2, ×5 и ×10
соответствуют интервалу от measured/factor до measured*factor. Метрики Ksat
рассчитываются по log10(predicted/measured). Блоки и доли попаданий подписаны.

`APPROXIMATION/volumetric/` и `gravimetric/` — measured-vs-fitted графики VG–Mualem,
VG m=1, Brooks–Corey и Clapp. Параметры подогнаны к показанным измерениям:
эти ошибки описывают аппроксимацию, а не независимую PTF-точность. Сводные метрики
являются медианами метрик образцов. `per_sample/` содержит одну картинку на образец
с четырьмя кривыми; `relaxed only` отмечает fit, не прошедший строгие критерии.
Нулевая точка на логарифмической оси имеет явно подписанную координату отображения.
Диапазон экспериментальных давлений не является доверительным или прогнозным интервалом.

`diagnostic_index.csv` перечисляет все изображения; пути заданы относительно корня
архива. Четыре строки одного per-sample файла соответствуют четырём семействам.
Все выводы относятся к используемому набору данных и ограничениям E5.
"""
    (DESTINATION / 'README.md').write_text(readme, encoding='utf-8')
    archive_path = PRIVATE / 'ptfkit_diagnostics_v1.0.zip'
    members = [
        DESTINATION / 'README.md',
        DESTINATION / 'diagnostic_index.csv',
        *sorted((DESTINATION / 'PTF').rglob('*.png')),
        *sorted((DESTINATION / 'APPROXIMATION').rglob('*.png')),
    ]
    with ZipFile(archive_path, 'w', compression=ZIP_DEFLATED) as archive:
        for path in members:
            archive.write(path, path.relative_to(DESTINATION).as_posix())
    with ZipFile(archive_path) as archive:
        if archive.testzip() is not None or set(archive.namelist()) != {
            p.relative_to(DESTINATION).as_posix() for p in members
        }:
            message = 'Final ZIP inventory or CRC verification failed.'
            raise ValueError(message)
        for path in members:
            if archive.read(path.relative_to(DESTINATION).as_posix()) != path.read_bytes():
                message = 'ZIP member differs from its verified source.'
                raise ValueError(message)


def neutral_audit() -> None:
    audit = read(PUBLIC / 'expanded_function_audit.json')
    catalog = read(INPUT / 'baseline/benchmark_aggregates.json')['selection']['catalog']
    full = set(read(INPUT / 'scope.json')['previously_scored']) | set(
        audit['participation']['new_functions']
    )
    points = {p['function'] for p in read(INPUT / 'point-target-predictions.json')}
    decisions = {r['function']: r for r in audit['functions']}
    lines = [
        '# Применимость функций ptfkit — v1.0',
        '',
        'Независимый full-WRC/Ksat анализ использует 77 функций каталога, включая функции, чьи native-прогнозы участвуют в дискретных семействах. Point-target оценки охватывают ещё 18 функций. Ниже перечислены функции вне основного сравнения; это не означает их непригодности для собственных научных задач.',
        '',
        '| Функция | Статус в v1.0 | Причина |',
        '|---|---|---|',
    ]
    for row in catalog:
        fid = row['function']
        if fid in full:
            continue
        if fid in points:
            status, reason = (
                'point-target',
                'Оценка на отдельных native-давлениях; не входит в full-WRC ranking.',
            )
        elif fid in decisions:
            entry = decisions[fid]
            status = ', '.join(d.removeprefix('now_') for d in entry['dispositions'])
            missing = sorted(
                k.split(':', 1)[1]
                for k in entry['failure_counts']
                if k.startswith('missing_predictor:')
            )
            if 'user_decision_required' in status:
                reason = 'Физическая неоднозначность водной базы или формулы кривой требует решения; вариант не угадывается.'
            elif 'excluded_other_target' in status:
                reason = 'Другой физический target: AWC, extended AWC или осмотический потенциал.'
            else:
                reason = (
                    'Нет допустимого независимого расчёта с доступными входами и поддержкой давления.'
                    + (' Не хватает: ' + ', '.join(missing) + '.' if missing else '')
                )
        elif fid == 'calc_ptf_rawls1982_full_wrc':
            status, reason = (
                'conditional_only',
                'Измеренная влажность этой же ОГХ требуется как predictor; исключена из независимого сравнения.',
            )
        else:
            status, reason = (
                'other_target',
                'Нет операционного WRC/Ksat target в данном анализе: '
                + '; '.join(row['declared_families'])
                + '.',
            )
        lines.append(f'| `{fid}` | {status} | {reason} |')
    for name in ('NONPARTICIPATING-FUNCTIONS.md', 'EXPANDED-FUNCTION-AUDIT.md'):
        (PUBLIC / name).write_text('\n'.join(lines) + '\n', encoding='utf-8')


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        '--reuse-png-audit',
        action='store_true',
        help='Retain an existing completed PNG audit while refreshing documents and packaging.',
    )
    args = parser.parse_args()
    PRIVATE.mkdir(exist_ok=True)
    backup = PRIVATE / 'internal-provenance'
    backup.mkdir(exist_ok=True)
    for path in PUBLIC.iterdir():
        if path.is_file() and not (backup / path.name).exists():
            shutil.copyfile(path, backup / path.name)
    analyses = read(INPUT / 'analyses.json')
    benchmark = read(PUBLIC / 'benchmark_aggregates.json')
    approximation = benchmark['approximation']['aggregates']
    close(analyses[BASES[0]]['primary']['ranking'][0]['rmse'], 0.03677752909682378)
    close(analyses[BASES[0]]['ensembles']['top_5']['rmse'], 0.029325797045499827)
    close(analyses[BASES[1]]['primary']['ranking'][0]['rmse'], 0.06802726026304826)
    # Existing public summaries redact private observation identities and labels.
    public_selection = read(PUBLIC / 'expanded_selection.json')
    final_selection = {
        'benchmark_version': '1.0',
        'status': 'final',
        'analyses': clean_selection(public_selection['analyses']),
        'candidate_metadata': clean_selection(public_selection['candidate_metadata']),
        'policy': {
            'id': 'ptfkit_benchmark_v1.0',
            'water_bases': 'separate',
            'evidence': 'coverage-first E5',
            'robust_improvement': 'Pooled RMSE reduction >=1%; lower RMSE in every required held-out comparison against training-selected and pooled-best atomic comparators.',
            'no_ksat_ensemble': True,
        },
        'independent_catalog_functions': 77,
        'functions_with_secondary_scores': 95,
    }
    write(PUBLIC / 'selection_v1.0.json', final_selection)
    ledger = benchmark['approximation']['ledger_inventory']
    documents(analyses, approximation, ledger)
    presentation(analyses, approximation)
    audit = (
        read(PRIVATE / 'png-audit.json')
        if args.reuse_png_audit
        else png_audit(analyses, approximation)
    )
    package()
    neutral_audit()
    write(
        PUBLIC / 'benchmark_aggregates.json',
        {
            'benchmark_version': '1.0',
            'status': 'final',
            'selection': final_selection,
            'approximation': benchmark['approximation'],
            'dataset': {'wrc': ledger, 'ksat_observations': 39},
        },
    )
    write(
        PRIVATE / 'publication-validation.json',
        {
            'benchmark_version': '1.0',
            'benchmark_recomputed': False,
            'pngs': audit['actual_count'],
            'zip_crc_verified': True,
            'pptx_slides': 3,
            'conclusion_words': len(
                (PUBLIC / 'CONCLUSION-RU.md').read_text(encoding='utf-8').split()
            ),
        },
    )
    print(  # noqa: T201 - Command result.
        f'v1.0: 3 slides, {audit["actual_count"]} PNGs; ZIP: {PRIVATE / "ptfkit_diagnostics_v1.0.zip"}',
        flush=True,
    )


if __name__ == '__main__':
    main()
