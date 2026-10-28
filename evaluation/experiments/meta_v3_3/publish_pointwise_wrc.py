"""Rebuild the v1.0 handoff from every compatible experimental WRC coordinate.

Private measurements, predictions and diagnostic images remain under local/.
Public files contain aggregate results and the Russian presentation only.
"""

# ruff: noqa: INP001, E501, RUF001, PLR0915, PLR2004 - Scientific publication layouts.

from __future__ import annotations

import csv
import hashlib
import json
import shutil
import textwrap
from collections import defaultdict
from pathlib import Path
from statistics import median
from zipfile import ZIP_DEFLATED, ZipFile

import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from evaluation.experiments.meta_v3_3 import model_selection as original
from evaluation.experiments.meta_v3_3 import pointwise_wrc as wrc
from evaluation.experiments.meta_v3_3 import render_diagnostics as render
from evaluation.experiments.meta_v3_3.final_approximation import predict, reconstruct


ROOT = Path(__file__).resolve().parents[3]
PUBLIC = Path(__file__).with_name('final_results')
PRIVATE = ROOT / 'local/meta-final-v1_0-pointwise'
INPUT = ROOT / 'local/meta-expanded-benchmark'
BASES = wrc.BASES


def read(path: Path):  # noqa: ANN201 - Scientific JSON roots vary.
    return json.loads(path.read_bytes())


def write_json(path: Path, value: object) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8'
    )


def write_csv(path: Path, rows: list[dict]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(
            {
                k: json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v
                for k, v in row.items()
            }
            for row in rows
        )


def route(name: str, metadata: dict, catalog: dict) -> str:
    if name.startswith('ensemble:'):
        return 'ensemble'
    value = metadata[name]['route_type']
    targets = catalog.get(name, {}).get('wrc_targets', [])
    if any('axis_input' in t or t.get('adapter') == 'logistic_pressure_scale_v1' for t in targets):
        return 'continuous'
    if value == 'discrete_family_WRC' or len(catalog.get(name, {}).get('wrc_targets', [])) >= 3:
        return 'discrete-family'
    if value == 'continuous':
        return 'continuous'
    return 'parameter-derived'


def annotation(row: dict, basis: str, kind: str) -> str:
    return (
        f'Route: {kind}\nWater basis: {basis}\n\n'
        f'Points: {row["n_predicted_points"]}/{row["n_total_points"]} '
        f'({row["point_coverage"]:.1%})\n'
        f'Samples: {row["n_predicted_samples"]}/{row["n_total_samples"]} '
        f'({row["sample_coverage"]:.1%})\n\n'
        f'Sample-balanced RMSE: {row["sample_balanced_rmse"]:.6f}\n'
        f'Pointwise RMSE: {row["rmse"]:.6f}\n'
        f'Pointwise MAE: {row["mae"]:.6f}\nPointwise bias: {row["bias"]:+.6f}\n\n'
        f'Pressure: {row["pressure_min_kpa"]:.5g} to {row["pressure_max_kpa"]:.5g} kPa'
    )


def ptf_plots(points: list[dict], analyses: dict, metadata: dict, catalog: dict) -> list[dict]:
    rows = []
    for basis, analysis in analyses.items():
        folder, unit, label = render.basis_info(basis)
        all_points = points + analysis['ensemble_predictions']
        for summary in analysis['individual'] + analysis['ensembles']:
            name = summary['model']
            group = [p for p in all_points if p['basis'] == basis and p['function'] == name]
            kind = route(name, metadata, catalog)
            fig, ax = render.canvas(
                '\n'.join(textwrap.wrap(name, 100)), annotation(summary, label.lower(), kind)
            )
            render.scatter(ax, group, f'water content ({unit})')
            fig.text(
                0.09,
                0.035,
                'All predictable experimental points; ensembles: source-held-out forecasts.',
                fontsize=9,
            )
            relative = render.save(fig, Path('PTF') / f'{folder}_WRC' / f'{render.slug(name)}.png')
            rows.append(
                {
                    'category': 'PTF',
                    'basis': basis,
                    'model': name,
                    'route_type': kind,
                    **summary,
                    'file_path': relative,
                }
            )
    return rows


def ksat_plots(points: list[dict], observations: list[dict]) -> tuple[list[dict], dict]:
    groups, rows, results = defaultdict(list), [], {}
    for point in points:
        if point['basis'] == 'ksat':
            groups[point['function']].append(point)
    total = sum(p['basis'] == 'ksat' for p in observations)
    for name, group in sorted(groups.items()):
        metrics = original.metrics(group, ksat=True)
        blocks = {
            block: original.metrics([p for p in group if p['block'] == block], ksat=True)
            for block in sorted({p['block'] for p in group})
        }
        results[name] = {**metrics, 'block_sensitivity': blocks, 'total': total}
        note = (
            f'N: {len(group)}/{total} ({len(group) / total:.1%})\n\n'
            f'RMSE(log10): {metrics["rmse"]:.6f}\nMAE(log10): {metrics["mae"]:.6f}\n'
            f'Bias(log10): {metrics["bias"]:+.6f}\n\n'
            + '\n'.join(f'Within x{n}: {metrics[f"factor_{n}"]:.1%}' for n in (2, 5, 10))
            + '\n\nBlocks (N; RMSE log10):\n'
            + '\n'.join(f'{b}: {s["observations"]}; {s["rmse"]:.4f}' for b, s in blocks.items())
        )
        fig, ax = render.canvas('\n'.join(textwrap.wrap(name, 100)), note)
        render.scatter(ax, group, 'Ksat (m/s)', ksat=True)
        path = render.save(fig, Path('PTF/Ksat') / f'{render.slug(name)}.png')
        rows.append(
            {
                'category': 'PTF',
                'basis': 'ksat',
                'model': name,
                'route_type': 'Ksat',
                **metrics,
                'n_total_points': total,
                'point_coverage': len(group) / total,
                'file_path': path,
            }
        )
    return rows, results


def secondary_plots(points: list[dict], ledger: list[dict]) -> list[dict]:
    """Retain native point targets as clearly identified secondary diagnostics."""
    groups, rows = defaultdict(list), []
    for point in points:
        groups[point['basis'], point['function']].append(point)
    for (basis, name), group in sorted(groups.items()):
        denominator = [p for p in ledger if p['basis'] == basis]
        exact = []
        for p in group:
            matches = [
                o
                for o in denominator
                if wrc.sample_key(o) == wrc.sample_key(p)
                and np.isclose(o['pressure_kpa'], p['pressure_kpa'], rtol=1e-12, atol=1e-14)
            ]
            if matches:
                exact.append({**p, **matches[0], 'predicted': p['predicted']})
        measured = wrc.score(exact, denominator)
        metrics = wrc.score(group, denominator)
        folder, unit, label = render.basis_info(basis)
        note = (
            f'Secondary fixed-pressure targets\nWater basis: {label}\n'
            f'Target pairs: {len(group)}\n'
            f'Direct experimental points: {len(exact)}/{len(denominator)} ({measured["point_coverage"]:.1%})\n'
            f'Sampled curves: {metrics["n_predicted_samples"]}/{metrics["n_total_samples"]} ({metrics["sample_coverage"]:.1%})\n\n'
            f'SB RMSE (all target pairs): {metrics["sample_balanced_rmse"]:.6f}\n'
            f'RMSE (all target pairs): {metrics["rmse"]:.6f}\nMAE: {metrics["mae"]:.6f}\nBias: {metrics["bias"]:+.6f}\n\n'
            f'Pressure: {metrics["pressure_min_kpa"]:.5g} to {metrics["pressure_max_kpa"]:.5g} kPa\n\n'
            'Other target pairs use interpolation\ninside the measured curve support.\nExcluded from main WRC ranking.'
        )
        fig, ax = render.canvas('\n'.join(textwrap.wrap(name, 100)), note)
        render.scatter(ax, group, f'water content ({unit})')
        path = render.save(
            fig, Path('PTF') / f'{folder}_WRC/point_targets' / f'{render.slug(name)}.png'
        )
        rows.append(
            {
                'category': 'SECONDARY',
                'basis': basis,
                'model': name,
                'route_type': 'fixed-pressure point-target',
                **metrics,
                'point_coverage': measured['point_coverage'],
                'n_predicted_points': len(exact),
                'n_reconstructed_target_pairs': len(group) - len(exact),
                'metric_scope': 'all secondary target pairs, including reconstructed observations',
                'file_path': path,
            }
        )
    return rows


def approximation_plots(fits: list[dict], ledger: list[dict]) -> list[dict]:
    rows = []
    for basis in BASES:
        folder, unit, label = render.basis_info(basis)
        for family in render.FAMILIES:
            records = [r for r in fits if r['basis'] == basis and r['family'] == family]
            eligible = [r for r in records if r['status'] != 'ineligible']
            accepted = [r for r in records if r['relaxed_accepted']]
            points = [p for r in accepted for p in r['points']]
            summary = wrc.score(points, [p for p in ledger if p['basis'] == basis])
            note = (
                f'Family: {render.NAMES[family]}\nWater basis: {label.lower()} ({unit})\n\n'
                f'Eligible points: {sum(r["n"] for r in eligible)}\nEligible samples: {len(eligible)}\n'
                f'Fit success (relaxed): {len(accepted)}/{len(eligible)}\n'
                f'Fit success (strict): {sum(r["strict_accepted"] for r in records)}/{len(eligible)}\n'
                f'Sample coverage (all): {len(accepted)}/{len(records)} ({len(accepted) / len(records):.1%})\n'
                f'Point coverage (all): {summary["n_predicted_points"]}/{summary["n_total_points"]} ({summary["point_coverage"]:.1%})\n\n'
                f'Median RMSE: {median(r["rmse"] for r in accepted):.6f}\n'
                f'Median MAE: {median(r["mae"] for r in accepted):.6f}\n'
                f'Median bias: {median(r["bias"] for r in accepted):+.6f}\n\n'
                f'Validated observed support: {summary["pressure_min_kpa"]:.5g} to {summary["pressure_max_kpa"]:.5g} kPa'
            )
            fig, ax = render.canvas(f'Measured vs fitted: {render.NAMES[family]} / {label}', note)
            render.scatter(ax, points, f'water content ({unit})', fitted=True)
            fig.text(
                0.09,
                0.035,
                'Same-data descriptive fits; relaxed fits include flagged boundary/rank cases.',
                fontsize=9,
            )
            path = render.save(fig, Path('APPROXIMATION') / folder / f'{family}.png')
            rows.append(
                {
                    'category': 'APPROXIMATION',
                    'basis': basis,
                    'model': family,
                    **summary,
                    'eligible_points': sum(r['n'] for r in eligible),
                    'eligible_samples': len(eligible),
                    'fit_success': len(accepted),
                    'median_rmse': median(r['rmse'] for r in accepted),
                    'median_mae': median(r['mae'] for r in accepted),
                    'median_bias': median(r['bias'] for r in accepted),
                    'file_path': path,
                }
            )
    by_sample = defaultdict(list)
    for record in fits:
        by_sample[record['basis'], record['source_id'], record['sample_id']].append(record)
    for key, records in sorted(by_sample.items()):
        accepted = [r for r in records if r['relaxed_accepted']]
        if not accepted:
            continue
        basis, source, sample = key
        _, unit, label = render.basis_info(basis)
        points = accepted[0]['points']
        low, high = min(p['pressure_kpa'] for p in points), max(p['pressure_kpa'] for p in points)
        zero = min(p['pressure_kpa'] for p in points if p['pressure_kpa'] > 0) / 3
        h = np.geomspace(max(low, zero), high, 300)
        if low == 0:
            h = np.concatenate(([0], h))
        note = (
            f'Water basis: {label}\nPoints: {len(points)}/{len(points)} (100%)\n'
            f'Sample coverage: 1/1 (100%)\nFit success: {len(accepted)}/{len(records)} families\n'
            f'Observed support: {low:.4g} to {high:.4g} kPa\n\n'
        )
        for record in accepted:
            note += (
                f'{render.NAMES[record["family"]]}\nRMSE {record["rmse"]:.5f}; MAE {record["mae"]:.5f}\n'
                f'Bias {record["bias"]:+.5f}; {"strict" if record["strict_accepted"] else "relaxed only"}\n\n'
            )
        fig, ax = render.canvas('\n'.join(textwrap.wrap(f'{sample} / {source}', 100)), note)
        ax.scatter(
            [p['pressure_kpa'] or zero for p in points],
            [p['observed'] for p in points],
            s=20,
            color='black',
            label='Measured',
        )
        for record in accepted:
            family = record['family']
            stored = predict(
                family,
                record['parameters'],
                np.array([p['pressure_kpa'] for p in record['points']]),
                1 if basis == BASES[0] else 5,
            )
            if not np.allclose(
                stored, [p['predicted'] for p in record['points']], rtol=1e-10, atol=1e-12
            ):
                msg = 'Saved approximation predictions differ from curve parameters.'
                raise ValueError(msg)
            y = predict(family, record['parameters'], h, 1 if basis == BASES[0] else 5)
            ax.plot(
                np.where(h == 0, zero, h),
                y,
                label=render.NAMES[family],
                color=render.COLORS[family],
            )
        ax.set(xscale='log', xlabel='Suction magnitude (kPa)', ylabel=f'Water content ({unit})')
        ax.legend(fontsize=8)
        fig.text(
            0.09,
            0.035,
            f'Same-data approximation. Zero suction display coordinate: {zero:.4g} kPa; curves stay inside observed support.',
            fontsize=8,
        )
        path = render.save(fig, Path('APPROXIMATION/per_sample') / f'{render.sample_stem(key)}.png')
        rows.append(
            {
                'category': 'APPROXIMATION',
                'basis': basis,
                'model': 'all_families',
                'sample': sample,
                'source': source,
                'fit_success': len(accepted),
                'file_path': path,
            }
        )
    return rows


def coverage_plot(analyses: dict) -> None:
    fig, axes = plt.subplots(1, 2, figsize=(14, 6))
    for ax, (basis, analysis) in zip(axes, analyses.items(), strict=True):
        nominations = analysis['nominations']
        frontier = [r for r in analysis['individual'] if r['model'] in nominations['pareto']]
        for kind, rows, color, marker in (
            ('Individual', analysis['individual'], '#287889', 'o'),
            ('Source-held-out ensemble', analysis['ensembles'], '#bb6740', '^'),
        ):
            ax.scatter(
                [100 * r['point_coverage'] for r in rows],
                [r['sample_balanced_rmse'] for r in rows],
                label=kind,
                color=color,
                marker=marker,
                s=35,
            )
        frontier.sort(key=lambda r: r['point_coverage'])
        ax.plot(
            [100 * r['point_coverage'] for r in frontier],
            [r['sample_balanced_rmse'] for r in frontier],
            color='#187044',
            label='Individual Pareto frontier',
        )
        for label, row in (
            ('Accuracy', nominations['best_accuracy']),
            ('Compromise', nominations['compromise']),
        ):
            ax.annotate(
                f'{label}: {row["point_coverage"]:.1%}\nRMSE={row["sample_balanced_rmse"]:.4f}',
                (100 * row['point_coverage'], row['sample_balanced_rmse']),
                xytext=(6, 10),
                textcoords='offset points',
                fontsize=8,
            )
        ax.set(
            title=basis.replace('_water_content', '').capitalize(),
            xlabel='Experimental point coverage (%)',
            ylabel='Sample-balanced RMSE',
            xlim=(0, 105),
        )
        ax.grid(visible=True, alpha=0.2)
        ax.legend(fontsize=8)
    fig.suptitle('ptfkit v1.0: error and coverage on all predictable experimental points')
    fig.tight_layout()
    for path in (
        PUBLIC / 'wrc_accuracy_coverage.png',
        PRIVATE / 'diagnostics/wrc_accuracy_coverage.png',
    ):
        fig.savefig(path, dpi=180, facecolor='white')
    plt.close(fig)


def publication_figures(fits: list[dict]) -> None:
    figures = PUBLIC / 'figures'
    for filename in ('figure_1_ptf_selection.png', 'figure_1_slide.png'):
        shutil.copyfile(PUBLIC / 'wrc_accuracy_coverage.png', figures / filename)
    fig, axes = plt.subplots(1, 2, figsize=(13, 6))
    for ax, basis in zip(axes, BASES, strict=True):
        values, labels = [], []
        for family in render.FAMILIES:
            records = [r for r in fits if r['basis'] == basis and r['family'] == family]
            accepted = [r for r in records if r['relaxed_accepted']]
            values.append(median(r['rmse'] for r in accepted))
            labels.append(f'{render.NAMES[family]}\n{len(accepted)}/{len(records)} samples')
        bars = ax.bar(labels, values, color=[render.COLORS[f] for f in render.FAMILIES])
        ax.bar_label(bars, fmt='%.4f', padding=3)
        ax.set(
            title=basis.replace('_water_content', '').capitalize(),
            ylabel='Median sample RMSE (accepted fits)',
        )
        ax.tick_params(axis='x', labelsize=8)
    fig.suptitle('ptfkit v1.0: same-data approximation error and sample coverage')
    fig.tight_layout()
    fig.savefig(figures / 'figure_2_wrc_approximation.png', dpi=180, facecolor='white')
    plt.close(fig)
    # These specific obsolete exports contradict the PNG-only publication.
    for filename in (
        'figure_1_ptf_selection.svg',
        'figure_1_slide.svg',
        'figure_2_wrc_approximation.svg',
    ):
        path = (figures / filename).resolve()
        if path.is_relative_to(PUBLIC.resolve()):
            path.unlink(missing_ok=True)


def describe(row: dict | None) -> str:
    if row is None:
        return 'No eligible model'
    return (
        f'`{row["model"]}`: sample-balanced RMSE {row["sample_balanced_rmse"]:.6f}; '
        f'pointwise RMSE {row["rmse"]:.6f}; points {row["n_predicted_points"]}/{row["n_total_points"]} '
        f'({row["point_coverage"]:.1%}); samples {row["n_predicted_samples"]}/{row["n_total_samples"]} ({row["sample_coverage"]:.1%})'
    )


def describe_ru(row: dict | None) -> str:
    if row is None:
        return 'Нет модели с таким покрытием'
    return (
        f'`{row["model"]}`: SB RMSE {row["sample_balanced_rmse"]:.6f}; '
        f'обычная RMSE {row["rmse"]:.6f}; точки {row["n_predicted_points"]}/{row["n_total_points"]} '
        f'({row["point_coverage"]:.1%}); образцы {row["n_predicted_samples"]}/{row["n_total_samples"]} '
        f'({row["sample_coverage"]:.1%})'
    )


def documents(analyses: dict, inventory: dict, ksat: dict, fits: list[dict]) -> None:  # noqa: C901 - Report sections.
    lines = [
        '# ptfkit v1.0: all experimental WRC points',
        '',
        f'The ledger contains {inventory["observations"]} experimental points, {inventory["samples"]} samples and {inventory["sources"]} sources. Volumetric: 1069 points / 77 samples; gravimetric: 632 / 65. Fitted workbook parameters and reconstructed target values are excluded.',
        '',
        'Every PTF is evaluated on all compatible experimental coordinates with available predictors and valid pressure support. No minimum number of observed points per sample or common cohort is imposed. Coverage denominators are the complete ledger for the respective water basis. Bias is prediction minus measurement. R squared is descriptive.',
        '',
        'Sample-balanced RMSE = sqrt(mean over samples of mean squared point error within each sample). Pointwise RMSE uses equal point weights. Sample-balanced MAE and bias are also supplied. Models have different evaluation subsets; rankings describe this dataset and do not establish causal superiority on a common population.',
        '',
    ]
    russian = [
        '# Выводы ptfkit v1.0',
        '',
        'Для каждой PTF использованы все экспериментальные точки ОГХ, которые она способна предсказать. Модели сравниваются по ошибке и покрытию; объёмная и массовая влажность анализируются отдельно.',
        '',
        'База: 1701 экспериментальная точка, 142 образца, 19 источников. Объёмная влажность: 1069 точек; массовая: 632. Подогнанные и интерполированные значения не считаются независимыми измерениями.',
        '',
    ]
    for basis, analysis in analyses.items():
        n = analysis['nominations']
        best_ensemble = analysis['ensembles'][0] if analysis['ensembles'] else None
        lines += [f'## {basis}', '', '**Best accuracy:** ' + describe(n['best_accuracy']), '']
        for threshold, row in n['high_coverage'].items():
            lines += [f'**Best individual at coverage >= {threshold}%:** ' + describe(row), '']
        lines += [
            '**Accuracy-coverage compromise:** ' + describe(n['compromise']),
            '',
            n['compromise_policy'],
            '',
            '**Best source-held-out ensemble:** ' + describe(best_ensemble),
            '',
            '| Model | Points predicted / total | Coverage % | Samples predicted / total | Sample coverage % | SB RMSE | RMSE | MAE | Bias |',
            '|---|---:|---:|---:|---:|---:|---:|---:|---:|',
        ]
        for row in analysis['individual'] + analysis['ensembles']:
            lines.append(
                f'| {row["model"]} | {row["n_predicted_points"]}/{row["n_total_points"]} | {100 * row["point_coverage"]:.2f} | {row["n_predicted_samples"]}/{row["n_total_samples"]} | {100 * row["sample_coverage"]:.2f} | {row["sample_balanced_rmse"]:.6f} | {row["rmse"]:.6f} | {row["mae"]:.6f} | {row["bias"]:+.6f} |'
            )
        threshold = analysis['high_coverage_comparison_threshold_percent']
        reference_label = (
            f'highest attained threshold ({threshold}%)'
            if threshold
            else 'maximum available coverage (below 25%)'
        )
        lines += ['', f'Ensemble comparison with the best individual at {reference_label}:', '']
        if not analysis['ensemble_comparisons']:
            lines += [
                '',
                'No individual reaches 25% coverage; no qualifying high-coverage comparison.',
            ]
        for comparison in analysis['ensemble_comparisons']:
            lines.append(
                f'- {comparison["ensemble"]}: common-point SB RMSE {comparison["ensemble_common"]["sample_balanced_rmse"]:.6f} versus {comparison["individual_common"]["sample_balanced_rmse"]:.6f}; N={comparison["n_common_points"]}; lower error: {comparison["lower_common_error"]}; coverage at least individual: {comparison["coverage_at_least_individual"]}.'
            )
        ensemble_note = ''
        if best_ensemble:
            comparison = next(
                (
                    c
                    for c in analysis['ensemble_comparisons']
                    if c['ensemble'] == best_ensemble['model']
                ),
                None,
            )
            if comparison:
                verdict = (
                    comparison['lower_common_error'] and comparison['coverage_at_least_individual']
                )
                scope = (
                    f'при пороге {threshold}%'
                    if threshold
                    else 'при максимальном доступном покрытии (<25%)'
                )
                ensemble_note = f'Ансамбль улучшает ошибку при не меньшем покрытии относительно индивидуальной модели {scope}: {"да" if verdict else "нет"}. Сравнение ошибок на общих точках: {comparison["ensemble_common"]["sample_balanced_rmse"]:.6f} против {comparison["individual_common"]["sample_balanced_rmse"]:.6f}.'
        lines += ['']
        russian += [
            f'## {"Объёмная" if basis == BASES[0] else "Массовая"} ОГХ',
            '',
            'Минимальная ошибка: ' + describe_ru(n['best_accuracy']),
            '',
            *[
                f'Лучшая индивидуальная модель при покрытии ≥{t}%: ' + describe_ru(r) + '\n'
                for t, r in n['high_coverage'].items()
            ],
            'Компромисс ошибки и покрытия (эвристический выбор на границе Парето): '
            + describe_ru(n['compromise']),
            '',
            'Лучший ансамбль по проверке с исключением источника: ' + describe_ru(best_ensemble),
            '',
            'Ошибки получены на различных подмножествах; малое покрытие ограничивает интерпретацию минимальной RMSE. Ансамбли выбирают участников только по обучающим источникам.',
            '',
        ]
        if ensemble_note:
            russian += [ensemble_note, '']
    lines += [
        '## Ksat',
        '',
        'Ksat remains a separate benchmark without ensembles. The original common-cohort and maximum-coverage recommendations and source/unit caveats are retained: Ferrer-Julia gleyic cambisol sand on the 30-sample common set; Dane-Puckett on all 39 samples. Per-function log10 RMSE, MAE, bias, factor fractions and block sensitivity are saved in ksat_metrics.json and shown in the regenerated PNG.',
        '',
        '## Approximation and artifacts',
        '',
        'Four approximation families are fitted separately to each water basis. These are same-data fits, never independent PTF observations. Aggregated PNG use all accepted fits per family, with eligible counts and relaxed/strict success shown explicitly; per-sample curves stay inside observed support.',
        '',
        'See wrc_accuracy_coverage.csv, ptf_selection_summary.csv, selection_v1.0.json, PNG-AUDIT.md and wrc_accuracy_coverage.png. Private ledger, predictions, diagnostic_index.csv, all diagnostic PNG and ZIP are in local/meta-final-v1_0-pointwise/. Reproduce with the commands in ../README.md.',
        '',
    ]
    lines += [
        '| Basis | Approximation family | Accepted / total samples | Median RMSE | Median MAE | Median bias |',
        '|---|---|---:|---:|---:|---:|',
    ]
    for basis in BASES:
        candidates = []
        for family in render.FAMILIES:
            records = [r for r in fits if r['basis'] == basis and r['family'] == family]
            accepted = [r for r in records if r['relaxed_accepted']]
            loss = median(r['rmse'] for r in accepted)
            candidates.append((loss, family))
            lines.append(
                f'| {basis} | {family} | {len(accepted)}/{len(records)} | {loss:.6f} | {median(r["mae"] for r in accepted):.6f} | {median(r["bias"] for r in accepted):+.6f} |'
            )
        best_loss, best_family = min(candidates)
        russian += [
            f'Аппроксимация ({basis}): минимальная медианная RMSE у {render.NAMES[best_family]}, {best_loss:.6f}; это подгонка к тем же измерениям.',
            '',
        ]
    lines += [
        '',
        'Legacy expanded_selection.json and point_target_selection.json retain historical cohort analyses. They are supplementary records, not the current primary WRC ranking.',
        '',
    ]
    russian += [
        'Ksat сохраняется отдельным: gleyic cambisol sand на общем наборе 30 образцов; Dane–Puckett при покрытии 39/39. Ансамбль Ksat не строится. Все выводы относятся к данному набору; лабораторные методы и допущения E5 ограничивают переносимость результатов.',
        '',
    ]
    (PUBLIC / 'FINAL-RESULTS.md').write_text('\n'.join(lines), encoding='utf-8')
    (PUBLIC / 'CONCLUSION-RU.md').write_text('\n'.join(russian), encoding='utf-8')
    methodology = """# Ансамбли ОГХ ptfkit v1.0

Объёмная и массовая влажность анализируются отдельно. Единица оценки — исходная
экспериментальная пара давления и влажности. Подогнанные параметры и
интерполированные target-значения в основной benchmark не включаются.

В каждом fold исключается целый источник. Только на остальных источниках PTF
ранжируются по sample-balanced RMSE: сначала средняя квадратичная ошибка внутри
каждого образца, затем среднее по образцам и квадратный корень. При равенстве
учитываются абсолютный sample-balanced bias и идентификатор функции. Каждая PTF
оценивается на всех доступных ей обучающих точках; их покрытия могут различаться.

Top-3 и top-5 выбирают до трёх/пяти лучших PTF с обучающими прогнозами. Для median
и all-available mean берутся все PTF с обучающими прогнозами. На каждой отложенной
точке используются только доступные участники выбранного состава: среднее или
медиана; один участник даёт свой прогноз; ноль участников означает пропуск.
Состав, fold и число участников каждой точки сохраняются в приватной таблице.
Источник целиком отсутствует в выборе состава; отложенные target-значения
используются только для итоговой оценки. Общая когорта не требуется.

Итоговые метрики рассчитаны по объединённым отложенным прогнозам. Покрытие
рассчитывается относительно полного experimental ledger своей водной базы.
Сравнения с индивидуальной PTF также приводятся на пересечении прогнозируемых
точек; различия покрытия всегда показываются отдельно. Это описательная проверка
на данном наборе источников, без утверждения статистической значимости.
Модели с узким покрытием могут иметь малую ошибку; поэтому сохраняются пороги
25%, 50%, 75% и граница Парето. Для практического компромисса выбирается минимальная
RMSE на границе Парето при самом высоком достижимом пороге (75%, затем 50%, затем
25%). Если ни один порог не достигнут, выбирается максимальное доступное покрытие
на границе Парето. Этот выбор не означает единственного научно оптимального решения.
Ksat оценивается отдельно; ансамбль Ksat не строится.
"""
    for filename in (
        'ENSEMBLE-METHODOLOGY-RU.md',
        'ENSEMBLE-METHODOLOGY.md',
        'Как делали ансамбль.md',
    ):
        (PUBLIC / filename).write_text(methodology, encoding='utf-8')
    write_json(PUBLIC / 'ksat_metrics.json', ksat)


def presentation(analyses: dict) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.5))
    box.text_frame.text = 'ptfkit v1.0: ошибка и покрытие экспериментальных ОГХ'
    box.text_frame.paragraphs[0].font.size = Pt(26)
    slide.shapes.add_picture(
        str(PUBLIC / 'wrc_accuracy_coverage.png'), Inches(0.3), Inches(0.9), width=Inches(8.2)
    )
    box = slide.shapes.add_textbox(Inches(8.7), Inches(1), Inches(4.2), Inches(6.2))
    frame = box.text_frame
    frame.word_wrap = True
    for basis, analysis in analyses.items():
        paragraph = frame.add_paragraph()
        paragraph.text = 'Объёмная ОГХ' if basis == BASES[0] else 'Массовая ОГХ'
        paragraph.font.size = Pt(19)
        paragraph.font.bold = True
        n = analysis['nominations']
        high = next(
            (n['high_coverage'][t] for t in ('75', '50', '25') if n['high_coverage'][t]), None
        )
        for label, row in (
            ('Минимальная ошибка', n['best_accuracy']),
            ('Высокое покрытие', high),
            ('Лучший ансамбль', analysis['ensembles'][0] if analysis['ensembles'] else None),
        ):
            p = frame.add_paragraph()
            p.text = (
                label
                + ': '
                + (
                    f'{row["model"]}\nSB RMSE {row["sample_balanced_rmse"]:.4f}; покрытие {row["point_coverage"]:.1%}'
                    if row
                    else 'нет модели ≥25%'
                )
            )
            p.font.size = Pt(12)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.2), Inches(6.6))
    box.text_frame.word_wrap = True
    for text in (
        'Как считали',
        '1701 экспериментальная точка: 1069 объёмных, 632 массовых; 142 образца, 19 источников.',
        'Для каждой PTF использованы все экспериментальные точки ОГХ, которые она способна предсказать; затем модели сравнивались по ошибке и покрытию.',
        'Sample-balanced RMSE даёт образцам равный вес. Покрытие показано от полного ledger каждой водной базы. Разные подмножества ограничивают сравнение точности.',
        'Ансамбли: top-3, top-5, медиана, среднее всех доступных. Состав выбирается только на обучающих источниках; отложенный источник полностью исключается. В каждой точке усредняются доступные участники.',
        'Ksat отдельно: gleyic cambisol sand на общем наборе 30 образцов; Dane–Puckett на 39/39. Сохраняются log10-метрики, множители ×2/×5/×10 и анализ блоков. Ансамбль Ksat не строится.',
        'Аппроксимации VG–Mualem, VG m=1, Brooks–Corey, Clapp: описательная подгонка к измерениям. Они не являются независимыми прогнозами PTF.',
    ):
        p = box.text_frame.add_paragraph()
        p.text = text
        p.font.size = Pt(22 if text == 'Как считали' else 18)
        p.space_after = Pt(14)
    for filename in ('ptfkit_benchmark_summary.pptx', 'ptfkit_benchmark_summary_RU_v1.0.pptx'):
        deck.save(str(PUBLIC / filename))


def audit_png(rows: list[dict], expected: set[str]) -> dict:
    directory = PRIVATE / 'diagnostics'
    actual = {p.relative_to(directory).as_posix() for p in directory.rglob('*.png')}
    invalid = []
    for relative in sorted(actual):
        try:
            with Image.open(directory / relative) as png:
                png.verify()
            with Image.open(directory / relative) as png:
                png.load()
                if png.format != 'PNG' or png.width < 1200 or png.height < 800:
                    invalid.append(relative)
        except OSError:
            invalid.append(relative)
    missing, stale = sorted(expected - actual), sorted(actual - expected)
    if missing or stale or invalid:
        msg = f'PNG audit failed: missing={missing}, stale={stale}, invalid={invalid}'
        raise ValueError(msg)
    report = {
        'expected_ptf': sum(p.startswith('PTF/') for p in expected),
        'actual_ptf': sum(p.startswith('PTF/') for p in actual),
        'expected_approximation': sum(p.startswith('APPROXIMATION/') for p in expected),
        'actual_approximation': sum(p.startswith('APPROXIMATION/') for p in actual),
        'missing': missing,
        'stale': stale,
        'invalid': invalid,
        'regenerated': sorted(actual),
        'index_rows': len(rows),
        'sha256': {
            p: hashlib.sha256((directory / p).read_bytes()).hexdigest() for p in sorted(actual)
        },
    }
    write_json(PRIVATE / 'png-audit.json', report)
    lines = [
        '# ptfkit v1.0 PNG audit',
        '',
        f'Expected / actual PTF PNG: {report["expected_ptf"]} / {report["actual_ptf"]}.',
        f'Expected / actual approximation PNG: {report["expected_approximation"]} / {report["actual_approximation"]}.',
        'One additional accuracy-coverage PNG.',
        '',
        'Missing: none. Stale/outdated: none. Invalid/corrupt: none.',
        '',
        'Every file was regenerated from current pointwise forecasts or current saved fits. Expected paths are independently enumerated from scored model IDs, all eight basis/family combinations and every sample with an accepted fit. The image index is checked against that set. PNG headers, complete decoding and dimensions are verified. SHA256 per image is saved in local/meta-final-v1_0-pointwise/png-audit.json.',
        '',
        'Regenerated files:',
        '',
        *['- ' + p for p in sorted(actual) if '/per_sample/' not in p],
        '',
        'Per-sample filenames contain private identities; the complete regenerated-file list is in local/meta-final-v1_0-pointwise/png-audit.json.',
        '',
    ]
    (PUBLIC / 'PNG-AUDIT.md').write_text('\n'.join(lines), encoding='utf-8')
    return report


def main() -> None:
    PRIVATE.mkdir(parents=True, exist_ok=True)
    (PRIVATE / 'diagnostics').mkdir(exist_ok=True)
    render.OUTPUT = PRIVATE / 'diagnostics'
    ledger, ledger_audit = reconstruct(ROOT / 'local/meta_v3.3.xlsx')
    raw = read(INPUT / 'predictions.json')
    historical_ksat = read(INPUT / 'baseline/selection_v1.0.json')['analyses']['ksat']
    points = wrc.align_predictions(raw, ledger)
    metadata = read(INPUT / 'candidate-metadata.json')
    frozen = read(ROOT / 'evaluation/datasets/meta_v3_3.model-selection-summary.json')
    catalog = {r['function']: r for r in frozen['catalog']}
    analyses = {
        basis: wrc.analyze(
            [p for p in points if p['basis'] == basis], [p for p in ledger if p['basis'] == basis]
        )
        for basis in BASES
    }
    fit_artifact = read(ROOT / 'local/meta-final-benchmark/approximation-fits.json')
    workbook_hash = hashlib.sha256((ROOT / 'local/meta_v3.3.xlsx').read_bytes()).hexdigest()
    if (
        fit_artifact['workbook_sha256'] != workbook_hash
        or frozen['input_sha256']['workbook'] != workbook_hash
    ):
        msg = 'Workbook fingerprint differs from the frozen inputs.'
        raise ValueError(msg)
    fits = fit_artifact['results']
    ledger_index = {p['id']: p for p in ledger}
    for record in fits:
        for p in record.get('points', []):
            if any(
                p[k] != ledger_index[p['id']][k]
                for k in ('basis', 'source_id', 'sample_id', 'observed', 'pressure_kpa')
            ):
                msg = 'Saved fit target differs from the experimental ledger.'
                raise ValueError(msg)
    write_json(
        PRIVATE / 'experimental-wrc-ledger.json',
        {'workbook_sha256': workbook_hash, 'audit': ledger_audit, 'observations': ledger},
    )
    write_csv(PRIVATE / 'experimental_wrc_ledger.csv', ledger)
    all_forecasts = points + [p for a in analyses.values() for p in a['ensemble_predictions']]
    write_csv(PRIVATE / 'wrc_pointwise_predictions.csv', all_forecasts)
    accuracy = [
        {'basis': b, 'kind': kind, 'route_type': route(r['model'], metadata, catalog), **r}
        for b, a in analyses.items()
        for kind, collection in (('individual', a['individual']), ('ensemble', a['ensembles']))
        for r in collection
    ]
    for name in ('wrc_accuracy_coverage.csv', 'ptf_selection_summary.csv'):
        write_csv(PUBLIC / name, accuracy)
    rows = ptf_plots(points, analyses, metadata, catalog)
    ksat_rows, ksat = ksat_plots(raw, read(ROOT / 'local/meta-diagnostic-replay/observations.json'))
    rows.extend(ksat_rows)
    secondary = secondary_plots(read(INPUT / 'point-target-predictions.json'), ledger)
    rows.extend(secondary)
    write_csv(PUBLIC / 'point_target_summary.csv', secondary)
    write_csv(
        PUBLIC / 'ptf_selection_summary.csv',
        accuracy
        + [
            {
                'basis': 'ksat',
                'kind': 'individual',
                'model': r['model'],
                'n_predicted_points': r['observations'],
                'n_total_points': r['n_total_points'],
                'point_coverage': r['point_coverage'],
                'rmse': r['rmse'],
                'mae': r['mae'],
                'bias': r['bias'],
                'metric_unit': 'log10_ratio',
                **{f'factor_{n}': r[f'factor_{n}'] for n in (2, 5, 10)},
            }
            for r in ksat_rows
        ],
    )
    rows.extend(approximation_plots(fits, ledger))
    coverage_plot(analyses)
    publication_figures(fits)
    expected = {
        f'PTF/{render.basis_info(b)[0]}_WRC/{render.slug(r["model"])}.png'
        for b, a in analyses.items()
        for r in a['individual'] + a['ensembles']
    }
    expected |= {f'PTF/Ksat/{render.slug(n)}.png' for n in ksat}
    expected |= {
        f'PTF/{render.basis_info(b)[0]}_WRC/point_targets/{render.slug(n)}.png'
        for b, n in {
            (p['basis'], p['function']) for p in read(INPUT / 'point-target-predictions.json')
        }
    }
    expected |= {
        f'APPROXIMATION/{render.basis_info(b)[0]}/{family}.png'
        for b in BASES
        for family in render.FAMILIES
    }
    expected |= {
        f'APPROXIMATION/per_sample/{render.sample_stem((r["basis"], r["source_id"], r["sample_id"]))}.png'
        for r in fits
        if r['relaxed_accepted']
    }
    indexed = {r['file_path'] for r in rows}
    if indexed != expected:
        msg = 'Diagnostic index differs from independently enumerated scientific outputs.'
        raise ValueError(msg)
    expected.add('wrc_accuracy_coverage.png')
    rows.append(
        {
            'category': 'SUMMARY',
            'model': 'accuracy_coverage',
            'file_path': 'wrc_accuracy_coverage.png',
        }
    )
    write_csv(PRIVATE / 'diagnostics/diagnostic_index.csv', rows)
    write_csv(PRIVATE / 'diagnostic_index.csv', rows)
    # The public index omits private sample rows; the complete one accompanies the ZIP.
    write_csv(
        PUBLIC / 'diagnostic_index.csv', [r for r in rows if '/per_sample/' not in r['file_path']]
    )
    audit = audit_png(rows, expected)
    public_analysis = {
        b: {k: v for k, v in a.items() if k != 'ensemble_predictions'} for b, a in analyses.items()
    }
    result = {
        'benchmark_version': 'v1.0',
        'method': 'all_experimental_points',
        'workbook_sha256': workbook_hash,
        'ledger_inventory': ledger_audit['inventory'],
        'analyses': public_analysis,
        'candidate_metadata': metadata,
        'ksat': ksat,
        'historical_ksat_analysis': historical_ksat,
        'policy': {
            'ranking': 'sample_balanced_rmse',
            'ensembles': 'leave_one_source_out_training_only',
            'coverage_denominator': 'all experimental observations and samples of the same water basis',
            'secondary_reconstructed_targets': 'excluded from primary experimental benchmark',
        },
    }
    write_json(PUBLIC / 'selection_v1.0.json', result)
    write_json(PRIVATE / 'selection_v1.0.json', result)
    write_json(
        PUBLIC / 'benchmark_aggregates.json',
        {
            'benchmark_version': 'v1.0',
            'status': 'pointwise_complete',
            'selection': result,
            'dataset': ledger_audit['inventory'],
            'approximation': read(
                ROOT / 'local/meta-final-benchmark/approximation-aggregates.json'
            ),
        },
    )
    documents(analyses, ledger_audit['inventory'], ksat, fits)
    presentation(analyses)
    readme = """# ptfkit v1.0 pointwise diagnostics

PTF/volumetric_WRC and PTF/gravimetric_WRC contain every scored individual
route and four source-held-out ensembles. Each PNG states error, point/sample
coverage, water basis and observed pressure range. Equal x/y limits and a 1:1
line apply. Ksat uses log axes and x2/x5/x10 bands, with separate log metrics
and block sensitivity. Ksat has no ensemble.

APPROXIMATION contains all eight family/basis aggregates and every sample with
an accepted fit. All family-specific accepted fits are used; strict and relaxed
success are distinguished. Fits describe measured data; they are not PTF
predictions. Per-sample fit curves stay within measured support.

wrc_accuracy_coverage.png highlights the individual Pareto frontier.
diagnostic_index.csv is the complete image index. png-audit.json in the parent
directory records independently expected counts, file integrity and SHA256.
All images were regenerated for the pointwise benchmark; no SVG is packaged.
Secondary fixed-pressure PNG distinguish direct experimental coverage from
reconstructed target pairs. Their errors refer to all secondary target pairs;
they never enter the primary experimental-point ranking or ensembles.
This directory contains private sample identities and must stay under local/.
"""
    (PRIVATE / 'diagnostics/README.md').write_text(readme, encoding='utf-8')
    archive = PRIVATE / 'ptfkit_diagnostics_v1.0.zip'
    with ZipFile(archive, 'w', compression=ZIP_DEFLATED) as bundle:
        for file in sorted((PRIVATE / 'diagnostics').rglob('*')):
            if file.is_file():
                bundle.write(file, file.relative_to(PRIVATE / 'diagnostics').as_posix())
    with ZipFile(archive) as bundle:
        if bundle.testzip() is not None:
            msg = 'Diagnostic archive CRC check failed.'
            raise ValueError(msg)
    write_json(
        PUBLIC / 'publication_v1.0_validation.json',
        {
            'method': result['method'],
            'workbook_sha256': workbook_hash,
            'png_counts': {k: v for k, v in audit.items() if isinstance(v, int)},
            'missing_png': [],
            'invalid_png': [],
            'stale_png': [],
            'zip_crc_valid': True,
            'private_output': PRIVATE.relative_to(ROOT).as_posix(),
        },
    )
    # Include current reports, figures and deck in the final public inventory.
    write_json(
        PUBLIC / 'artifact-hashes.json',
        {
            p.relative_to(PUBLIC).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(PUBLIC.rglob('*'))
            if p.is_file() and p.name != 'artifact-hashes.json'
        },
    )
    print(  # noqa: T201 - Command result.
        f'Pointwise v1.0 complete: {len(ledger)} observations; {len(expected)} PNG; {archive}',
        flush=True,
    )


if __name__ == '__main__':
    main()
