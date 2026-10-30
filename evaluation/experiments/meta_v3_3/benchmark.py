"""Reproduce workbook-wide accounting, catalog exclusions and reviewed native evaluation."""

# ruff: noqa: INP001 - Standalone research entry point.

from __future__ import annotations

import argparse
import hashlib
import json
import runpy
from collections import Counter
from pathlib import Path
from typing import cast

from openpyxl import load_workbook
from ptfkit_eval import (
    audit_predictor_workbook,
    export_normalized_retention_dataset,
    load_dataset_mapping,
    normalize_retention_workbook,
    read_retention_run_report,
)
from ptfkit_eval.artifacts import artifact_value
from ptfkit_eval.documents import canonical_json, document_hash, plain

from evaluation.experiments.meta_v3_3.class_comparison import run as run_class_comparison
from evaluation.experiments.meta_v3_3.ensembles import run as run_ensembles
from evaluation.experiments.meta_v3_3.function_audit import audit as audit_functions
from evaluation.experiments.meta_v3_3.function_audit import render as render_function_audit
from evaluation.experiments.meta_v3_3.function_audit import target_pools
from evaluation.experiments.meta_v3_3.reporting import render_report
from evaluation.experiments.meta_v3_3.scope import workbook_scope
from evaluation.experiments.meta_v3_3.source_audit import priority_source_audit
from evaluation.experiments.meta_v3_3.summaries import (
    approximation_summary,
    blocker_impact,
    catalog_support,
    count_only_checkpoint,
    eligibility_matrix,
    function_priorities,
    reconcile_rows,
    region_metrics,
    source_readiness,
)
from ptfkit import get_manifest


def _write(output: Path, name: str, document: object) -> None:
    with (output / name).open('xb') as stream:
        stream.write(canonical_json(document))


def _function_coverage(matrix: list[dict], coverage: dict) -> list[dict]:
    records = []
    for function in coverage['functions']:
        rows = [r for r in matrix if r['function'] == function['id']]
        evaluated = {r['master_key'] for r in rows if r['status'] == 'eligible_and_evaluated'}
        keys = {r['master_key'] for r in rows}
        records.append(
            {
                'function': function['id'],
                'evaluated_samples': len(evaluated),
                'eligible_samples': len(evaluated),
                'blocked_samples': len(keys - evaluated),
                'reason_counts_by_sample_target': dict(
                    sorted(Counter(reason for r in rows for reason in r['reasons']).items())
                ),
            }
        )
    return records


def run(  # noqa: PLR0915 - One run transaction.
    workbook: Path, paper: Path, output: Path, *, checkpoint: Path | None = None
) -> None:
    """Reuse existing audited normalization, predictor evaluation and fitting implementations."""
    root = Path(__file__).resolve().parents[3]
    experiments = root / 'evaluation/experiments'
    mapping = load_dataset_mapping(root / 'evaluation/datasets/meta_v3_3.yaml')
    coverage = runpy.run_path(str(experiments / 'catalog_coverage/audit.py'))['audit'](root)
    manifest = get_manifest()
    generated = json.loads((root / 'targets/ptfkit-py/src/ptfkit/_manifest.json').read_bytes())
    if document_hash(manifest) != document_hash(generated):
        message = 'Installed and generated manifests must agree before benchmarking.'
        raise ValueError(message)
    normalized = normalize_retention_workbook(workbook, mapping)
    predictor_result = audit_predictor_workbook(workbook, mapping, manifest=manifest)
    predictors = cast('dict', artifact_value(predictor_result))
    normalization = cast('dict', artifact_value(normalized))
    rows = reconcile_rows(normalization)
    followup_path = Path(__file__).with_name('source-followup.json')
    followup = json.loads(followup_path.read_bytes())
    with workbook.open('rb') as stream:
        book = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = book[normalized.summary.sheet_name]
            raw_sources = {
                number: values[0]
                for number, values in enumerate(
                    sheet.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True), 2
                )
            }
            scope = workbook_scope(
                book,
                cast('dict', plain(mapping.document)),
                rows,
                {r['sample_key'] for r in predictors['eligibility']},
            )
            priority_audit = priority_source_audit(book, followup)
            master_name = next(
                s['name']
                for s in cast('dict', plain(mapping.document))['worksheets']
                if s['id'] == 'samples_master'
            )
            candidate_values = {
                f'{master_name}!{number}': {
                    chr(65 + index): value for index, value in enumerate(values)
                }
                for number, values in enumerate(
                    book[master_name].iter_rows(min_row=2, max_col=24, values_only=True), 2
                )
            }
        finally:
            book.close()
    for row in rows:
        row['raw_source'] = raw_sources[row['row']]
    sources = source_readiness(cast('dict', plain(mapping.document)), rows)
    followup_index = {record['source_id']: record for record in followup['sources']}
    master_followup = {r['raw_source']: r for r in followup['additional_master_blocks']}
    for block in scope['unreviewed_master_sources']:
        block['followup'] = master_followup.get(block['raw_source'])
    for source in sources:
        source['followup'] = followup_index.get(source['source_id'])
        source['canonicalization_claim_status'] = (
            'exactly_derived' if source['normalized_observations'] else 'unresolved'
        )
    output.mkdir(parents=True, exist_ok=False)
    datasets = []
    for name, quantity in (
        ('volume', 'volumetric_water_content'),
        ('mass', 'gravimetric_water_content'),
    ):
        exported = export_normalized_retention_dataset(
            normalized,
            dataset_id=f'meta-v3-3-reviewed-{name}',
            material_basis='unspecified',
            water_quantity=quantity,
        )
        dataset = cast('dict', plain(exported.dataset.document))
        datasets.append(dataset)
        _write(output, f'canonical-{name}.json', dataset)
    _write(output, 'row-accounting.json', rows)
    _write(output, 'normalization.json', normalization['summary'])
    _write(output, 'predictor-audit.json', predictors)
    _write(output, 'implemented-functions.json', coverage)
    _write(output, 'workbook-scope.json', scope)
    _write(output, 'source-followup.json', followup)
    _write(output, 'priority-source-audit.json', priority_audit)
    review = root / 'evaluation/datasets/meta_v3_3.clapp-source-review.json'
    artifacts = runpy.run_path(str(experiments / 'panina_clapp/run.py'))['run'](
        output / 'canonical-volume.json', review, paper, output / 'native-clapp'
    )
    read_retention_run_report(artifacts)
    runpy.run_path(str(experiments / 'panina_clapp/report.py'))['publish'](
        artifacts, review, output / 'native-decision'
    )
    native = json.loads((output / 'native-decision/summary.json').read_bytes())
    comparison_artifacts, comparison = run_class_comparison(
        output / 'native-clapp', output / 'native-class-comparison'
    )
    ensemble = run_ensembles(output / 'native-class-comparison', output / 'ensembles')
    overlay = json.loads(review.read_bytes())
    for source in sources:
        selected = [
            s for s in overlay['selected_samples'] if s['source_id'] == source['master_source']
        ]
        if selected:
            source['reviewed_profile_count'] = len({s['profile_id'] for s in selected})
            source['reviewed_overlay'] = {
                'sample_count': len(selected),
                'physical_claim_status': 'directly_reported',
                'material_claim_status': 'exactly_derived_from_reviewed_source_descriptions',
                'review_sha256': hashlib.sha256(review.read_bytes()).hexdigest(),
            }
    points = json.loads((artifacts / 'point-decisions.json').read_bytes())
    evaluation_plan = json.loads((comparison_artifacts / 'evaluation-plan.json').read_bytes())
    functions = {
        candidate['id']: candidate['function'] for candidate in evaluation_plan['candidates']
    }
    eligibility = json.loads((comparison_artifacts / 'sample-eligibility.json').read_bytes())
    evaluated = [
        {**r, 'function': functions[r['candidate_id']], 'target': 'volumetric_water_content'}
        for r in eligibility
        if r['eligible_point_count']
    ]
    support = catalog_support(cast('dict', plain(manifest)))
    _write(output, 'catalog-support.json', support)
    matrix = eligibility_matrix(coverage, predictors, datasets, evaluated, support)
    function_audit = audit_functions(
        coverage, matrix, cast('dict', plain(mapping.document)), support, candidate_values
    )
    function_audit['target_pools'] = target_pools(coverage, support, ensemble)
    _write(output, 'function-audit.json', function_audit)
    (output / 'function-audit.md').write_text(
        render_function_audit(function_audit, ensemble), encoding='utf-8', newline='\n'
    )
    impact = blocker_impact(matrix)
    priorities = function_priorities(coverage, matrix)
    _write(output, 'function-priorities.json', priorities)
    _write(output, 'blocker-impact.json', impact)
    _write(output, 'eligibility-matrix.json', matrix)
    _write(
        output, 'exclusions.json', [r for r in matrix if r['status'] != 'eligible_and_evaluated']
    )
    runpy.run_path(str(experiments / 'measured_retention/fit.py'))['main'](
        [str(output / 'canonical-volume.json'), '--output', str(output / 'approximation-fits.json')]
    )
    fits = json.loads((output / 'approximation-fits.json').read_bytes())
    approximations = approximation_summary(fits)
    unsupported = [
        {
            'source_id': sample['source_id'],
            'sample_id': sample['sample_id'],
            'family': family['family'],
            'status': 'unsupported_target',
            'reason': 'registered_family_requires_volumetric_water_content_no_density_conversion',
        }
        for sample in datasets[1]['samples']
        for family in approximations['families']
    ]
    _write(output, 'unsupported-approximation-targets.json', unsupported)
    diagnostics = region_metrics([p['comparison'] for p in points if p['eligible']])
    summary = {
        'schema_version': 1,
        'protocol': 'meta_v3_3_benchmark_v1',
        'status': 'source_review_incomplete',
        'workbook_sha256': normalized.summary.sha256,
        'mapping_sha256': document_hash(mapping.document),
        'manifest_sha256': document_hash(manifest),
        'followup_sha256': hashlib.sha256(followup_path.read_bytes()).hexdigest(),
        'implementation_sha256': {
            path.relative_to(root).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in [
                *Path(__file__).parent.glob('*.py'),
                Path(__file__).with_name('report-template.md'),
                experiments / 'measured_retention/fit.py',
                experiments / 'panina_clapp/run.py',
                experiments / 'panina_clapp/report.py',
            ]
        },
        'workbook_scope': {key: value for key, value in scope.items() if key != 'rows'},
        'row_accounting': normalization['summary'],
        'sources': sources,
        'functions': _function_coverage(matrix, coverage),
        'matrix_record_count': len(matrix),
        'blocker_impact': impact,
        'priority_source_audit': [
            {
                'raw_source': source['raw_source'],
                'master_row_count': source['master_row_count'],
                'original_row_count': source['original_row_count'],
                'exact_id_count': sum(
                    r['status'] == 'exact_literal_id' for r in source['crosswalk']
                ),
                'ordered_vector_match_count': sum(
                    r['ordered_values_equal'] is True for r in source['crosswalk']
                ),
                'duplicate_vector_group_count': len(source['duplicate_master_vectors']),
                'unique_nonempty_vector_count': source['unique_nonempty_vector_count'],
                'nonempty_reported_water_vector_count': source[
                    'nonempty_reported_water_vector_count'
                ],
                'reused_water_multiset_group_count': len(source['reused_water_value_multisets']),
                'released_observation_count': source['released_observation_count'],
                'remaining_claims': source['remaining_claims'],
            }
            for source in priority_audit['sources']
        ],
        'native_evaluation': native,
        'class_comparison': comparison,
        'function_audit': function_audit,
        'ensemble_comparison': ensemble,
        'function_priorities': priorities,
        'approximations': approximations,
        'pressure_diagnostics': diagnostics,
        'remaining_blockers': sorted(
            (s for s in sources if s['excluded_observations']),
            key=lambda s: (-s['excluded_observations'], s['source_id']),
        ),
    }
    _write(output, 'source-readiness.json', sources)
    _write(output, 'approximation-summary.json', approximations)
    _write(output, 'pressure-diagnostics.json', diagnostics)
    selection = {
        name: native[name]
        for name in (
            'selected_in_fold',
            'selection_stability',
            'strategy_outer_performance',
            'final_recommended_atomic',
            'final_recommended_ensemble',
            'ensemble_execution',
        )
    }
    selection['sources'] = [
        {
            'source_id': source['source_id'],
            'status': 'insufficient_evidence',
            'reason': 'one_reviewed_profile_and_two_atomic_candidates'
            if source.get('reviewed_overlay')
            else 'no_source_reviewed_compatible_predictions',
            'final_recommended_atomic': None,
            'final_recommended_ensemble': None,
        }
        for source in sources
    ]
    for name in ('final_recommended_atomic', 'final_recommended_ensemble', 'selection_stability'):
        selection[name] = comparison['selection']
    _write(output, 'selection.json', selection)
    _write(
        output,
        'uncertainty.json',
        {
            'performance': [c['performance_uncertainty'] for c in native['cohorts']],
            'approximation': {
                'status': 'conditional_local_linear_iid_student_t',
                'coverage_qualification': 'not_validated',
                'artifact': 'approximation-fits.json',
            },
            'ptf_individual_prediction_interval': {'status': 'not_estimated'},
        },
    )
    _write(output, 'summary.json', summary)
    (output / 'report.md').write_text(render_report(summary), encoding='utf-8', newline='\n')
    # Publish the inventory last: absence of this file means the run is incomplete.
    inventory = {
        p.relative_to(output).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(output.rglob('*'))
        if p.is_file()
    }
    _write(output, 'artifact-hashes.json', inventory)
    if checkpoint is not None:
        with checkpoint.open('xb') as stream:
            stream.write(canonical_json(count_only_checkpoint(summary)))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=Path, required=True)
    parser.add_argument('--paper', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--checkpoint', type=Path)
    arguments = parser.parse_args()
    run(arguments.workbook, arguments.paper, arguments.output, checkpoint=arguments.checkpoint)
