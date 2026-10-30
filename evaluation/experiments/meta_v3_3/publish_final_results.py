"""Render completed calculations, then create the compact benchmark presentation."""

# ruff: noqa: INP001, E501, RUF001, ISC004, PLR2004, PERF401, PLR0917 - Scientific report and fixed slide layout.

from __future__ import annotations

import hashlib
import importlib.metadata
import json
import math
from pathlib import Path

import matplotlib as mpl


mpl.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from evaluation.experiments.meta_v3_3 import prepare_selection_report as selection


OUTPUT = selection.OUTPUT
ROOT = selection.ROOT
FAMILIES = (
    'van_genuchten_mualem_v1',
    'van_genuchten_m1_v1',
    'brooks_corey_v1',
    'clapp1978_gradual_air_entry_v1',
)
NAMES = dict(zip(FAMILIES, ('VG–Mualem', 'VG m=1', 'Brooks–Corey', 'Clapp'), strict=True))
COLORS = dict(zip(FAMILIES, ('#3876a6', '#16766d', '#c58833', '#9368a0'), strict=True))
BASES = ('volumetric_water_content', 'gravimetric_water_content')
FOOTER = 'Exploratory benchmark, meta_v3.3; results are dataset-specific.'


def save_figure(fig: plt.Figure, name: str) -> None:
    for extension in ('png', 'svg'):
        fig.savefig(
            OUTPUT / 'figures' / f'{name}.{extension}',
            dpi=220,
            facecolor='white',
            metadata={'Date': None} if extension == 'svg' else None,
        )
    plt.close(fig)


def approximation_figure(aggregates: dict) -> None:
    plt.rcParams.update(
        {
            'font.size': 11,
            'axes.spines.top': False,
            'axes.spines.right': False,
            'svg.fonttype': 'none',
            'svg.hashsalt': 'ptfkit',
        }
    )
    fig, axes = plt.subplots(2, 2, figsize=(13.5, 5.4), gridspec_kw={'width_ratios': [1, 1.4]})
    for row, basis in enumerate(BASES):
        aggregate = aggregates[basis]
        summaries = {r['family']: r for r in aggregate['summary']}
        unit = 'm³/m³' if row == 0 else 'kg/kg'
        label = 'Volumetric' if row == 0 else 'Gravimetric'
        distributions = [aggregate['families'][f]['rmse_distribution'] for f in FAMILIES]
        ax = axes[row, 0]
        boxes = ax.boxplot(
            distributions,
            patch_artist=True,
            tick_labels=[NAMES[f] for f in FAMILIES],
            widths=0.55,
            flierprops={'markersize': 2},
            medianprops={'color': '#162c3d', 'linewidth': 1.5},
        )
        for patch, family in zip(boxes['boxes'], FAMILIES, strict=True):
            patch.set_facecolor(COLORS[family])
            patch.set_alpha(0.6)
        ax.set_yscale('log')
        ax.set_ylabel(f'Per-sample RMSE ({unit})')
        n = aggregate['summary'][0]['paired_relaxed_samples']
        ax.set_title(f'{label}: paired cohort N={n}', loc='left', fontsize=12, pad=20)
        for i, family in enumerate(FAMILIES, 1):
            r = summaries[family]
            ax.text(
                i,
                1.01,
                f'{r["strict_successes"]}/{r["eligible_samples"]}',
                transform=ax.get_xaxis_transform(),
                ha='center',
                fontsize=9,
                color=COLORS[family],
            )
        ax.yaxis.grid(visible=True, alpha=0.15)
        ax.tick_params(axis='x', labelsize=9)
        right = axes[row, 1]
        support_counts = []
        for family in FAMILIES:
            bins = [b for b in aggregate['families'][family]['bins'] if b['samples']]
            x = [
                math.log10(math.sqrt(b['pressure_min_kpa'] * b['pressure_max_kpa']))
                if b['pressure_min_kpa'] > 0
                else math.log10(b['pressure_max_kpa'] / 2)
                for b in bins
            ]
            center = [b['median_residual'] for b in bins]
            lower, upper = (
                [b['empirical_error_lower_95'] for b in bins],
                [b['empirical_error_upper_95'] for b in bins],
            )
            right.fill_between(x, lower, upper, color=COLORS[family], alpha=0.10)
            right.plot(
                x,
                center,
                marker='o',
                markersize=3,
                linewidth=1.6,
                color=COLORS[family],
                label=NAMES[family],
            )
            support_counts = [b['samples'] for b in bins]
        right.axhline(0, color='#4d5d68', linewidth=0.6)
        right.set_title('Median residual + empirical 95% residual range', loc='left', fontsize=12)
        right.set_ylabel(f'Fitted − observed ({unit})')
        right.set_xlabel('log₁₀ suction (kPa); zero included in first bin', fontsize=10)
        right.grid(visible=True, alpha=0.12)
        right.text(
            0.02,
            0.03,
            f'Observed bins: {min(support_counts)}–{max(support_counts)} samples/bin',
            transform=right.transAxes,
            fontsize=9,
        )
        if row == 0:
            right.legend(ncols=4, fontsize=8, loc='upper center', frameon=False)
    fig.tight_layout(rect=(0.005, 0.06, 0.995, 0.98), h_pad=2.1, w_pad=2)
    fig.text(
        0.02,
        0.015,
        'Fractions above boxes: strict successes / eligible. Relaxed descriptive success: 100% for every family.\n'
        'Same samples and observations within each basis. Ribbons describe fitted-point errors; sample-bootstrap summary intervals are in JSON.',
        fontsize=8.5,
    )
    save_figure(fig, 'figure_2_wrc_approximation')


def selection_slide_figure(data: dict) -> None:
    fig, axes = plt.subplots(1, 2, figsize=(13.4, 3.65))
    configurations = [
        ('volumetric_water_content', 'Volumetric WRC: N=17, 495 points', 'RMSE (m³/m³)'),
        ('ksat', 'Ksat common cohort: N=30', 'RMSE (log₁₀ Ksat)'),
    ]
    for ax, (basis, title, xlabel) in zip(axes, configurations, strict=True):
        analysis = data['analyses'][basis]
        records = analysis['primary']['ranking'][:5]
        labels = [
            r['function']
            .replace('calc_ptf_varallyai1982_', '')
            .replace('calc_ptf_ferrerjulia2004_', '')
            .replace('calc_ptf_', '')
            .replace('_', ' ')
            for r in records
        ]
        values = [r['rmse'] for r in records]
        colors = ['#16766d'] + ['#9eafbc'] * 4
        if basis != 'ksat':
            labels.append('Mean ensemble')
            values.append(analysis['ensembles']['mean']['rmse'])
            colors.append('#c58833')
        ax.barh(range(len(values)), values, color=colors, height=0.65)
        ax.set_yticks(range(len(values)), labels, fontsize=11)
        ax.invert_yaxis()
        for i, value in enumerate(values):
            ax.text(value + 0.015 * max(values), i, f'{value:.4f}', va='center', fontsize=11)
        ax.set_xlim(0, max(values) * 1.2)
        ax.set_title(title, loc='left', fontsize=13)
        ax.set_xlabel(xlabel, fontsize=12)
        ax.xaxis.grid(visible=True, alpha=0.15)
        ax.set_axisbelow(True)
    fig.tight_layout(rect=(0, 0.02, 1, 1), w_pad=2.5)
    save_figure(fig, 'figure_1_slide')


def fnum(value: float | None, digits: int = 5) -> str:
    return 'not estimable' if value is None else f'{value:.{digits}g}'


def final_report(data: dict) -> str:
    a = data['approximation']['aggregates']
    prior = selection.report(data['selection']).split('## Analytical WRC approximation:')[0]
    prior = prior.replace(
        '# ptfkit benchmark results — selection complete; approximation pending',
        '# ptfkit benchmark results',
    )
    start, end = prior.index('**Status:**'), prior.index('## Dataset and interpretation')
    prior = (
        prior[:start]
        + 'The existing PTF selection is retained. Four analytical families were independently fitted across the complete usable ledger; the presentation follows completed calculations and figures.\n\n'
        + prior[end:]
    )
    lines = [
        prior,
        '![PTF selection](figures/figure_1_ptf_selection.png)',
        '',
        '## Analytical approximation of experimental WRC',
        '',
        '**Preferred overall descriptive family: van Genuchten with m=1, separately for both water bases.** '
        'This is fitted-curve approximation, not independent PTF prediction. Of 142 samples, 141 have enough distinct pressures: '
        '76 volumetric and 65 gravimetric. One volumetric curve has too few points. All four families converge to nondegenerate descriptive fits on these same samples (564 eligible fits).',
        '',
        '| Basis | Family | Strict / eligible | Relaxed / eligible | Median RMSE | Median MAE | Median ΔAICc | Median ΔBIC |',
        '|---|---|---:|---:|---:|---:|---:|---:|',
    ]
    for basis in BASES:
        for r in a[basis]['summary']:
            lines.append(
                f'| {"Volume (m³/m³)" if basis == BASES[0] else "Mass (kg/kg)"} | {NAMES[r["family"]]} | {r["strict_successes"]}/{r["eligible_samples"]} | {r["relaxed_successes"]}/{r["eligible_samples"]} | {r["median_rmse"]:.6f} | {r["median_mae"]:.6f} | {fnum(r["median_delta_aicc"], 4)} | {fnum(r["median_delta_bic"], 4)} |'
            )
    lines.extend(
        [
            '',
            'RMSE/MAE medians use the identical relaxed-success sample/point cohorts within each basis. '
            'Strict acceptance is a sensitivity analysis, not a filter silently applied to the plotted distributions. '
            'VG m=1 also leads median RMSE on the strict common cohorts. ΔAICc/ΔBIC are differences from the best family on each sample; '
            'their medians do not identify the same criterion as median RMSE. In particular, the simpler Clapp model is competitive under volumetric AICc.',
            '',
            'The sample-bootstrap median paired RMSE difference (VG m=1 minus VG–Mualem) is:',
            '',
        ]
    )
    for basis in BASES:
        comparison = next(
            r for r in a[basis]['paired_comparisons'] if r['comparator'] == FAMILIES[0]
        )
        interval = comparison['bootstrap_median_difference_95']
        lines.append(
            f'- {"Volume" if basis == BASES[0] else "Mass"}: {comparison["median_paired_rmse_difference"]:.6f}, bootstrap 95% [{interval[0]:.6f}, {interval[1]:.6f}] (N={comparison["samples"]}).'
        )
    lines.extend(
        [
            '',
            'These are exploratory within-dataset, sample-level resamples. Source/profile dependence is not resolved, '
            'and the near-zero volumetric upper limit supports a modest advantage rather than universal superiority.',
            '',
            '### Acceptance and parameter diagnostics',
            '',
            'Each curve uses deterministic bounded multistart least squares in the original water unit, with ordered residual/saturated water coordinates. '
            'Brooks–Corey receives additional starts in every observed-pressure partition. Strict acceptance requires optimizer convergence, '
            'finite physical predictions, n>p, no critical imposed-bound hit, nonconstant fitted values, full column-normalized Jacobian rank and condition ≤10¹⁰. '
            'A residual-water value of zero is allowed as a physical nested model; it is still recorded as a boundary hit. Relaxed success retains boundary/conditioning flags '
            'and is used only to describe the achievable curve fit. Bounds and all rejection counts are explicit in JSON.',
            '',
            'Volumetric amplitudes are bounded by 1 m³/m³. Gravimetric amplitudes use the same registered shapes with an explicit homogeneous amplitude rescaling '
            'and a numerical upper box of 5 kg/kg; this is not a physical universal maximum. No density conversion or water-basis relabeling is used. '
            'The Clapp constraint 0<b<11.5 is preserved. The four registered families are the complete comparison roster.',
            '',
            'AICc/BIC count p curve parameters plus estimated residual variance (K=p+1). AICc is undefined for n≤K+1; '
            f'common AICc coverage is {a[BASES[0]]["summary"][0]["aicc_paired_samples"]} volume and {a[BASES[1]]["summary"][0]["aicc_paired_samples"]} mass samples. '
            'Information criteria are descriptive Gaussian approximations; iid residual regularity and boundary assumptions have not been validated. '
            'R² is stored descriptively and never determines acceptance. Conditional local-linear parameter intervals are retained for suitable strict, interior fits; '
            'they are not prediction intervals, and intervals crossing physical bounds are flagged.',
            '',
            'Typical strict interior VG m=1 fits (closest to the strict median RMSE among fits with intervals) are retained as private '
            '[volumetric](../../../../local/meta-final-benchmark/parameter-uncertainty-example-volumetric_water_content-van_genuchten_m1_v1.json) and '
            '[gravimetric](../../../../local/meta-final-benchmark/parameter-uncertainty-example-gravimetric_water_content-van_genuchten_m1_v1.json) examples, with parameters, units and conditional 95% intervals.',
            '',
            '### Pressure support, error and practical use',
            '',
            '| Basis | VG m=1 strict support min–max (kPa) | Central 5–95% support (kPa) | Strict samples |',
            '|---|---:|---:|---:|',
        ]
    )
    for basis in BASES:
        r = next(r for r in a[basis]['summary'] if r['family'] == FAMILIES[1])
        lines.append(
            f'| {"Volume" if basis == BASES[0] else "Mass"} | {r["strict_support_min_kpa"]:.5g}–{r["strict_support_max_kpa"]:.6g} | {r["strict_support_p05_kpa"]:.5g}–{r["strict_support_p95_kpa"]:.6g} | {r["strict_successes"]} |'
        )
    lines.extend(
        [
            '',
            'These are observed/support ranges with equal sample weighting for percentiles. They neither establish uniform performance throughout the interval '
            'nor authorize extrapolation. The very large dry-end pressures are retained workbook values and require source/protocol caution.',
            '',
            '| Basis / region (kPa) | Lowest numerical MAE family | Samples / points | Median AE | 95th percentile AE | Bias | Pairwise superiority established? |',
            '|---|---|---:|---:|---:|---:|---|',
        ]
    )
    for basis in BASES:
        for region in ('wet', 'intermediate', 'dry'):
            best = a[basis]['region_lowest_mae'][region]
            r = a[basis]['families'][best]['regions'][region]
            comparisons = a[basis]['region_paired_comparisons'][region]
            clear = len(comparisons) == 3 and all(
                c['bootstrap_mean_difference_95'][1] < 0 for c in comparisons
            )
            label = {'wet': '[0,10)', 'intermediate': '[10,1500)', 'dry': '[1500,∞)'}[region]
            lines.append(
                f'| {"Volume" if basis == BASES[0] else "Mass"}, {region} {label} | {NAMES[best]} | {r["samples"]} / {r["observations"]} | {r["median_absolute_error"]:.5f} | {r["p95_absolute_error"]:.5f} | {r["bias"]:+.5f} | {"All three exploratory paired intervals below zero" if clear else "No clear separation from all alternatives"} |'
            )
    lines.extend(
        [
            '',
            'Regional winners minimize equal-sample mean absolute error; median AE and tail errors describe the same cohort. '
            'Numerical regional differences alone do not justify model switching. Detailed results for every family/region and the paired bootstrap differences are in JSON. '
            'For VG m=1, volume errors are smallest in the wet region (<10 kPa; N=76, median AE 0.00344, 95th-percentile AE 0.01164 m³/m³). '
            'Mass errors are smallest in the observed dry region (≥1500 kPa; N=37, median AE 0.00477, 95th-percentile AE 0.01483 kg/kg). '
            'The per-family lowest observed MAE region and its N/tail error are recorded in CSV; regional support differs, so these are descriptive reliability patterns. '
            'Read practical trust from support, sample counts and regional errors together; no arbitrary absolute-error threshold or single confidence range is imposed.',
            '',
            '![Approximation families](figures/figure_2_wrc_approximation.png)',
            '',
            'The pressure ribbons are equal-sample empirical 2.5–97.5% fitted-point residual quantiles. Separately, 1,000 whole-sample bootstrap draws per bin '
            'provide 95% intervals for the median residual, bias, MAE and RMSE. All points of a selected sample stay together. '
            'These fitted-point residuals do not validate future-sample predictive coverage. The first pressure bin includes zero; it is displayed at a finite log coordinate. '
            'Parameter intervals, bootstrap uncertainty of summaries and observed pressure support are distinct quantities.',
            '',
            '## Deliverables and reproducibility',
            '',
            '- `ptf_selection_summary.csv`: 60 existing selection/ensemble rows, with explicit cohorts and units.',
            '- `wrc_approximation_summary.csv`: eight family/basis rows, success, paired errors, information criteria, boundary diagnostics and strict/relaxed support.',
            '- `benchmark_aggregates.json`: full original selection aggregates, new approximation aggregates, pressure bins, paired comparisons, policies and provenance.',
            '- `figures/`: standalone PNG/SVG Figure 1, Figure 2 and a larger-label Figure 1 variant for slides.',
            '- `ptfkit_benchmark_summary.pptx`: two main slides and one methods/support appendix.',
            '- Private `local/meta-final-benchmark/`: observation ledger, exclusions, all 568 sample/family decisions, parameters, point residuals and conditional parameter intervals. Original data are not committed.',
            '',
            'The workbook SHA-256 matches the pinned mapping and previous selection. Reconstructed observation counts, sample counts, basis totals and pressure extrema '
            'match the prior ledger; 92 stored fitted workbook parameters remain excluded from observations. Workbook-origin and E5 applicability limitations remain unchanged. '
            'No new PTF was unlocked or recalibrated, and Part A was not rerun.',
            '',
            'Reproduce with the source Python packages on `PYTHONPATH` (`targets/ptfkit-eval/src`, `targets/ptfkit-py/src` and the repository root): '
            '`python -m evaluation.experiments.meta_v3_3.final_approximation`; then '
            '`python -m evaluation.experiments.meta_v3_3.publish_final_results`. The rendering command only runs after the aggregate calculations exist. '
            'Exact runtime versions and reproduction source hashes are in JSON.',
            '',
            FOOTER,
            '',
        ]
    )
    return '\n'.join(lines)


def text_box(
    slide: object,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    size: float = 14,
    color: str = '253746',
    *,
    bold: bool = False,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.01)
    for index, line in enumerate(text.split('\n')):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = 'Aptos'
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(3)


def new_slide(presentation: Presentation, title: str, number: int) -> object:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    text_box(slide, 0.4, 0.22, 12.5, 0.5, title, 25, bold=True)
    text_box(slide, 0.4, 7.14, 12, 0.2, FOOTER, 9, '657782')
    text_box(slide, 12.5, 7.14, 0.3, 0.2, str(number), 9, '657782')
    return slide


def presentation(data: dict) -> None:
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.3333), Inches(7.5)
    slide = new_slide(p, 'Выбор педотрансферных функций для ОГХ и Ksat', 1)
    text_box(
        slide,
        0.42,
        0.81,
        12.4,
        0.3,
        'Ledger: 1,701 WRC points / 142 samples; 39 Ksat observations. Comparison cohorts are shown separately.',
        12,
    )
    cards = [
        ('Volumetric WRC', '0.0421', 'Varallyai chernozem B · N=17'),
        ('Gravimetric WRC', '0.0899', 'Chakraborty eq4 · N=8'),
        ('Ksat, maximum coverage', '0.718', 'Dane–Puckett · N=39 · log10 RMSE'),
        ('WRC mean ensemble', '−6.2%', 'Pooled gain; no robust held-out gain'),
    ]
    for i, (label, value, detail) in enumerate(cards):
        x = 0.45 + i * 3.2
        text_box(slide, x, 1.22, 3.05, 0.25, label, 11, '657782')
        text_box(slide, x, 1.48, 3.05, 0.43, value, 27, '16766D', bold=True)
        text_box(slide, x, 1.96, 3.05, 0.38, detail, 10)
    slide.shapes.add_picture(
        str(OUTPUT / 'figures/figure_1_slide.png'), Inches(0.15), Inches(2.4), width=Inches(13.0)
    )
    text_box(
        slide,
        0.5,
        6.08,
        12.3,
        0.78,
        'Varallyai B is the pooled volumetric candidate; Dane–Puckett maximizes Ksat coverage.\nKsat common-cohort winner: gleyic cambisol sand (RMSE 0.270). Different cohorts cannot be directly ranked.',
        15,
    )
    slide = new_slide(p, 'Аппроксимация экспериментальной ОГХ', 2)
    a = data['approximation']['aggregates']
    volume = next(r for r in a[BASES[0]]['summary'] if r['family'] == FAMILIES[1])
    mass = next(r for r in a[BASES[1]]['summary'] if r['family'] == FAMILIES[1])
    text_box(
        slide,
        0.42,
        0.85,
        12.4,
        0.55,
        f'VG m=1: lowest median RMSE in both bases — {volume["median_rmse"]:.4f} m³/m³ (N=76); {mass["median_rmse"]:.4f} kg/kg (N=65).',
        17,
        '16766D',
        bold=True,
    )
    slide.shapes.add_picture(
        str(OUTPUT / 'figures/figure_2_wrc_approximation.png'),
        Inches(0.2),
        Inches(1.48),
        width=Inches(12.9),
    )
    text_box(
        slide,
        0.45,
        6.72,
        12.4,
        0.36,
        'Volume: smallest errors below 10 kPa; regional VG differences are unclear. Mass: VG m=1 leads in all regions.\nObserved support is in the appendix; residual ribbons describe fitted-point errors, not future-sample prediction coverage.',
        10,
    )
    slide = new_slide(p, 'Methods, acceptance and observed pressure support', 3)
    text_box(slide, 0.5, 1.05, 6.1, 0.4, 'Strict vs descriptive success', 20, '16766D', bold=True)
    text_box(
        slide,
        0.5,
        1.65,
        6.0,
        2.2,
        'Strict: convergence; physical finite curve; n>p; no critical bound; full Jacobian rank; condition ≤10¹⁰.\nZero residual water is admissible and still flagged.\nDescriptive: converged physical curve; boundary and rank warnings retained.\nAll four families: 100% descriptive success on eligible samples.',
        16,
    )
    text_box(slide, 6.9, 1.05, 5.9, 0.4, 'VG m=1: observed support', 20, '16766D', bold=True)
    support_text = []
    for r, label in [(volume, 'Volume'), (mass, 'Mass')]:
        support_text.append(
            f'{label}: strict {r["strict_successes"]}/{r["eligible_samples"]}\nMin–max: {r["strict_support_min_kpa"]:.3g}–{r["strict_support_max_kpa"]:.3g} kPa\nCentral 5–95%: {r["strict_support_p05_kpa"]:.3g}–{r["strict_support_p95_kpa"]:.3g} kPa'
        )
    text_box(slide, 6.9, 1.65, 5.8, 2.3, '\n\n'.join(support_text), 15)
    text_box(slide, 0.5, 4.3, 12.2, 0.4, 'What the intervals mean', 20, '16766D', bold=True)
    text_box(
        slide,
        0.5,
        4.9,
        12.1,
        1.55,
        'Empirical 95% ribbons: fitted-point residual distribution with equal sample weights.\nBootstrap: whole-sample draws, 1,000 per pressure bin; summary intervals retained in JSON.\nParameter intervals: conditional local-linear estimates, separate from residual bands.\nMass basis: explicit homogeneous shape extension; no assumed bulk-density conversion.',
        16,
    )
    text_box(
        slide,
        0.5,
        6.58,
        12.1,
        0.28,
        'Limits: E5 workbook semantics; unknown methods; sparse profile IDs; source dependence; extreme dry-end pressures; exploratory selection.',
        10,
    )
    p.core_properties.title = 'ptfkit benchmark summary'
    p.core_properties.subject = 'Dataset-specific PTF selection and analytical WRC approximation'
    p.core_properties.author = 'ptfkit'
    p.save(OUTPUT / 'ptfkit_benchmark_summary.pptx')


def main() -> None:
    path = OUTPUT / 'benchmark_aggregates.json'
    data = json.loads(path.read_bytes())
    if 'aggregates' not in data['approximation']:
        message = 'Complete the full-ledger approximation calculations before publishing.'
        raise ValueError(message)
    selection.figure(data['selection'])
    selection_slide_figure(data['selection'])
    approximation_figure(data['approximation']['aggregates'])
    (OUTPUT / 'FINAL-RESULTS.md').write_text(final_report(data), encoding='utf-8')
    presentation(data)
    data['status'] = 'complete'
    data['runtime'] = {
        name: importlib.metadata.version(name)
        for name in (
            'numpy',
            'scipy',
            'openpyxl',
            'matplotlib',
            'python-pptx',
            'pyyaml',
            'jsonschema',
        )
    }
    sources = [
        Path(__file__),
        Path(__file__).with_name('final_approximation.py'),
        Path(__file__).with_name('selection_audit.py'),
        Path(__file__).with_name('prepare_selection_report.py'),
        ROOT / 'targets/ptfkit-eval/src/ptfkit_eval/families.py',
        ROOT / 'targets/ptfkit-eval/src/ptfkit_eval/adapters.py',
        ROOT / 'targets/ptfkit-eval/src/ptfkit_eval/normalization.py',
        ROOT / 'evaluation/datasets/meta_v3_3.yaml',
    ]
    data['reproduction_implementation_sha256'] = {
        p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in sources
    }
    data['presentation'] = {
        'status': 'complete',
        'slides': 3,
        'main_slides': 2,
        'appendix_slides': 1,
    }
    path.write_text(
        json.dumps(data, indent=2, sort_keys=True, allow_nan=False) + '\n', encoding='utf-8'
    )
    inventory = {
        p.relative_to(OUTPUT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(OUTPUT.rglob('*'))
        if p.is_file() and p.name != 'artifact-hashes.json'
    }
    (OUTPUT / 'artifact-hashes.json').write_text(
        json.dumps(inventory, indent=2, sort_keys=True) + '\n', encoding='utf-8'
    )


if __name__ == '__main__':
    main()
