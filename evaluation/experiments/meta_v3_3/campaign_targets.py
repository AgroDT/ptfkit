"""Measured-range target reconstruction and exact-identity density discovery."""

# ruff: noqa: INP001 - Experiment-specific retained-data context.

from __future__ import annotations

import math
from collections import defaultdict
from typing import TYPE_CHECKING

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from ptfkit_eval.documents import document_hash
from ptfkit_eval.families import evaluate_retention_family
from ptfkit_eval.harmonization import EvidenceTier, EvidenceValue, reconstruct_water
from ptfkit_eval.water import within_water_bounds

from evaluation.experiments.meta_v3_3 import exploratory_inputs as previous
from evaluation.experiments.meta_v3_3.campaign_inputs import density, evidence


if TYPE_CHECKING:
    from pathlib import Path


def discover_density(workbook: Path, rows: dict, mapping: dict) -> dict:  # noqa: C901 - Exact identity and origin gates.
    """Read every header-mapped duplicate view once, requiring exact source/sample keys."""
    master = next(s for s in mapping['worksheets'] if s['id'] == 'samples_master')
    header = next(c['source_header'] for c in master['columns'] if c['role'] == 'dry_bulk_density')
    source_header = next(c['source_header'] for c in master['columns'] if c['role'] == 'source')

    def normalize(value: object) -> str:
        return ''.join(str(value).casefold().split())

    candidates = defaultdict(list)
    book = load_workbook(workbook, read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(workbook, read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in book:
            values = iter(sheet.iter_rows(values_only=True))
            head = next(values)
            if 'sample_id' not in head or source_header not in head:
                continue
            columns = [i for i, h in enumerate(head) if normalize(h) == normalize(header)]
            if len(columns) != 1:
                continue
            identity_column, source_column = head.index('sample_id'), head.index(source_header)
            column = columns[0]
            formula_rows = formulas[sheet.title].iter_rows(min_row=2, values_only=True)
            for number, (value, formula_row) in enumerate(
                zip(values, formula_rows, strict=True), 2
            ):
                key = (value[source_column], value[identity_column])
                numeric = value[column]
                if isinstance(formula_row[column], str) and formula_row[column].startswith('='):
                    continue
                if (
                    key in rows
                    and isinstance(numeric, (int, float))
                    and not isinstance(numeric, bool)
                    and math.isfinite(numeric)
                    and numeric > 0
                ):
                    candidates[key].append(
                        (float(numeric), f'{sheet.title}!{get_column_letter(column + 1)}{number}')
                    )
    finally:
        book.close()
        formulas.close()
    result = {
        'linked_samples': 0,
        'conflicting_samples': 0,
        'rules': 'exact_source_and_sample_identity_no_borrowing',
    }
    for key, row in rows.items():
        if 'Q' in row.get('formulas', {}):
            continue
        available = candidates[key]
        if previous.raw_number(row, 'Q') is not None:
            continue
        if len({value for value, _ in available}) > 1:
            result['conflicting_samples'] += 1
        elif available:
            row['linked_density'] = EvidenceValue(
                available[0][0],
                EvidenceTier.APPROXIMATE,
                'same_sample_density_link_v1',
                tuple(sorted({cell for _, cell in available})),
                (
                    ('assumed_unit', 'gram_per_cubic_centimeter'),
                    ('identity_match', 'exact_source_and_sample'),
                    ('analytical_comparability', 'assumed'),
                ),
            )
            result['linked_samples'] += 1
    return result


def record_formulas(workbook: Path, rows: dict) -> None:
    """Retain formula origin, particularly to reject cached density averages."""
    by_sheet = defaultdict(dict)
    for row in rows.values():
        by_sheet[row['worksheet']][row['row']] = row
    book = load_workbook(workbook, read_only=True, data_only=False, keep_links=False)
    try:
        for name, selected in by_sheet.items():
            for number, values in enumerate(
                book[name].iter_rows(min_row=2, max_col=22, values_only=True), 2
            ):
                if number in selected:
                    selected[number]['formulas'] = {
                        get_column_letter(i): value
                        for i, value in enumerate(values, 1)
                        if isinstance(value, str) and value.startswith('=')
                    }
    finally:
        book.close()


def measured_series(dataset: dict, sample: dict) -> list[tuple[float, EvidenceValue]]:
    return [
        (
            o['suction_kpa'],
            EvidenceValue(
                o['value'],
                EvidenceTier.STRICT,
                'reviewed_observation_v1',
                tuple(o['source_cells']),
                (
                    ('observation_id', o['id']),
                    ('origin', 'measured'),
                    ('suction_kpa', o['suction_kpa']),
                    ('original_water_basis', dataset['target']),
                ),
            ),
        )
        for o in dataset['observations']
        if (o['source_id'], o['sample_id']) == (sample['source_id'], sample['sample_id'])
        and o['origin'] == 'measured'
        and o['review_status'] == 'confirmed'
        and not o['quality_flags']
    ]


def accepted_fit_index(artifact: dict, datasets: list[dict]) -> dict:
    """Accept only fitted VG/BC records whose experimental coordinates match exactly."""
    index = {}
    observations = {}
    for dataset in datasets:
        if dataset['target']['quantity'] != 'volumetric_water_content':
            continue
        for sample in dataset['samples']:
            observations[sample['source_id'], sample['sample_id']] = sorted(
                (p, previous.numeric(v), dict(v.details)['observation_id'])
                for p, v in measured_series(dataset, sample)
            )
    for row in artifact['results']:
        fit = row['fit']
        family = fit['plan']['family_id']
        key = (row['source_id'], row['sample_id'])
        if fit['status'] != 'fitted' or family not in (
            'van_genuchten_mualem_v1',
            'brooks_corey_v1',
        ):
            continue
        fitted = sorted(
            (p['suction_kpa'], p['water_content'], p['reference']) for p in fit['observations']
        )
        if fitted != observations.get(key):
            continue
        identity = (*key, family)
        if identity in index:
            message = 'Duplicate accepted fit identity.'
            raise ValueError(message)
        index[identity] = {
            'family': family,
            'parameters': {p['name']: p['value'] for p in fit['parameters']},
            'fit_sha256': document_hash(fit),
            'pressure_range_kpa': [fitted[0][0], fitted[-1][0]],
            'status': fit['status'],
        }
    return index


def reconstruct(  # noqa: PLR0911 - Explicit non-extrapolating reconstruction dispositions.
    dataset: dict,
    sample: dict,
    pressure: float,
    route: str,
    fits: dict,
    *,
    interpolate: bool = True,
) -> tuple[EvidenceValue | None, str | None]:
    points = measured_series(dataset, sample)
    if not points:
        return None, 'target_observations_absent'
    try:
        interpolated = reconstruct_water(
            points, pressure, water_quantity=dataset['target']['quantity']
        )
    except ValueError:
        return None, 'invalid_or_ambiguous_wrc'
    if interpolated is None:
        return None, 'target_pressure_outside_measured_range'
    if route == 'log_pressure_linear_interpolation_v1':
        if interpolated.tier == EvidenceTier.RECONSTRUCTED:
            if not interpolate:
                return None, 'target_requires_interpolation'
            return EvidenceValue(
                interpolated.value,
                EvidenceTier.RECONSTRUCTED,
                route,
                interpolated.sources,
                (
                    *interpolated.details,
                    (
                        'formula',
                        'theta = theta_lo + (theta_hi-theta_lo)*log10(p/p_lo)/log10(p_hi/p_lo)',
                    ),
                    ('original_water_basis', dataset['target']),
                ),
                interpolated.parents,
            ), None
        return EvidenceValue(
            interpolated.value,
            interpolated.tier,
            interpolated.policy,
            interpolated.sources,
            (
                *interpolated.details,
                ('requested_pressure_kpa', pressure),
                ('exact_observation_existed', True),
                ('original_water_basis', dataset['target']),
            ),
            interpolated.parents,
        ), None
    fit = fits.get((sample['source_id'], sample['sample_id'], route))
    if fit is None:
        return None, 'accepted_fit_absent_or_observation_mismatch'
    lower, upper = fit['pressure_range_kpa']
    if not lower <= pressure <= upper:
        return None, 'target_pressure_outside_measured_range'
    value = float(evaluate_retention_family(fit['family'], fit['parameters'], [pressure])[0])
    exact = next((previous.numeric(v) for p, v in points if p == pressure), None)
    return EvidenceValue(
        value,
        EvidenceTier.RECONSTRUCTED,
        'accepted_fit_target_v1',
        tuple(sorted({s for _, p in points for s in p.sources})),
        (
            ('accepted_fit', fit),
            ('requested_pressure_kpa', pressure),
            ('original_water_basis', dataset['target']),
            ('interpolation_alternative', previous.numeric(interpolated)),
            ('fit_minus_interpolation', value - previous.numeric(interpolated)),
            ('exact_observation', exact),
            ('origin', 'reconstructed'),
        ),
        tuple(p for _, p in points),
    ), None


def convert_target(
    value: EvidenceValue,
    original: dict,
    required: dict,
    row: dict | None,
    registration: dict,
    *,
    density_conversion: bool = True,
) -> tuple[EvidenceValue | None, str | None]:
    result = value
    if original['quantity'] != required['quantity']:
        rho = density(row, registration) if row else None
        if rho is None:
            return None, 'same_sample_density_missing'
        if not density_conversion:
            return None, 'density_conversion_policy_disabled'
        factor = registration['water_density_g_cm3'] / previous.numeric(rho)
        if original['quantity'] == 'gravimetric_water_content':
            factor = 1 / factor
        result = evidence(
            previous.numeric(value) * factor,
            'same_sample_water_basis_conversion_v1',
            (value, rho),
            factor=factor,
            original_contract=original,
            required_contract=required,
            water_density_g_cm3=registration['water_density_g_cm3'],
        )
    if original['material_basis'] != required['material_basis']:
        result = evidence(
            previous.numeric(result),
            'meta_target_material_v1',
            (result,),
            original_contract=original,
            required_contract=required,
        )
    if not within_water_bounds(previous.numeric(result), required['quantity']):
        return None, 'density_converted_target_outside_physical_bounds'
    return result, None


def density_audit(datasets: list[dict], rows: dict, registration: dict) -> list[dict]:
    groups = defaultdict(
        lambda: {
            'mass_basis_samples': 0,
            'mass_basis_points': 0,
            'same_sample_density_available': 0,
            'convertible': 0,
            'missing_density': 0,
        }
    )
    for dataset in datasets:
        if dataset['target']['quantity'] != 'gravimetric_water_content':
            continue
        for sample in dataset['samples']:
            group = groups[sample['source_id']]
            points = measured_series(dataset, sample)
            row = rows.get((sample['source_id'], sample['sample_id']))
            rho = density(row, registration) if row else None
            group['mass_basis_samples'] += 1
            group['mass_basis_points'] += len(points)
            group['same_sample_density_available'] += rho is not None
            group['missing_density'] += rho is None
            if rho and any(
                within_water_bounds(
                    previous.numeric(value)
                    * previous.numeric(rho)
                    / registration['water_density_g_cm3'],
                    'volumetric_water_content',
                )
                for _, value in points
            ):
                group['convertible'] += 1
    return [{'source_id': source, **values} for source, values in sorted(groups.items())]
