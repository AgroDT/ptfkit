"""Coverage-first second-chance pass over exactly the three requested groups."""

# ruff: noqa: INP001, E501, C901, PLR0912, PLR0915, PLR2004 - Explicit scientific audit.

from __future__ import annotations

import hashlib
import json
import math
import shutil
import warnings
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from ptfkit import get_manifest
from ptfkit_eval.documents import plain

from evaluation.experiments.meta_v3_3 import campaign_contracts as contracts
from evaluation.experiments.meta_v3_3 import expanded_predictors as predictors
from evaluation.experiments.meta_v3_3 import expanded_routes as routes
from evaluation.experiments.meta_v3_3 import model_selection as original
from evaluation.experiments.meta_v3_3 import workbook_semantic as semantic
from evaluation.experiments.meta_v3_3.final_approximation import json_write
from evaluation.experiments.meta_v3_3.replay_diagnostic_predictions import restored_datasets
from evaluation.experiments.meta_v3_3.summaries import catalog_support


ROOT = Path(__file__).resolve().parents[3]
PUBLIC = Path(__file__).parent / 'final_results'
OUTPUT = ROOT / 'local/meta-expanded-benchmark'
BASES = ('volumetric_water_content', 'gravimetric_water_content', 'ksat')


def read(path: Path) -> Any:  # noqa: ANN401 - JSON artifacts contain object and array roots.
    return json.loads(path.read_bytes())


def eligible_catalog(baseline: dict) -> tuple[set[str], set[str]]:
    included = set().union(*(set(baseline['selection']['analyses'][b]['coverage']) for b in BASES))
    eligible = {
        r['function']
        for r in baseline['selection']['catalog']
        if r['function'] not in included
        and not r['executions']
        and (r['wrc_disposition'] != 'not_WRC_output' or r['ksat_targets'])
    }
    if len(eligible) != 119:
        message = 'The frozen three-group roster must contain 64+17+38 functions.'
        raise ValueError(message)
    return included, eligible


def build_context() -> tuple[dict, list[dict], list[dict], dict]:
    workbook = ROOT / 'local/meta_v3.3.xlsx'
    manifest = plain(get_manifest())
    if not isinstance(manifest, dict):
        message = 'The native manifest must be an object.'
        raise TypeError(message)
    context: dict = {
        'root': ROOT,
        'datasets': restored_datasets(workbook),
        'catalog': contracts.load_catalog(ROOT),
        'manifest': manifest,
        'old_registration': read(Path(__file__).with_name('exploratory-policies.json')),
    }
    context['support'] = catalog_support(context['manifest'])
    book = load_workbook(workbook, data_only=True, keep_links=False)
    formulas = load_workbook(workbook, data_only=False, keep_links=False)
    try:
        audit, observations, rows = original.rescan(context, workbook, book, formulas)
        # Scan all textual cells/comments for requested alias families; retain hits
        # as evidence, never assign a numeric value without a same-record column.
        terms = (
            'cec',
            'cation',
            'altitude',
            'precip',
            'slope',
            'gravel',
            'stoniness',
            'gmd',
            'median sand',
            'particle density',
            'porosity',
            'nitrogen',
        )
        audit['expanded_alias_scan'] = [
            {
                'sheet': s.title,
                'cell': c.coordinate,
                'text': str(c.value),
                'comment': c.comment.text if c.comment else None,
            }
            for s in book
            for cells in s
            for c in cells
            if any(
                t in (str(c.value) + ' ' + (c.comment.text if c.comment else '')).casefold()
                for t in terms
            )
        ]
    finally:
        book.close()
        formulas.close()
    return context, observations, rows, audit


def matched_row(observation: dict, rows: list[dict]) -> dict | None:
    candidates = [
        r
        for r in rows
        if (r['source_id'], r['sample_id']) == (observation['source_id'], observation['sample_id'])
    ]
    if observation['basis'] == 'ksat':
        candidates = [
            r
            for r in candidates
            if r['properties'].get('ksat') and observation['id'] in r['properties']['ksat'].sources
        ]
    elif observation.get('predictor_workbook_row'):
        candidates = [
            r
            for r in candidates
            if f'{r["worksheet"]}!{r["row"]}' == observation['predictor_workbook_row']
        ]
    else:
        candidates = [r for r in candidates if r['sheet_index'] == 0]
    return candidates[0] if len(candidates) == 1 else None


def run_predictions() -> None:
    baseline_path = OUTPUT / 'baseline/benchmark_aggregates.json'
    if not baseline_path.exists():
        (OUTPUT / 'baseline').mkdir(parents=True)
        for path in PUBLIC.iterdir():
            if path.is_file():
                shutil.copyfile(path, OUTPUT / 'baseline' / path.name)
    baseline = read(baseline_path)
    if (
        hashlib.sha256((ROOT / 'local/meta_v3.3.xlsx').read_bytes()).hexdigest()
        != baseline['selection']['input_sha256']['workbook']
    ):
        message = 'Workbook differs from the frozen comparison.'
        raise ValueError(message)
    included, eligible = eligible_catalog(baseline)
    context, observations, rows, workbook_audit = build_context()
    catalog = {f['id']: f for f in context['catalog']}
    isolated = {
        r['function']
        for r in baseline['selection']['catalog']
        if r['function'] in eligible
        and r['wrc_disposition'] == 'isolated_pressure_outputs_not_full_WRC'
        and not (r['wrc_targets'] or r['ksat_targets'])
    }
    registered = routes.registry(context['catalog'], isolated, context)
    json_write(OUTPUT / 'family-adapter-registry.json', registered)
    json_write(PUBLIC / 'expanded_wrc_registry.json', registered)
    samples = {
        (s['source_id'], s['sample_id']): s for d in context['datasets'] for s in d['samples']
    }
    failures = defaultdict(Counter)
    examples = defaultdict(list)
    native_records, added, point_scores = [], [], []
    metadata = {}
    cache = {}

    def native(fid: str, row: dict) -> tuple[dict, dict]:
        key = fid, row['worksheet'], row['row']
        if key in cache:
            result = cache[key]
            if isinstance(result, Exception):
                raise result
            return result
        function = catalog[fid]
        sample = samples.get((row['source_id'], row['sample_id']), {'features': []})
        values, proofs, reasons = predictors.resolve(function, sample, row, context)
        # Never satisfy measured-water inputs from the experimental target curve.
        if reasons:
            failures[fid].update(reasons)
            error = ValueError('unresolved_inputs')
            cache[key] = error
            raise error
        try:
            with warnings.catch_warnings():
                warnings.simplefilter('ignore', RuntimeWarning)
                output = contracts.native_outputs(function, values, {})
        except (ValueError, TypeError, ArithmeticError) as error:
            failures[fid]['native_evaluation:' + type(error).__name__ + ':' + str(error)] += 1
            if len(examples[fid]) < 3:
                examples[fid].append(
                    {
                        'worksheet': row['worksheet'],
                        'row': row['row'],
                        'inputs': values,
                        'error': str(error),
                    }
                )
            cache[key] = error
            raise
        cache[key] = output, proofs
        return output, proofs

    def target_value(fid: str, target: dict, row: dict) -> tuple[float, dict]:
        output, proofs = native(fid, row)
        unit = (
            'mass_fraction' if target['quantity'].startswith('gravimetric') else 'volume_fraction'
        )
        value = contracts.convert_value(output[target['field']], target['unit'], unit)
        if not math.isfinite(value) or value < 0 or (unit == 'volume_fraction' and value > 1):
            message = 'nonphysical_native_water_output'
            raise ValueError(message)
        return value, proofs

    obs_groups = defaultdict(list)
    for observation in observations:
        if observation['basis'] != 'ksat':
            obs_groups[
                observation['source_id'], observation['sample_id'], observation['basis']
            ].append(observation)
    # Evaluate family members once per exact predictor record, retaining native outputs.
    for family in registered['families'].values():
        fid = family['id']
        metadata[fid] = {
            'route_type': 'discrete_family_WRC',
            'members': sorted({m['function'] for m in family['members']}),
            'source': family['source'],
            'native_basis': family['basis'],
            'mechanism': 'Fixed-pressure family reconstruction',
        }
        for identity, group in sorted(obs_groups.items()):
            matched = [matched_row(o, rows) for o in group]
            unique = {(r['worksheet'], r['row']) for r in matched if r is not None}
            if len(unique) != 1 or any(r is None for r in matched):
                continue
            row = matched[0]
            if row is None:
                continue
            coordinates = {}
            native_sources = []
            for member in family['members']:
                member_id, target = member['function'], member['target']
                try:
                    value, proofs = target_value(member_id, target, row)
                except (ValueError, TypeError, ArithmeticError) as error:
                    if str(error) != 'unresolved_inputs':
                        failures[member_id][str(error)] += 1
                    continue
                pressure = target['pressure_kpa']
                if pressure in coordinates:
                    message = 'Family registry mixes alternative outputs at the same pressure.'
                    raise ValueError(message)
                coordinates[pressure] = value
                native_sources.append(member_id)
                native_records.append(
                    {
                        'family': fid,
                        'function': member_id,
                        'source_id': identity[0],
                        'sample_id': identity[1],
                        'basis': family['basis'],
                        'pressure_kpa': pressure,
                        'predicted': value,
                        'predictor_record': f'{row["worksheet"]}!{row["row"]}',
                        'predictor_evidence': proofs,
                    }
                )
                # Secondary point benchmark reconstructs observations only inside
                # each measured curve's support, never outside it or across samples.
                try:
                    observed = routes.interpolate_native(
                        [(o['pressure_kpa'], o['observed']) for o in group], pressure
                    )
                    prediction, conversion = predictors.water_convert(
                        value, family['basis'], identity[2], row
                    )
                    point_scores.append(
                        {
                            **group[0],
                            'id': f'point:{identity!r}:{pressure:.12g}',
                            'function': member_id,
                            'observed': observed,
                            'predicted': prediction,
                            'pressure_kpa': pressure,
                            'route_type': 'point_target',
                            'target_reconstruction': {
                                'method': 'exact or log-pressure linear interpolation',
                                'source_observation_ids': [o['id'] for o in group],
                            },
                            'water_basis_conversion': conversion,
                            'water_conditioned': False,
                            'minimum_tier': 'E5',
                        }
                    )
                except (ValueError, ArithmeticError):
                    pass
            if len(coordinates) < 3:
                for member in family['members']:
                    failures[member['function']]['insufficient_native_pressure_points'] += 1
                continue
            nonmonotone = any(
                coordinates[b] > coordinates[a] + 1e-12
                for a, b in zip(sorted(coordinates)[:-1], sorted(coordinates)[1:], strict=True)
            )
            for o in group:
                try:
                    value = routes.interpolate_native(list(coordinates.items()), o['pressure_kpa'])
                    value, conversion = predictors.water_convert(
                        value, family['basis'], o['basis'], row
                    )
                    added.append(
                        {
                            **o,
                            'function': fid,
                            'predicted': value,
                            'route_type': 'discrete_family_WRC',
                            'water_conditioned': False,
                            'minimum_tier': 'E5',
                            'native_pressure_support_kpa': sorted(coordinates),
                            'native_members_available': sorted(set(native_sources)),
                            'native_member_functions_used': sorted(
                                {
                                    m['function']
                                    for m in family['members']
                                    if m['target']['pressure_kpa']
                                    in routes.supporting_pressures(
                                        list(coordinates), o['pressure_kpa']
                                    )
                                }
                            ),
                            'nonmonotone_native_predictions': nonmonotone,
                            'water_basis_conversion': conversion,
                        }
                    )
                except (ValueError, ArithmeticError):
                    pass
    # Parameter-derived routes use predicted anchors only. Saxton helper inversions
    # compose with the source's existing base PTF instead of measured water inputs.
    for fid, adapter in registered['adapters'].items():
        parent = adapter.get('parent_function', fid)
        metadata[fid] = {
            'route_type': 'parameter_derived_WRC',
            'members': sorted({fid, parent}),
            'source': catalog[fid]['source'],
            'native_basis': 'volumetric_water_content',
            'mechanism': 'Retention-parameter adapters',
            'adapter': adapter,
        }
        for o in observations:
            if o['basis'] == 'ksat':
                continue
            row = matched_row(o, rows)
            if row is None:
                failures[fid]['ambiguous_or_missing_predictor_record'] += 1
                continue
            try:
                output, proofs = native(parent, row)
                if parent != fid:
                    output = {**output, 'theta_s': 1 - output['normal_density'] / 2.65}
                value = routes.parameter_curve(fid, output, o['pressure_kpa'])
                value, conversion = predictors.water_convert(
                    value, 'volumetric_water_content', o['basis'], row
                )
                added.append(
                    {
                        **o,
                        'function': fid,
                        'predicted': value,
                        'route_type': 'parameter_derived_WRC',
                        'water_conditioned': False,
                        'minimum_tier': 'E5',
                        'predictor_evidence': proofs,
                        'water_basis_conversion': conversion,
                        'predicted_parameters': output,
                    }
                )
            except (ValueError, TypeError, ArithmeticError) as error:
                failures[fid][str(error)] += 1
    # Retry original executable WRC/Ksat routes after recovering actual predictors.
    failed_ids = {
        r['function']
        for r in baseline['selection']['catalog']
        if r['function'] in eligible and (r['wrc_targets'] or r['ksat_targets'])
    }
    for fid in sorted(failed_ids):
        function = catalog[fid]
        water, conductivity = original.candidate_targets(function, context)
        metadata.setdefault(
            fid,
            {
                'route_type': 'continuous'
                if water and 'axis_input' in water[0]
                else 'parameter_derived_WRC',
                'members': [fid],
                'source': function['source'],
                'mechanism': 'Predictor recovery',
            },
        )
        for o in observations:
            targets = conductivity if o['basis'] == 'ksat' else water
            if not targets or (fid in registered['adapters'] and o['basis'] != 'ksat'):
                continue
            target = targets[0]
            row = matched_row(o, rows)
            if row is None:
                failures[fid]['ambiguous_or_missing_predictor_record'] += 1
                continue
            sample = samples.get((o['source_id'], o['sample_id']), {'features': []})
            axis = {target['axis_input']} if 'axis_input' in target else set()
            values, proofs, reasons = predictors.resolve(function, sample, row, context, axis)
            if reasons:
                failures[fid].update(reasons)
                continue
            try:
                if axis:
                    binding = next(
                        b for b in function['inputs'] if b['name'] == target['axis_input']
                    )
                    axis_value = contracts.axis_value({**target, 'pressure_kpa': o['pressure_kpa']})
                    if not within_domain(binding, values, axis_value):
                        failures[fid]['pressure_outside_declared_domain'] += 1
                        continue
                with warnings.catch_warnings():
                    warnings.simplefilter('ignore', RuntimeWarning)
                    value = (
                        semantic.prediction(
                            function, {**target, 'pressure_kpa': o['pressure_kpa']}, values, context
                        )
                        if o['basis'] == 'ksat'
                        or len(water) == 1
                        or axis
                        or target.get('adapter') in original.CURVES
                        else original.discrete_prediction(
                            function, water, o['pressure_kpa'], values, context
                        )
                    )
                conversion = None
                if o['basis'] != 'ksat':
                    value, conversion = predictors.water_convert(
                        value, target['quantity'], o['basis'], row
                    )
                if (
                    not math.isfinite(value)
                    or value < 0
                    or (o['basis'] == 'ksat' and value == 0)
                    or (o['basis'] == 'volumetric_water_content' and value > 1)
                ):
                    failures[fid]['nonpositive_or_nonfinite_or_nonphysical_output'] += 1
                    continue
                added.append(
                    {
                        **o,
                        'function': fid,
                        'predicted': value,
                        'route_type': metadata[fid]['route_type']
                        if o['basis'] != 'ksat'
                        else 'Ksat',
                        'water_conditioned': False,
                        'minimum_tier': 'E5',
                        'predictor_evidence': proofs,
                        'water_basis_conversion': conversion,
                    }
                )
            except (ValueError, TypeError, ArithmeticError) as error:
                failures[fid]['native_or_adapter:' + type(error).__name__ + ':' + str(error)] += 1
                if len(examples[fid]) < 3:
                    examples[fid].append(
                        {
                            'worksheet': row['worksheet'],
                            'row': row['row'],
                            'pressure_kpa': o['pressure_kpa'],
                            'inputs': values,
                            'error': str(error),
                        }
                    )
    old = [
        p
        for p in read(ROOT / 'local/meta-diagnostic-replay/predictions.json')
        if not p['water_conditioned']
    ]
    for point in old:
        fid = point['function']
        function = catalog[fid]
        metadata.setdefault(
            fid,
            {
                'route_type': 'continuous'
                if point['adapter'].startswith('native_scalar')
                else 'parameter_derived_WRC',
                'members': [fid],
                'source': function['source'],
                'mechanism': 'Previously scored',
            },
        )
        point['route_type'] = metadata[fid]['route_type'] if point['basis'] != 'ksat' else 'Ksat'
    # A function may have multiple output fields at one pressure; a secondary
    # function/pressure/sample observation is retained exactly once.
    point_scores = list(
        {
            (p['function'], p['basis'], p['source_id'], p['sample_id'], p['pressure_kpa']): p
            for p in point_scores
        }.values()
    )
    all_points = old + added
    if any(
        p['function'] == 'calc_ptf_rawls1982_full_wrc' or p.get('water_conditioned')
        for p in all_points
    ):
        message = 'Conditional predictions must never enter independent scoring.'
        raise ValueError(message)
    json_write(OUTPUT / 'predictions.json', all_points)
    json_write(OUTPUT / 'point-target-predictions.json', point_scores)
    json_write(OUTPUT / 'native-member-predictions.json', native_records)
    json_write(OUTPUT / 'candidate-metadata.json', metadata)
    json_write(OUTPUT / 'failures.json', {k: dict(v) for k, v in failures.items()})
    json_write(OUTPUT / 'failure-examples.json', dict(examples))
    json_write(OUTPUT / 'workbook-reaudit.json', workbook_audit)
    json_write(
        OUTPUT / 'scope.json',
        {
            'previously_scored': sorted(included),
            'reconsidered': sorted(eligible),
            'old_prediction_count': len(old),
            'new_prediction_count': len(added),
            'point_prediction_count': len(point_scores),
        },
    )
    print(  # noqa: T201 - Command result.
        f'Expanded predictions: {len(old)} baseline + {len(added)} new; {len(point_scores)} secondary points.',
        flush=True,
    )


def within_domain(binding: dict, values: dict, value: float) -> bool:
    return predictors.inputs.legacy_domain(value, binding) and predictors.within_named_domain(
        binding.get('domain') or '', {**values, binding['name']: value, 'value': value}
    )


if __name__ == '__main__':
    run_predictions()
