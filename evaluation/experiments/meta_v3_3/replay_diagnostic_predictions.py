"""Recover diagnostic forecasts from the pinned workbook and persisted review.

This is a diagnostic replay, not a replacement for the historical run inventory.
The unavailable frozen resource directories and source PDF are not fabricated.
The persisted source review supplies the reviewed classes and profile identities.
"""

# ruff: noqa: INP001, E501 - Standalone scientific artifact recovery.

from __future__ import annotations

import copy
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from ptfkit import get_manifest
from ptfkit_eval import (
    export_normalized_retention_dataset,
    load_dataset_mapping,
    load_retention_dataset,
    normalize_retention_workbook,
)
from ptfkit_eval.documents import plain

from evaluation.experiments.meta_v3_3 import campaign_contracts as contracts
from evaluation.experiments.meta_v3_3 import model_selection as selection
from evaluation.experiments.meta_v3_3.class_comparison import add_class_coordinates
from evaluation.experiments.meta_v3_3.final_approximation import json_write
from evaluation.experiments.meta_v3_3.summaries import catalog_support


ROOT = Path(__file__).resolve().parents[3]
PUBLIC = Path(__file__).parent / 'final_results'
OUTPUT = ROOT / 'local/meta-diagnostic-replay'


def restored_datasets(workbook: Path) -> list[dict]:
    """Reapply the recorded review to the original canonical normalization."""
    normalized = normalize_retention_workbook(
        workbook, load_dataset_mapping(ROOT / 'evaluation/datasets/meta_v3_3.yaml')
    )
    datasets = [
        plain(
            export_normalized_retention_dataset(
                normalized,
                dataset_id=f'meta-v3-3-reviewed-{name}',
                material_basis='unspecified',
                water_quantity=quantity,
            ).dataset.document
        )
        for name, quantity in (
            ('volume', 'volumetric_water_content'),
            ('mass', 'gravimetric_water_content'),
        )
    ]
    review_path = ROOT / 'evaluation/datasets/meta_v3_3.clapp-source-review.json'
    review = json.loads(review_path.read_bytes())
    canonical = load_retention_dataset(datasets[0])
    if canonical.sha256 != review['canonical_dataset_sha256']:
        msg = 'Canonical volume dataset differs from the persisted source review.'
        raise ValueError(msg)
    reviewed = copy.deepcopy(datasets[0])
    selected = {(r['source_id'], r['sample_id']): r for r in review['selected_samples']}
    reviewed['samples'] = []
    for sample in datasets[0]['samples']:
        identity = sample['source_id'], sample['sample_id']
        if identity not in selected:
            continue
        decision = selected[identity]
        updated = copy.deepcopy(sample)
        updated['profile_id'] = decision['profile_id']
        updated['features'] = [
            {
                'id': 'texture',
                'kind': 'categorical',
                'value': decision['texture'],
                'review_status': 'confirmed',
                'origin': 'measured',
                'category_system': decision['category_system'],
                'source_file': review['paper']['id'],
                'source_cells': decision['source_locations'],
            }
        ]
        reviewed['samples'].append(updated)
    if len(reviewed['samples']) != review['eligible_count']:
        msg = 'Persisted reviewed identities are not all present.'
        raise ValueError(msg)
    reviewed['observations'] = [
        o for o in reviewed['observations'] if (o['source_id'], o['sample_id']) in selected
    ]
    reviewed['dataset_id'] = 'meta-v3-3-panina-gray-forest-clapp-reviewed'
    reviewed['target']['material_basis'] = 'mineral_soil'
    reviewed['method'] = (
        'Reviewed Panina 2015 mineral gray forest profile; Table 17 texture classes.'
    )
    reviewed['source_files'].extend(
        [
            {key: review['paper'][key] for key in ('id', 'filename', 'sha256')},
            {
                'id': 'clapp_source_review',
                'filename': review_path.name,
                'sha256': hashlib.sha256(review_path.read_bytes()).hexdigest(),
            },
        ]
    )
    reviewed = add_class_coordinates(reviewed)
    for dataset in datasets:
        for field in ('samples', 'observations'):
            dataset[field] = [
                r for r in dataset[field] if (r['source_id'], r['sample_id']) not in selected
            ]
    return [*datasets, reviewed]


def main() -> None:
    workbook = ROOT / 'local/meta_v3.3.xlsx'
    # Publication files evolve; replay is pinned to the historical checkpoint.
    frozen = json.loads(
        (ROOT / 'evaluation/datasets/meta_v3_3.model-selection-summary.json').read_bytes()
    )
    digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
    if digest != frozen['input_sha256']['workbook']:
        msg = 'Workbook differs from the final benchmark.'
        raise ValueError(msg)
    if OUTPUT.exists():
        msg = 'Diagnostic replay output must be new.'
        raise ValueError(msg)
    manifest = plain(get_manifest())
    if manifest != json.loads((ROOT / 'targets/ptfkit-py/src/ptfkit/_manifest.json').read_bytes()):
        msg = 'Installed and generated manifests differ.'
        raise ValueError(msg)
    context = {
        'root': ROOT,
        'datasets': restored_datasets(workbook),
        'manifest': manifest,
        'catalog': contracts.load_catalog(ROOT),
        'support': catalog_support(manifest),
        'old_registration': json.loads(
            Path(__file__).with_name('exploratory-policies.json').read_bytes()
        ),
    }
    book = load_workbook(workbook, data_only=True, keep_links=False)
    formulas = load_workbook(workbook, data_only=False, keep_links=False)
    try:
        audit, observations, rows = selection.rescan(context, workbook, book, formulas)
    finally:
        book.close()
        formulas.close()
    points, decisions = selection.predict(context, observations, rows)
    analyses = {}
    for basis in ('volumetric_water_content', 'gravimetric_water_content', 'ksat'):
        analyses[basis] = selection.analyze(
            [p for p in points if p['basis'] == basis and not p['water_conditioned']],
            ksat=basis == 'ksat',
        )
    conditioned = [p for p in points if p['water_conditioned']]
    for basis in sorted({p['basis'] for p in conditioned}):
        analyses[basis + '_water_conditioned'] = selection.analyze(
            [p for p in conditioned if p['basis'] == basis]
        )
    for basis in ('volumetric_water_content', 'gravimetric_water_content'):
        analyses[basis + '_reviewed_targets_only'] = selection.analyze(
            [
                p
                for p in points
                if p['basis'] == basis and not p['water_conditioned'] and not p.get('assumptions')
            ]
        )
    # Keep the exact historical scoring implementation and independently compare
    # every published analysis, including compatible cohorts and held-out results.
    public = selection.public_summary({'analyses': analyses})['analyses']
    if public != frozen['analyses']:
        OUTPUT.mkdir(parents=True)
        json_write(OUTPUT / 'replayed-analyses.json', public)
        msg = 'Replayed analyses differ from the final benchmark; inspect before rendering.'
        raise ValueError(msg)
    OUTPUT.mkdir(parents=True)
    for name, value in (
        ('predictions', points),
        ('observations', observations),
        ('audit', audit),
        ('decisions', decisions),
        ('analyses', analyses),
    ):
        json_write(OUTPUT / f'{name}.json', value)
    json_write(
        OUTPUT / 'validation.json',
        {
            'workbook_sha256': digest,
            'all_published_analyses_identical': True,
            'historical_frozen_directories_available': False,
            'source_pdf_rechecked': False,
            'review_policy': 'Reapplied persisted source-review JSON; no new source interpretation.',
            'predictions': len(points),
        },
    )
    print(  # noqa: T201 - Command result.
        f'Recovered {len(points)} forecasts; all published analyses match exactly.', flush=True
    )


if __name__ == '__main__':
    main()
