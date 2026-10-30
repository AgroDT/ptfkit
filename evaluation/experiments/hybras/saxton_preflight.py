"""Run a target-free native numerical preflight; do not grant benchmark eligibility."""

# ruff: noqa: INP001 - Standalone raw-predictor computational audit.

from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
from collections import Counter
from importlib.metadata import version
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

import ptfkit_eval
from ptfkit import _ptfkit, get_manifest, saxton2006
from ptfkit_eval import (
    FunctionCurveCandidate,
    evaluate_candidate,
    load_evaluation_plan,
    load_retention_dataset,
)
from ptfkit_eval.documents import document_hash


_ARCHIVE_SHA256 = 'e9164c2f74e9ad75033b99248b0eee7ba426230150a499a6e6b3a285eebcf618'
_DATASET_SHA256 = '7ee2c78e07da2ffe97b987e776620ce275390b1e24939a63552f08f3039fd665'
_CLAY_LIMIT_PERCENT = 60
_ORGANIC_MATTER_LIMIT_PERCENT = 8
_PERCENT_SCALE = 100
_WET_BRANCH_MAXIMUM_KPA = 33


def evaluate_raw_predictors(sand: object, clay: object, organic_matter: object) -> dict:
    """Convert reported texture percentages only; OM remains a mass percentage."""
    inputs = {'sand_percent': sand, 'clay_percent': clay, 'organic_matter_percent': organic_matter}
    if not all(
        isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)
        for v in inputs.values()
    ):
        return {
            'raw_inputs': inputs,
            'status': 'not_executed',
            'reasons': ['missing_or_nonnumeric_predictor'],
        }
    sand, clay, organic_matter = float(sand), float(clay), float(organic_matter)
    reasons = []
    if (
        not 0 <= sand <= _PERCENT_SCALE
        or not 0 <= clay <= _PERCENT_SCALE
        or sand + clay > _PERCENT_SCALE
    ):
        reasons.append('invalid_texture_percentages')
    if clay > _CLAY_LIMIT_PERCENT:
        reasons.append('clay_above_60_percent')
    if not 0 <= organic_matter <= _ORGANIC_MATTER_LIMIT_PERCENT:
        reasons.append('organic_matter_outside_0_8_percent')
    if reasons:
        return {'raw_inputs': inputs, 'status': 'not_executed', 'reasons': reasons}
    result = saxton2006.calc_ptf_saxton2006(
        sand=sand / _PERCENT_SCALE, clay=clay / _PERCENT_SCALE, organic_matter=organic_matter
    )
    outputs = {name: float(value) for name, value in result._asdict().items()}
    if not all(math.isfinite(value) for value in outputs.values()):
        reasons.append('nonfinite_native_output')
        outputs = {name: value if math.isfinite(value) else None for name, value in outputs.items()}
    if all(math.isfinite(v) for v in [result.theta_1500, result.theta_33, result.theta_s]) and not (
        0 < result.theta_1500 < result.theta_33 < result.theta_s <= 1
    ):
        reasons.append('water_anchor_order_or_physical_bounds')
    if math.isfinite(result.air_entry_tension) and not (
        0 < result.air_entry_tension < _WET_BRANCH_MAXIMUM_KPA
    ):
        reasons.append('air_entry_outside_wet_branch')
    if any(
        math.isfinite(value) and value <= 0
        for value in [
            result.retention_a,
            result.retention_b,
            result.conductivity_lambda,
            result.normal_density,
        ]
    ):
        reasons.append('nonpositive_curve_or_density_parameter')
    return {
        'raw_inputs': inputs,
        'status': 'executed_preflight_only',
        'reasons': reasons,
        'native_outputs': outputs,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('archive', type=Path)
    parser.add_argument('dataset', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    raw = args.archive.read_bytes()
    document = json.loads(args.dataset.read_bytes())
    dataset = load_retention_dataset(document)
    if hashlib.sha256(raw).hexdigest() != _ARCHIVE_SHA256 or dataset.sha256 != _DATASET_SHA256:
        parser.error('The reviewed official archive and positive-head cohort are required.')
    selected = {int(s['sample_id'].removeprefix('hybras-horizon-')): s for s in document['samples']}
    with ZipFile(BytesIO(raw)) as archive:
        workbook = load_workbook(
            BytesIO(archive.read('HYBRAS_COMPLETE_V1.xlsx')), read_only=True, data_only=True
        )
        try:
            rows = [
                (number, row)
                for number, row in enumerate(
                    workbook['HYBRAS'].iter_rows(min_row=4, values_only=True), start=4
                )
                if row[0] in selected
            ]
        finally:
            workbook.close()
    if len(rows) != len(selected):
        parser.error('Every reviewed horizon must join exactly to the source workbook.')
    args.output.mkdir(parents=True, exist_ok=False)
    candidate = _candidate()
    counts = Counter()
    with (args.output / 'preflight.jsonl').open('xb') as stream:
        for number, row in rows:
            result = evaluate_raw_predictors(row[42], row[40], row[54])
            if result['status'] == 'executed_preflight_only':
                result.update(_curve_preflight(candidate, result['raw_inputs']))
                counts[result['curve_status']] += 1
            result.update(
                sample_id=selected[row[0]]['sample_id'],
                source_id=selected[row[0]]['source_id'],
                source_cells=[f'HYBRAS!AO{number}', f'HYBRAS!AQ{number}', f'HYBRAS!BC{number}'],
            )
            counts[result['status']] += 1
            for reason in result['reasons']:
                counts[reason] += 1
            stream.write(
                (
                    json.dumps(result, sort_keys=True, allow_nan=False, separators=(',', ':'))
                    + '\n'
                ).encode('utf-8')
            )
    summary = {
        'schema_version': 1,
        'scope': 'raw_predictor_numerical_preflight_not_external_accuracy',
        'function': 'calc_ptf_saxton2006',
        'dataset_sha256': dataset.sha256,
        'archive_sha256': _ARCHIVE_SHA256,
        'samples': len(rows),
        'counts': dict(counts),
        'native_eligibility': 'not_granted',
        'blockers': [
            'fine_earth_target_basis_not_bound',
            'predictor_protocol_and_origin_review_pending',
            'normal_density_and_applicability_review_pending',
        ],
        'script_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        'environment': {
            'python': platform.python_version(),
            **{name: version(name) for name in ('ptfkit', 'ptfkit-eval', 'numpy', 'openpyxl')},
        },
        'manifest_sha256': document_hash(get_manifest()),
        'evaluator_source_sha256': document_hash(
            {
                p.relative_to(Path(ptfkit_eval.__file__).parent).as_posix(): hashlib.sha256(
                    p.read_bytes()
                ).hexdigest()
                for p in sorted(Path(ptfkit_eval.__file__).parent.rglob('*'))
                if p.suffix in {'.py', '.json'}
            }
        ),
        'wrapper_sha256': hashlib.sha256(Path(saxton2006.__file__).read_bytes()).hexdigest(),
        'native_binary_sha256': hashlib.sha256(Path(_ptfkit.__file__).read_bytes()).hexdigest(),
        'preflight_sha256': hashlib.sha256(
            (args.output / 'preflight.jsonl').read_bytes()
        ).hexdigest(),
    }
    (args.output / 'summary.json').write_text(
        json.dumps(summary, indent=2, allow_nan=False) + '\n', encoding='utf-8', newline='\n'
    )


def _candidate() -> FunctionCurveCandidate:
    plan = load_evaluation_plan(
        {
            'schema_version': 1,
            'ptfkit_manifest_schema_version': 1,
            'candidates': [
                {
                    'id': 'saxton_preflight',
                    'kind': 'function_curve',
                    'function': 'calc_ptf_saxton2006',
                    'axis': {
                        'quantity': 'matric_potential',
                        'unit': 'kilopascal',
                        'convention': 'magnitude',
                        'transform': 'log10',
                    },
                    'target': {
                        'quantity': 'volumetric_water_content',
                        'unit': 'volume_fraction',
                        'material_basis': 'fine_earth',
                        'denominator_basis': 'bulk_volume',
                    },
                    'prediction': {'kind': 'model_adapter', 'id': 'saxton2006_h_theta'},
                }
            ],
        }
    )
    candidate = plan.candidates[0]
    if not isinstance(candidate, FunctionCurveCandidate):
        message = 'The native preflight requires the registered Saxton function candidate.'
        raise TypeError(message)
    return candidate


def _curve_preflight(candidate: FunctionCurveCandidate, inputs: dict) -> dict:
    try:
        curve = evaluate_candidate(
            candidate,
            [0, 33, 1500],
            {
                'sand': inputs['sand_percent'] / _PERCENT_SCALE,
                'clay': inputs['clay_percent'] / _PERCENT_SCALE,
                'organic_matter': inputs['organic_matter_percent'],
            },
        )
    except (ValueError, ArithmeticError) as error:
        return {
            'curve_status': 'curve_preflight_failed',
            'curve_error': f'{type(error).__name__}: {error}',
        }
    return {
        'curve_status': 'curve_preflight_passed',
        'curve_points': [
            {'pressure_kpa': p.axis_value, 'predicted_volume_fraction': p.value}
            for p in curve.points
        ],
    }


if __name__ == '__main__':
    main()
