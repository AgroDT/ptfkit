"""Import positive-head HYBRAS observations for retention approximation."""

# ruff: noqa: INP001 - Standalone reviewed-data adapter.

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from ptfkit_eval import RetentionSeriesPoint, diagnose_retention_series, load_retention_dataset
from ptfkit_eval.documents import canonical_json


ARCHIVE_SHA256 = 'e9164c2f74e9ad75033b99248b0eee7ba426230150a499a6e6b3a285eebcf618'
WORKBOOK_SHA256 = 'a185fc02e5a46d46bfaeea705c2e37a0f4459278819b72ccbef3955c0874f4c9'


def build_dataset(archive_bytes: bytes) -> tuple[dict, dict]:
    """Preserve reported cm heads; do not recover rounded nominal kPa labels."""
    if hashlib.sha256(archive_bytes).hexdigest() != ARCHIVE_SHA256:
        message = 'HYBRAS adapter requires the reviewed official archive checksum.'
        raise ValueError(message)
    with ZipFile(BytesIO(archive_bytes)) as archive:
        workbook_bytes = archive.read('HYBRAS_COMPLETE_V1.xlsx')
    if hashlib.sha256(workbook_bytes).hexdigest() != WORKBOOK_SHA256:
        message = 'HYBRAS workbook checksum differs from the reviewed source.'
        raise ValueError(message)
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    try:
        horizons = {
            row[0]: (number, row)
            for number, row in enumerate(
                workbook['HYBRAS'].iter_rows(min_row=4, values_only=True), start=4
            )
            if isinstance(row[0], int)
        }
        points = list(workbook['RAWRET_lines'].iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    samples = []
    metadata = []
    identities = {}
    for code, (number, row) in horizons.items():
        references = [row[36], row[37]]
        source_id = 'hybras-study-' + hashlib.sha256(canonical_json(references)).hexdigest()[:16]
        sample_id = f'hybras-horizon-{code}'
        identities[code] = (source_id, sample_id)
        samples.append(
            {
                'source_id': source_id,
                'sample_id': sample_id,
                'profile_id': f'hybras-profile-{row[5]}',
                'features': [],
            }
        )
        metadata.append(
            {
                'sample_id': sample_id,
                'source_id': source_id,
                'horizon_row': number,
                'original_profile_id': row[5],
                'references': references,
                'retention_method': row[124],
                'retention_path': row[127],
                'sample_type': row[129],
                'source_comments': list(row[38:40]),
            }
        )
    observations = []
    lineage = []
    omitted = []
    for number, row in enumerate(points, start=2):
        if row[0] is None:
            continue
        if row[3] == 0:
            omitted.append(
                {
                    'row': number,
                    'reason': 'derived_porosity'
                    if 'porosidade total=' in str(row[5])
                    else 'saturation_deferred',
                }
            )
            continue
        source_id, sample_id = identities[row[0]]
        observation_id = f'hybras-raw-row-{number}'
        ambiguous_pressure = '1atm=1020' in str(row[6])
        observations.append(
            {
                'id': observation_id,
                'source_id': source_id,
                'sample_id': sample_id,
                # Standard cmH2O conversion. Source rounding/conversion remains in lineage.
                'suction_kpa': row[3] * 0.0980665,
                'value': row[4],
                'origin': 'measured',
                'review_status': 'blocked' if ambiguous_pressure else 'confirmed',
                'quality_flags': ['pressure_conversion_review_required']
                if ambiguous_pressure
                else [],
                'source_file': 'hybras-workbook',
                'source_cells': [f'RAWRET_lines!A{number}:H{number}'],
            }
        )
        lineage.append(
            {
                'observation_id': observation_id,
                'reported_head_cm': row[3],
                'raw_curve_id': row[1],
                'measurement_flag': row[2],
                'comments': list(row[5:8]),
            }
        )
    document = {
        'schema_version': 1,
        'dataset_id': 'hybras-v1-positive-head-volume',
        'kind': 'measured',
        'target': {
            'quantity': 'volumetric_water_content',
            'unit': 'volume_fraction',
            'material_basis': 'unspecified',
            'denominator_basis': 'bulk_volume',
        },
        'method': (
            'HYBRAS raw positive-head drying observations of undisturbed samples; '
            'all zero-head points deferred. Reported rounded cm heads preserved; '
            'digitization and replicate averaging comments retained in provenance sidecar. '
            'No predictor bindings or external PTF eligibility granted.'
        ),
        'source_files': [
            {
                'id': 'hybras-workbook',
                'filename': 'HYBRAS_COMPLETE_V1.xlsx',
                'sha256': WORKBOOK_SHA256,
            }
        ],
        'samples': samples,
        'observations': observations,
    }
    dataset = load_retention_dataset(document)
    provenance = {
        'archive_sha256': ARCHIVE_SHA256,
        'canonical_dataset_sha256': dataset.sha256,
        'samples': metadata,
        'observations': lineage,
        'omitted_zero_head': omitted,
        'limitations': [
            'Material basis and predictor protocols not bound for native PTF scoring.',
            'Measured targets may be digitized or replicate averages; source comments preserved.',
            'Profile IDs group horizons, but do not establish independence between profiles.',
            'Study reference strings define source groups; calibration overlap not yet reviewed.',
            'Zero-head measurements deferred, including reported measured saturation.',
            (
                'Source comments declaring 1atm=1020cm require original-pressure review; '
                'do not repair by inference.'
            ),
        ],
    }
    return document, provenance


def confirmed_subset(document: dict) -> dict:
    """Keep only reviewed pressures for descriptive fitting, preserving the full import."""
    subset = json.loads(json.dumps(document))
    subset['dataset_id'] = 'hybras-v1-reviewed-positive-head-volume'
    subset['observations'] = [
        point
        for point in subset['observations']
        if point['review_status'] == 'confirmed' and not point['quality_flags']
    ]
    identities = {(point['source_id'], point['sample_id']) for point in subset['observations']}
    subset['samples'] = [
        sample
        for sample in subset['samples']
        if (sample['source_id'], sample['sample_id']) in identities
    ]
    subset['method'] += (
        ' Ambiguous atm-to-cm conversion records excluded; full import retains blockers.'
    )
    load_retention_dataset(subset)
    return subset


def summarize_import(document: dict, provenance: dict) -> dict:
    """Report noncorrecting physical findings and the exact imported population."""
    grouped = defaultdict(list)
    for point in document['observations']:
        grouped[point['sample_id']].append(
            RetentionSeriesPoint(point['suction_kpa'], point['value'], point['id'])
        )
    qualities = [diagnose_retention_series(points) for points in grouped.values()]
    return {
        'schema_version': 1,
        'dataset_id': document['dataset_id'],
        'canonical_dataset_sha256': load_retention_dataset(document).sha256,
        'archive_sha256': ARCHIVE_SHA256,
        'samples': len(document['samples']),
        'profiles': len({sample['profile_id'] for sample in document['samples']}),
        'source_reference_groups': len({sample['source_id'] for sample in document['samples']}),
        'positive_head_observations': len(document['observations']),
        'pressure_conversion_blocked_observations': sum(
            point['review_status'] == 'blocked' for point in document['observations']
        ),
        'confirmed_subset_sha256': load_retention_dataset(confirmed_subset(document)).sha256,
        'confirmed_subset_samples': len(confirmed_subset(document)['samples']),
        'confirmed_subset_observations': len(confirmed_subset(document)['observations']),
        'deferred_zero_head': dict(
            sorted(Counter(item['reason'] for item in provenance['omitted_zero_head']).items())
        ),
        'curves_with_physical_diagnostics': sum(bool(q.diagnostics) for q in qualities),
        'diagnostics': dict(
            sorted(Counter(item.code for q in qualities for item in q.diagnostics).items())
        ),
        'point_count_distribution': {
            str(count): total
            for count, total in sorted(Counter(q.point_count for q in qualities).items())
        },
        'native_ptf_eligibility': 'not_granted',
        'material_basis': 'unspecified',
        'pressure_policy': 'reported_cm_times_0.0980665_no_snap_to_approximate_kPa_labels',
        'parameter_uncertainty_status': 'not_qualified_for_empirical_coverage',
        'limitations': provenance['limitations'],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('archive', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    dataset, provenance = build_dataset(args.archive.read_bytes())
    args.output.mkdir(parents=True, exist_ok=False)
    for name, document in [
        ('dataset.json', dataset),
        ('fitting-dataset.json', confirmed_subset(dataset)),
        ('provenance.json', provenance),
        ('import-summary.json', summarize_import(dataset, provenance)),
    ]:
        with (args.output / name).open('x', encoding='utf-8', newline='\n') as stream:
            json.dump(document, stream, indent=2, ensure_ascii=False, allow_nan=False)
            stream.write('\n')


if __name__ == '__main__':
    main()
