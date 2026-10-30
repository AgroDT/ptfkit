"""Inventory the official HYBRAS archive without granting evaluation eligibility."""

# ruff: noqa: INP001 - Standalone source-data inventory.

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


def audit_archive(path: Path) -> dict:
    """Keep raw points, derived saturation and fitted tables distinguishable."""
    archive_bytes = path.read_bytes()
    with ZipFile(BytesIO(archive_bytes)) as archive:
        workbook_bytes = archive.read('HYBRAS_COMPLETE_V1.xlsx')
        workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
        try:
            horizons = [
                row
                for row in workbook['HYBRAS'].iter_rows(min_row=4, values_only=True)
                if isinstance(row[0], int)
            ]
            points = [
                row
                for row in workbook['RAWRET_lines'].iter_rows(min_row=2, values_only=True)
                if row[0] is not None
            ]
        finally:
            workbook.close()
    codes = [row[0] for row in horizons]
    if len(set(codes)) != len(codes):
        message = 'Duplicate horizon code'
        raise ValueError(message)
    if {row[0] for row in points} - set(codes):
        message = 'Raw point has no horizon metadata'
        raise ValueError(message)
    zero_points = [row for row in points if row[3] == 0]
    # The source repeats this saturation provenance on every point of a curve.
    # It applies to h=0 only; positive-head points must not be dropped with it.
    derived_zero = [row for row in zero_points if 'porosidade total=' in str(row[5])]
    return {
        'schema_version': 1,
        'status': 'inventory_only_not_evaluation_eligible',
        'source_url': 'https://rigeo.sgb.gov.br/items/edfbb5ff-604c-43c6-940b-18fc6fa53caf',
        'archive_sha256': hashlib.sha256(archive_bytes).hexdigest(),
        'workbook_sha256': hashlib.sha256(workbook_bytes).hexdigest(),
        'horizons': len(horizons),
        'distinct_profile_ids': len({row[5] for row in horizons if row[5] is not None}),
        'missing_profile_ids': sum(row[5] is None for row in horizons),
        'raw_table_points': len(points),
        'zero_head_points': len(zero_points),
        'zero_head_derived_from_porosity': len(derived_zero),
        'positive_head_points': sum(
            isinstance(row[3], (int, float)) and row[3] > 0 for row in points
        ),
        'raw_flags': dict(sorted(Counter(str(row[2]) for row in points).items())),
        'raw_head_unit': 'centimeter_water',
        'raw_water_unit': 'volume_fraction',
        'excluded_from_observation_import': [
            'HYDRAULIC_PROPS estimated water contents',
            'CVG_PARAMETERS contributed fit parameters',
            'derived van Genuchten parameters',
            'zero-head values derived from porosity',
        ],
        'pending_review': [
            'Per-point digitization, averaging and original pressure conversion provenance',
            'Profile identity and source-study grouping for nested validation',
            'Texture classification, material basis and predictor method compatibility',
            'Potential overlap with PTF calibration datasets',
            'Redistribution permissions before publishing raw observations',
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('archive', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    result = audit_archive(args.archive)
    with args.output.open('x', encoding='utf-8', newline='\n') as stream:
        json.dump(result, stream, indent=2, ensure_ascii=False, allow_nan=False)
        stream.write('\n')


if __name__ == '__main__':
    main()
