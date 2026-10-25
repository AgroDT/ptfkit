"""Recover explicitly reported controlled kPa regimes from rounded HYBRAS heads."""

# ruff: noqa: INP001 - Standalone source-reviewed data adapter.

from __future__ import annotations

import argparse
import hashlib
import io
import json
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from ptfkit_eval import load_retention_dataset
from ptfkit_eval.documents import canonical_json


ARCHIVE_SHA256 = 'e9164c2f74e9ad75033b99248b0eee7ba426230150a499a6e6b3a285eebcf618'
DATASET_SHA256 = '7ee2c78e07da2ffe97b987e776620ce275390b1e24939a63552f08f3039fd665'

# Exact source method cells, not a free-text numeric parser or nearest-grid rule.
PRESSURE_PROTOCOLS = {
    'mesa de tens\u00e3o na tens\u00e3o de 6 kPa e panela de press\u00e3o nas tens\u00f5es '
    'de 10, 33, 100 e 1500 kPa.': (6, 10, 33, 100, 1500),
    'mesa de tens\u00e3o para as tens\u00f5es de 1, 3, 5, 8 e 10 kPa e panela de press\u00e3o '
    'para as tens\u00f5es de 33, 100, 500 e 1500 kPa.': (1, 3, 5, 8, 10, 33, 100, 500, 1500),
    'mesa de tens\u00e3o para as tens\u00f5es de 2, 4 e 6 kPa e panela de press\u00e3o '
    'para as tens\u00f5es de 10, 33, 100, 500 e 1500 kPa.': (2, 4, 6, 10, 33, 100, 500, 1500),
    'unidades de suc\u00e7\u00e3o funis de Buchner para as suc\u00e7\u00f5es de 1, 2, 4, 6, 8 e '
    '10 kPa e panela de press\u00e3o para as suc\u00e7\u00f5es de 33, 60, 100, 500 e 1500 kPa.': (
        1,
        2,
        4,
        6,
        8,
        10,
        33,
        60,
        100,
        500,
        1500,
    ),
    'unidades de suc\u00e7\u00e3o funis de Buchner para as tens\u00f5es de 1, 2, 4, 6, 8 e '
    '10 kPa e panela de press\u00e3o para as tens\u00f5es de 33, 60, 100, 500 e 1500 kPa.': (
        1,
        2,
        4,
        6,
        8,
        10,
        33,
        60,
        100,
        500,
        1500,
    ),
}


def controlled_pressure(method: object, comment: object, head_cm: float) -> float | None:
    """Require exact protocol, original kPa declaration and unique rounded-head identity."""
    if (
        not isinstance(method, str)
        or method not in PRESSURE_PROTOCOLS
        or not isinstance(comment, str)
        or 'originalmente em kPa.' not in comment
        or '(1kPa=10.20 cmh2o)' not in comment
    ):
        return None
    matches = [p for p in PRESSURE_PROTOCOLS[method] if float(f'{p * 10.20:.3g}') == head_cm]
    return float(matches[0]) if len(matches) == 1 else None


def restore(document: dict, workbook: object) -> tuple[dict, dict, dict]:
    """Produce a separate coordinate view with exact target and source-cell checks."""
    sheet_rows = list(workbook['HYBRAS'].iter_rows(values_only=True))
    horizons = {row[0]: (number, row) for number, row in enumerate(sheet_rows[3:], start=4)}
    raw_rows = list(workbook['RAWRET_lines'].iter_rows(values_only=True))
    raw_points = {f'hybras-raw-row-{n}': (n, r) for n, r in enumerate(raw_rows[1:], start=2)}
    samples = {(s['source_id'], s['sample_id']): s for s in document['samples']}
    restored = []
    lineage = []
    exclusions: Counter[str] = Counter()
    for point in document['observations']:
        raw_number, raw = raw_points[point['id']]
        number, row = horizons[raw[0]]
        expected_source = (
            'hybras-study-' + hashlib.sha256(canonical_json([row[36], row[37]])).hexdigest()[:16]
        )
        identity = (expected_source, f'hybras-horizon-{raw[0]}')
        if (
            identity != (point['source_id'], point['sample_id'])
            or samples[identity]['profile_id'] != f'hybras-profile-{row[5]}'
            or raw[4] != point['value']
            or raw[3] * 0.0980665 != point['suction_kpa']
            or point['source_cells'] != [f'RAWRET_lines!A{raw_number}:H{raw_number}']
        ):
            message = 'Canonical observation differs from the reviewed source row.'
            raise ValueError(message)
        pressure = controlled_pressure(row[124], raw[6], raw[3])
        if pressure is None:
            exclusions['protocol_or_original_unit_not_confirmed'] += 1
            continue
        wide = [i for i in range(64, 95) if sheet_rows[2][i] == raw[3]]
        if len(wide) != 1 or row[wide[0]] != raw[4]:
            exclusions['wide_long_value_or_coordinate_mismatch'] += 1
            continue
        restored.append({**point, 'suction_kpa': pressure})
        lineage.append(
            {
                'observation_id': point['id'],
                'source_id': point['source_id'],
                'sample_id': point['sample_id'],
                'reported_head_cm': raw[3],
                'parent_suction_kpa': point['suction_kpa'],
                'controlled_pressure_kpa': pressure,
                'method_source_cell': f'HYBRAS!DU{number}',
                'wide_water_source_cell': f'HYBRAS!{get_column_letter(wide[0] + 1)}{number}',
                'raw_source_cells': point['source_cells'],
                'method': row[124],
                'pressure_comment': raw[6],
            }
        )
    identities = {(p['source_id'], p['sample_id']) for p in restored}
    subset = json.loads(canonical_json(document))
    subset['samples'] = [
        s for s in subset['samples'] if (s['source_id'], s['sample_id']) in identities
    ]
    subset['observations'] = restored
    subset['dataset_id'] = 'hybras-v1-explicit-controlled-kpa-volume'
    subset['method'] = (
        'Explicit source-controlled kPa regimes from exact HYBRAS method cells, original-unit '
        'comments and matching wide/long water values. Separate view; no nearest-pressure '
        'snapping or interpolation. Material and predictor eligibility remain unbound.'
    )
    review = {
        'schema_version': 1,
        'scope': 'source_coordinate_restoration_not_native_eligibility',
        'archive_sha256': ARCHIVE_SHA256,
        'parent_dataset_sha256': DATASET_SHA256,
        'script_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        'protocols': PRESSURE_PROTOCOLS,
        'lineage': lineage,
    }
    summary = {
        'schema_version': 1,
        'scope': review['scope'],
        'parent_dataset_sha256': DATASET_SHA256,
        'archive_sha256': ARCHIVE_SHA256,
        'script_sha256': review['script_sha256'],
        'samples': len(subset['samples']),
        'profiles': len({s['profile_id'] for s in subset['samples']}),
        'sources': len({s['source_id'] for s in subset['samples']}),
        'observations': len(restored),
        'pressure_counts_kpa': dict(
            sorted(Counter(str(p['suction_kpa']) for p in restored).items())
        ),
        'exclusions': dict(exclusions),
        'native_eligibility': 'not_granted',
        'limitations': [
            'Targets and identities stay unchanged in this separate pressure-coordinate view.',
            'Absent, ranged and unreviewed method lists do not permit original-unit inversion.',
            'Atmosphere-based, MPa-based and cm-based protocols are outside this first view.',
            'Material, predictor and function applicability still require independent binding.',
            'Historical fits and accuracy results belong to the unchanged rounded-head datasets.',
        ],
    }
    return subset, review, summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('archive', type=Path)
    parser.add_argument('dataset', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    raw = args.archive.read_bytes()
    if hashlib.sha256(raw).hexdigest() != ARCHIVE_SHA256:
        message = 'The reviewed HYBRAS archive checksum is required.'
        raise ValueError(message)
    dataset = load_retention_dataset(args.dataset)
    if dataset.sha256 != DATASET_SHA256:
        message = 'The reviewed positive-head dataset checksum is required.'
        raise ValueError(message)
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        workbook = openpyxl.load_workbook(
            io.BytesIO(archive.read('HYBRAS_COMPLETE_V1.xlsx')), read_only=True, data_only=True
        )
    try:
        subset, review, summary = restore(json.loads(args.dataset.read_bytes()), workbook)
    finally:
        workbook.close()
    review_raw = canonical_json(review)
    review_hash = hashlib.sha256(review_raw).hexdigest()
    subset['source_files'].append(
        {'id': 'hybras-pressure-review', 'filename': 'pressure-review.json', 'sha256': review_hash}
    )
    validated = load_retention_dataset(subset)
    summary['dataset_sha256'] = validated.sha256
    summary['review_sha256'] = review_hash
    args.output.mkdir(parents=True, exist_ok=False)
    (args.output / 'dataset.json').write_bytes(canonical_json(subset))
    (args.output / 'pressure-review.json').write_bytes(review_raw)
    (args.output / 'summary.json').write_bytes(canonical_json(summary))


if __name__ == '__main__':
    main()
