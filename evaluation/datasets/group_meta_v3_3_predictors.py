"""Partition all workbook sample records by explicitly available predictors.

Run with the prepared ptfkit-eval Python environment from the repository root.
"""

# ruff: noqa: INP001, C901, PLR0912, PLR0915, PLR2004, E501, PERF401 - Standalone workbook export with explicit layouts.

from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / 'local/meta_v3.3.xlsx'
OUTPUT = ROOT / 'local/meta-v3_3-all-predictor-subsets'
FEATURES = (
    ('texture_label', 'L'),
    ('physical_clay_lt_10_um', 'M'),
    ('clay', 'N'),
    ('silt_1_50_um', 'O'),
    ('sand_gt_50_um', 'P'),
    ('bulk_density', 'Q'),
    ('organic_carbon', 'R'),
    ('humus', 'S'),
    ('ph', 'T'),
    ('porosity', 'U'),
    ('particle_density', 'V'),
)


def numeric(value: object) -> bool:
    return isinstance(value, (float, int)) and not isinstance(value, bool) and math.isfinite(value)


def present(value: object) -> bool:
    if numeric(value):
        return True
    return isinstance(value, str) and value.strip() not in ('', '-', '?', 'NA', 'N/A')


def write_json(path: Path, value: object) -> None:
    def encode(item: object) -> str:
        if isinstance(item, (date, datetime)):
            return item.isoformat()
        message = f'Unsupported workbook value: {type(item).__name__}'
        raise TypeError(message)

    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False, default=encode) + '\n',
        encoding='utf-8',
    )


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    (OUTPUT / 'subsets').mkdir(exist_ok=True)
    book = load_workbook(INPUT, data_only=True, keep_links=False)
    formulas = load_workbook(INPUT, data_only=False, keep_links=False)
    records, audit = [], []
    # Each wide view remains independent, including the view lacking cached IDs.
    for sheet_index in (0, 3, 4, 5, 6):
        sheet = book.worksheets[sheet_index]
        clay_upper, previous_source = 1, None
        for row in range(2, sheet.max_row + 1):
            source, number = sheet[f'C{row}'].value, sheet[f'B{row}'].value
            shifted = sheet_index == 6
            offset = 3 if shifted else 0
            clay_column = 14 + offset
            if source != previous_source:
                clay_upper, previous_source = 1, source
            label = sheet.cell(row, clay_column).value
            if isinstance(label, str) and '< 2' in label:
                clay_upper = 2
            if not isinstance(source, str) or not source.strip():
                status = 'empty_or_no_source'
            elif number == 0.1:
                status = 'source_control'
            elif not numeric(number) or number <= 0 or int(number) != number:
                status = 'review_noninteger_or_missing_sample_number'
            else:
                status = 'sample_record'
            audit.append({'worksheet': sheet.title, 'row': row, 'status': status})
            if status != 'sample_record':
                continue
            values, evidence, derived = {}, {}, {}
            for feature_name, column in FEATURES:
                index = ord(column) - ord('A') + 1 + offset
                cell = sheet.cell(row, index)
                value = cell.value
                name = f'clay_lt_{clay_upper}_um' if feature_name == 'clay' else feature_name
                if name == 'texture_label':
                    available = isinstance(value, str) and present(value)
                else:
                    available = numeric(value)
                if not available:
                    continue
                proof = {'cell': f'{sheet.title}!{cell.coordinate}', 'value': value}
                formula = formulas[sheet.title][cell.coordinate]
                if formula.data_type == 'f':
                    derived[name] = {**proof, 'formula': formula.value}
                else:
                    values[name] = value
                    evidence[name] = proof
            record = {
                'record_id': f'sheet{sheet_index}_row{row}',
                'worksheet': sheet.title,
                'row': row,
                'source_id': source,
                'sample_number': number,
                'sample_id': f'{source}_{int(number)}',
                'cached_sample_id': sheet[f'A{row}'].value,
                'predictors': values,
                'predictor_evidence': evidence,
                'formula_predictors': derived,
                'record_kind': 'wide_sample',
            }
            records.append(record)

    # PSD rows are separate records: do not transfer their values to another view.
    sheet = book.worksheets[2]
    headers = {}
    for row in range(1, sheet.max_row + 1):
        source, number = sheet[f'B{row}'].value, sheet[f'C{row}'].value
        is_sample = isinstance(source, str) and present(source) and numeric(number) and number > 0
        if not is_sample:
            headers = {c.column: str(c.value) for c in sheet[row][3:] if present(c.value)}
            if row > 1:
                audit.append(
                    {'worksheet': sheet.title, 'row': row, 'status': 'psd_header_or_unresolved'}
                )
            continue
        values, evidence, derived = {}, {}, {}
        for cell in sheet[row][3:]:
            if not numeric(cell.value):
                continue
            # Literal block labels are retained; unlabeled fractions stay unassigned.
            name = (
                f'psd_column_{cell.column_letter}:{headers.get(cell.column, "unresolved_header")}'
            )
            proof = {'cell': f'{sheet.title}!{cell.coordinate}', 'value': cell.value}
            formula = formulas[sheet.title][cell.coordinate]
            if formula.data_type == 'f':
                derived[name] = {**proof, 'formula': formula.value}
            else:
                values[name], evidence[name] = cell.value, proof
        records.append(
            {
                'record_id': f'sheet2_row{row}',
                'worksheet': sheet.title,
                'row': row,
                'source_id': source,
                'sample_number': number,
                'sample_id': sheet[f'A{row}'].value,
                'cached_sample_id': sheet[f'A{row}'].value,
                'predictors': values,
                'predictor_evidence': evidence,
                'formula_predictors': derived,
                'record_kind': 'psd_sample',
            }
        )
        audit.append({'worksheet': sheet.title, 'row': row, 'status': 'sample_record'})

    # Retain long-sheet identities absent from every predictor-bearing view.
    known = {(r['source_id'], r['sample_number']) for r in records}
    long_only = defaultdict(list)
    sheet = book.worksheets[1]
    for row in range(2, sheet.max_row + 1):
        source, number = sheet[f'B{row}'].value, sheet[f'C{row}'].value
        if isinstance(source, str) and present(source) and numeric(number) and number > 0:
            status = 'long_observation_identity'
            if (source, number) not in known:
                long_only[source, number].append(row)
        else:
            status = 'long_unresolved_identity'
        audit.append({'worksheet': sheet.title, 'row': row, 'status': status})
    for (source, number), rows in sorted(long_only.items()):
        records.append(
            {
                'record_id': f'sheet1_identity_row{rows[0]}',
                'worksheet': sheet.title,
                'row': rows[0],
                'source_id': source,
                'sample_number': number,
                'sample_id': sheet[f'A{rows[0]}'].value,
                'cached_sample_id': sheet[f'A{rows[0]}'].value,
                'predictors': {},
                'predictor_evidence': {},
                'formula_predictors': {},
                'record_kind': 'long_only_identity',
                'observation_rows': rows,
            }
        )
    for row in range(2, book.worksheets[7].max_row + 1):
        audit.append(
            {'worksheet': book.worksheets[7].title, 'row': row, 'status': 'source_metadata'}
        )

    groups = defaultdict(list)
    for record in records:
        signature = sorted(record['predictors'])
        group = 'predictors_' + hashlib.sha256(json.dumps(signature).encode()).hexdigest()[:12]
        record['subset_id'] = group
        groups[group].append(record)
        # Raw source row preserves targets, context, formulas and excluded cells.
        sheet = book[record['worksheet']]
        record['raw_cells'] = {
            c.coordinate: {
                'value': c.value,
                'formula': formulas[sheet.title][c.coordinate].value
                if formulas[sheet.title][c.coordinate].data_type == 'f'
                else None,
            }
            for c in sheet[record['row']]
            if c.value is not None
        }

    summary, memberships = [], []
    workbook = Workbook()
    workbook.remove(workbook.active)
    index = workbook.create_sheet('Summary')
    index.append(['subset_id', 'n_records', 'n_source_sample_keys', 'predictors'])
    for group, members in sorted(groups.items()):
        names = sorted(members[0]['predictors'])
        row = {
            'subset_id': group,
            'n_records': len(members),
            'n_source_sample_keys': len({(r['source_id'], r['sample_number']) for r in members}),
            'predictors': '; '.join(names),
        }
        summary.append(row)
        index.append(list(row.values()))
        fields = [
            'record_id',
            'worksheet',
            'row',
            'source_id',
            'sample_number',
            'sample_id',
            'record_kind',
            *names,
        ]
        table = [
            {**{k: r[k] for k in fields if k not in names}, **r['predictors']} for r in members
        ]
        write_csv(OUTPUT / 'subsets' / f'{group}.csv', table, fields)
        write_json(OUTPUT / 'subsets' / f'{group}.json', {'predictors': names, 'records': members})
        target = workbook.create_sheet(group)
        target.append(fields)
        for item in table:
            target.append([item[k] for k in fields])
        memberships.extend(
            {
                k: r[k]
                for k in (
                    'record_id',
                    'subset_id',
                    'worksheet',
                    'row',
                    'source_id',
                    'sample_number',
                    'sample_id',
                    'record_kind',
                )
            }
            for r in members
        )
    for target in workbook:
        target.freeze_panes = 'A2'
        target.auto_filter.ref = target.dimensions
    workbook.save(OUTPUT / 'predictor_subsets.xlsx')
    write_csv(OUTPUT / 'subset_summary.csv', summary, list(summary[0]))
    write_csv(OUTPUT / 'sample_membership.csv', memberships, list(memberships[0]))
    write_csv(OUTPUT / 'row_audit.csv', audit, ['worksheet', 'row', 'status'])
    expected = sum(s.max_row - 1 for s in book)
    if len(audit) != expected or len({(a['worksheet'], a['row']) for a in audit}) != expected:
        message = 'Workbook rows were lost or duplicated in the audit.'
        raise ValueError(message)
    if not len(memberships) == len(records) == len({r['record_id'] for r in records}):
        message = 'Predictor records were lost or duplicated in the partition.'
        raise ValueError(message)
    review = []
    for item in audit:
        if item['status'] == 'review_noninteger_or_missing_sample_number':
            sheet = book[item['worksheet']]
            review.append(
                {
                    **item,
                    'cached_cells': {
                        c.coordinate: c.value for c in sheet[item['row']] if c.value is not None
                    },
                    'formulas': {
                        c.coordinate: c.value
                        for c in formulas[sheet.title][item['row']]
                        if c.data_type == 'f'
                    },
                }
            )
    write_json(OUTPUT / 'review_records.json', review)
    validation = {
        'workbook_sha256': hashlib.sha256(INPUT.read_bytes()).hexdigest(),
        'subsets': len(groups),
        'records': len(records),
        'source_sample_keys': len({(r['source_id'], r['sample_number']) for r in records}),
        'record_kinds': dict(Counter(r['record_kind'] for r in records)),
        'rows_audited': expected,
        'row_statuses': dict(Counter(a['status'] for a in audit)),
        'all_records_assigned_once': True,
    }
    write_json(OUTPUT / 'validation.json', validation)
    (OUTPUT / 'README.md').write_text(
        '# ptfkit workbook predictor subsets\n\n'
        f'{len(records)} records in {len(groups)} disjoint availability subsets. '
        'See predictor_subsets.xlsx, subset_summary.csv and subsets/*.csv (or JSON with cell evidence and raw source rows).\n\n'
        'All workbook sheets are audited, regardless of usable targets. Wide views and PSD rows remain separate records; '
        'repeated source/sample keys are not independent samples. Long-only identities have no assigned predictors. '
        'Review row_audit.csv and review_records.json before treating noninteger/missing-number records as samples. '
        'Source-control rows and formula-only empty padding are excluded. No target availability filter is applied.\n\n'
        'Signatures use explicitly stored numeric predictors and literal texture labels. Formula-derived predictors '
        'are retained separately in JSON and do not define availability. No imputation, PSD closure, organic conversions '
        'or cross-row unions are performed. PSD headers remain literal, including unresolved headers. '
        'Depth and other context remain in raw_cells; they are not part of the signature. '
        'Values retain workbook scales and units; zero and out-of-domain numeric values are retained. '
        'Availability does not establish physical compatibility or applicability of a PTF.\n\n'
        'For validation, keep duplicate views, all pressures and overlapping identities in the same fold; '
        'review source aliases and profile identities before splitting. These are availability groups, not train/test folds.\n\n'
        'Reproduce: targets/ptfkit-eval/.venv/Scripts/python.exe evaluation/datasets/group_meta_v3_3_predictors.py\n',
        encoding='utf-8',
    )
    book.close()
    formulas.close()
    print(json.dumps(validation, indent=2))  # noqa: T201 - Command result.


if __name__ == '__main__':
    main()
